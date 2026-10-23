#!/usr/bin/env python3
"""DPM Segregated web wiring: page, templates, command line, real run, download name."""
import base64
import io
import os
import tempfile
import time
from pathlib import Path
from types import SimpleNamespace

TMP = Path(tempfile.mkdtemp())
os.environ["DW_AUTH_DB"] = str(TMP / "users.db")
os.environ["DW_JOB_DIR"] = str(TMP / "jobs")
os.environ["DW_ADMIN_USER"] = "boss"
os.environ["DW_ADMIN_PASS"] = "boss-password-1"

from openpyxl import load_workbook  # noqa: E402

import app as A  # noqa: E402
import chunked  # noqa: E402

A.app.config["TESTING"] = True
client = A.app.test_client()
token = base64.b64encode(b"boss:boss-password-1").decode()
client.environ_base["HTTP_AUTHORIZATION"] = f"Basic {token}"
client.environ_base["HTTP_ACCEPT"] = "text/html"

assert A.TOOLS["segregated"][2] == ("dpm_segregated.py",)
assert "segregated" not in A.KVB_TOOLS
assert chunked._NAMA_TOOL["segregated"] == "Segregated"


def wait(location):
    job_id = location.rstrip("/").rsplit("/", 1)[-1]
    for _ in range(3000):
        state = A._read_state(job_id)
        if state and state["state"] in ("done", "failed"):
            return job_id, state
        time.sleep(0.01)
    raise AssertionError("job timeout")


def daily(day, extra=0):
    return (f"Login,Name,Currency,Deposit,Commission,Balance,Equity\n"
            f"1001,Ann,USD,{100 + extra},-1,{1000 + day},{1010 + day}\n"
            f"1002,Bob,JPY,10000,-100,{50000 + day},{50500 + day}\n").encode()


def fx(days):
    rows = ["date,Currency,rate"]
    for d in days:
        rows += [f"2026.09.{d:02d} 00:00:00,USD,", f"2026.09.{d:02d} 00:00:00,JPY,100"]
    return ("\n".join(rows) + "\n").encode()


# Page: DPM only, chunked always, FX input required, template links.
page = client.get("/tool/segregated")
assert page.status_code == 200
html = page.data.decode()
assert "var DIRECT = 0" in html
assert 'id="deals-file"' in html and "multiple required" in html
equity = html[html.index('id="equity-file"'):]
equity = equity[:equity.index(">")]
assert "required" in equity and ".csv" in equity, equity
assert "/tool/segregated/template/segregated" in html
assert "/tool/segregated/template/fx" in html
assert "Segregated" in client.get("/").data.decode()
kvb = client.get("/kvb/tool/segregated")
assert kvb.status_code in (302, 303, 404), kvb.status_code

# Templates download and the daily one runs through the engine as-is.
for nama in ("segregated", "fx"):
    r = client.get(f"/tool/segregated/template/{nama}")
    assert r.status_code == 200 and r.data[:2] == b"PK", nama
    (TMP / f"t-{nama}.xlsx").write_bytes(r.data)
assert client.get("/tool/segregated/template/mt5").status_code == 404
wb = load_workbook(TMP / "t-segregated.xlsx", read_only=True)
assert next(wb["Segregated"].iter_rows(values_only=True))[:4] == ("Date", "Login", "Name", "Currency")
wb.close()

# Validation: one file can never be daily + FX.
one = client.post("/tool/segregated", data={"file": (io.BytesIO(daily(1)), "2026-09-01.csv")},
                  content_type="multipart/form-data")
assert one.status_code in (400, 422) and b"Client Equity FX file is required" in one.data, one.status_code

# Command line built by the app (fake subprocess).
real_run = A.subprocess.run
commands = []


def fake_run(command, **_kwargs):
    commands.append(command)
    Path(command[command.index("-o") + 1]).write_bytes(b"PK fake xlsx")
    return SimpleNamespace(returncode=0, stdout="=== DPM Segregated DONE ===\nPERIOD: 2026-09\n", stderr="")


try:
    A.subprocess.run = fake_run
    r = client.post("/tool/segregated", data={"file": [
        (io.BytesIO(daily(1)), "2026-09-01.csv"),
        (io.BytesIO(daily(2)), "2026-09-02.csv"),
        (io.BytesIO(fx([1, 2])), "Client Equity - FX.csv"),
    ]}, content_type="multipart/form-data")
    assert r.status_code == 303, (r.status_code, r.data[:400])
    job_id, state = wait(r.headers["Location"])
    assert state["state"] == "done", state
    cmd = commands[-1]
    assert Path(cmd[1]).name == "dpm_segregated.py"
    assert "--delete-input" in cmd and "--allow-mt4" not in cmd
    assert sum(1 for c in cmd if c.endswith(".csv")) == 3
    assert Path(cmd[cmd.index("-o") + 1]).name == "Segregated.xlsx"
    assert state["name"] == "Segregated - Sep 2026.xlsx", state.get("name")
finally:
    A.subprocess.run = real_run

# Real run end to end.
r = client.post("/tool/segregated", data={"file": [
    (io.BytesIO(daily(1)), "2026-09-01.csv"),
    (io.BytesIO(daily(2, extra=5)), "2026-09-02.csv"),
    (io.BytesIO(daily(3)), "20260903.csv"),
    (io.BytesIO(fx([1, 2, 3])), "Client Equity - FX.csv"),
]}, content_type="multipart/form-data")
assert r.status_code == 303, (r.status_code, r.data[:400])
job_id, state = wait(r.headers["Location"])
assert state["state"] == "done", state
assert state["name"] == "Segregated - Sep 2026.xlsx"
assert "PERIOD: 2026-09" in state["log"]
job_dir = A._job_dir(job_id)
assert not list(job_dir.glob("*.csv")), list(job_dir.iterdir())
dl = client.get(f"/job/{job_id}/download")
assert dl.status_code == 200
assert "Segregated - Sep 2026.xlsx" in dl.headers.get("Content-Disposition", "")
assert "spreadsheetml" in dl.headers["Content-Type"]
out = TMP / "out.xlsx"
out.write_bytes(dl.data)
wb = load_workbook(out, read_only=True)
assert wb.sheetnames[:3] == ["Monthly Summary", "Daily", "Verifikasi"], wb.sheetnames
rows = list(wb["Monthly Summary"].iter_rows(values_only=True))
head = rows[0]
ann = next(r for r in rows[1:] if str(r[0]) == "1001")
bob = next(r for r in rows[1:] if str(r[0]) == "1002")
assert ann[head.index("Deposit")] == 305
assert ann[head.index("Balance")] == 1003 and ann[head.index("Equity")] == 1013
assert bob[head.index("Deposit (USD)")] == 300
assert bob[head.index("Balance (USD)")] == 500.03
assert len(list(wb["Daily"].iter_rows(values_only=True))) == 7
wb.close()

# Failure page shows the STOP from the engine (no FX file among the uploads).
r = client.post("/tool/segregated", data={"file": [
    (io.BytesIO(daily(1)), "2026-09-01.csv"),
    (io.BytesIO(daily(2)), "2026-09-02.csv"),
]}, content_type="multipart/form-data")
assert r.status_code == 303
job_id, state = wait(r.headers["Location"])
assert state["state"] == "failed" and "Client Equity FX" in (state["error"] or ""), state
assert client.get(f"/job/{job_id}").status_code == 422

print("OK test_segregated_web")
