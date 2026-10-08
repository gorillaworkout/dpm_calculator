#!/usr/bin/env python3
"""PL by Desk — monthly P&L by sales desk for DPM Malaysia.

FinanceOS settlement (PDF, or the team's converted Excel) supplies column B and
the equity balances. The Metabase workbook supplies the per-country totals, the
desk mapping and the targets. Nothing in those inputs is hardcoded.

    python3 pl_desk.py <financeos.pdf|xlsx> <metabase.xlsx> -o "PL by Desk.xlsx"

Refuses with a message and writes no workbook when the settlement or the
Metabase export fails the checks in validate_document / load_metabase.
"""
import argparse
import calendar
import re
import shutil
import subprocess
import sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CENT = Decimal("0.01")
TOLERANCE_PER_CHILD = Decimal("0.005")
VARIANCE_LIMIT = 1000
RAW_SHEET = "Raw Metabase"
MANUAL_REBATE = "Manual Rebate (Reward)"
MANUAL_INPUT_FLOOR = 8
MANUAL_INPUT_EXTRA = 4
NUM_FMT = '#,##0.00;[Red]\\(#,##0.00\\);\\-'
PCT_FMT = "0.00%"

# Company profile is the only place a second entity (KVB later) needs to land.
COMPANIES = {
    "dpm": {
        "legal_name": "DPM Malaysia Sdn. Bhd.",
        "currency": "CONSOLIDATED",
        "fallback_desk": "GEM",
    },
}

REQUIRED_LABELS = (
    "Closed P/L", "Floating P/L Changes", "Swap", "Swap (deals)", "Commissions",
    "CRM Swap Refund", "Rebates (calculated by system)", "Rebate",
    "MT5 Trading Account Changes", "Negative Balance Compensation", "Other Adjustment",
    "Rewards Wallet Changes", "Reward Other Promotions", "Rewards (Clients' Wallet)",
    "Manual Rebate (Reward)", "Currency Rate Diff(USD)", "Rounding Adjustment",
)
REQUIRED_SECTIONS = ("A", "B", "C")
KNOWN_SECTIONS = {"Trade", "Cash Movement", "Adjustment"}
# August 2026 vocabulary. A label that is not in this set is reported as new;
# it is still written under its parent. Manual-rebate descriptions are free
# text and are not matched against this set.
KNOWN_LABELS = {
    "Closed P/L", "Floating P/L Changes", "Swap", "Swap (deals)", "Commissions",
    "CRM Swap Refund", "Deposit to MT5 Trading Account", "Withdrawal MT5 Trading Account",
    "Withdrawal from Clients' Wallet",
    "Clients' Wallet transfer to MT5 trading account (test)",
    "Clients' Wallet transfer to MT5 trading account",
    "MT5 trading account received from Clients' Wallet",
    "Rebates (calculated by system)", "Rebate", "Rollback Rebate", "Test",
    "MT5 Trading Account Changes", "Profit Sharing", "Negative Balance Compensation",
    "Other Adjustment", "Adjustment: Inactive", "Adjustment: CRM-D",
    "Adjustment: Deposit Stage Credit Stage", "Adjustment: SO adj",
    "Adjustment: Written Off", "Dividend", "Rewards Wallet Changes", "Swap Free",
    "Reward TransferOut", "Reward Adjustment", "Reward Expiration",
    "Reward Other Promotions", "Welcome Bonus", "Switch Bonus", "Deposit Bonus",
    "Deposit Stage Credit", "Cash Back Bonus", "Loyalty Program",
    "Live Lucky Draw Promotion", "Adj: Welcome Bonus", "Adj: Trade Gold Win Gold",
    "Adj: Expiration", "Adj: Lunar New Year Promo", "Adj: Deduction of Deposit Credit Stage",
    "Others", "Rewards (Clients' Wallet)", "TransferOut", "TransferIn", "Deduction",
    "Manual Rebate (Reward)", "MT5 Transfer IN", "MT5 Transfer OUT",
    "Profit by sharing (Wallet's Profit Sharing)", "Currency Rate Diff(USD)",
    "System: FX translation of the opening balance", "Rounding Adjustment",
    "2026-08-21 [USD]",
}
# Yellow, and kept out of the Adjustment subtotal. They stay on the sheet so the
# layout still shows them, with column B blank.
EXCLUDED_FROM_SUBTOTAL = {
    "MT5 Transfer IN", "MT5 Transfer OUT",
    "Profit by sharing (Wallet's Profit Sharing)",
}
# Yellow leaf inside an otherwise calculated group.
YELLOW_LEAVES = {"Profit Sharing"}
METRIC_BY_LABEL = {
    "Closed P/L": ("closed_pnl_usd", 1),
    "Commissions": ("commission_usd", 1),
    "Rebate": ("rebate_usd", 1),
}
METRIC_FIELDS = (
    "closed_pnl_usd", "commission_usd", "rebate_usd",
    "deposit_usd", "withdrawal_usd", "ndp_usd", "lots",
)
REQUIRED_META_COLUMNS = (
    "data_date", "country_code", "deposit_usd", "withdrawal_usd", "ndp_usd",
    "lots", "closed_pnl_usd", "commission_usd", "rebate_usd",
)
TARGET_ALIASES = {
    "closed_pnl_usd": "closed_pnl_usd", "closed p/l": "closed_pnl_usd",
    "commission_usd": "commission_usd", "commissions": "commission_usd",
    "rebate_usd": "rebate_usd", "rebate": "rebate_usd",
    "deposit_usd": "deposit_usd", "deposit": "deposit_usd",
    "withdrawal_usd": "withdrawal_usd", "withdrawal": "withdrawal_usd",
    "ndp_usd": "ndp_usd", "ndp": "ndp_usd",
    "lots": "lots", "volume": "lots",
}

AMT = r"\(?-?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?\)?"
RE_AMT_END = re.compile(rf"^(?P<label>.*?)\s+(?P<amt>{AMT})$")
RE_SECTION = re.compile(r"^(?P<code>[A-Z])\.\)\s+(?P<name>.+?)$")
RE_SUBTOTAL = re.compile(rf"^Sub-Total:\s*(?P<amt>{AMT})$")
RE_EQUITY = re.compile(
    rf"^(?P<cat>Total|MT5|Wallet)\s+Previous:\s*(?P<p>{AMT})\s+Current:\s*(?P<c>{AMT})\s+Change:\s*(?P<d>{AMT})$"
)
RE_TOTAL_MOVE = re.compile(rf"^Total Movement:\s*(?P<amt>{AMT})$")
RE_DIFF = re.compile(rf"\(Diff\):\s*(?P<amt>{AMT})$")
RE_PERIOD = re.compile(r"(\d{4}-\d{2}-\d{2})\D+(\d{4}-\d{2}-\d{2})")
RE_EXCEL_SECTION = re.compile(r"^([A-Z])\.\)\s+(.+?)\s+Sub-Total$")

FILL_NAVY = PatternFill("solid", fgColor="17365D")
FILL_SECTION = PatternFill("solid", fgColor="1F4E78")
FILL_HEAD = PatternFill("solid", fgColor="5B9BD5")
FILL_META = PatternFill("solid", fgColor="002060")
FILL_GREEN = PatternFill("solid", fgColor="00B050")
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
FILL_ORANGE = PatternFill("solid", fgColor="FFC000")
FILL_DIFF = PatternFill("solid", fgColor="FFEB9C")
FILL_VARIANCE = PatternFill("solid", fgColor="FFC7CE")
FONT_WHITE = Font(color="FFFFFF", bold=True)
FONT_WHITE_TITLE = Font(color="FFFFFF", bold=True, size=14)
FONT_GREEN = Font(color="008000")
FONT_GREEN_B = Font(color="008000", bold=True)
FONT_MUTED = Font(color="666666", bold=True)
FONT_DIFF = Font(color="9C6500", bold=True)
FONT_VARIANCE = Font(color="9C0006")
FONT_BOLD = Font(bold=True)


class PlDeskError(Exception):
    """The run is refused. No workbook is written."""


def money(value):
    """Parse a FinanceOS amount. Brackets are negative. Result is a cent-rounded float."""
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    quant = Decimal(text.replace(",", "")).quantize(CENT, rounding=ROUND_HALF_UP)
    if negative:
        quant = -quant
    return float(quant)


def q2(value):
    return float(Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP))


def fmt_amt(value):
    if value is None:
        return ""
    number = q2(value)
    sign = "-" if number < 0 else ""
    return sign + f"{abs(number):,.2f}"


def put_text(cell, value):
    """Store text as text, including values Excel would otherwise treat as formulas."""
    if value is None:
        cell.value = None
        return
    text = str(value)
    cell.value = text
    if text[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"
        cell.quotePrefix = True


def fix_text(text):
    # PDFKit writes UTF-16 through a WinAnsi font. The arrow and the Diff label
    # come out garbled; amounts are unaffected.
    text = text.replace("!\u2019", "->").replace("!\x92", "->")
    return text


def read_pdf_lines(path):
    """Return (lines, engine). lines are {page, x0, text}. Raise when there is no text layer."""
    path = Path(path)
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None
    if pdfplumber is not None:
        try:
            with pdfplumber.open(path) as pdf:
                lines = []
                for page_number, page in enumerate(pdf.pages, 1):
                    for line in page.extract_text_lines(keep_blank_chars=True) or []:
                        text = fix_text(re.sub(r"\s{2,}", "  ", line["text"].strip()))
                        if text:
                            lines.append({"page": page_number, "x0": round(line["x0"], 1), "text": text})
        except Exception as exc:
            raise PlDeskError(
                "This PDF could not be read. If it is a scan, upload the original "
                f"FinanceOS PDF (it has a text layer) or the converted Excel. ({exc})"
            ) from exc
        if not lines:
            raise PlDeskError(
                "This PDF has no text layer. It looks scanned. Upload the original "
                "FinanceOS PDF, or the team's converted Excel workbook."
            )
        return lines, "pdfplumber"
    if not shutil.which("pdftotext"):
        raise PlDeskError(
            "pdfplumber is not installed, and pdftotext was not found. "
            "Install pdfplumber (pip). poppler-utils provides the optional pdftotext fallback."
        )
    try:
        result = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            capture_output=True, text=True, check=True,
        )
    except subprocess.CalledProcessError as exc:
        raise PlDeskError("pdftotext could not read this PDF.") from exc
    lines = []
    for page_number, page in enumerate(result.stdout.split("\f"), 1):
        for raw in page.splitlines():
            if raw.strip():
                lead = len(raw) - len(raw.lstrip(" "))
                text = fix_text(re.sub(r"\s{2,}", "  ", raw.strip()))
                lines.append({"page": page_number, "x0": float(lead), "text": text})
    if not lines:
        raise PlDeskError(
            "This PDF has no text layer. It looks scanned. Upload the original "
            "FinanceOS PDF, or the team's converted Excel workbook."
        )
    return lines, "pdftotext"


def _levels(lines):
    xs = []
    for x in sorted({line["x0"] for line in lines if line["text"].startswith("- ")}):
        if not xs or x - xs[-1] > 1:
            xs.append(x)
    return xs


def parse_pdf_lines(lines, engine="pdfplumber"):
    """Turn extracted lines into a settlement document. Does not apply the fail-safe checks."""
    if not lines or "reconciliation" not in lines[0]["text"].casefold():
        title = lines[0]["text"] if lines else ""
        raise PlDeskError(
            "Unexpected title "
            f"{title!r}. Expected a FinanceOS Trading Account Reconciliation."
        )
    levels = _levels(lines)

    def level_of(x):
        if not levels:
            raise PlDeskError("The PDF has no indented settlement lines to read.")
        return 1 + min(range(len(levels)), key=lambda i: abs(levels[i] - x))

    doc = {
        "engine": engine,
        "pages": max(line["page"] for line in lines),
        "header": {"title": lines[0]["text"]},
        "equity": {},
        "sections": [],
        "total_movement": None,
        "diff": None,
        "warnings": [],
    }
    section = None
    pending = None
    rate_idx = None
    for idx, line in enumerate(lines[1:], 1):
        text = line["text"]
        if pending is not None and (
            text.startswith("- ") or RE_SECTION.match(text) or RE_SUBTOTAL.match(text)
            or RE_TOTAL_MOVE.match(text)
        ):
            doc["warnings"].append(f"label without amount: {pending['label']!r}")
            pending = None
        if pending is not None:
            matched = RE_AMT_END.match(text)
            if matched and not text.startswith("- "):
                pending["label"] = (pending["label"] + " " + matched["label"]).strip()
                pending["amount"] = money(matched["amt"])
                pending["wrapped"] = True
                doc["warnings"].append(
                    f"wrapped label joined: '{pending['label'][:80]}' (p{line['page']})"
                )
                pending = None
                continue
            pending["label"] += " " + text
            continue
        period = RE_PERIOD.search(text) if text.startswith("Period:") else None
        if period:
            doc["header"]["period_from"], doc["header"]["period_to"] = period.group(1), period.group(2)
            continue
        if text.startswith("Currency:"):
            doc["header"]["currency"] = text.split(":", 1)[1].strip()
            continue
        if section is None and not doc["equity"] and idx == 1:
            doc["header"]["company"] = text.strip()
            continue
        if text.startswith("First day this currency"):
            doc["header"]["rate_note"] = text
            rate_idx = idx
            continue
        if rate_idx == idx - 1 and text != "Equity / Balance" and not RE_EQUITY.match(text):
            doc["header"]["rate_note"] = doc["header"].get("rate_note", "") + " " + text
            rate_idx = idx
            continue
        if text == "Equity / Balance":
            continue
        equity = RE_EQUITY.match(text)
        if equity:
            doc["equity"][equity["cat"]] = {key: money(equity[key]) for key in ("p", "c", "d")}
            continue
        section_match = RE_SECTION.match(text)
        if section_match:
            section = {
                "code": section_match["code"], "name": section_match["name"].strip(),
                "subtotal": None, "items": [],
            }
            doc["sections"].append(section)
            continue
        subtotal = RE_SUBTOTAL.match(text)
        if subtotal:
            if section is None:
                raise PlDeskError("A Sub-Total appears before its section.")
            section["subtotal"] = money(subtotal["amt"])
            continue
        total = RE_TOTAL_MOVE.match(text)
        if total:
            doc["total_movement"] = money(total["amt"])
            continue
        diff = RE_DIFF.search(text)
        if diff:
            doc["diff"] = money(diff["amt"])
            continue
        if text.startswith("- "):
            if section is None:
                raise PlDeskError(f"A settlement line appears before any section: {text}")
            body = text[2:]
            matched = RE_AMT_END.match(body)
            item = {
                "label": (matched["label"] if matched else body).strip(),
                "amount": money(matched["amt"]) if matched else None,
                "level": level_of(line["x0"]),
                "page": line["page"],
                "wrapped": False,
                "children": [],
            }
            section["items"].append(item)
            if matched is None:
                pending = item
            continue
        doc["warnings"].append(f"unparsed line p{line['page']}: {text!r}")
    if pending is not None:
        doc["warnings"].append(f"label without amount: {pending['label']!r}")
    for section in doc["sections"]:
        section["tree"] = _tree(section["items"])
    return doc


def _tree(items):
    roots, stack = [], []
    for item in items:
        while stack and stack[-1]["level"] >= item["level"]:
            stack.pop()
        (stack[-1]["children"] if stack else roots).append(item)
        stack.append(item)
    return roots


def walk(nodes, parents=()):
    for node in nodes:
        yield parents, node
        yield from walk(node["children"], parents + (node["label"],))


def parse_pdf(path, company="dpm"):
    lines, engine = read_pdf_lines(path)
    doc = parse_pdf_lines(lines, engine)
    doc["source"] = str(path)
    validate_document(doc, company)
    return doc


def _excel_amount(value, label):
    if isinstance(value, str) and value[:1] == "=":
        raise PlDeskError(
            "The converted FinanceOS workbook has formulas in the amount column. "
            "Paste values, then upload it again."
        )
    if value is None or value == "":
        raise PlDeskError(f"A settlement line has no amount: {label!r}.")
    if isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise PlDeskError(f"Amount for {label!r} is not a number.")
    try:
        return money(value) if isinstance(value, str) else q2(value)
    except Exception as exc:
        raise PlDeskError(f"Amount for {label!r} is not a number.") from exc


def parse_finance_excel(path, company="dpm"):
    """Read the team's manual FinanceOS conversion (values + indent) into the same document."""
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = _finance_excel_sheet(workbook)
        rows = []
        for cells in sheet.iter_rows():
            label_cell = cells[0] if cells else None
            rows.append((
                int(round(label_cell.alignment.indent or 0)) if label_cell is not None else 0,
                [cell.value for cell in cells[:4]],
            ))
    finally:
        workbook.close()
    doc = {
        "engine": "excel", "pages": None, "source": str(path),
        "header": {}, "equity": {}, "sections": [], "total_movement": None,
        "diff": None, "warnings": [],
    }
    title = next((vals[0] for _i, vals in rows if isinstance(vals[0], str) and "reconciliation" in vals[0].casefold()), None)
    if not title:
        raise PlDeskError(
            "This workbook is not a FinanceOS settlement. Expected a sheet titled "
            "Trading Account Reconciliation."
        )
    doc["header"]["title"] = str(title).strip()
    for _indent, vals in rows:
        label = str(vals[0]).strip() if vals[0] is not None else ""
        if label == "Company" and vals[1]:
            doc["header"]["company"] = str(vals[1]).strip()
        elif label == "Currency" and vals[1]:
            doc["header"]["currency"] = str(vals[1]).strip()
        elif label == "Period" and vals[1]:
            period = RE_PERIOD.search(str(vals[1]))
            if period:
                doc["header"]["period_from"], doc["header"]["period_to"] = period.group(1), period.group(2)
        elif label == "Rate note" and vals[1]:
            doc["header"]["rate_note"] = str(vals[1]).strip()
        elif label in ("Total", "MT5", "Wallet") and _is_number(vals[1]) and _is_number(vals[2]) and _is_number(vals[3]):
            doc["equity"][label] = {"p": q2(vals[1]), "c": q2(vals[2]), "d": q2(vals[3])}
    section = None
    for indent, vals in rows:
        label = str(vals[0]).strip() if vals[0] is not None else ""
        if not label:
            continue
        section_match = RE_EXCEL_SECTION.match(label)
        if section_match:
            section = {
                "code": section_match.group(1),
                "name": section_match.group(2).strip(),
                "subtotal": _excel_amount(vals[1], label),
                "items": [],
            }
            doc["sections"].append(section)
            continue
        if label == "Total Movement":
            doc["total_movement"] = _excel_amount(vals[1], label)
            section = None
            continue
        if label == "Diff" or label.endswith("(Diff)"):
            doc["diff"] = _excel_amount(vals[1], label)
            continue
        if section is not None and indent >= 1:
            section["items"].append({
                "label": label,
                "amount": _excel_amount(vals[1], label),
                "level": indent,
                "page": None,
                "wrapped": False,
                "children": [],
            })
    for section in doc["sections"]:
        section["tree"] = _tree(section["items"])
    validate_document(doc, company)
    return doc


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _finance_excel_sheet(workbook):
    ranked = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        blob = " ".join(
            str(cell.value) for row in sheet.iter_rows(max_row=20, max_col=2) for cell in row
            if cell.value is not None
        )
        if "reconciliation" in blob.casefold() or "Sub-Total" in blob:
            ranked.append(sheet)
    if not ranked:
        raise PlDeskError(
            "This workbook is not a FinanceOS settlement. Expected a sheet titled "
            "Trading Account Reconciliation."
        )
    return ranked[0]


def _company_profile(company):
    try:
        return COMPANIES[company]
    except KeyError as exc:
        raise PlDeskError(
            f"Company {company!r} is not configured. PL by Desk is available for DPM first."
        ) from exc


def _calendar_month(period_from, period_to):
    try:
        start = date.fromisoformat(period_from)
        end = date.fromisoformat(period_to)
    except (TypeError, ValueError) as exc:
        raise PlDeskError("The FinanceOS period is missing or not a pair of dates.") from exc
    last = calendar.monthrange(start.year, start.month)[1]
    if start.day != 1 or end != date(start.year, start.month, last):
        raise PlDeskError(
            f"The FinanceOS period {period_from} to {period_to} is not one calendar month."
        )
    return start


def validate_document(doc, company="dpm"):
    """Raise PlDeskError when the settlement cannot be trusted. Rounding stays a warning."""
    profile = _company_profile(company)
    header = doc["header"]
    title = header.get("title") or ""
    if "trading account reconciliation" not in title.casefold():
        raise PlDeskError(
            f"Unexpected title {title!r}. Expected a FinanceOS Trading Account Reconciliation."
        )
    legal = (header.get("company") or "").strip()
    if legal.casefold() != profile["legal_name"].casefold():
        raise PlDeskError(
            f"Unexpected company {legal!r}. This report expects {profile['legal_name']!r}."
        )
    currency = (header.get("currency") or "").strip()
    allowed = {profile["currency"].upper(), f"{profile['currency'].upper()} USD"}
    if currency.upper() not in allowed:
        raise PlDeskError(
            f"Unexpected currency {currency!r}. Only {profile['currency']} USD is in scope."
        )
    _calendar_month(header.get("period_from"), header.get("period_to"))
    if set(doc["equity"]) < {"Total", "MT5", "Wallet"}:
        raise PlDeskError("The equity block is missing (Total, MT5 and Wallet are required).")
    present = {section["code"] for section in doc["sections"]}
    missing_sections = [code for code in REQUIRED_SECTIONS if code not in present]
    if missing_sections:
        raise PlDeskError("A required section is missing: " + ", ".join(missing_sections) + ".")
    problems = []
    for row in reconciliation_checks(doc):
        if row["status"] != "FAIL":
            continue
        if row.get("diff") is not None:
            problems.append(
                f"{row['check']} (diff {fmt_amt(row['diff'])}, tolerance {fmt_amt(row.get('tolerance'))})"
            )
        else:
            problems.append(row["check"])
    if problems:
        extra = f" ({len(problems)} problems)" if len(problems) > 1 else ""
        raise PlDeskError("FinanceOS reconciliation failed: " + problems[0] + extra)


def reconciliation_checks(doc):
    """Sub-totals, group totals, equity identities and the required-label list."""
    results = []

    def check(name, reported, expected, children):
        if reported is None or expected is None:
            results.append({"check": name, "status": "FAIL", "diff": None})
            return
        diff = q2(Decimal(str(reported)) - Decimal(str(expected)))
        tolerance = float(TOLERANCE_PER_CHILD * Decimal(children))
        if abs(diff) < 0.005:
            status = "OK"
        elif abs(diff) <= tolerance + 1e-9:
            status = "ROUNDING"
        else:
            status = "FAIL"
        results.append({
            "check": name, "reported": q2(reported), "sum": q2(expected),
            "diff": diff, "tolerance": q2(tolerance) if tolerance else 0.0, "status": status,
            "children": children,
        })

    for section in doc["sections"]:
        for _parents, node in walk(section["tree"]):
            if node["amount"] is None:
                results.append({
                    "check": f"amount missing for {node['label']!r}",
                    "status": "FAIL", "diff": None,
                })
        if section["subtotal"] is None:
            results.append({
                "check": f"{section['code']}.) {section['name']} Sub-Total is missing",
                "status": "FAIL", "diff": None,
            })
            continue
        roots = section["tree"]
        check(
            f"{section['code']}.) {section['name']} Sub-Total",
            section["subtotal"],
            sum(node["amount"] or 0 for node in roots),
            len(roots),
        )
        for _parents, node in walk(section["tree"]):
            if not node["children"]:
                continue
            check(
                f"{section['code']}: {node['label']}",
                node["amount"],
                sum(child["amount"] or 0 for child in node["children"]),
                len(node["children"]),
            )
    if doc["total_movement"] is None:
        results.append({"check": "Total Movement is missing", "status": "FAIL", "diff": None})
    else:
        check(
            "Total Movement = sections",
            doc["total_movement"],
            sum(section["subtotal"] or 0 for section in doc["sections"]),
            len(doc["sections"]),
        )
    equity = doc["equity"]
    if {"Total", "MT5", "Wallet"} <= set(equity):
        for key, label in (("p", "Previous"), ("c", "Current"), ("d", "Change")):
            check(
                f"Equity {label}: MT5 + Wallet = Total",
                equity["Total"][key],
                equity["MT5"][key] + equity["Wallet"][key],
                2,
            )
        for category in ("Total", "MT5", "Wallet"):
            check(
                f"Equity {category}: Current - Previous = Change",
                equity[category]["d"],
                equity[category]["c"] - equity[category]["p"],
                2,
            )
        if doc["diff"] is None:
            results.append({"check": "Diff is missing", "status": "FAIL", "diff": None})
        else:
            check(
                "Diff = Total Movement - Equity Change",
                doc["diff"],
                doc["total_movement"] - equity["Total"]["d"],
                2,
            )
    labels = {node["label"] for section in doc["sections"] for _p, node in walk(section["tree"])}
    for required in REQUIRED_LABELS:
        if required not in labels:
            results.append({"check": f"required label {required!r} is missing", "status": "FAIL", "diff": None})
    return results


def find_label(doc, label):
    for section in doc["sections"]:
        for _parents, node in walk(section["tree"]):
            if node["label"] == label:
                return node
    return None


def new_labels(doc):
    found = []
    for section in doc["sections"]:
        if section["name"] not in KNOWN_SECTIONS:
            found.append({
                "section": section["code"], "label": section["name"],
                "amount": section["subtotal"], "where": "section",
            })
        for parents, node in walk(section["tree"]):
            if node["label"] == MANUAL_REBATE or MANUAL_REBATE in parents:
                continue
            if node["label"] not in KNOWN_LABELS:
                found.append({
                    "section": section["code"], "label": node["label"],
                    "amount": node["amount"],
                    "where": " > ".join(parents) or section["name"],
                })
    return found


def test_entries(doc):
    found = []
    for section in doc["sections"]:
        for parents, node in walk(section["tree"]):
            if "test" in node["label"].casefold():
                found.append({
                    "section": section["code"], "label": node["label"],
                    "amount": node["amount"], "where": " > ".join(parents),
                })
    return found


# ----------------------------------------------------------------- metabase
def _norm_header(value):
    return str(value).strip().lower() if value is not None else ""


def _sheet_rows(sheet):
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def _header_index(header):
    return {_norm_header(name): index for index, name in enumerate(header) if _norm_header(name)}


def _looks_like_metabase(headers):
    """The desk-rebate export, even when a required metric column is missing.

    The sample workbook also carries an Outcome sheet, so a broken export must
    still be recognised as Metabase. load_metabase then names the missing column.
    """
    return {"data_date", "country_code"} <= set(headers)


def classify_workbook(path):
    """Return 'metabase', 'finance' or None. Metabase wins when both shapes are present."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        metabase = False
        finance = False
        for name in workbook.sheetnames:
            rows = _sheet_rows(workbook[name])
            if not rows:
                continue
            headers = set(_header_index(rows[0]))
            if _looks_like_metabase(headers):
                metabase = True
            blob = " ".join(str(cell) for row in rows[:25] for cell in row[:4] if cell is not None)
            if "reconciliation" in blob.casefold() or "Sub-Total" in blob:
                finance = True
        if metabase:
            return "metabase"
        if finance:
            return "finance"
        return None
    finally:
        workbook.close()


def identify_inputs(paths):
    finance, metabase = [], []
    for path in paths:
        path = Path(path)
        if not path.is_file():
            raise PlDeskError(f"File not found: {path}")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            finance.append(path)
            continue
        if suffix not in (".xlsx", ".xlsm"):
            raise PlDeskError(
                f"{path.name} is not a FinanceOS PDF or an Excel workbook."
            )
        kind = classify_workbook(path)
        if kind == "metabase":
            metabase.append(path)
        elif kind == "finance":
            finance.append(path)
        else:
            raise PlDeskError(
                f"{path.name} is not a FinanceOS settlement or a Metabase desk-rebate export."
            )
    if len(finance) != 1 or len(metabase) != 1:
        raise PlDeskError(
            "Upload two files: the FinanceOS settlement (PDF, or the converted Excel) "
            "and the Metabase workbook that contains the desk-rebate export, Desks and Target."
        )
    return finance[0], metabase[0]


def _parse_data_date(value, where):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise PlDeskError(f"data_date on row {where} is not a date.")
    if isinstance(value, (int, float)):
        text = f"{int(value):08d}"
    elif isinstance(value, str):
        text = value.strip()
    else:
        raise PlDeskError(f"data_date on row {where} is not a date.")
    try:
        if re.fullmatch(r"\d{8}", text):
            return date(int(text[:4]), int(text[4:6]), int(text[6:8]))
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return date.fromisoformat(text)
    except ValueError as exc:
        raise PlDeskError(f"data_date on row {where} is not a date.") from exc
    raise PlDeskError(f"data_date on row {where} is not a date.")


def _parse_metric(value, field, where):
    if value is None or value == "":
        raise PlDeskError(f"{field} for {where} is not a number.")
    if isinstance(value, bool):
        raise PlDeskError(f"{field} for {where} is not a number.")
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return money(value)
        except Exception as exc:
            raise PlDeskError(f"{field} for {where} is not a number.") from exc
    raise PlDeskError(f"{field} for {where} is not a number.")


def _find_sheet(workbook, names):
    wanted = {name.casefold() for name in names}
    for sheet in workbook.sheetnames:
        if sheet.strip().casefold() in wanted:
            return workbook[sheet]
    return None


def load_metabase(path, company="dpm"):
    profile = _company_profile(company)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {name: _sheet_rows(workbook[name]) for name in workbook.sheetnames}
    finally:
        workbook.close()
    candidates = []
    for name, rows in tables.items():
        if not rows:
            continue
        headers = _header_index(rows[0])
        if set(REQUIRED_META_COLUMNS) <= set(headers):
            candidates.append((name, rows, headers))
    if not candidates:
        missing = ", ".join(REQUIRED_META_COLUMNS)
        raise PlDeskError(f"The Metabase export is missing required columns ({missing}).")
    preferred = [item for item in candidates if item[0].strip().casefold().startswith("query result")]
    chosen = preferred or candidates
    if len(chosen) > 1 and not (len(preferred) == 1):
        raise PlDeskError("More than one sheet looks like the Metabase desk-rebate export.")
    sheet_name, rows, headers = chosen[0]
    records = []
    months = set()
    seen = {}
    for offset, row in enumerate(rows[1:], start=2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        code = row[headers["country_code"]] if headers["country_code"] < len(row) else None
        if code is None or str(code).strip() == "":
            raise PlDeskError(f"Row {offset} of the Metabase export has no country_code.")
        if isinstance(code, float) and not isinstance(code, bool):
            raise PlDeskError(f"country_code on row {offset} is not a country code.")
        country = str(code).strip().upper()
        if country in seen:
            raise PlDeskError(
                f"Country {country} appears more than once in the Metabase export "
                f"(rows {seen[country]} and {offset})."
            )
        seen[country] = offset
        when = _parse_data_date(row[headers["data_date"]] if headers["data_date"] < len(row) else None, offset)
        months.add((when.year, when.month))
        values = {"country_code": country, "data_date": when}
        for field in METRIC_FIELDS:
            raw = row[headers[field]] if headers[field] < len(row) else None
            values[field] = _parse_metric(raw, field, country)
        original = []
        for index, _name in enumerate(rows[0]):
            original.append(row[index] if index < len(row) else None)
        values["_original"] = original
        records.append(values)
    if not records:
        raise PlDeskError("The Metabase export has no country rows.")
    if len(months) != 1:
        listed = ", ".join(f"{year}-{month:02d}" for year, month in sorted(months))
        raise PlDeskError(
            f"The Metabase export covers more than one month ({listed}). It must be a single month."
        )
    desks_sheet = _find_sheet_in(tables, ("desks",))
    if desks_sheet is None:
        raise PlDeskError("The Metabase workbook has no Desks sheet.")
    desk_names, explicit, duplicates = _read_desks(desks_sheet, profile["fallback_desk"])
    if duplicates:
        country, names = duplicates[0]
        raise PlDeskError(
            f"Country {country} sits in more than one desk ({', '.join(names)})."
        )
    fallback = next((name for name in desk_names if name.casefold() == profile["fallback_desk"].casefold()), None)
    if fallback is None:
        raise PlDeskError(
            f"The Desks sheet has no {profile['fallback_desk']} column. "
            "Unmapped countries are assigned there, so that desk is required."
        )
    names = _read_country_names(_find_sheet_in(tables, ("country code", "country codes")))
    auto = []
    for record in records:
        desk = explicit.get(record["country_code"])
        if desk is None:
            desk = fallback
            auto.append(record["country_code"])
        record["desk"] = desk
        record["country_name"] = names.get(record["country_code"], "")
    targets, recognised, target_cells = _read_targets(_find_sheet_in(tables, ("target",)), desk_names)
    copies = {
        "Desks": desks_sheet,
        "Target": _find_sheet_in(tables, ("target",)) or [],
        "Country code": _find_sheet_in(tables, ("country code", "country codes")) or [],
    }
    return {
        "sheet": sheet_name,
        "header": list(rows[0]),
        "rows": records,
        "month": months.pop(),
        "desks": desk_names,
        "auto_gem": auto,
        "fallback_desk": fallback,
        "targets": targets,
        "targets_recognised": recognised,
        "target_cells": target_cells,
        "copies": copies,
        "names_missing_sheet": _find_sheet_in(tables, ("country code", "country codes")) is None,
    }


def _find_sheet_in(tables, names):
    wanted = {name.casefold() for name in names}
    for name, rows in tables.items():
        if name.strip().casefold() in wanted:
            return rows
    return None


def _read_desks(rows, fallback_name):
    header_at = None
    for index, row in enumerate(rows[:6]):
        texts = [str(cell).strip() for cell in row if isinstance(cell, str) and cell.strip()]
        if len(texts) >= 2:
            header_at = index
            break
    if header_at is None:
        raise PlDeskError("The Desks sheet has no desk headers.")
    header = rows[header_at]
    columns = []
    seen_names = {}
    for index, cell in enumerate(header):
        if not isinstance(cell, str) or not cell.strip():
            continue
        name = cell.strip()
        if name.casefold() in seen_names:
            raise PlDeskError(f"Desk name {name} appears twice on the Desks sheet.")
        seen_names[name.casefold()] = name
        columns.append((index, name))
    explicit = {}
    duplicates = []
    owners = {}
    for row in rows[header_at + 1:]:
        for index, name in columns:
            if index >= len(row) or row[index] is None:
                continue
            if isinstance(row[index], (int, float)) and not isinstance(row[index], bool):
                continue
            code = str(row[index]).strip().upper()
            if not code or code.casefold() in ("desk", "country", "country code"):
                continue
            owners.setdefault(code, [])
            if name not in owners[code]:
                owners[code].append(name)
            explicit[code] = name
    for code, names in owners.items():
        if len(names) > 1:
            duplicates.append((code, names))
    return [name for _i, name in columns], explicit, duplicates


def _read_country_names(rows):
    names = {}
    if not rows:
        return names
    for offset, row in enumerate(rows):
        if not row or row[0] is None:
            continue
        code = str(row[0]).strip()
        if offset == 0 and code.casefold() in ("country_code", "code", "country code"):
            continue
        name = row[1] if len(row) > 1 and row[1] is not None else ""
        names[code.upper()] = str(name).strip()
    return names


def _read_targets(rows, desk_names):
    """Recognise a sheet with desk names across the header and a metric in column A."""
    if not rows:
        return {}, False, {}
    wanted = {name.casefold(): name for name in desk_names}
    header_at = None
    columns = {}
    for index, row in enumerate(rows[:8]):
        found = {}
        for col, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().casefold() in wanted:
                found[col] = wanted[cell.strip().casefold()]
        if len(found) >= 2:
            header_at = index
            columns = found
            break
    if header_at is None:
        return {}, False, {}
    targets = {field: {desk: None for desk in desk_names} for field in METRIC_FIELDS}
    cells = {}
    for offset, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        if not row or row[0] is None:
            continue
        key = TARGET_ALIASES.get(str(row[0]).strip().casefold())
        if not key:
            continue
        for col, desk in columns.items():
            cells[(key, desk)] = (offset, col + 1)
            if col >= len(row) or row[col] is None or row[col] == "":
                continue
            targets[key][desk] = _parse_metric(row[col], f"target {key}", desk)
    return targets, True, cells


def _same_month(doc, metabase):
    start = _calendar_month(doc["header"].get("period_from"), doc["header"].get("period_to"))
    year, month = metabase["month"]
    if (start.year, start.month) != (year, month):
        raise PlDeskError(
            f"FinanceOS period is {doc['header']['period_from']} to {doc['header']['period_to']}, "
            f"but the Metabase data_date month is {year}-{month:02d}."
        )
    return f"{start.year}-{start.month:02d}"


# ----------------------------------------------------------------- outcome rows
def _sum_formula(row_numbers):
    if not row_numbers:
        return "=0"
    if len(row_numbers) == 1:
        return f"=B{row_numbers[0]}"
    if row_numbers[-1] - row_numbers[0] + 1 == len(row_numbers):
        return f"=SUM(B{row_numbers[0]}:B{row_numbers[-1]})"
    return "=" + "+".join(f"B{number}" for number in row_numbers)


def _add_tree(add, node, section_code, force_blank):
    is_cash = section_code == "B"
    if node["label"] == MANUAL_REBATE and not force_blank:
        spec = add({
            "role": "group", "label": node["label"], "level": node["level"],
            "yellow": False, "pdf_amount": node["amount"],
        })
        count = max(MANUAL_INPUT_FLOOR, len(node["children"]) + MANUAL_INPUT_EXTRA)
        kids = [
            add({"role": "manual", "label": "", "level": node["level"] + 1, "orange": True})
            for _ in range(count)
        ]
        spec["kids"] = kids
        spec["manual_source"] = list(node["children"])
        return spec
    if force_blank or node["label"] in EXCLUDED_FROM_SUBTOTAL:
        spec = add({
            "role": "leaf", "label": node["label"], "level": node["level"],
            "yellow": True, "hidden": is_cash, "pdf_amount": node["amount"],
        })
        for child in node["children"]:
            _add_tree(add, child, section_code, True)
        return spec
    if node["children"]:
        spec = add({
            "role": "group", "label": node["label"], "level": node["level"],
            "yellow": is_cash, "hidden": is_cash, "pdf_amount": node["amount"],
        })
        spec["kids"] = [_add_tree(add, child, section_code, False) for child in node["children"]]
        return spec
    yellow = is_cash or node["label"] in YELLOW_LEAVES
    metric = METRIC_BY_LABEL.get(node["label"])
    return add({
        "role": "leaf", "label": node["label"], "level": node["level"],
        "yellow": yellow, "hidden": is_cash,
        "value": None if yellow else node["amount"],
        "pdf_amount": node["amount"],
        "metric": metric[0] if metric else None,
        "sign": metric[1] if metric else 1,
    })


def build_outcome_specs(doc):
    """Row specs for the settlement body. Row numbers are filled in afterwards."""
    rows = []

    def add(spec):
        spec.setdefault("kids", [])
        spec.setdefault("value", None)
        spec.setdefault("metric", None)
        spec.setdefault("sign", 1)
        spec.setdefault("yellow", False)
        spec.setdefault("orange", False)
        spec.setdefault("hidden", False)
        spec.setdefault("level", 0)
        spec.setdefault("pdf_amount", None)
        spec.setdefault("b_formula", None)
        spec.setdefault("role", "leaf")
        rows.append(spec)
        return spec

    sections = {}
    for section in doc["sections"]:
        spec = add({
            "role": "section",
            "label": f"{section['code']}.) {section['name']}  Sub-Total",
        })
        sections[section["code"]] = spec
        kids = []
        for node in section["tree"]:
            child = _add_tree(add, node, section["code"], False)
            if node["label"] not in EXCLUDED_FROM_SUBTOTAL:
                kids.append(child)
        spec["kids"] = kids
    add({"role": "blank"})
    total = add({"role": "total", "label": "Total Movement"})
    diff = add({"role": "diff", "label": "Diff"})
    add({"role": "blank"})
    gross = add({"role": "gp", "label": "Gross profit"})
    margin = add({"role": "margin", "label": "Gross profit margin"})
    add({"role": "blank"})
    deposit_pdf = find_label(doc, "Deposit to MT5 Trading Account")
    withdrawal_pdf = find_label(doc, "Withdrawal MT5 Trading Account")
    deposit = add({
        "role": "metric", "label": "Deposit", "metric": "deposit_usd", "sign": 1,
        "value": None if deposit_pdf is None else deposit_pdf["amount"],
        "pdf_amount": None if deposit_pdf is None else deposit_pdf["amount"],
    })
    withdrawal = add({
        "role": "metric", "label": "Withdrawal", "metric": "withdrawal_usd", "sign": -1,
        "value": None if withdrawal_pdf is None else withdrawal_pdf["amount"],
        "pdf_amount": None if withdrawal_pdf is None else withdrawal_pdf["amount"],
    })
    ndp = add({"role": "metric", "label": "NDP", "metric": "ndp_usd", "sign": 1})
    volume = add({"role": "metric", "label": "Volume", "metric": "lots", "sign": 1})
    for number, spec in enumerate(rows, start=14):
        spec["row"] = number
    for spec in rows:
        if spec["role"] in ("section", "group") and spec["kids"]:
            spec["b_formula"] = _sum_formula([kid["row"] for kid in spec["kids"]])
    total["b_formula"] = "=" + "+".join(f"B{sections[code]['row']}" for code in sections)
    diff["b_formula"] = f"=D9-B{total['row']}"
    rebate = next((spec for spec in rows if spec["label"] == "Rebate"), None)
    if rebate is None:
        raise PlDeskError("FinanceOS reconciliation failed: required label 'Rebate' is missing.")
    gross["b_formula"] = f"=-B{sections['A']['row']}-B{sections['C']['row']}-B{rebate['row']}"
    margin["b_formula"] = f"=IF(B{sections['A']['row']}=0,\"\",B{gross['row']}/-B{sections['A']['row']})"
    ndp["b_formula"] = f"=B{deposit['row']}+B{withdrawal['row']}"
    ndp["c_formula"] = f"=C{deposit['row']}+C{withdrawal['row']}"
    return {
        "rows": rows,
        "sections": sections,
        "deposit": deposit,
        "withdrawal": withdrawal,
        "ndp": ndp,
        "volume": volume,
        "rebate": rebate,
        "gross": gross,
        "total": total,
    }


# ----------------------------------------------------------------- workbook
def _style_amount(cell, formula=None, value=None, fmt=NUM_FMT, font=None, fill=None):
    if formula is not None:
        cell.value = formula
    elif value is not None:
        cell.value = value
    cell.number_format = fmt
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill


def _write_outcome(workbook, doc, metabase, body, columns):
    sheet = workbook.active
    sheet.title = "Outcome"
    profile_name = doc["header"]["company"]
    period = f"{doc['header']['period_from']} to {doc['header']['period_to']}"
    put_text(sheet.cell(1, 1), "Monthly Trading Account Reconciliation")
    sheet.cell(1, 1).font = FONT_WHITE_TITLE
    sheet.cell(1, 1).fill = FILL_NAVY
    put_text(sheet.cell(2, 1), f"Company: {profile_name}")
    sheet.cell(2, 1).font = FONT_MUTED
    put_text(sheet.cell(3, 1), f"Currency: {doc['header'].get('currency', 'CONSOLIDATED')}")
    sheet.cell(3, 1).font = FONT_MUTED
    put_text(sheet.cell(4, 1), f"Period: {period}")
    sheet.cell(4, 1).font = FONT_MUTED
    if doc["header"].get("rate_note"):
        put_text(sheet.cell(5, 1), doc["header"]["rate_note"])
        sheet.cell(5, 1).font = Font(color="666666", italic=True, size=9)
        sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=5)
    desks = metabase["desks"]
    for index, desk in enumerate(desks):
        start = 6 + index * 3
        put_text(sheet.cell(6, start), desk)
        sheet.cell(6, start).font = FONT_BOLD
        sheet.cell(6, start).alignment = Alignment(horizontal="center")
        sheet.merge_cells(start_row=6, start_column=start, end_row=6, end_column=start + 2)
        for offset, title in enumerate(("Target", "Actual", "Completion %")):
            put_text(sheet.cell(7, start + offset), title)
            sheet.cell(7, start + offset).font = FONT_BOLD
            sheet.cell(7, start + offset).alignment = Alignment(horizontal="center")
    check_col = 6 + len(desks) * 3
    put_text(sheet.cell(7, check_col), "Desks − C")
    sheet.cell(7, check_col).font = FONT_BOLD
    for col, title, fill in (
        (1, "Category", FILL_HEAD), (2, "Previous", FILL_HEAD),
        (3, "Current", FILL_HEAD), (4, "Change", FILL_HEAD),
    ):
        put_text(sheet.cell(8, col), title)
        sheet.cell(8, col).fill = fill
        sheet.cell(8, col).font = FONT_WHITE
    equity = doc["equity"]
    put_text(sheet.cell(9, 1), "Total")
    _style_amount(sheet.cell(9, 2), formula="=B10+B11", font=FONT_BOLD)
    _style_amount(sheet.cell(9, 3), formula="=C10+C11", font=FONT_BOLD)
    _style_amount(sheet.cell(9, 4), formula="=C9-B9", font=FONT_BOLD)
    put_text(sheet.cell(10, 1), "MT5")
    _style_amount(sheet.cell(10, 2), value=equity["MT5"]["p"], font=FONT_GREEN)
    _style_amount(sheet.cell(10, 3), value=equity["MT5"]["c"], font=FONT_GREEN)
    _style_amount(sheet.cell(10, 4), formula="=C10-B10")
    put_text(sheet.cell(11, 1), "Wallet")
    _style_amount(sheet.cell(11, 2), value=equity["Wallet"]["p"], font=FONT_GREEN)
    _style_amount(sheet.cell(11, 3), value=equity["Wallet"]["c"], font=FONT_GREEN)
    _style_amount(sheet.cell(11, 4), formula="=C11-B11")
    headers = (
        (1, "Description", FILL_HEAD), (2, "Amount", FILL_HEAD),
        (3, "Metabase", FILL_META), (4, "Variance", FILL_HEAD),
        (5, "Metabase description", FILL_GREEN),
    )
    for col, title, fill in headers:
        put_text(sheet.cell(13, col), title)
        sheet.cell(13, col).fill = fill
        sheet.cell(13, col).font = FONT_WHITE
    put_text(sheet.cell(13, check_col), "Desks − C")
    sheet.cell(13, check_col).font = FONT_WHITE
    sheet.cell(13, check_col).fill = FILL_HEAD
    last_data = 1 + len(metabase["rows"])
    desk_header_cell = {}
    for index, desk in enumerate(desks):
        desk_header_cell[desk] = f"{get_column_letter(6 + index * 3)}$6"
    for spec in body["rows"]:
        row = spec["row"]
        if spec["role"] == "blank":
            continue
        label_cell = sheet.cell(row, 1)
        put_text(label_cell, spec["label"])
        if spec["level"] > 1:
            label_cell.alignment = Alignment(indent=spec["level"] - 1)
        _paint_label(label_cell, spec)
        _write_amount_cell(sheet, spec)
        if spec["metric"]:
            _write_metric_cells(
                sheet, spec, columns, last_data, desks, desk_header_cell,
                body, check_col, metabase,
            )
        if spec["hidden"]:
            sheet.row_dimensions[row].hidden = True
    last_row = body["rows"][-1]["row"]
    sheet.conditional_formatting.add(
        f"D14:D{last_row}",
        FormulaRule(
            formula=[f"AND(ISNUMBER(D14),ABS(D14)>{VARIANCE_LIMIT})"],
            fill=FILL_VARIANCE, font=FONT_VARIANCE,
        ),
    )
    sheet.column_dimensions["A"].width = 62
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 22
    for col in range(6, check_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 14
    sheet.column_dimensions[get_column_letter(check_col)].width = 16
    sheet.freeze_panes = "A14"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = "Monthly Trading Account Reconciliation"
    return check_col


def _paint_label(cell, spec):
    if spec["role"] == "section":
        cell.fill = FILL_SECTION
        cell.font = FONT_WHITE
    elif spec["orange"]:
        cell.fill = FILL_ORANGE
        cell.font = FONT_GREEN
    elif spec["yellow"]:
        cell.fill = FILL_YELLOW
        cell.font = FONT_GREEN
    elif spec["role"] == "diff":
        cell.fill = FILL_DIFF
        cell.font = FONT_DIFF
    elif spec["role"] in ("total", "gp", "margin"):
        cell.font = FONT_BOLD
    elif spec["role"] in ("leaf", "group", "metric"):
        cell.font = FONT_GREEN_B if spec["role"] == "group" else FONT_GREEN


def _write_amount_cell(sheet, spec):
    cell = sheet.cell(spec["row"], 2)
    if spec["orange"]:
        cell.fill = FILL_ORANGE
        cell.number_format = NUM_FMT
        return
    font = FONT_WHITE if spec["role"] == "section" else None
    if spec["role"] == "diff":
        font = FONT_DIFF
        cell.fill = FILL_DIFF
    elif spec["role"] in ("group", "leaf", "metric") and spec["role"] != "section":
        font = FONT_GREEN
    if spec["b_formula"]:
        fmt = PCT_FMT if spec["role"] == "margin" else NUM_FMT
        _style_amount(cell, formula=spec["b_formula"], fmt=fmt, font=font or FONT_BOLD)
        if spec["role"] == "section":
            cell.font = FONT_WHITE
            cell.fill = FILL_SECTION
        return
    if spec["value"] is not None:
        _style_amount(cell, value=spec["value"], font=font or FONT_GREEN)


def _write_metric_cells(sheet, spec, columns, last_data, desks, desk_header_cell, body, check_col, metabase):
    row = spec["row"]
    field = spec["metric"]
    letter = columns[field]
    data_range = f"'{RAW_SHEET}'!${letter}$2:${letter}${last_data}"
    desk_letter = columns["desk"]
    desk_range = f"'{RAW_SHEET}'!${desk_letter}$2:${desk_letter}${last_data}"
    if spec.get("c_formula"):
        c_formula = spec["c_formula"]
    elif spec["sign"] < 0:
        c_formula = f"=-SUM({data_range})"
    else:
        c_formula = f"=SUM({data_range})"
    _style_amount(sheet.cell(row, 3), formula=c_formula)
    if field != "lots":
        if spec["value"] is None and not spec["b_formula"]:
            variance = f'=IF(B{row}="","",B{row}-C{row})'
        else:
            variance = f"=B{row}-C{row}"
        _style_amount(sheet.cell(row, 4), formula=variance)
    put_text(sheet.cell(row, 5), field)
    actual_cells = []
    for index, desk in enumerate(desks):
        target_col = 6 + index * 3
        actual_col = target_col + 1
        completion_col = target_col + 2
        target_ref = _target_formula(metabase, field, desk)
        if target_ref:
            _style_amount(sheet.cell(row, target_col), formula=target_ref)
        actual_ref = _actual_formula(
            spec, body, desk, index, data_range, desk_range, desk_header_cell,
        )
        _style_amount(sheet.cell(row, actual_col), formula=actual_ref)
        target_cell = f"{get_column_letter(target_col)}{row}"
        actual_cell = f"{get_column_letter(actual_col)}{row}"
        _style_amount(
            sheet.cell(row, completion_col),
            formula=f'=IF(OR({target_cell}="",{target_cell}=0),"",{actual_cell}/{target_cell})',
            fmt=PCT_FMT,
        )
        actual_cells.append(actual_cell)
    check = "=" + "+".join(actual_cells) + f"-C{row}"
    _style_amount(sheet.cell(row, check_col), formula=check)


def _target_formula(metabase, field, desk):
    """Point at the copied Target sheet when its layout was recognised."""
    if not metabase["targets_recognised"]:
        return None
    found = metabase["target_cells"].get((field, desk))
    if not found:
        return None
    ref = f"Target!{get_column_letter(found[1])}{found[0]}"
    return f'=IF({ref}="","",{ref})'


def _actual_formula(spec, body, desk, index, data_range, desk_range, desk_header_cell):
    header = desk_header_cell[desk]
    if spec["metric"] == "ndp_usd":
        deposit_row = body["deposit"]["row"]
        withdrawal_row = body["withdrawal"]["row"]
        actual_col = get_column_letter(6 + index * 3 + 1)
        return f"={actual_col}{deposit_row}+{actual_col}{withdrawal_row}"
    if spec["sign"] < 0:
        return f'=-SUMIFS({data_range},{desk_range},{header})'
    return f'=SUMIFS({data_range},{desk_range},{header})'


def _write_parsed(workbook, doc, checks):
    sheet = workbook.create_sheet("FinanceOS (parsed)")
    for col, title, width in (
        (1, "Section", 14), (2, "Level", 10), (3, "Group", 42),
        (4, "Label", 78), (5, "Amount", 18), (6, "Page", 10), (7, "Note", 42),
    ):
        put_text(sheet.cell(1, col), title)
        sheet.cell(1, col).font = FONT_WHITE
        sheet.cell(1, col).fill = FILL_SECTION
        sheet.column_dimensions[get_column_letter(col)].width = width
    put_text(sheet.cell(2, 1), "Header")
    put_text(sheet.cell(2, 4), doc["header"].get("title"))
    put_text(sheet.cell(3, 1), "Company")
    put_text(sheet.cell(3, 4), doc["header"].get("company"))
    put_text(sheet.cell(4, 1), "Currency")
    put_text(sheet.cell(4, 4), doc["header"].get("currency"))
    put_text(sheet.cell(5, 1), "Period")
    put_text(sheet.cell(5, 4), f"{doc['header'].get('period_from')} to {doc['header'].get('period_to')}")
    put_text(sheet.cell(6, 1), "Engine")
    put_text(sheet.cell(6, 4), doc.get("engine"))
    row = 8
    for category in ("Total", "MT5", "Wallet"):
        equity = doc["equity"][category]
        put_text(sheet.cell(row, 1), "Equity")
        put_text(sheet.cell(row, 4), category)
        _style_amount(sheet.cell(row, 5), value=equity["c"])
        put_text(sheet.cell(row, 7), f"Previous {fmt_amt(equity['p'])}; Change {fmt_amt(equity['d'])}")
        row += 1
    row += 1
    for section in doc["sections"]:
        put_text(sheet.cell(row, 1), section["code"])
        put_text(sheet.cell(row, 4), f"{section['code']}.) {section['name']}  Sub-Total")
        _style_amount(sheet.cell(row, 5), value=section["subtotal"])
        sheet.cell(row, 4).font = FONT_BOLD
        row += 1
        for parents, node in walk(section["tree"]):
            put_text(sheet.cell(row, 1), section["code"])
            sheet.cell(row, 2, node["level"])
            put_text(sheet.cell(row, 3), " > ".join(parents))
            put_text(sheet.cell(row, 4), node["label"])
            if node["amount"] is not None:
                _style_amount(sheet.cell(row, 5), value=node["amount"])
            if node.get("page"):
                sheet.cell(row, 6, node["page"])
            notes = []
            if node.get("wrapped"):
                notes.append("wrapped label joined")
            if "test" in node["label"].casefold():
                notes.append("test entry, kept in the total")
            if node["label"] in EXCLUDED_FROM_SUBTOTAL or node["label"] in YELLOW_LEAVES:
                notes.append("left blank on Outcome")
            if section["code"] == "B":
                notes.append("Cash Movement amount left blank on Outcome")
            if MANUAL_REBATE in parents:
                notes.append("manual rebate, not copied to Outcome")
            put_text(sheet.cell(row, 7), "; ".join(dict.fromkeys(notes)))
            row += 1
    row += 1
    put_text(sheet.cell(row, 4), "Total Movement")
    if doc["total_movement"] is not None:
        _style_amount(sheet.cell(row, 5), value=doc["total_movement"])
    row += 1
    put_text(sheet.cell(row, 4), "Diff")
    if doc["diff"] is not None:
        _style_amount(sheet.cell(row, 5), value=doc["diff"])
    row += 2
    for col, title in enumerate(("Check", "Reported", "Sum of children", "Diff", "Tolerance", "Status"), start=1):
        put_text(sheet.cell(row, col), title)
        sheet.cell(row, col).font = FONT_WHITE
        sheet.cell(row, col).fill = FILL_SECTION
    for item in checks:
        row += 1
        put_text(sheet.cell(row, 1), item["check"])
        if item.get("reported") is not None:
            _style_amount(sheet.cell(row, 2), value=item["reported"])
            _style_amount(sheet.cell(row, 3), value=item["sum"])
            _style_amount(sheet.cell(row, 4), value=item["diff"])
            _style_amount(sheet.cell(row, 5), value=item.get("tolerance"))
        put_text(sheet.cell(row, 6), item["status"])
        if item["status"] == "ROUNDING":
            sheet.cell(row, 6).fill = FILL_YELLOW
        elif item["status"] == "FAIL":
            sheet.cell(row, 6).fill = FILL_VARIANCE
    sheet.freeze_panes = "A2"
    return sheet


def _normalized_header(header):
    names = [("" if name is None else name) for name in header]
    while names and names[-1] == "":
        names.pop()
    names.extend(["desk", "country_name"])
    return names


def _column_letters(header):
    names = _normalized_header(header)
    columns = {}
    for index, name in enumerate(names):
        key = _norm_header(name)
        if key in METRIC_FIELDS or key in ("country_code", "data_date", "desk", "country_name"):
            columns[key] = get_column_letter(index + 1)
    missing = [field for field in METRIC_FIELDS if field not in columns]
    if missing or "desk" not in columns:
        raise PlDeskError("The Metabase export is missing required columns (" + ", ".join(missing) + ").")
    return columns


def _write_raw(workbook, metabase):
    sheet = workbook.create_sheet(RAW_SHEET)
    header = _normalized_header(metabase["header"])
    for col, name in enumerate(header, start=1):
        put_text(sheet.cell(1, col), name)
        sheet.cell(1, col).font = FONT_WHITE
        sheet.cell(1, col).fill = FILL_SECTION
    desk_col = header.index("desk") + 1
    name_col = header.index("country_name") + 1
    for row_index, record in enumerate(metabase["rows"], start=2):
        original = list(record["_original"][: len(header) - 2])
        while len(original) < len(header) - 2:
            original.append(None)
        for col, value in enumerate(original, start=1):
            cell = sheet.cell(row_index, col)
            if isinstance(value, str):
                put_text(cell, value)
            elif isinstance(value, datetime):
                cell.value = value.date()
                cell.number_format = "DD MMM YYYY"
            elif _is_number(value):
                cell.value = value
                cell.number_format = NUM_FMT
            elif value is not None:
                put_text(cell, value)
        put_text(sheet.cell(row_index, desk_col), record["desk"])
        put_text(sheet.cell(row_index, name_col), record["country_name"])
        if record["country_code"] in metabase["auto_gem"]:
            sheet.cell(row_index, desk_col).fill = FILL_YELLOW
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{1 + len(metabase['rows'])}"
    for col in range(1, len(header) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 18
    sheet.column_dimensions[get_column_letter(name_col)].width = 42


def _write_copy(workbook, title, rows):
    sheet = workbook.create_sheet(title)
    if not rows:
        put_text(sheet.cell(1, 1), "(empty in the upload)")
        return sheet
    width = max(len(row) for row in rows)
    for r, row in enumerate(rows, start=1):
        for c in range(width):
            value = row[c] if c < len(row) else None
            cell = sheet.cell(r, c + 1)
            if isinstance(value, str):
                put_text(cell, value)
            elif isinstance(value, datetime):
                cell.value = value
            else:
                cell.value = value
    sheet.freeze_panes = "A2"
    return sheet


def _write_check(workbook, doc, metabase, body, checks, variances):
    sheet = workbook.create_sheet("Check")
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 88
    sheet.column_dimensions["C"].width = 22
    sheet.column_dimensions["D"].width = 22
    sheet.column_dimensions["E"].width = 22
    put_text(sheet.cell(1, 1), "Check")
    sheet.cell(1, 1).font = Font(bold=True, size=14)
    put_text(sheet.cell(2, 1), "Read this sheet before using Outcome.")
    row = 4
    put_text(sheet.cell(row, 1), "Note")
    put_text(sheet.cell(row, 2), (
        "Deposit and Withdrawal are compared with Metabase in the block under Gross profit, "
        "not inside Cash Movement. Cash Movement column B stays blank (yellow), including "
        "Rebates. Gross profit still references the blank Rebate cell, so it starts deducting "
        "rebates when that cell is filled in. Wallet withdrawals are not part of the "
        "Metabase comparison; the Withdrawal line is the MT5 withdrawal."
    ))
    sheet.row_dimensions[row].height = 48
    sheet.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2
    rounding = [item for item in checks if item["status"] == "ROUNDING"]
    row = _check_heading(sheet, row, "Rounding")
    if not rounding:
        put_text(sheet.cell(row, 2), "None")
        row += 1
    for item in rounding:
        put_text(sheet.cell(row, 1), "Rounding")
        put_text(sheet.cell(row, 2), f"{item['check']}: diff {fmt_amt(item['diff'])} (tolerance {fmt_amt(item.get('tolerance'))})")
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
    row += 1
    row = _check_heading(sheet, row, "Diff")
    put_text(sheet.cell(row, 1), "Diff")
    pdf_diff = doc["diff"]
    if pdf_diff is not None and abs(q2(pdf_diff)) >= 0.01:
        put_text(sheet.cell(row, 2), f"FinanceOS Diff is {fmt_amt(pdf_diff)} (not zero).")
        sheet.cell(row, 1).fill = FILL_YELLOW
    else:
        put_text(sheet.cell(row, 2), "FinanceOS Diff is zero.")
    row += 1
    put_text(sheet.cell(row, 1), "Diff")
    put_text(sheet.cell(row, 2), (
        "Outcome Diff will not be zero: Cash Movement column B is blank, and Profit Sharing, "
        "MT5 Transfer IN, MT5 Transfer OUT and Wallet Profit Sharing are blank and outside the subtotal."
    ))
    sheet.cell(row, 1).fill = FILL_YELLOW
    row += 2
    row = _check_heading(sheet, row, "Wrapped label")
    wrapped = [warning for warning in doc["warnings"] if warning.startswith("wrapped label")]
    other = [warning for warning in doc["warnings"] if not warning.startswith("wrapped label")]
    if not wrapped:
        put_text(sheet.cell(row, 2), "None")
        row += 1
    for warning in wrapped:
        put_text(sheet.cell(row, 1), "Wrapped label")
        put_text(sheet.cell(row, 2), warning)
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
    for warning in other:
        put_text(sheet.cell(row, 1), "Unparsed line")
        put_text(sheet.cell(row, 2), warning)
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
    row += 1
    row = _check_heading(sheet, row, "New label")
    fresh = new_labels(doc)
    if not fresh:
        put_text(sheet.cell(row, 2), "None")
        row += 1
    for item in fresh:
        put_text(sheet.cell(row, 1), "New label")
        put_text(sheet.cell(row, 2), f"{item['section']}: {item['label']} ({fmt_amt(item['amount'])}) under {item['where']}")
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
    row += 1
    row = _check_heading(sheet, row, "Test entry")
    tests = test_entries(doc)
    if not tests:
        put_text(sheet.cell(row, 2), "None")
        row += 1
    for item in tests:
        put_text(sheet.cell(row, 1), "Test entry")
        put_text(sheet.cell(row, 2), f"{item['label']} ({fmt_amt(item['amount'])}) — flagged, not excluded")
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
    row += 1
    row = _check_heading(sheet, row, "Assigned to GEM")
    if not metabase["auto_gem"]:
        put_text(sheet.cell(row, 2), "None. Every country is on the Desks sheet.")
        row += 1
    else:
        put_text(sheet.cell(row, 1), "Assigned to GEM")
        put_text(sheet.cell(row, 2), (
            f"{len(metabase['auto_gem'])} countries are not on the Desks sheet and were assigned "
            f"to {metabase['fallback_desk']}. Desk totals then include them."
        ))
        sheet.cell(row, 1).fill = FILL_YELLOW
        row += 1
        for col, title in enumerate(("Country", "Name", "Closed P/L", "Deposit"), start=1):
            put_text(sheet.cell(row, col), title)
            sheet.cell(row, col).font = FONT_BOLD
        row += 1
        by_code = {record["country_code"]: record for record in metabase["rows"]}
        ordered = sorted(metabase["auto_gem"], key=lambda code: abs(by_code[code]["closed_pnl_usd"]), reverse=True)
        for code in ordered:
            record = by_code[code]
            put_text(sheet.cell(row, 1), code)
            put_text(sheet.cell(row, 2), record["country_name"])
            _style_amount(sheet.cell(row, 3), value=record["closed_pnl_usd"])
            _style_amount(sheet.cell(row, 4), value=record["deposit_usd"])
            row += 1
    row += 1
    row = _check_heading(sheet, row, "Variance over 1,000")
    if not variances:
        put_text(sheet.cell(row, 2), "None")
        row += 1
    else:
        for col, title in enumerate(("Category", "Line", "FinanceOS", "Metabase", "Variance"), start=1):
            put_text(sheet.cell(row, col), title)
            sheet.cell(row, col).font = FONT_BOLD
        row += 1
        for item in variances:
            put_text(sheet.cell(row, 1), "Variance over 1,000")
            put_text(sheet.cell(row, 2), item["line"])
            _style_amount(sheet.cell(row, 3), value=item["finance"])
            _style_amount(sheet.cell(row, 4), value=item["metabase"])
            _style_amount(sheet.cell(row, 5), value=item["variance"])
            sheet.cell(row, 1).fill = FILL_VARIANCE
            row += 1
    row += 1
    row = _check_heading(sheet, row, "Blank targets")
    put_text(sheet.cell(row, 1), "Blank targets")
    if not metabase["targets_recognised"] or not _any_target(metabase):
        put_text(sheet.cell(row, 2), (
            "Targets are blank, so Completion % is blank. On the Target sheet put metric names "
            "in column A (Closed P/L, Commissions, Rebate, Deposit, Withdrawal, NDP, Volume) "
            "and the desk names across the header row, then run the report again."
        ))
    else:
        blanks = _blank_targets(metabase)
        put_text(sheet.cell(row, 2), "Blank target cells: " + (", ".join(blanks) if blanks else "None"))
    sheet.cell(row, 1).fill = FILL_YELLOW
    sheet.cell(row, 2).alignment = Alignment(wrap_text=True)
    row += 2
    row = _check_heading(sheet, row, "Desk check")
    put_text(sheet.cell(row, 2), "Each formula is the sum of the desk Actual columns minus column C. It should be 0.")
    row += 1
    check_col = 6 + len(metabase["desks"]) * 3
    letter = get_column_letter(check_col)
    for spec in body["rows"]:
        if not spec["metric"]:
            continue
        put_text(sheet.cell(row, 1), "Desk check")
        put_text(sheet.cell(row, 2), spec["label"])
        sheet.cell(row, 3).value = f"=Outcome!{letter}{spec['row']}"
        sheet.cell(row, 3).number_format = NUM_FMT
        row += 1
    manual = next((spec for spec in body["rows"] if spec["label"] == MANUAL_REBATE), None)
    row += 1
    row = _check_heading(sheet, row, "Manual rebate")
    count = len(manual["manual_source"]) if manual else 0
    slots = len(manual["kids"]) if manual else 0
    put_text(sheet.cell(row, 1), "Manual rebate")
    put_text(sheet.cell(row, 2), (
        f"The PDF has {count} manual rebate lines. They are not copied onto Outcome. "
        f"{slots} orange input rows are left blank and the group total sums them. "
        "The descriptions and amounts are on FinanceOS (parsed), under Manual Rebate (Reward)."
    ))
    sheet.cell(row, 2).alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[row].height = 32
    sheet.auto_filter.ref = None
    return sheet


def _check_heading(sheet, row, title):
    put_text(sheet.cell(row, 1), title)
    sheet.cell(row, 1).font = FONT_WHITE
    sheet.cell(row, 1).fill = FILL_SECTION
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    return row + 1


def _any_target(metabase):
    return any(
        value is not None
        for by_desk in metabase["targets"].values()
        for value in by_desk.values()
    )


def _blank_targets(metabase):
    blanks = []
    for field, by_desk in metabase["targets"].items():
        for desk, value in by_desk.items():
            if value is None:
                blanks.append(f"{desk} {field}")
    return blanks


def _metric_total(metabase, field, sign=1):
    """Cent-rounded sum. Metabase exports extra decimals; the report shows 2."""
    total = sign * sum(record[field] for record in metabase["rows"])
    return float(Decimal(str(total)).quantize(CENT, rounding=ROUND_HALF_UP))


def _variances(doc, metabase, body):
    """Only lines whose FinanceOS amount is filled in. Blank yellow rows are not flagged."""
    pairs = []
    closed = find_label(doc, "Closed P/L")
    commissions = find_label(doc, "Commissions")
    if closed and closed["amount"] is not None:
        pairs.append(("Closed P/L", closed["amount"], _metric_total(metabase, "closed_pnl_usd")))
    if commissions and commissions["amount"] is not None:
        pairs.append(("Commissions", commissions["amount"], _metric_total(metabase, "commission_usd")))
    if body["deposit"]["value"] is not None:
        pairs.append(("Deposit", body["deposit"]["value"], _metric_total(metabase, "deposit_usd")))
    if body["withdrawal"]["value"] is not None:
        pairs.append((
            "Withdrawal", body["withdrawal"]["value"],
            _metric_total(metabase, "withdrawal_usd", sign=-1),
        ))
    if body["deposit"]["value"] is not None and body["withdrawal"]["value"] is not None:
        finance_ndp = q2(body["deposit"]["value"] + body["withdrawal"]["value"])
        meta_ndp = q2(
            _metric_total(metabase, "deposit_usd") + _metric_total(metabase, "withdrawal_usd", sign=-1)
        )
        pairs.append(("NDP", finance_ndp, meta_ndp))
    flagged = []
    for line, finance, meta in pairs:
        variance = float(Decimal(str(finance - meta)).quantize(CENT, rounding=ROUND_HALF_UP))
        if abs(variance) > VARIANCE_LIMIT:
            flagged.append({
                "line": line, "finance": finance, "metabase": meta, "variance": variance,
            })
    return flagged


def _assert_desk_totals(metabase):
    desks = metabase["desks"]
    for field in METRIC_FIELDS:
        raw = sum(record[field] for record in metabase["rows"])
        by_desk = {desk: 0.0 for desk in desks}
        for record in metabase["rows"]:
            by_desk[record["desk"]] += record[field]
        if abs(sum(by_desk.values()) - raw) > 0.02:
            raise PlDeskError(
                f"Desk totals for {field} do not reconcile to the Metabase total "
                f"({fmt_amt(sum(by_desk.values()))} vs {fmt_amt(raw)})."
            )


def _safe_remove(path):
    try:
        Path(path).unlink()
    except OSError:
        pass


def build_report(paths, output, company="dpm"):
    """Write the workbook. On any refused check, leave `output` uncreated."""
    output = Path(output)
    finance_path, metabase_path = identify_inputs(list(paths))
    if finance_path.suffix.lower() == ".pdf":
        doc = parse_pdf(finance_path, company)
    else:
        doc = parse_finance_excel(finance_path, company)
    metabase = load_metabase(metabase_path, company)
    period = _same_month(doc, metabase)
    checks = reconciliation_checks(doc)
    if any(item["status"] == "FAIL" for item in checks):
        raise PlDeskError("FinanceOS reconciliation failed.")
    _assert_desk_totals(metabase)
    body = build_outcome_specs(doc)
    variances = _variances(doc, metabase, body)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".partial.xlsx")
    workbook = Workbook()
    try:
        columns = _column_letters(metabase["header"])
        _write_outcome(workbook, doc, metabase, body, columns)
        _write_raw(workbook, metabase)
        _write_parsed(workbook, doc, checks)
        _write_copy(workbook, "Desks", metabase["copies"]["Desks"])
        _write_copy(workbook, "Target", metabase["copies"]["Target"])
        _write_copy(workbook, "Country code", metabase["copies"]["Country code"])
        _write_check(workbook, doc, metabase, body, checks, variances)
        _order_sheets(workbook)
        workbook.save(temporary)
    except Exception:
        _safe_remove(temporary)
        raise
    finally:
        workbook.close()
    temporary.replace(output)
    equity = doc["equity"]
    change = q2((equity["MT5"]["c"] + equity["Wallet"]["c"]) - (equity["MT5"]["p"] + equity["Wallet"]["p"]))
    trade = next(section["subtotal"] for section in doc["sections"] if section["code"] == "A")
    return {
        "period": period,
        "trade_subtotal": trade,
        "closed_pl": find_label(doc, "Closed P/L")["amount"],
        "equity_change": change,
        "ndp": _metric_total(metabase, "ndp_usd"),
        "output": str(output),
    }


def _order_sheets(workbook):
    order = ["Outcome", "FinanceOS (parsed)", RAW_SHEET, "Desks", "Target", "Country code", "Check"]
    for offset, title in enumerate(order):
        current = workbook.sheetnames.index(title)
        workbook.move_sheet(workbook[title], offset=offset - current)


def main(argv=None):
    parser = argparse.ArgumentParser(description="PL by Desk — monthly FinanceOS vs Metabase report")
    parser.add_argument("files", nargs="+", help="FinanceOS PDF or Excel, and the Metabase workbook")
    parser.add_argument("-o", "--output", required=True, help="result .xlsx")
    parser.add_argument("--company", default="dpm", help="company profile (default: dpm)")
    args = parser.parse_args(argv)
    output = Path(args.output)
    try:
        summary = build_report(args.files, output, company=args.company)
    except PlDeskError as exc:
        print(exc, file=sys.stderr)
        _safe_remove(output)
        return 2
    except Exception:
        print("PL by Desk failed.", file=sys.stderr)
        import traceback
        traceback.print_exc()
        _safe_remove(output)
        return 1
    print(f"PERIOD: {summary['period']}")
    print(f"Trade Sub-Total: {summary['trade_subtotal']:.2f}")
    print(f"Closed P/L: {summary['closed_pl']:.2f}")
    print(f"NDP: {summary['ndp']:.2f}")
    print(f"Equity change: {summary['equity_change']:.2f}")
    print(f"Output: {output.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
