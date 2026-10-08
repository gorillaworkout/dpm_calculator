#!/usr/bin/env python3
"""PL by Desk: parser refusals and the August 2026 end-to-end report.

    python3 test_pl_desk.py

The August fixtures live outside the repo (private financial files). The test
looks for them under the cloud-agent uploads directory.
"""
import io
import os
import re
import shutil
import tempfile
import time
from pathlib import Path

import pypdfium2 as pdfium
from openpyxl import Workbook, load_workbook

UPLOADS = Path("/home/ubuntu/.cursor/projects/workspace/uploads")
PDF = UPLOADS / "087abc019a138ade51dea1f103695e4671438ce6865273e326b4c993f43cd353_d196.pdf"
META = UPLOADS / "473f6ca6e9dd2fd09f8c1f5e763368769f77f010f04ab850b5f22e7b3c58c377_ea69.xlsx"
EXCEL = UPLOADS / "92eb037e2b8a257cabf17fadf0e8c49f435650278aa3ec45250a3be21cdff61b_1970.xlsx"
for fixture in (PDF, META, EXCEL):
    assert fixture.is_file(), fixture

import pl_desk as P

TMP = Path(tempfile.mkdtemp(prefix="pl-desk-test-"))
AMT_TAIL = re.compile(r"(\(?-?(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}\)?)(?!.*\d)")


def money_tail(text):
    matched = AMT_TAIL.search(text)
    assert matched, text
    return matched, P.money(matched.group(1))


def replace_tail(text, value):
    matched, _current = money_tail(text)
    return text[:matched.start(1)] + P.fmt_amt(value) + text[matched.end(1):]


def base_lines():
    lines, engine = P.read_pdf_lines(PDF)
    assert engine == "pdfplumber"
    return [dict(line) for line in lines]


LINES = base_lines()


def clone():
    return [dict(line) for line in LINES]


def first(lines, pred):
    for index, line in enumerate(lines):
        if pred(line["text"]):
            return index, line
    raise AssertionError("line not found")


def expect_parse_error(lines, fragment):
    doc = P.parse_pdf_lines(lines)
    try:
        P.validate_document(doc)
    except P.PlDeskError as exc:
        assert fragment in str(exc), str(exc)
        return str(exc)
    raise AssertionError(f"expected a refusal containing {fragment!r}")


def absent(path):
    path = Path(path)
    assert not path.exists(), path
    assert not path.with_name(path.stem + ".partial.xlsx").exists()


def expect_build_error(paths, fragment):
    output = TMP / f"refused-{abs(hash(fragment)) % 10**8}.xlsx"
    try:
        P.build_report(paths, output)
    except P.PlDeskError as exc:
        assert fragment in str(exc), str(exc)
        absent(output)
        return str(exc)
    raise AssertionError(f"expected a refusal containing {fragment!r}")


def rows_named(sheet, label):
    return [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == label]


def fill_rgb(cell):
    color = cell.fill.fgColor
    return (color.rgb or "") if color is not None else ""


# ---------------------------------------------------------------- parser
assert P.money("(5,276,735.09)") == -5276735.09
assert P.money("-5,276,735.09") == -5276735.09
assert P.money("10.00") == 10.0

bracket = clone()
index, line = first(bracket, lambda text: text.startswith("- Closed P/L"))
assert "-5,276,735.09" in line["text"]
line["text"] = line["text"].replace("-5,276,735.09", "(5,276,735.09)", 1)
doc = P.parse_pdf_lines(bracket)
P.validate_document(doc)
assert P.find_label(doc, "Closed P/L")["amount"] == -5276735.09

missing = clone()
index, line = first(missing, lambda text: text.startswith("- Commissions "))
line["text"] = line["text"].replace("Commissions", "Commission", 1)
expect_parse_error(missing, "Commissions")

broken = clone()
index, line = first(broken, lambda text: text.startswith("Sub-Total:"))
_matched, current = money_tail(line["text"])
line["text"] = replace_tail(line["text"], current + 1)
message = expect_parse_error(broken, "Sub-Total")
assert "diff" in message

dropped = clone()
index, line = first(dropped, lambda text: text.startswith("- Swap (deals)"))
del dropped[index]
message = expect_parse_error(dropped, "Sub-Total")
assert "-5,406.04" in message

nested = clone()
index, line = first(nested, lambda text: text.startswith("- Welcome Bonus "))
del nested[index]
expect_parse_error(nested, "Reward Other Promotions")

lost = clone()
index, line = first(lost, lambda text: text.startswith("- Closed P/L"))
line["text"] = "- Closed P/L"
expect_parse_error(lost, "Closed P/L")

fresh = clone()
index, line = first(fresh, lambda text: text.startswith("- Closed P/L"))
fresh.insert(index + 1, {"page": line["page"], "x0": line["x0"], "text": "- Extra Fee  -10.00"})
for pred in (
    lambda text: text.startswith("Sub-Total:"),
    lambda text: text.startswith("Total Movement:"),
    lambda text: "(Diff):" in text,
):
    _i, target = first(fresh, pred)
    _matched, current = money_tail(target["text"])
    target["text"] = replace_tail(target["text"], round(current - 10, 2))
doc = P.parse_pdf_lines(fresh)
P.validate_document(doc)
added = [item["label"] for item in P.new_labels(doc)]
assert added == ["Extra Fee"], added
assert P.find_label(doc, "Extra Fee")["amount"] == -10.0

blank_pdf = TMP / "scanned.pdf"
document = pdfium.PdfDocument.new()
document.new_page(595, 842)
document.save(str(blank_pdf))
try:
    P.build_report([blank_pdf, META], TMP / "scanned-out.xlsx")
except P.PlDeskError as exc:
    assert "text layer" in str(exc), str(exc)
else:
    raise AssertionError("a scanned PDF must be refused")
absent(TMP / "scanned-out.xlsx")


def edited_meta(name, edit):
    dest = TMP / name
    shutil.copy(META, dest)
    workbook = load_workbook(dest)
    edit(workbook)
    workbook.save(dest)
    workbook.close()
    return dest


def query(workbook):
    return workbook["Query result-Desk Rebate"]


def duplicate_country(workbook):
    sheet = query(workbook)
    last = sheet.max_row + 1
    for col in range(1, sheet.max_column + 1):
        sheet.cell(last, col).value = sheet.cell(2, col).value


def two_desks(workbook):
    workbook["Desks"].cell(9, 6).value = "VN"


def wrong_month(workbook):
    sheet = query(workbook)
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value is not None:
            sheet.cell(row, 1).value = 20260701


def drop_column(workbook):
    sheet = query(workbook)
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(1, col).value == "closed_pnl_usd":
            sheet.cell(1, col).value = "closed_pnl"
            return
    raise AssertionError("closed_pnl_usd header missing")


def text_amount(workbook):
    query(workbook).cell(2, 4).value = "n/a"


expect_build_error([PDF, edited_meta("dup.xlsx", duplicate_country)], "appears more than once")
expect_build_error([PDF, edited_meta("two.xlsx", two_desks)], "more than one desk")
expect_build_error([PDF, edited_meta("month.xlsx", wrong_month)], "2026-07")
expect_build_error([PDF, edited_meta("col.xlsx", drop_column)], "missing required columns")
expect_build_error([PDF, edited_meta("text.xlsx", text_amount)], "not a number")
expect_build_error([META, META], "Upload two files")

book = Workbook()
cell = book.active["A1"]
P.put_text(cell, "=1+1")
assert cell.value == "=1+1" and cell.data_type == "s" and cell.quotePrefix
quote_path = TMP / "quote.xlsx"
book.save(quote_path)
book.close()
reloaded = load_workbook(quote_path)
assert reloaded.active["A1"].value == "=1+1"
assert reloaded.active["A1"].data_type == "s"
assert reloaded.active["A1"].quotePrefix
reloaded.close()


# ---------------------------------------------------------------- August report
def inspect(summary, label, wrapped=True):
    assert summary["period"] == "2026-08", summary
    assert summary["trade_subtotal"] == -4628109.21, summary
    assert summary["closed_pl"] == -5276735.09, summary
    assert summary["ndp"] == 4566872.28, summary
    assert summary["equity_change"] == 398488.72, summary
    workbook = load_workbook(summary["output"])
    try:
        assert workbook.sheetnames == [
            "Outcome", "FinanceOS (parsed)", "Raw Metabase", "Desks", "Target",
            "Country code", "Check",
        ]
        outcome = workbook["Outcome"]
        assert outcome["A1"].value == "Monthly Trading Account Reconciliation"
        assert "Daily" not in outcome["A1"].value
        change = (outcome["C10"].value + outcome["C11"].value) - (outcome["B10"].value + outcome["B11"].value)
        assert round(change, 2) == 398488.72

        trade = rows_named(outcome, "A.) Trade  Sub-Total")
        assert trade == [14], trade
        trade_cell = outcome.cell(14, 2)
        assert trade_cell.value.startswith("=SUM(") and trade_cell.data_type == "f"
        cash_header = rows_named(outcome, "B.) Cash Movement  Sub-Total")[0]
        trade_values = [
            outcome.cell(row, 2).value
            for row in range(15, cash_header)
            if isinstance(outcome.cell(row, 2).value, (int, float))
        ]
        assert round(sum(trade_values), 2) == -4628109.21

        closed_row = rows_named(outcome, "Closed P/L")[0]
        assert outcome.cell(closed_row, 2).value == -5276735.09
        assert outcome.cell(closed_row, 2).data_type == "n"
        assert outcome.cell(closed_row, 4).value == f"=B{closed_row}-C{closed_row}"
        assert outcome.cell(closed_row, 4).data_type == "f"
        assert "SUMIFS" in outcome.cell(closed_row, 7).value
        assert "F$6" in outcome.cell(closed_row, 7).value
        assert outcome.cell(closed_row, 21).value.endswith(f"-C{closed_row}")

        commission_row = rows_named(outcome, "Commissions")[0]
        deposit_row = rows_named(outcome, "Deposit")[0]
        withdrawal_row = rows_named(outcome, "Withdrawal")[0]
        ndp_row = rows_named(outcome, "NDP")[0]
        assert outcome.cell(ndp_row, 2).value == f"=B{deposit_row}+B{withdrawal_row}"
        assert outcome.cell(ndp_row, 3).value == f"=C{deposit_row}+C{withdrawal_row}"
        assert "-" not in outcome.cell(ndp_row, 2).value
        assert outcome.cell(ndp_row, 7).value == f"=G{deposit_row}+G{withdrawal_row}"

        rebate_row = rows_named(outcome, "Rebate")[0]
        assert outcome.cell(rebate_row, 2).value is None
        assert outcome.cell(rebate_row, 4).value == f'=IF(B{rebate_row}="","",B{rebate_row}-C{rebate_row})'
        gross_row = rows_named(outcome, "Gross profit")[0]
        adjustment = rows_named(outcome, "C.) Adjustment  Sub-Total")[0]
        assert outcome.cell(gross_row, 2).value == f"=-B14-B{adjustment}-B{rebate_row}"

        sharing = rows_named(outcome, "Profit Sharing")[0]
        assert outcome.cell(sharing, 2).value is None
        assert fill_rgb(outcome.cell(sharing, 1)).endswith("FFFF00")
        for excluded in (
            "MT5 Transfer IN", "MT5 Transfer OUT",
            "Profit by sharing (Wallet's Profit Sharing)",
        ):
            row = rows_named(outcome, excluded)[0]
            assert outcome.cell(row, 2).value is None, excluded
            assert f"B{row}" not in outcome.cell(adjustment, 2).value

        cash_deposit = rows_named(outcome, "Deposit to MT5 Trading Account")[0]
        assert outcome.cell(cash_deposit, 2).value is None
        assert outcome.row_dimensions[cash_deposit].hidden

        manual = rows_named(outcome, "Manual Rebate (Reward)")[0]
        formula = outcome.cell(manual, 2).value
        matched = re.fullmatch(r"=SUM\(B(\d+):B(\d+)\)", formula)
        assert matched, formula
        start, end = int(matched.group(1)), int(matched.group(2))
        assert end - start + 1 == 11
        for row in range(start, end + 1):
            assert outcome.cell(row, 1).value in (None, "")
            assert outcome.cell(row, 2).value is None
            assert fill_rgb(outcome.cell(row, 1)).endswith("FFC000")
            assert fill_rgb(outcome.cell(row, 2)).endswith("FFC000")

        rules = list(outcome.conditional_formatting._cf_rules.values())
        assert rules and "ABS(D14)>1000" in rules[0][0].formula[0]

        raw = workbook["Raw Metabase"]
        headers = [raw.cell(1, col).value for col in range(1, raw.max_column + 1)]
        desk_col = headers.index("desk") + 1
        code_col = headers.index("country_code") + 1
        ndp_col = headers.index("ndp_usd") + 1
        closed_col = headers.index("closed_pnl_usd") + 1
        commission_col = headers.index("commission_usd") + 1
        by_code = {}
        ndp_total = 0.0
        commission_total = 0.0
        closed_total = 0.0
        for row in range(2, raw.max_row + 1):
            code = raw.cell(row, code_col).value
            if not code:
                continue
            by_code[code] = raw.cell(row, desk_col).value
            ndp_total += raw.cell(row, ndp_col).value
            commission_total += raw.cell(row, commission_col).value
            closed_total += raw.cell(row, closed_col).value
            if code == "EG":
                assert fill_rgb(raw.cell(row, desk_col)).endswith("FFFF00")
        assert by_code["EG"] == "GEM"
        assert by_code["VN"] == "VN"
        assert "Thai " not in set(by_code.values())
        assert round(ndp_total, 2) == 4566872.28
        assert round(closed_total, 2) == -5273899.21
        finance_commission = outcome.cell(commission_row, 2).value
        commission_variance = round(finance_commission - round(commission_total, 2), 2)
        assert abs(commission_variance) == 34.84
        assert abs(commission_variance) < 1000

        check = workbook["Check"]
        flagged = [
            check.cell(row, 2).value
            for row in range(1, check.max_row + 1)
            if check.cell(row, 1).value == "Variance over 1,000"
        ]
        assert "Closed P/L" in flagged
        assert "Commissions" not in flagged
        closed_check = next(
            row for row in range(1, check.max_row + 1)
            if check.cell(row, 1).value == "Variance over 1,000" and check.cell(row, 2).value == "Closed P/L"
        )
        assert check.cell(closed_check, 3).value == -5276735.09
        assert check.cell(closed_check, 4).value == -5273899.21
        assert check.cell(closed_check, 5).value == -2835.88
        text = " ".join(
            str(check.cell(row, col).value or "")
            for row in range(1, check.max_row + 1)
            for col in range(1, 3)
        )
        assert "test" in text.casefold()
        assert "Extra Fee" not in text
        if wrapped:
            assert "wrapped label" in text
        assert "Assigned to GEM" in text
        assert "EG" in text
        assert "Targets are blank" in text
        desk_checks = [
            check.cell(row, 3).value
            for row in range(1, check.max_row + 1)
            if check.cell(row, 1).value == "Desk check" and check.cell(row, 2).value == "Closed P/L"
        ]
        assert desk_checks == [f"=Outcome!U{closed_row}"]
        print(f"OK {label}")
    finally:
        workbook.close()


inspect(P.build_report([PDF, META], TMP / "aug.pdf.xlsx"), "pdf end to end")
inspect(P.build_report([EXCEL, META], TMP / "aug.xlsx.xlsx"), "excel fallback", wrapped=False)


# ---------------------------------------------------------------- web wiring
os.environ["DW_AUTH_DB"] = str(TMP / "users.db")
os.environ["DW_JOB_DIR"] = str(TMP / "jobs")
os.environ["DW_ADMIN_USER"] = "tester"
os.environ["DW_ADMIN_PASS"] = "tester-password-1"

import app as A  # noqa: E402

assert "pl-desk" not in A.KVB_TOOLS
assert A.TOOLS["pl-desk"][2] == ("pl_desk.py",)
A.app.config["TESTING"] = True
client = A.app.test_client()
client.environ_base["HTTP_AUTHORIZATION"] = "Basic " + __import__("base64").b64encode(
    b"tester:tester-password-1").decode()

page = client.get("/tool/pl-desk")
assert page.status_code == 200
assert b'id="finance-file"' in page.data and b'id="metabase-file"' in page.data
assert b'name="bulan"' not in page.data and b'name="saldo_jw"' not in page.data
assert b"var DIRECT = 80 * 1024 * 1024" in page.data
assert client.get("/kvb/tool/pl-desk").status_code == 404
kvb_home = client.get("/kvb")
assert kvb_home.status_code == 200 and b"PL by Desk" not in kvb_home.data

one = client.post("/tool/pl-desk", data={"file": (io.BytesIO(PDF.read_bytes()), "aug.pdf")})
assert one.status_code == 400 and b"two files" in one.data

with PDF.open("rb") as finance, META.open("rb") as metabase:
    posted = client.post("/tool/pl-desk", data={
        "file": [(finance, "FinanceOS Aug 2026.pdf"), (metabase, "desk-rebate.xlsx")],
    })
assert posted.status_code == 303, posted.status_code
job_id = posted.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
state = None
for _ in range(240):
    state = A._read_state(job_id)
    if state and state["state"] in ("done", "failed"):
        break
    time.sleep(0.25)
assert state and state["state"] == "done", state
assert state["name"] == "PL by Desk - Aug 2026.xlsx"
download = client.get(f"/job/{job_id}/download")
assert download.status_code == 200, download.status_code
assert download.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
assert download.data[:2] == b"PK"
assert b"Aug 2026" in download.headers["Content-Disposition"].encode()

print("OK pl desk")
