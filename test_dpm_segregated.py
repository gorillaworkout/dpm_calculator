"""DPM Segregated monthly roll-up on a synthetic month: python3 test_dpm_segregated.py"""
import csv
import datetime
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

import dpm_segregated as S

HEADER = ["Login", "Name", "Currency", "Deposit", "Withdrawal", "Closed P/L", "Commission",
          "Volume", "Margin Level %", "Prev Balance", "Balance", "Equity"]
SEPT = [datetime.date(2026, 9, d) for d in range(1, 31)]
LAST = SEPT[-1]


def jpy_rate(day):
    return 100.0 + day.day          # a different rate every day


def rows_for(day):
    """1001 USD; 1002 JPY; 1003 gone on the last day; 1004 Equity 0 on the last day."""
    d = day.day
    out = [
        ["1001", "Alice", "USD", 10, -1, 2.5, -0.5, 0.1, 350, 1000 + d - 1, 1000 + d, 1000 + d + 0.5],
        ["1002", "=Bob", "JPY", 1000, 0, 500, -100, 1, "", 50000 + d - 1, 50000 + d, 51000 + d],
    ]
    if day != LAST:
        out.append(["1003", "Cara", "USD", 1, 0, 0, 0, 0, "", 10, 10, 10])
    out.append(["1004", "Dan", "USD", 0, 0, 1, 0, 0, "", 5, 5, 0 if day == LAST else 7])
    return out


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_xlsx(path, header, rows, title=None, sheet="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    if title:
        ws.append([title])
        ws.append([])
    ws.append(header)
    for r in rows:
        ws.append([float(v) if isinstance(v, (int, float)) else v for v in r])
    wb.save(path)


def write_fx(path, days=SEPT, skip=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    for d in days:
        ws.append([d, "USD", None])
        if (d, "JPY") not in skip:
            ws.append([d, "JPY", jpy_rate(d)])
    wb.save(path)


def month(tmp, skip_fx=(), extra=None):
    files = []
    for i, d in enumerate(SEPT):
        if i % 2:
            p = tmp / f"Segregated {d:%Y-%m-%d}.csv"
            write_csv(p, HEADER, rows_for(d))
        else:
            p = tmp / f"segregated_{d:%Y%m%d}.xlsx"
            write_xlsx(p, HEADER, rows_for(d), title=f"Segregated {d}" if i == 0 else None)
        files.append(p)
    fx = tmp / "Client Equity - FX.xlsx"
    write_fx(fx, skip=skip_fx)
    return files + [fx] + (extra or [])


def sheet_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]


def cells(ws):
    head = [c.value for c in ws[1]]
    return head, {r[0].value: {head[i]: c for i, c in enumerate(r)} for r in ws.iter_rows(min_row=2)}


def verif(wb):
    return {r[0]: r[1] for r in wb["Verifikasi"].iter_rows(min_row=3, values_only=True) if r and r[0]}


def yellow(cell):
    return cell.fill is not None and cell.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00")


def stop(paths, out, needle):
    try:
        S.proses(paths, out)
    except SystemExit as exc:
        assert needle in str(exc), exc
        return
    raise AssertionError(f"expected STOP: {needle}")


# ---------------------------------------------------------------- unit bits
assert S.tanggal_nama_file("Segregated 2026-09-01.csv") == datetime.date(2026, 9, 1)
assert S.tanggal_nama_file("seg_20260902.xlsx") == datetime.date(2026, 9, 2)
assert S.tanggal_nama_file("seg 03-09-2026.csv") == datetime.date(2026, 9, 3)
assert S.tanggal_nama_file("1 Sept.xlsx", 2026) == datetime.date(2026, 9, 1)
assert S.tanggal_nama_file("1 Sept.xlsx") is None
assert S.tanggal_nama_file("report.csv") is None
assert S.tanggal_sel(20260901) == datetime.date(2026, 9, 1)
assert S.tanggal_sel("2026.09.01 00:00:00") == datetime.date(2026, 9, 1)
assert S.tanggal_sel(46266) == datetime.date(2026, 9, 1)
assert S.kunci(" Closed P/L ") == "CLOSEDPL"

with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    # ------------------------------------------------------------ full month
    paths = month(tmp)
    out = tmp / "out" / "Segregated.xlsx"
    r = S.proses(paths, out)
    assert r["_periode"] == "2026-09"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Monthly Summary", "Daily", "Verifikasi"], wb.sheetnames
    head, summ = cells(wb["Monthly Summary"])
    assert head[:6] == ["Account", "Currency", "Name", "Days", "First date", "Last date"], head
    assert "Margin Level %" not in head          # a level is never summed
    assert head[-1] == "Check"
    assert [h for h in head if h.endswith("(USD)")] == [
        "Deposit (USD)", "Withdrawal (USD)", "Closed P/L (USD)", "Commission (USD)",
        "Prev Balance (USD)", "Balance (USD)", "Equity (USD)"], head
    assert "Volume (USD)" not in head
    a = {k: c.value for k, c in summ["1001"].items()}
    assert a["Days"] == 30 and a["Deposit"] == 300 and a["Withdrawal"] == -30
    assert a["Closed P/L"] == 75 and a["Commission"] == -15 and abs(a["Volume"] - 3) < 1e-9
    assert a["Balance"] == 1030 and a["Equity"] == 1030.5      # last day, not summed
    assert a["Prev Balance"] == 1000                           # first day's opening
    assert a["Balance (USD)"] == 1030 and a["Deposit (USD)"] == 300
    assert a["Check"] is None
    b = {k: c.value for k, c in summ["1002"].items()}
    assert b["Name"] == "=Bob" and summ["1002"]["Name"].data_type == "s"
    assert b["Deposit"] == 30000
    assert b["Deposit (USD)"] == round(sum(1000 / jpy_rate(d) for d in SEPT), 2)
    assert b["Closed P/L (USD)"] == round(sum(500 / jpy_rate(d) for d in SEPT), 2)
    assert b["Balance"] == 50030 and b["Balance (USD)"] == round(50030 / jpy_rate(LAST), 2)
    assert b["Equity (USD)"] == round(51030 / jpy_rate(LAST), 2)
    c = summ["1003"]
    assert c["Days"].value == 29 and c["Last date"].value.date() == datetime.date(2026, 9, 29)
    assert c["Balance"].value is None and yellow(c["Balance"]) and yellow(c["Balance (USD)"])
    assert "not in the last day's file" in c["Check"].value and yellow(c["Check"])
    d = summ["1004"]
    assert d["Equity"].value == 0 and yellow(d["Equity"]) and d["Balance"].value == 5
    assert not yellow(d["Balance"]) and "Equity is 0" in d["Check"].value

    dh, daily = sheet_dicts(wb["Daily"])
    assert dh[:3] == ("Date", "Account", "Currency"), dh
    assert len(daily) == 4 * 30 - 1
    assert daily[0]["Date"].date() == datetime.date(2026, 9, 1)
    assert daily[-1]["Date"].date() == LAST
    j = [x for x in daily if x["Account"] == "1002" and x["Date"].date() == datetime.date(2026, 9, 5)][0]
    assert j["Deposit (USD)"] == round(1000 / 105, 2) and j["Balance (USD)"] == round(50005 / 105, 2)
    assert j["Margin Level %"] is None and daily[0]["Margin Level %"] == 350

    v = verif(wb)
    assert v["Account column"] == "Login" and v["Currency column"] == "Currency"
    assert v["Last day (Balance / Equity taken from)"] == "2026-09-30"
    assert "Balance, Equity" == v["Last day's value (not summed)"]
    assert v["First day's value (not summed)"] == "Prev Balance"
    assert "Volume" in v["Summed, not converted (counts / volume / already USD)"]
    assert "Margin Level %" in v["Numeric, not summed (prices / rates / levels): Daily only"]
    assert "1003" in v["Accounts missing on the last day (Balance / Equity blank)"]
    assert "1004" in v["Equity blank or 0 on the last day"]
    assert "Balance blank or 0 on the last day" not in v   # 1003 is listed as missing instead
    assert v["Daily files read"] == 30
    assert "Days of Sep 2026 with no file" not in v
    assert "Missing FX rate (USD left blank)" not in v

    # ------------------------------------------------------------ FX required
    stop(paths[:-1], tmp / "x.xlsx", "Client Equity FX file is required")

    # ------------------------------------------------------------ missing rate for one day
    sub = tmp / "missing"
    sub.mkdir()
    paths = month(sub, skip_fx={(datetime.date(2026, 9, 10), "JPY")})
    out = sub / "o.xlsx"
    S.proses(paths, out)
    wb = load_workbook(out)
    head, summ = cells(wb["Monthly Summary"])
    assert summ["1002"]["Deposit (USD)"].value is None and yellow(summ["1002"]["Deposit (USD)"])
    assert summ["1002"]["Balance (USD)"].value == round(50030 / jpy_rate(LAST), 2)
    assert summ["1001"]["Deposit (USD)"].value == 300
    assert "FX rate missing" in summ["1002"]["Check"].value
    v = verif(wb)
    assert v["Missing FX rate (USD left blank)"] == "2026-09-10 JPY: 1 rows"
    ws = wb["Daily"]
    dh = [c.value for c in ws[1]]
    bad = [row for row in ws.iter_rows(min_row=2)
           if row[1].value == "1002" and row[0].value.date() == datetime.date(2026, 9, 10)][0]
    assert bad[dh.index("Deposit (USD)")].value is None and yellow(bad[dh.index("Deposit (USD)")])

    # ------------------------------------------------------------ duplicate date, identical file, gaps
    sub = tmp / "dup"
    sub.mkdir()
    write_csv(sub / "a 2026-09-01.csv", HEADER, rows_for(SEPT[0]))
    write_csv(sub / "b 2026-09-01.csv", HEADER, rows_for(SEPT[0])[:1])
    write_csv(sub / "copy.csv", HEADER, rows_for(SEPT[0])[:1])   # same bytes as b
    (sub / "copy.csv").write_bytes((sub / "b 2026-09-01.csv").read_bytes())
    write_csv(sub / "c 2026-09-03.csv", HEADER, rows_for(SEPT[2]))
    write_fx(sub / "fx.csv.xlsx")
    S.proses(sorted(sub.iterdir()), sub / "o.xlsx")
    v = verif(load_workbook(sub / "o.xlsx"))
    assert "a 2026-09-01.csv, b 2026-09-01.csv" in v["Same date in more than one file (rows were added together)"]
    assert "'copy.csv' is byte-for-byte the same as 'b 2026-09-01.csv'" in v["Identical files skipped"] \
        or "'b 2026-09-01.csv' is byte-for-byte the same as 'copy.csv'" in v["Identical files skipped"]
    assert v["Days of Sep 2026 with no file"].startswith("2, 4, 5")

    # ------------------------------------------------------------ date column, header under a title,
    # different column order and names, missing currency column, bad number
    sub = tmp / "content"
    sub.mkdir()
    h2 = ["Report Date", "Account ID", "Profit", "Swap", "Balance", "Equity", "Profit (USD)"]
    write_xlsx(sub / "day one.xlsx", h2, [[datetime.date(2026, 9, 1), 7, 5, -1, 100, 101, 999]],
               title="Daily report")
    write_csv(sub / "day two.csv", list(reversed(h2)),
              [list(reversed(["2026-09-02", "7", "1,250.50", "N/A", "", "", "999"]))])
    write_fx(sub / "fx.xlsx")
    S.proses(sorted(sub.iterdir()), sub / "o.xlsx")
    wb = load_workbook(sub / "o.xlsx")
    head, summ = cells(wb["Monthly Summary"])
    row = {k: c.value for k, c in summ["7"].items()}
    assert row["Profit"] == 1255.5 and row["Profit (USD)"] == 1255.5   # recalculated, not 1998
    assert "Swap" not in head                                            # spoiled by 'N/A'
    assert row["Balance"] is None and "Balance blank" in row["Check"]
    v = verif(wb)
    assert "No currency column" in v and v["Account column"] == "Account ID"
    assert "Swap not summed" in " ".join(v) and "'N/A'" in v["Column Swap not summed"]
    assert v["USD columns in the files: recalculated from the base column"] == "Profit (USD)"
    assert v["Date taken from"] == "Date column"

    # ------------------------------------------------------------ year-less name takes the FX year
    sub = tmp / "sept"
    sub.mkdir()
    write_xlsx(sub / "1 Sept.xlsx", HEADER, rows_for(SEPT[0]))
    write_fx(sub / "fx.xlsx")
    r = S.proses(sorted(sub.iterdir()), sub / "o.xlsx")
    assert r["Period"].startswith("2026-09-01 to 2026-09-01")

    # ------------------------------------------------------------ no date anywhere
    sub = tmp / "nodate"
    sub.mkdir()
    write_csv(sub / "report.csv", HEADER, rows_for(SEPT[0]))
    write_fx(sub / "fx.xlsx")
    stop(sorted(sub.iterdir()), sub / "o.xlsx", "cannot tell which day 'report.csv' is")

    # ------------------------------------------------------------ unknown file
    write_csv(sub / "other.csv", ["a", "b"], [[1, 2]])
    (sub / "report.csv").rename(sub / "r 2026-09-01.csv")
    stop(sorted(sub.iterdir()), sub / "o.xlsx", "'other.csv'")

    # ------------------------------------------------------------ CLI, --delete-input, PERIOD line
    sub = tmp / "cli"
    sub.mkdir()
    paths = month(sub)
    res = subprocess.run([sys.executable, str(Path(S.__file__).resolve()), *map(str, paths),
                          "-o", str(sub / "Segregated.xlsx"), "--delete-input"],
                         capture_output=True, text=True)
    assert res.returncode == 0, res.stderr
    assert "PERIOD: 2026-09" in res.stdout and "CHECK: Accounts missing on the last day" in res.stdout
    assert all(not p.exists() for p in paths) and (sub / "Segregated.xlsx").exists()
    assert not list(sub.glob(".segregated-*")), "spool file left behind"

    # ------------------------------------------------------------ volume smoke: 31 days x 3,000 accounts
    sub = tmp / "big"
    sub.mkdir()
    oct_days = [datetime.date(2026, 10, d) for d in range(1, 32)]
    for d in oct_days:
        write_csv(sub / f"seg {d:%Y%m%d}.csv", HEADER,
                  [[str(10000 + k), f"N{k}", "JPY" if k % 3 == 0 else "USD", 1, 0, 1, 0, 0.01, "",
                    100, 100 + d.day, 100 + d.day] for k in range(3000)])
    write_fx(sub / "fx.xlsx", days=oct_days)
    r = S.proses(sorted(sub.iterdir()), sub / "o.xlsx")
    assert r["Daily rows"] == 93000 and r["Accounts (Monthly Summary rows)"] == 3000
    assert "Days of Oct 2026 with no file" not in r

print("test_dpm_segregated: OK")
