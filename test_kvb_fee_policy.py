"""KVB multi-fee consolidation must never import DPM EXTRA_FEES.

Self-contained: the private 'Template D&W.xlsx' is not in the repo, so the empty
template is generated from the synthetic source with buat_template.py.
"""
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
import hitung_dw

HERE = Path(__file__).parent


def source(path):
    wb = Workbook()
    d = wb.active
    assert d is not None
    d.title = "D"
    d.append(["Currency", "Paid Date", "Transaction", "Payment Gateway", "USD"])
    d.append(["USD", datetime(2026, 8, 1), 100, "TEST", 100])
    w = wb.create_sheet("W")
    w.append(["Currency", "Completed Date", "Transaction", "Payment Gateway", "USD", "Status"])
    x = wb.create_sheet("Xero")
    x.append(["Date", "USD"])
    x.append([datetime(2026, 8, 1), 1])
    for name in ("D&W TD FEE", "D&W TD FEE -add."):
        fee = wb.create_sheet(name)
        fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
        fee.append(["USD", "TEST", 0.01, 0.02])
    wb.save(path)


def run(*args):
    p = subprocess.run([sys.executable, *map(str, args)], cwd=HERE, capture_output=True, text=True)
    assert p.returncode == 0, p.stdout + p.stderr
    return p.stdout


def prepare(src, tpl, dst, *extra):
    run(HERE / "isi_template.py", src, "-t", tpl, "--bulan", "2026-08", *extra, "-o", dst)


def fee_keys(path):
    wb = load_workbook(path, data_only=True)
    keys = set(hitung_dw.load_fee_table(wb, pakai_extra_fees=False))
    wb.close()
    return keys


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    src = tmp / "source.xlsx"
    dpm = tmp / "dpm.xlsx"
    kvb = tmp / "kvb.xlsx"
    tpl = tmp / "Template D&W.xlsx"
    source(src)
    run(HERE / "buat_template.py", src, "-o", tpl)
    prepare(src, tpl, dpm)
    prepare(src, tpl, kvb, "--tanpa-extra-fees")
    assert ("VND", "BECKPAY") in fee_keys(dpm), "DPM must retain fallback fees"
    assert ("VND", "BECKPAY") not in fee_keys(kvb), "KVB must not receive DPM fallback fees"

print("OK KVB multi-fee policy")
