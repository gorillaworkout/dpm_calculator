#!/usr/bin/env python3
"""Terjemahkan export "KVB Plus" jadi bentuk yang dikenali isi_template.py dan
hitung_dw.py (format DPM) -- TANPA mengubah kedua file itu sama sekali.

Alur lengkap (dikonfirmasi user 1 Okt 2026, opsi B: jangan sentuh mesin DPM):
    python3 siapkan_kvb.py   "KVB Plus.xlsx"        -o terjemahan.xlsx
    python3 isi_template.py  terjemahan.xlsx --bulan 2026-08 -o template.xlsx
    python3 hitung_dw.py     template.xlsx --period 2026-08  -o hasil.xlsx

KVB dan DPM memakai sheet D/W/Xero/Fund Transfer Table/tabel fee dengan
JUMLAH & URUTAN KOLOM yang nyaris sama persis -- cuma beda NAMA header, nama
sheet, dan istilah 'Source Name' untuk rebate withdrawal. isi_template.py dan
hitung_dw.py SUDAH mencocokkan kolom berdasarkan nama header (bukan posisi),
jadi begitu nama-nama ini diterjemahkan ke istilah DPM, sisanya jalan apa
adanya lewat pipeline yang sudah ada -- lihat perbandingan lengkap yang sudah
dibahas dengan user.

Script ini HANYA mengganti NAMA (header kolom, nama sheet) dan MENORMALISASI
nilai kolom 'Source Name' di sheet W (lihat normalisasi_source_name()).
Angka/nilai transaksi tidak disentuh sama sekali.
"""
import argparse
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook


def norm(v):
    return " ".join(str(v if v is not None else "").strip().upper().split())


def cari_sheet(wb, kandidat):
    for x in wb.sheetnames:
        if norm(x) in kandidat:
            return x
    return None


# --- Alias nama kolom: {header KVB (dinormalisasi) -> header DPM} ------------
# Kolom yang TIDAK disebut di sini dibiarkan apa adanya -- sudah sama persis,
# atau cuma beda huruf besar/kecil (dicocokkan case-insensitive di hilir).
D_ALIAS = {
    "SOURCE.NAME": "Source Name",
    "SALES_UID": "Sales UID",
    "SALES_NAME": "Sales Name",
    "SALES_DEPT": "Sales Dept",
    "SALES_AREA": "Sales Area",
    "USER_ID": "User ID",
    "ACCOUNT": "Account",
    "METHOD": "Currency",
    "MT_ORDER": "MT Order",
    "TICKET": "Ticket",
    "USD (CRM)": "USD",
    "RATE": "Rate",
    "TRANSACTION": "Transaction",
    "RECEIVE_COMPANY_ACCOUNT": "Payment Gateway",
    "REFERENACE": "Reference",      # salah ketik di sumber KVB, bukan di sini
    "HASH": "Hash",
}

W_ALIAS = {
    "SOURCE.NAME": "Source Name",
    "PAYMENT CURRENCY": "Currency",
    "USD (CRM)": "USD",
    "RATE": "Rate",
    "CHARGES": "Charges",
    "TRANSACTION": "Transaction",
    "STATUS": "Status",
}

FEE_ALIAS = {
    "PAYMENT CHANNEL": "Payment Gateway",
}

D_SUMBER = ("D", "DEPOSIT", "DEPOSITS")
W_SUMBER = ("W", "WITHDRAWAL", "WITHDRAWALS")
XERO_SUMBER = ("XERO",)
FTT_SUMBER = ("TRANSFER", "FUND TRANSFER TABLE", "FUND TRANSFER")
FEE_SUMBER = ("HANDING FEE", "HANDLING FEE", "D&W FEE")

# Sheet W KVB menulis Source Name dengan PREFIX BULAN, mis. "07-2026 Withdrawal"
# dan "07-2026 Wallet withdrawal" -- padahal DPM menulisnya bersih: "Withdrawal"
# / "Rebate Withdrawal". Rumus MTOATD (mtoatd_spec.py, dipanggil dari
# hitung_dw.py) membedakan MT4出金 vs Wallet出金 dari TEKS PERSIS itu
# (case-sensitive, exact match, bukan "startswith"). TANPA normalisasi ini,
# SEMUA baris W dari KVB gagal cocok di kedua baris itu -- Wallet出金 dan
# MT4出金 akan kosong total, padahal datanya ada.
_PREFIX_BULAN = re.compile(r"^\d{1,2}-\d{4}\s+")


def normalisasi_source_name(v):
    """'07-2026 Withdrawal' -> 'Withdrawal' ; '08-2026 Wallet withdrawal' ->
    'Rebate Withdrawal'. Pola yang tidak dikenal dibiarkan apa adanya (supaya
    kelihatan di hasil kalau ternyata ada varian baru, bukan diam-diam salah)."""
    if v in (None, ""):
        return v
    teks = str(v).strip()
    tanpa_prefix = _PREFIX_BULAN.sub("", teks).strip()
    kunci = tanpa_prefix.lower()
    if kunci == "wallet withdrawal":
        return "Rebate Withdrawal"
    if kunci == "withdrawal":
        return "Withdrawal"
    return teks


def salin_dengan_alias(ws_src, wb_out, nama_sheet, alias, kolom_sumber_sn=None):
    """Salin SELURUH isi ws_src ke sheet baru 'nama_sheet' di wb_out. Header
    baris 1 diganti sesuai 'alias' (kunci = norm(header asli)); header yang
    tidak ada di 'alias' disalin apa adanya. Nilai baris TIDAK disentuh,
    KECUALI kolom Source Name (indeks 0-based 'kolom_sumber_sn') yang lewat
    normalisasi_source_name() -- cuma dipakai untuk sheet W."""
    ws_out = wb_out.create_sheet(nama_sheet)
    rows = ws_src.iter_rows(values_only=True)
    header = next(rows)
    ws_out.append([alias.get(norm(h), h) for h in header])
    n = 0
    for row in rows:
        if kolom_sumber_sn is not None and len(row) > kolom_sumber_sn:
            row = list(row)
            row[kolom_sumber_sn] = normalisasi_source_name(row[kolom_sumber_sn])
        ws_out.append(row)
        n += 1
    return n


def salin_apa_adanya(ws_src, wb_out, nama_sheet):
    """Salin isi sheet TANPA mengubah apa pun -- dipakai untuk Xero dan
    Fund Transfer Table, yang pembacanya (hitung_dw.py) sudah mencari kolom
    lewat NAMA header di dalam sheet, bukan posisi, jadi kolom tambahan KVB
    (mis. 'Xero Rate', duplikat 'USD (Xero)') aman diikutkan apa adanya."""
    ws_out = wb_out.create_sheet(nama_sheet)
    n = 0
    for row in ws_src.iter_rows(values_only=True):
        ws_out.append(row)
        n += 1
    return max(n - 1, 0)


def proses(path_in: Path, path_out: Path) -> dict:
    wb_src = openpyxl.load_workbook(path_in, read_only=True, data_only=True)

    nama_d = cari_sheet(wb_src, D_SUMBER)
    nama_w = cari_sheet(wb_src, W_SUMBER)
    nama_xero = cari_sheet(wb_src, XERO_SUMBER)
    nama_ftt = cari_sheet(wb_src, FTT_SUMBER)
    nama_fee = cari_sheet(wb_src, FEE_SUMBER)

    hilang = [label for label, nama in (
        ("D", nama_d), ("W", nama_w), ("Xero", nama_xero),
        ("Fund Transfer Table / Transfer", nama_ftt),
        ("D&W FEE / Handing Fee", nama_fee),
    ) if nama is None]
    if hilang:
        wb_src.close()
        sys.exit(f"STOP: sheet berikut tidak ditemukan di '{path_in.name}': {', '.join(hilang)}\n"
                 f"Sheet yang ada di file ini: {wb_src.sheetnames}\n"
                 "Pastikan ini memang export KVB Plus (Transfer/D/W/Xero/Handing Fee), "
                 "bukan file lain.")

    wb_out = Workbook()
    wb_out.remove(wb_out.active)      # buang sheet kosong bawaan Workbook()

    n_d = salin_dengan_alias(wb_src[nama_d], wb_out, "D", D_ALIAS)

    # Indeks kolom Source Name di sheet W HARUS dicari sebelum header diganti
    # nama (masih header asli KVB waktu dicari).
    header_w = next(wb_src[nama_w].iter_rows(max_row=1, values_only=True))
    i_sn_w = next((i for i, h in enumerate(header_w) if norm(h) == "SOURCE.NAME"), None)
    n_w = salin_dengan_alias(wb_src[nama_w], wb_out, "W", W_ALIAS, kolom_sumber_sn=i_sn_w)

    n_xero = salin_apa_adanya(wb_src[nama_xero], wb_out, "Xero")
    n_ftt = salin_apa_adanya(wb_src[nama_ftt], wb_out, "Fund Transfer Table")
    n_fee = salin_dengan_alias(wb_src[nama_fee], wb_out, "D&W FEE", FEE_ALIAS)

    path_out.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(path_out)
    wb_src.close()

    ringkas = {"d": n_d, "w": n_w, "xero": n_xero, "ftt": n_ftt, "fee": n_fee,
               "source_name_dinormalisasi": i_sn_w is not None}

    print(f"Sumber   : {path_in.name}")
    print(f"  sheet 'D'                  (dari '{nama_d}')   : {n_d:,} baris")
    tanda_sn = (" (Source Name dinormalisasi: prefix bulan dibuang, "
                "'Wallet withdrawal' -> 'Rebate Withdrawal')" if i_sn_w is not None else
                " -- !! kolom Source Name TIDAK KETEMU: MTOATD Wallet出金 akan kosong semua")
    print(f"  sheet 'W'                  (dari '{nama_w}')   : {n_w:,} baris{tanda_sn}")
    print(f"  sheet 'Xero'                (dari '{nama_xero}')              : {n_xero:,} baris")
    print(f"  sheet 'Fund Transfer Table' (dari '{nama_ftt}')              : {n_ftt:,} baris")
    print(f"  sheet 'D&W FEE'             (dari '{nama_fee}')          : {n_fee:,} baris")
    print(f"Simpan   : {path_out}")
    return ringkas


def main():
    ap = argparse.ArgumentParser(
        description="Terjemahkan export KVB Plus jadi bentuk yang dikenali "
                    "isi_template.py / hitung_dw.py (tanpa mengubah keduanya)")
    ap.add_argument("input", help="workbook export KVB Plus (sheet Transfer/D/W/Xero/Handing Fee)")
    ap.add_argument("-o", "--output", help="file hasil (default: <input>-kvb-diterjemahkan.xlsx)")
    args = ap.parse_args()

    src = Path(args.input).expanduser()
    if not src.is_file():
        sys.exit(f"STOP: file tidak ditemukan: {src}")
    dst = Path(args.output).expanduser() if args.output else \
        src.with_name(f"{src.stem}-kvb-diterjemahkan.xlsx")

    proses(src, dst)


if __name__ == "__main__":
    main()
