"""KVB D&W: Excel-serial dates survive the month filter, 1970 dates are reported,
and gateway names that differ only in case/whitespace become one gateway
(fee lookup, D&W Report grouping, detail rows)."""
import datetime
import importlib.util
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
PY = sys.executable


def muat(nama):
    spec = importlib.util.spec_from_file_location(nama + "_t", HERE / f"{nama}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


hd = muat("hitung_dw")
it = muat("isi_template")
sk = muat("siapkan_kvb")

# ---------------------------------------------------------------- unit checks
assert it.tanggal(46240) == datetime.date(2026, 8, 6)
assert it.tanggal(46240.5) == datetime.date(2026, 8, 6)
assert it.tanggal(datetime.datetime(2026, 8, 6, 13, 0)) == datetime.date(2026, 8, 6)
assert it.tanggal(25569) is None and it.tanggal_tidak_sah(25569)          # 1970-01-01
assert it.tanggal(datetime.datetime(1970, 1, 1)) is None
assert it.tanggal_tidak_sah(datetime.datetime(1970, 1, 1))
assert not it.tanggal_tidak_sah(datetime.datetime(2026, 8, 1))
assert it.tanggal("abc") is None and not it.tanggal_tidak_sah("abc")
assert hd.as_date(46240) == datetime.date(2026, 8, 6)
assert hd.as_date(20260806) == datetime.date(2026, 8, 6)                  # YYYYMMDD tetap
assert hd.as_date(25569) is None
assert sk.serial_ke_tanggal(46240) == datetime.datetime(2026, 8, 6)
assert sk.serial_ke_tanggal(25569) == 25569                                # tidak diubah
assert sk.tanggal_aneh(25569) and sk.tanggal_aneh(datetime.datetime(1970, 1, 1))

peta = sk.peta_gateway_kanon(["Chippay", "ChipPay ", "ChipPay", "USDT-ERC20 ", "USDT-ERC20",
                              "Foo  Bar", "foo bar", "foo bar"], ["ChipPay"])
assert sk.kanon_gateway("chippay", peta) == "ChipPay"
assert sk.kanon_gateway("  CHIPPAY  ", peta) == "ChipPay"
assert sk.kanon_gateway("USDT-ERC20 ", peta) == "USDT-ERC20"
assert sk.kanon_gateway("FOO   BAR", peta) == "foo bar"                   # paling sering

# fee lookup: spelling variants hit the same fee row
fee_wb = Workbook()
ws = fee_wb.active
ws.title = "D&W FEE"
ws.append(["Payment Gateway", "Currency", "Deposit", "Withdrawal"])
ws.append(["ChipPay", "CNY", 0.004, 0.007])
ws.append(["USDT-ERC20", "USDT", 0.01, 0.02])
tabel = hd.load_fee_table(fee_wb, pakai_extra_fees=False)
for g in ("Chippay", "ChipPay ", "CHIPPAY", " chip pay "):
    komp, _f, key, _how = hd.lookup_fee(tabel, "CNY", g, "withdrawal")
    assert komp and abs(komp["pct"] - 0.007) < 1e-12, (g, komp)
komp, *_ = hd.lookup_fee(tabel, "USDT", "USDT-ERC20 ", "deposit")
assert komp and abs(komp["pct"] - 0.01) < 1e-12
assert hd.peta_nama_gateway({"Chippay": 5, "CHIPPAY ": 1}, tabel)["CHIPPAY"] == "ChipPay"
assert hd.peta_nama_gateway({"Zpay": 1, "ZPAY": 3})["ZPAY"] == "ZPAY"
assert hd.kunci_channel("CNY", "Chippay ") == hd.kunci_channel("CNY", "ChipPay")

# report grouping merges pivot keys
rep = {"records": [], "pivot": {
    ("withdrawal", "CNY", "ChipPay"): {"crm_usd": 10.0, "tx": 70.0, "xero_usd": 10.0,
                                      "charges": 0.0, "n": 3, "urut": 0},
    ("withdrawal", "CNY", "Chippay"): {"crm_usd": 1.0, "tx": 7.0, "xero_usd": 1.0,
                                      "charges": 0.0, "n": 1, "urut": 1},
    ("withdrawal", "CNY", "ChipPay "): {"crm_usd": 2.0, "tx": 14.0, "xero_usd": 2.0,
                                       "charges": 0.0, "n": 1, "urut": 2}}}
hd.satukan_nama_gateway(rep, tabel)
assert list(rep["pivot"]) == [("withdrawal", "CNY", "ChipPay")]
assert rep["pivot"][("withdrawal", "CNY", "ChipPay")]["tx"] == 91.0

# ---------------------------------------------------------------- end to end
tmp = Path(tempfile.mkdtemp())
src = Workbook()
src.remove(src.active)
t = src.create_sheet("Transfer")
t.append(["汇出"] + [None] * 7 + ["收款"])
t.append(["汇款日期", "汇出部门", "支付渠道", "汇出币种", "Xero Rate", "金额", "USD (Xero)", "手续费",
          "收款日期", "收款部门", "收款人", "收入币种", "金额", "手续费", "USD (Xero)"])
t.append([46238, "KVB", "Chippay ", "CNY", None, 1000, None, None,
          46238, "KVB", "J Wallet", "USDT", 140, None, None])
d = src.create_sheet("D")
d.append(["Source.Name", "user_id", "account", "Paid Date", "Settlement Date", "method", "ticket",
          "USD (CRM)", "rate", "transaction", "receive_company_account"])
d.append(["08-2026 Deposit", 1, 1, datetime.datetime(2026, 8, 3), datetime.datetime(2026, 8, 3),
          "CNY", 1, 140, 0.14, 1000, "ChipPay"])
d.append(["08-2026 Deposit", 2, 2, 46239, 46239, "CNY", 2, 140, 0.14, 1000, "chippay "])
w = src.create_sheet("W")
w.append(["Source.Name", "User ID", "Account", "Apply Date", "Payment Currency", "Payment Gateway",
          "Reference", "USD (CRM)", "rate", "charges", "transaction", "Completed Date",
          "Settlement Date", "status"])
w.append(["08-2026 Withdrawal", 1, 1, 46240, "CNY", "ChipPay", 101, 140, 0.14, 0, 1000,
          46240, 46240, "finish"])                                         # serial Aug
w.append(["08-2026 Withdrawal", 2, 2, 46241, "CNY", "Chippay", 102, 70, 0.14, 0, 500,
          46241, 46241, "finish"])                                         # serial Aug
w.append(["08-2026 Withdrawal", 3, 3, 46242, "CNY", "ChipPay ", 103, 14, 0.14, 0, 100,
          datetime.datetime(2026, 8, 8), 46242, "finish"])                 # real date
w.append(["08-2026 Withdrawal", 4, 4, 46242, "USDT", "USDT-ERC20 ", 104, 50, 1, 0, 50,
          46242, 46242, "finish"])
w.append(["08-2026 Withdrawal", 5, 5, 46242, "USDT", "USDT-ERC20", 105, 25, 1, 0, 25,
          46242, 46242, "finish"])
w.append(["07-2026 Withdrawal", 6, 6, 46220, "CNY", "ChipPay", 106, 14, 0.14, 0, 100,
          datetime.datetime(2026, 7, 17), 46220, "finish"])                # July -> dropped
w.append(["08-2026 Withdrawal", 7, 7, 46243, "CNY", "ChipPay", 107, 14, 0.14, 0, 100,
          datetime.datetime(1970, 1, 1), 46243, "finish"])                 # 1970
w.append(["08-2026 Withdrawal", 8, 8, 46243, "CNY", "ChipPay", 108, 14, 0.14, 0, 100,
          25569, 46243, "finish"])                                         # serial 1970
x = src.create_sheet("Xero")
x.append(["Date", "CNY", "USDT", "USD"])
for day in range(1, 32):
    x.append([datetime.datetime(2026, 8, day), 7.0, 1.0, 1.0])
f = src.create_sheet("Handing Fee")
f.append(["Payment Channel", "Currency", "Deposit", "Withdrawal"])
f.append(["ChipPay", "CNY", 0.004, 0.007])
f.append(["USDT-ERC20", "USDT", 0.01, 0.02])
src.save(tmp / "kvb.xlsx")


def jalan(*args):
    r = subprocess.run([PY, *map(str, args)], cwd=HERE, capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    return r.stdout


out1 = jalan("siapkan_kvb.py", tmp / "kvb.xlsx", "-o", tmp / "t.xlsx")
assert "'ChipPay' / 'ChipPay ' / 'Chippay' / 'chippay ' -> 'ChipPay'" in out1, out1
assert "1970" in out1 and "baris 8" in out1 and "baris 9" in out1, out1
jalan("buat_template.py", tmp / "t.xlsx", "-o", tmp / "tpl0.xlsx")
out2 = jalan("isi_template.py", tmp / "t.xlsx", "-t", tmp / "tpl0.xlsx", "--bulan", "2026-08",
             "--tanpa-extra-fees", "-o", tmp / "tpl.xlsx")
assert "sheet W: 5 baris" in out2, out2
assert "2 baris bertanggal di luar 1990-2100" in out2, out2
tw = load_workbook(tmp / "tpl.xlsx", read_only=True)
assert "Verifikasi Tanggal" in tw.sheetnames
ver = list(tw["Verifikasi Tanggal"].iter_rows(min_row=3, values_only=True))
assert [r[1] for r in ver] == [8, 9] and all(r[0] == "W" for r in ver), ver
tw.close()

jalan("hitung_dw.py", tmp / "tpl.xlsx", "--period", "2026-08", "-o", tmp / "hasil.xlsx",
      "--tanpa-extra-fees", "--usd-shadow", "CNY:CHIPPAY,EP,UPAY,ZPAY",
      "--fee-currency-suffix", "--warn-high-pct", "--cny-rate-check")
hw = load_workbook(tmp / "hasil.xlsx", read_only=True, data_only=True)
det = hw["D&W Detail"].iter_rows(values_only=True)
h = list(next(det))
i_t, i_g, i_c, i_tx, i_fee = (h.index(k) for k in
                              ("Type", "Payment Gateway", "Currency", "Transaction", "Handling Fee"))
wd = [r for r in det if r[i_t] == "withdrawal"]
assert len(wd) == 5, wd
assert {r[i_g] for r in wd} == {"ChipPay", "USDT-ERC20"}, {r[i_g] for r in wd}
for r in wd:
    rate = 0.007 if r[i_c] == "CNY" else 0.02
    assert r[i_fee] is not None and abs(r[i_fee] - r[i_tx] * rate) < 1e-9, r
rpt = [r for r in hw["D&W Report"].iter_rows(values_only=True)]
gws = [r[1] for r in rpt if r[1] and str(r[1]).strip().upper() in ("CHIPPAY", "USDT-ERC20")]
assert gws.count("ChipPay") == 2 and gws.count("USDT-ERC20") == 1, gws    # 1 D + 1 W
assert not any(g != g.strip() or g == "Chippay" for g in gws)
hw.close()
print("KVB serial dates, 1970 report, gateway case/space merge, fee lookup: OK")
