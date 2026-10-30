#!/usr/bin/env python3
"""v1.32 KVB: CNY Rate | Balance (USDT) | Rate source on Payment Channel Balance.

Fixtures are synthetic and built in a temp folder. Run: python3 test_kvb_cny_balance_usdt.py
"""
import datetime
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import hitung_dw as h

ROOT = Path(__file__).resolve().parent
FLAGS = ["--tanpa-extra-fees", "--usd-shadow", "CNY:CHIPPAY,EP,UPAY,ZPAY",
         "--fee-currency-suffix", "--warn-high-pct", "--cny-rate-check"]
D1, D2, D3, D4 = (datetime.datetime(2026, 7, d) for d in (1, 2, 3, 4))


def _buku(path, rate_sheet="CNY RATE", rate_rows=None, headers=("Date", "TD", "Rate")):
    wb = Workbook()
    d = wb.active
    d.title = "D"
    d.append(["Source Name", "Currency", "Paid Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket"])
    d.append(["Deposit", "CNY", D1, 6900, "ChipPay", 1000, 6.9, "C1"])
    d.append(["Deposit", "CNY", D3, 6900, "ChipPay", 1000, 6.9, "C3"])
    d.append(["Deposit", "CNY", D2, 690, "EP", 100, 6.9, "E2"])
    d.append(["Deposit", "CNY", D2, 690, "OddPay", 100, 6.9, "O2"])
    d.append(["Deposit", "PKR", D2, 2780, "PA", 10, 278, "P2"])
    w = wb.create_sheet("W")
    w.append(["Source Name", "Currency", "Completed Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket", "Status"])
    w.append(["Withdrawal", "CNY", D4, 690, "ChipPay", 100, 6.9, "CW4", "finish"])
    x = wb.create_sheet("Xero")
    x.append(["Date", "CNY", "PKR", "USD"])
    for day in range(1, 32):
        x.append([datetime.datetime(2026, 7, day), 6.9, 278, 1])
    fee = wb.create_sheet("D&W FEE")
    fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
    for g in ("ChipPay", "EP", "OddPay"):
        fee.append(["CNY", g, 0, 0])
    fee.append(["PKR", "PA", 0, 0])
    if rate_sheet:
        cr = wb.create_sheet(rate_sheet)
        cr.append(list(headers))
        for row in (rate_rows if rate_rows is not None else _rates_default()):
            cr.append(list(row))
    wb.save(path)


def _rates_default():
    return [
        ("2026-07-02", "chip pay", 7.0),          # string date, spaced/lower TD
        (D4, "CHIPPAY", 8.0),
        (D1, " ep ", 0),                          # zero -> blank + flagged
        (D2, "Ep", "abc"),                        # text -> blank + flagged
        (D3, "EP", 6.5),
        (D3, "GhostPay", 7.1),                    # matches no gateway in the run
    ]


def _jalan(src, dst, *extra, opening=None):
    perintah = [sys.executable, str(ROOT / "hitung_dw.py"), str(src), "-o", str(dst),
                "--period", "2026-07", *FLAGS, *extra]
    if opening:
        perintah += ["--channel-opening", str(opening)]
    proc = subprocess.run(perintah, cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr[-3000:] + proc.stdout[-3000:]
    return proc.stdout


def _pcb(book):
    ws = book["Payment Channel Balance"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    out = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        item = dict(zip(header, r))
        if isinstance(item.get("Date"), datetime.datetime):
            item["Date"] = item["Date"].date()
        out.append(item)
    return header, out


def _teks(ws):
    return "\n".join(str(v) for row in ws.iter_rows(values_only=True) for v in row
                     if isinstance(v, str))


def _cari(rows, cur, gw, tgl):
    hasil = [r for r in rows if r["Currency"] == cur and h.norm(r["Payment Gateway"]) == gw
             and r["Date"] == tgl]
    assert hasil, (cur, gw, tgl)
    return hasil[0]


def test_lookup_rules_unit():
    wb = Workbook()
    ws = wb.active
    ws.title = "  cny   rate "
    ws.append([" date ", "td", "RATE "])
    for row in _rates_default():
        ws.append(list(row))
    ws.append([D3, "EP", 6.6])                   # conflicting duplicate -> first kept
    ws.append([D2, "ZPAY", 0.15])                # < 1 -> used, noted
    st = h.load_cny_rate(wb)
    assert st["dibaca"] and st["sheet"] == "  cny   rate "
    st["t0"] = datetime.date(2026, 7, 1)
    k = h.kunci_channel("CNY", "ChipPay")
    d = datetime.date
    assert h.kurs_cny_pcb(k, d(2026, 7, 2))[::2] == (7.0, "exact")
    assert h.kurs_cny_pcb(k, d(2026, 7, 3)) == (7.0, "CNY RATE 2026-07-02 (carried)", "carried")
    r, src, cara = h.kurs_cny_pcb(k, d(2026, 7, 1))
    assert (r, cara) == (7.0, "carried-back") and "no earlier rate" in src
    r, src, cara = h.kurs_cny_pcb(k, d(2026, 6, 30), pembuka=True)
    assert (r, cara) == (7.0, "opening") and "first rate in the period" in src, src
    ep = h.kunci_channel("CNY", "EP")
    assert h.kurs_cny_pcb(ep, d(2026, 7, 1))[::2] == (None, "bad")
    assert h.kurs_cny_pcb(ep, d(2026, 7, 2))[::2] == (None, "bad")
    assert h.kurs_cny_pcb(ep, d(2026, 7, 3))[0] == 6.5
    assert h.kurs_cny_pcb(h.kunci_channel("CNY", "OddPay"), d(2026, 7, 3))[2] == "none"
    assert [x["ignored"] for x in st["dobel"]] == [6.6]
    assert [x["td"] for x in st["kecil"]] == ["ZPAY"]
    assert sum(1 for m in st["masalah"] if "rate is 0" in m["why"]) == 1
    assert sum(1 for m in st["masalah"] if "not a number" in m["why"]) == 1
    h.reset_cny_rate()
    assert h.kurs_cny_pcb(k, d(2026, 7, 2))[2] == "nosheet"


def test_pcb_columns_on_real_run():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src, dst, opening = tmp / "kvb.xlsx", tmp / "out.xlsx", tmp / "open.txt"
        opening.write_text("ChipPay\tCNY\t1400\nChipPay\tUSD(CNY)\t200\n", encoding="utf-8")
        _buku(src, rate_sheet=" Cny Rate ", headers=(" DATE", "td ", " rate"))
        log = _jalan(src, dst, "--cny-balance-usdt", opening=opening)
        assert "Balance (USDT)" in log, log[-2000:]
        book = load_workbook(dst, data_only=True)
        header, rows = _pcb(book)
        i_bal = header.index("Balance")
        assert header[i_bal + 1:i_bal + 5] == ["CNY Rate", "Balance (USDT)", "Rate source",
                                               "Status"], header
        ws = book["Payment Channel Balance"]
        assert ws.cell(1, i_bal + 2).fill.fgColor.rgb.endswith("ACB9CA")
        assert ws.cell(1, i_bal + 3).fill.fgColor.rgb.endswith(h.RPT_FILL_HEAD2)
        d = datetime.date

        op = _cari(rows, "CNY", "CHIPPAY", d(2026, 6, 30))
        assert op["CNY Rate"] == 7.0 and abs(op["Balance (USDT)"] - 200) < 1e-9, op
        assert "Opening" in op["Rate source"], op
        r1 = _cari(rows, "CNY", "CHIPPAY", d(2026, 7, 1))
        assert r1["CNY Rate"] == 7.0 and "carried back" in r1["Rate source"], r1
        assert abs(r1["Balance (USDT)"] - r1["Balance"] / 7.0) < 1e-9
        r2 = _cari(rows, "CNY", "CHIPPAY", d(2026, 7, 2))
        assert r2["Rate source"] == "CNY RATE 2026-07-02", r2
        r3 = _cari(rows, "CNY", "CHIPPAY", d(2026, 7, 3))
        assert r3["CNY Rate"] == 7.0 and "(carried)" in r3["Rate source"], r3
        r4 = _cari(rows, "CNY", "CHIPPAY", d(2026, 7, 4))
        assert r4["CNY Rate"] == 8.0 and abs(r4["Balance (USDT)"] - r4["Balance"] / 8) < 1e-9
        r31 = _cari(rows, "CNY", "CHIPPAY", d(2026, 7, 31))
        assert r31["CNY Rate"] == 8.0 and "(carried)" in r31["Rate source"]

        e1 = _cari(rows, "CNY", "EP", d(2026, 7, 1))
        assert e1["CNY Rate"] is None and e1["Balance (USDT)"] is None
        assert "not a usable rate" in e1["Rate source"], e1
        e2 = _cari(rows, "CNY", "EP", d(2026, 7, 2))
        assert e2["CNY Rate"] is None and "'abc'" in e2["Rate source"], e2
        e3 = _cari(rows, "CNY", "EP", d(2026, 7, 3))
        assert e3["CNY Rate"] == 6.5

        o2 = _cari(rows, "CNY", "ODDPAY", d(2026, 7, 2))
        assert o2["CNY Rate"] is None and o2["Balance (USDT)"] is None
        assert o2["Rate source"] == "No CNY RATE rows for this gateway"

        for r in rows:
            if r["Currency"] != "CNY":
                assert (r["CNY Rate"], r["Balance (USDT)"], r["Rate source"]) == \
                    (None, None, None), r
        assert any(r["Currency"] == "USD" and h.norm(r["Payment Gateway"]) == "CHIPPAY"
                   for r in rows), "USD shadow block missing"

        hilang = _teks(book["Missing Data"])
        assert "8. CNY RATE SHEET" in hilang
        assert "CNY OddPay: no CNY RATE rows" in hilang, hilang[-3000:]
        assert "GhostPay" in hilang
        assert "rate is 0" in hilang and "rate is not a number" in hilang
        assert "carried back" in hilang
        legenda = _teks(book["Legend"])
        assert "CNY BALANCE IN USDT" in legenda
        assert " Cny Rate " in book.sheetnames   # the source sheet is kept as uploaded
        book.close()


def test_missing_sheet_keeps_running():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src, dst = tmp / "kvb.xlsx", tmp / "out.xlsx"
        _buku(src, rate_sheet="Sheet1")             # right headers, wrong sheet name
        log = _jalan(src, dst, "--cny-balance-usdt")
        assert "PERINGATAN" in log and "Sheet1" in log, log[-2000:]
        book = load_workbook(dst, data_only=True)
        header, rows = _pcb(book)
        assert "Balance (USDT)" in header
        cny = [r for r in rows if r["Currency"] == "CNY"]
        assert cny and all(r["CNY Rate"] is None and r["Balance (USDT)"] is None for r in cny)
        hilang = _teks(book["Missing Data"])
        assert "has no 'CNY RATE' sheet" in hilang
        assert "Sheet 'Sheet1' has the headers Date | TD | Rate" in hilang
        book.close()


def test_flag_off_keeps_old_layout():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src, dst = tmp / "kvb.xlsx", tmp / "out.xlsx"
        _buku(src)
        _jalan(src, dst)
        book = load_workbook(dst, data_only=True)
        header, _rows = _pcb(book)
        assert header == h.PCB_KOLOM, header
        assert "8. CNY RATE SHEET" not in _teks(book["Missing Data"])
        assert "CNY BALANCE IN USDT" not in _teks(book["Legend"])
        book.close()


def test_siapkan_and_isi_template_carry_cny_rate():
    import importlib
    sk = importlib.import_module("siapkan_kvb")
    it = importlib.import_module("isi_template")
    wb = Workbook()
    wb.active.title = "D"
    ws = wb.create_sheet("cny rate")
    ws.append(["Date", "TD", "Rate"])
    ws.append([D1, "ChipPay", 6.9])
    assert sk.cari_sheet_cny_rate(wb) == "cny rate"
    wb2 = Workbook()
    assert sk.cari_sheet_cny_rate(wb2) is None
    out = Workbook()
    it.salin_cny_rate(wb, out)
    assert "CNY RATE" in out.sheetnames
    rows = list(out["CNY RATE"].iter_rows(values_only=True))
    assert rows[0] == ("Date", "TD", "Rate") and rows[1][1:] == ("ChipPay", 6.9), rows
    kosong = Workbook()
    it.salin_cny_rate(wb2, kosong)
    assert "CNY RATE" not in kosong.sheetnames


if __name__ == "__main__":
    for nama, fn in sorted(globals().items()):
        if nama.startswith("test_") and callable(fn):
            fn()
            print("ok", nama)
