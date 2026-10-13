"""FX / rate file detection by header only: python3 test_deal_segregator_fx_detect.py"""
import csv
import datetime
import io
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from deal_segregator import _adalah_file_fx, proses

HEADERS = ["Deal", "Position", "Login", "Group", "Country", "Time", "Type", "Entry", "Symbol",
           "Volume", "Commission", "Fee", "Swap", "Profit", "Currency"]
USD_HEADERS = ["Commission (USD)", "Fee (USD)", "Swap (USD)", "Profit (USD)"]


def write_deals(path):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerow(["1", "P1", "100", r"real\DPMKT-15", "MY", "2026.07.01 01:02:03.004",
                    "buy", "out", "EURUSD", "1", "-2", "-1", "-4", "20", "JPY"])


def write_fx_xlsx(path, sheet, header):
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws.append(["nothing here"])
    ws2 = wb.create_sheet(sheet)
    ws2.append(header)
    ws2.append([datetime.date(2026, 7, 1), "JPY", 2])
    wb.save(path)
    wb.close()


def write_fx_csv(path):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow([" Date ", "CURRENCY", "Rate"])
        w.writerow(["2026-07-01", "JPY", "2"])


def check_converted(out):
    with zipfile.ZipFile(out) as z:
        wb = load_workbook(io.BytesIO(z.read("Deals - Daily.xlsx")))
        ws = wb["Daily"]
        headers = [c.value for c in ws[1]]
        assert headers[-4:] == USD_HEADERS, headers
        row = {headers[i]: c.value for i, c in enumerate(ws[2])}
        assert row["Commission (USD)"] == -1, row
        assert row["Fee (USD)"] == -0.5, row
        assert row["Swap (USD)"] == -2, row
        assert row["Profit (USD)"] == 10, row


def run(fx, tmp):
    deals, out = tmp / "deals.csv", tmp / f"{fx.stem}.zip"
    write_deals(deals)
    assert not _adalah_file_fx(deals), "Deals CSV must not be detected as FX"
    assert _adalah_file_fx(fx), f"{fx.name} must be detected as FX"
    proses([deals, fx], out)
    check_converted(out)


with TemporaryDirectory() as tmp:
    tmp = Path(tmp)

    # 1. Capital 'Rate' on 'Sheet1' (the reported CRM workbook).
    fx = tmp / "crm rate.xlsx"
    write_fx_xlsx(fx, "Sheet1", ["date", "Currency", "Rate"])
    run(fx, tmp)

    # 2. Existing 'Query result' format.
    fx = tmp / "Client Equity - FX.xlsx"
    write_fx_xlsx(fx, "Query result", ["date", "Currency", "rate"])
    run(fx, tmp)

    # 3. CSV FX file with odd case / spacing.
    fx = tmp / "rates.csv"
    write_fx_csv(fx)
    run(fx, tmp)

    # 4. FX only -> clean STOP.
    fx = tmp / "only fx.xlsx"
    write_fx_xlsx(fx, "Sheet1", ["date", "Currency", "Rate"])
    try:
        proses([fx], tmp / "fx_only.zip")
    except SystemExit as exc:
        msg = str(exc)
        assert msg.startswith("STOP: no Deals History file"), msg
        assert "only fx.xlsx" in msg, msg
    else:
        raise AssertionError("FX-only upload must STOP")

print("OK test_deal_segregator_fx_detect")
