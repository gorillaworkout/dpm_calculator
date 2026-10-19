#!/usr/bin/env python3
"""MT4 Raw Report opened and re-saved in Excel.

Excel pads every row to the same column count, so the title line becomes
`Raw Report for ... ,,,,,,` (or `;;;;;;` in a semicolon locale), the header and
the data rows get trailing empty columns, fields with a comma are quoted when
the file is saved with ',', and the footer/total lines are padded too.

The resaved file must give exactly the same result as the original: closed
deals, USD/USC Profit, Commission, and Volume, the Daily and Monthly sheets.

The real KVB reports are not in the repo. They are used when present
(override with MT4_REAL_RAW, MT4_FULL_RAW, MT4_ACCOUNT_TYPES); the committed
fixtures always run.

python3 test_deal_segregator_mt4_excel_resave.py
"""
import csv
import io
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import deal_segregator as ds
from deal_segregator import proses

ROOT = Path(__file__).resolve().parent
FIX = ROOT / "fixtures" / "kvb"
FIX_RAW = FIX / "KVB_MT4_RawReport_2026-10-04.csv"
FIX_TRAILING = FIX / "KVB_MT4_RawReport_2026-07-01_trailing_delim.csv"
FIX_CURRENCY = FIX / "KVB_MT4_Deals_2026-07-01_sample.csv"
FIX_TYPES = FIX / "KVB_MT4_Account_Types_sample.csv"

REAL_RAW = Path(os.environ.get(
    "MT4_REAL_RAW", "/workspace/kvb-mt4/KVB_MT4_RawReport_2026-10-04.csv"))
FULL_RAW = Path(os.environ.get(
    "MT4_FULL_RAW",
    "/home/box/sand-data/agents/24ca3b67-52c5-4d69-aa5a-b9f914126756/attachments/"
    "8763a29c84b737de10c89efbab54f9bf0788f65d8fe1f4c6dc7015076bbd477c.csv"))
REAL_TYPES = Path(os.environ.get(
    "MT4_ACCOUNT_TYPES", "/workspace/kvb-mt4/KVB_MT4_Account_Types.xlsx"))

TITLE = "Raw Report for 'abcc' from 2026.08.04 to 2026.08.04"
_DESIMAL = re.compile(r"^-?\d+\.\d+$")
_ANGKA = re.compile(r"^-?\d+(?:\.\d+)?$")


# ----------------------------------------------------------------- Excel resave
def _baca_sumber(src):
    raw = Path(src).read_bytes()
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252"), "cp1252"


def _pemisah_sumber(teks):
    baris = [b for b in teks.splitlines() if b.strip()]
    contoh = baris[1] if len(baris) > 1 else baris[0]
    return ";" if contoh.count(";") > contoh.count(",") else ","


def excel_resave_csv(src, dst, pemisah, lebar_extra=0, rapikan_angka=False):
    """Like Excel 'Save As CSV': every row padded to one column count, fields
    that contain the delimiter (or a quote) are quoted, CRLF line ends.
    rapikan_angka drops trailing zeros the way Excel shows numbers ("0.00" -> "0")."""
    teks, enc = _baca_sumber(src)
    rows = list(csv.reader(teks.splitlines(), delimiter=_pemisah_sumber(teks)))
    lebar = max(len(r) for r in rows) + lebar_extra
    with open(dst, "w", encoding=enc, newline="") as f:
        w = csv.writer(f, delimiter=pemisah, lineterminator="\r\n")
        for r in rows:
            if rapikan_angka:
                r = [c.rstrip("0").rstrip(".") if _DESIMAL.match(c) else c for c in r]
            w.writerow(r + [""] * (lebar - len(r)))
    return lebar


def excel_resave_xlsx(src, dst):
    """Like Excel 'Save As .xlsx': numbers become numeric cells."""
    teks, _ = _baca_sumber(src)
    rows = list(csv.reader(teks.splitlines(), delimiter=_pemisah_sumber(teks)))
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append([float(c) if _ANGKA.match(c) else (c if c != "" else None) for c in r])
    wb.save(dst)


# ----------------------------------------------------------------- run + compare
def _hasil(zip_path):
    out = {}
    with zipfile.ZipFile(zip_path) as z:
        for nama in z.namelist():
            wb = load_workbook(io.BytesIO(z.read(nama)), read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
                if ws.title == "Verifikasi":
                    # file names differ between the original and the resave
                    rows = [r for r in rows
                            if r and r[0] is not None and str(r[0]).startswith("MT4")]
                out[(nama, ws.title)] = rows
            wb.close()
    return out


def _verifikasi(hasil):
    rows = next(v for (n, t), v in hasil.items() if t == "Verifikasi" and "Daily" in n)
    return {str(r[0]): (r[1] if len(r) > 1 else None) for r in rows}


def _per_mata_uang(hasil):
    rows = next(v for (n, t), v in hasil.items() if t == "Daily")
    head = list(rows[0])
    i = {k: head.index(k) for k in ("Deals", "Volume", "Commission", "Profit", "Currency")}
    tot = {}
    for r in rows[1:]:
        t = tot.setdefault(r[i["Currency"]], {"Deals": 0, "Volume": 0.0,
                                              "Commission": 0.0, "Profit": 0.0})
        for k in ("Deals", "Volume", "Commission", "Profit"):
            t[k] += r[i[k]] or 0
    return {c: {k: round(v, 2) for k, v in t.items()} for c, t in tot.items()}


def _jalankan(path, types, d, label):
    out = Path(d) / f"{label}.zip"
    proses([path], out, allow_mt4=True, mt4_accounts=types)
    return _hasil(out)


def _bandingkan_resave(src, types, nama="MT4_RawReport_20260804.csv"):
    """Original vs comma resave vs semicolon resave. -> (original result, per currency)."""
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        asli = d / "asli" / nama
        asli.parent.mkdir()
        shutil.copy(src, asli)
        ref = _jalankan(asli, types, d, "asli")
        ref_v = _verifikasi(ref)
        ref_cur = _per_mata_uang(ref)
        assert ref_v["MT4 closed buy/sell kept"] > 0
        for label, pemisah, angka in (("koma", ",", True), ("titikkoma", ";", False)):
            p = d / label / nama
            p.parent.mkdir()
            excel_resave_csv(src, p, pemisah, rapikan_angka=angka)
            judul = p.read_text(encoding=_baca_sumber(src)[1]).splitlines()[0]
            assert judul.endswith(pemisah * 3), judul  # the padded title is really there
            info = ds._info_mt4(p)
            assert info and info["awal"] is not None, f"{label}: title not detected"
            assert info["pemisah"] == pemisah
            res = _jalankan(p, types, d, label)
            v = _verifikasi(res)
            for k in ("MT4 closed buy/sell kept", "MT4 Total Profit (closed buy/sell)",
                      "MT4 Total Commission (closed buy/sell)",
                      "MT4 Total Volume (closed buy/sell)",
                      "MT4 Total Profit raw (before USC ÷100)",
                      "MT4 footer/summary lines dropped",
                      "MT4 balance/credit rows dropped (same rule as MT5 non-trade Entry)"):
                assert v[k] == ref_v[k], (label, k, v[k], ref_v[k])
            assert _per_mata_uang(res) == ref_cur, (label, _per_mata_uang(res), ref_cur)
            assert v == ref_v, label
            assert res == ref, f"{label}: workbook differs from the original"
        return ref_v, ref_cur


# ----------------------------------------------------------------------- tests
def test_title_with_padding_is_recognised():
    for teks in (TITLE, TITLE + ",,,,,,", TITLE + ";;;;;;;", TITLE + "\t\t\t",
                 TITLE + " ,, ", '"' + TITLE + '",,,', "\ufeff" + TITLE + ";;",
                 '"Raw Report for \'a,b\' from 2026.08.04 to 2026.08.04",,,'):
        m = ds._cocok_judul_mt4(teks)
        assert m and m.group(1) == "2026.08.04", teks
    for teks in ("", "Deal;Login;Open Time", TITLE + ",x,,", "Raw Report for abcc",
                 "Raw Report for 'abcc' from 2026.08.04 to 2026.08.04 extra"):
        assert not ds._cocok_judul_mt4(teks), teks


def test_header_trailing_columns():
    h = ["Deal", "Login", "Comment"]
    assert ds._rapikan_header_mt4(h) == (h, None)
    assert ds._rapikan_header_mt4(h + [""]) == (h, None)
    # July layout: exactly two empty cells, currency under the second
    assert ds._rapikan_header_mt4(h + ["", ""]) == (h + ["", ""], 4)
    # Excel padding: more empty cells, no data sample -> no currency column
    assert ds._rapikan_header_mt4(h + [""] * 5) == (h, None)
    # with data: currency only where the rows really say USD/USC
    assert ds._rapikan_header_mt4(h + [""] * 5, [["1", "2", "c", "", "USD", "", ""]]) \
        == (h + ["", ""], 4)
    assert ds._rapikan_header_mt4(h + ["", ""], [["1", "2", "c;d", "e", "", ""]]) == (h, None)


def test_fixture_resave_matches_original():
    v, cur = _bandingkan_resave(FIX_RAW, FIX_TYPES)
    assert v["MT4 closed buy/sell kept"] > 0 and cur


def test_trailing_fixture_resave_matches_original():
    """';' inside Comment + trailing ';' + short footer, then padded by Excel."""
    v, _ = _bandingkan_resave(FIX_TRAILING, FIX_TYPES)
    assert v["MT4 closed buy/sell kept"] == 5
    assert round(v["MT4 Total Profit (closed buy/sell)"], 2) == 77.84


def test_header_first_currency_layout_padded():
    """Header-first July export with the currency column, padded wider by Excel:
    the currency cross-check column must still be found."""
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        asli = d / "a" / "deals.csv"
        asli.parent.mkdir()
        shutil.copy(FIX_CURRENCY, asli)
        info_asli = ds._info_mt4(asli)
        ref = _jalankan(asli, FIX_TYPES, d, "a")
        for label, pemisah in (("k", ","), ("s", ";")):
            p = d / label / "deals.csv"
            p.parent.mkdir()
            excel_resave_csv(FIX_CURRENCY, p, pemisah, lebar_extra=3)
            info = ds._info_mt4(p)
            assert info["currency_index"] == info_asli["currency_index"] is not None
            assert _jalankan(p, FIX_TYPES, d, label) == ref, label


def test_xlsx_resave_matches_original():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        asli = d / "a" / "MT4.csv"
        asli.parent.mkdir()
        shutil.copy(FIX_RAW, asli)
        ref = _jalankan(asli, FIX_TYPES, d, "a")
        x = d / "MT4.xlsx"
        excel_resave_xlsx(FIX_RAW, x)
        assert ds._info_mt4(x)["awal"] is not None
        assert _per_mata_uang(_jalankan(x, FIX_TYPES, d, "x")) == _per_mata_uang(ref)
        # A1 holding the whole padded title line as text (CSV pasted into Excel)
        wb = load_workbook(x)
        wb.active["A1"] = TITLE + ";;;;;;;;"
        x2 = d / "MT4_a1.xlsx"
        wb.save(x2)
        info = ds._info_mt4(x2)
        assert info and str(info["awal"]) == "2026-08-04"


def test_padded_blank_line_before_header():
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        p = d / "MT4.csv"
        excel_resave_csv(FIX_RAW, p, ",")
        lines = p.read_text(encoding="utf-8").split("\r\n")
        lines.insert(1, "," * (lines[0].count(",")))
        p.write_text("\r\n".join(lines), encoding="utf-8")
        asli = d / "a" / "MT4.csv"
        asli.parent.mkdir()
        shutil.copy(FIX_RAW, asli)
        assert _per_mata_uang(_jalankan(p, FIX_TYPES, d, "p")) == \
            _per_mata_uang(_jalankan(asli, FIX_TYPES, d, "a"))


def test_browser_sniff_recognises_padded_title():
    """templates/tool.html: the regex used before upload (no account-type file)."""
    html = (ROOT / "templates" / "tool.html").read_text(encoding="utf-8")
    m = re.search(r"var MT4_TITLE = /(.+)/i;", html)
    assert m, "MT4_TITLE regex not found in tool.html"
    rx = re.compile(m.group(1), re.IGNORECASE)
    for teks in (TITLE, TITLE + ",,,,", TITLE + ";;;;", TITLE + "\t\t", '"' + TITLE + '",,'):
        assert rx.search(teks), teks
    assert not rx.search("Deal,Login,Open Time")


def _real(path):
    if not path.exists() or not REAL_TYPES.exists():
        print(f"SKIP: {path} or {REAL_TYPES} not present")
        return False
    return True


def test_real_raw_report_resave():
    if not _real(REAL_RAW):
        return
    v, cur = _bandingkan_resave(REAL_RAW, REAL_TYPES)
    print("  real 2026-10-04:", v["MT4 closed buy/sell kept"], cur)


def test_full_raw_report_resave():
    if not _real(FULL_RAW):
        return
    v, cur = _bandingkan_resave(FULL_RAW, REAL_TYPES)
    assert v["MT4 closed buy/sell kept"] == 24325
    assert round(v["MT4 Total Profit (closed buy/sell)"], 2) == -286498.73
    print("  full 2026-07-01:", v["MT4 closed buy/sell kept"], cur)


if __name__ == "__main__":
    for nama, f in list(globals().items()):
        if nama.startswith("test_") and callable(f):
            f()
            print("OK", nama)
    print("OK")
