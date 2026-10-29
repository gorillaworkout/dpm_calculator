#!/usr/bin/env python3
"""End-to-end: a Deals History file uploaded in pieces must produce a complete
workbook. Run: python3 test_chunked_e2e.py

Self-contained: the ~12.5 MB Deals History CSV is synthetic (testkit.py), so the
test runs everywhere and forces several 5 MB pieces.
"""
import base64
import io
import time
import zipfile

from openpyxl import load_workbook

import testkit as T

TEST_ROOT = T.hermetic_env("tester", "tester-password-1")

from app import app  # noqa: E402

T.assert_isolated(TEST_ROOT)
LOGINS = tuple(range(100000, 101400))          # 1,400 accounts x 31 days x 3 deals
FIXTURE = T.write_deals_csv(TEST_ROOT / "31 Jul_Deals History synthetic.csv", logins=LOGINS)
EXPECTED = T.deals_expected(LOGINS)

app.config["TESTING"] = True
c = app.test_client()
c.environ_base["HTTP_AUTHORIZATION"] = "Basic " + base64.b64encode(
    b"tester:tester-password-1").decode()

CHUNK = 5 * 1024 * 1024        # small on purpose: forces many pieces
data = FIXTURE.read_bytes()
print(f"fixture {len(data)/1048576:.1f} MB -> {-(-len(data)//CHUNK)} chunks")

uid = c.post("/upload/begin", data={"bytes": str(len(data))}).get_json()["id"]
seq = 0
for off in range(0, len(data), CHUNK):
    r = c.post("/upload/chunk", data={
        "id": uid, "name": FIXTURE.name, "index": "0", "seq": str(seq),
        "chunk": (io.BytesIO(data[off:off + CHUNK]), "part"),
    }, content_type="multipart/form-data")
    assert r.status_code == 200, (seq, r.status_code, r.get_json())
    seq += 1

r = c.post("/tool/segregate", data={"upload_id": uid},
           content_type="multipart/form-data")
assert r.status_code == 303, (r.status_code, r.data[:500])
job_id = r.headers["Location"].rstrip("/").split("/")[-1]

deadline = time.time() + 900
while time.time() < deadline:
    s = c.get(f"/job/{job_id}")
    if s.status_code == 422:
        raise AssertionError(f"job failed: {s.data[:800]}")
    if b"/download" in s.data:
        break
    time.sleep(2)
else:
    raise AssertionError("job did not finish in 15 minutes")

d = c.get(f"/job/{job_id}/download")
assert d.status_code == 200, d.status_code
assert d.data[:2] == b"PK", "the result is not a ZIP/xlsx file"
assert seq >= 3, f"only {seq} chunks; the fixture must span several pieces"
# Nothing may be lost or duplicated across piece boundaries.
with zipfile.ZipFile(io.BytesIO(d.data)) as archive:
    wb = load_workbook(io.BytesIO(archive.read("Deals - Daily.xlsx")), read_only=True)
    try:
        rows = sum(1 for _ in wb["Daily"].iter_rows(values_only=True)) - 1
    finally:
        wb.close()
assert rows == EXPECTED["daily_rows"], (rows, EXPECTED)
print(f"OK chunked e2e — {len(d.data):,} bytes result from {seq} chunks")
