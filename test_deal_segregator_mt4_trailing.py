#!/usr/bin/env python3
"""MT4 Raw Report (July 2026): every data row ends with ';', the header does not,
some Comments contain ';', and the footer rows are shorter than the header.

python3 test_deal_segregator_mt4_trailing.py
"""
import io
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from deal_segregator import proses

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "fixtures" / "kvb" / "KVB_MT4_RawReport_2026-07-01_trailing_delim.csv"


def verifikasi(zip_path):
    with zipfile.ZipFile(zip_path) as z:
        nama = next(n for n in z.namelist() if n.endswith("Daily.xlsx"))
        wb = load_workbook(io.BytesIO(z.read(nama)), read_only=True, data_only=True)
        out = {}
        for row in wb["Verifikasi"].iter_rows(values_only=True):
            if row and row[0] is not None:
                out[str(row[0])] = row[1] if len(row) > 1 else None
        daily = list(wb["Daily"].iter_rows(values_only=True))
        return out, daily


def _peta_kosong():
    """Header-only account-type file: MT4 requires one; every login is then USD."""
    import tempfile as _tf
    from pathlib import Path as _P
    p = _P(_tf.mkdtemp()) / "account_types.csv"
    p.write_text("account,account_type\n", encoding="utf-8")
    return p


def run(path):
    out = path.with_suffix(".zip")
    proses([path], out, allow_mt4=True, mt4_accounts=_peta_kosong())
    return verifikasi(out)


def test_trailing_delimiter_parses():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "MT4_RawReport_20260701.csv"
        shutil.copy(SRC, p)
        v, daily = run(p)
        assert v["MT4 closed buy/sell kept"] == 5
        assert v["MT4 balance/credit rows dropped (same rule as MT5 non-trade Entry)"] == 1
        assert v["MT4 cancelled pending orders dropped"] == 1
        assert v["MT4 footer/summary lines dropped"] == 5
        # footer "Profit:" line in the excerpt
        assert round(v["MT4 Total Profit (closed buy/sell)"], 2) == 77.84
        assert round(v["MT4 balance/credit Profit excluded"], 2) == -213.00


def test_wide_row_with_bad_login_still_stops():
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "bad.csv"
        lines = SRC.read_text(encoding="utf-8").splitlines()
        lines.insert(3, "92000001;abc;2026.06.30 11:10:45;sell;xauusd;0.01;1;0;0;"
                        "2026.07.01 01:14:30;1;;;;;;0;0;0;0;1.00;1;x;y;")
        p.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
        try:
            run(p)
        except SystemExit as e:
            assert "more columns than the header" in str(e)
        else:
            raise AssertionError("expected STOP")


if __name__ == "__main__":
    test_trailing_delimiter_parses()
    test_wide_row_with_bad_login_still_stops()
    print("OK")
