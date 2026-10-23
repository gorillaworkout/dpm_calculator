#!/usr/bin/env python3
"""DPM Segregated: roll a month of daily segregated files into one workbook.

Usage:
    python3 dpm_segregated.py day01.csv day02.xlsx ... "Client Equity - FX.xlsx" -o out.xlsx

Input
  * One file per day (.csv, .xlsx or .xlsm), all in the same layout. Columns are
    found by header name, case/space/punctuation-insensitive, so column order and
    extra columns do not matter. Required: an account column (Login, Account,
    Account ID, ...). The day comes from a Date column in the file, otherwise
    from the file name (2026-09-01, 20260901, 01-09-2026, "1 Sep 2026", ...).
  * The month's Client Equity FX file (date, currency, rate). REQUIRED. It is
    detected by its header, the same reader Deal Segregator uses.

Output (one .xlsx)
  * Monthly Summary: one row per account (and currency). Every numeric column
    is summed over the month, EXCEPT Balance/Equity columns, which take the last
    day's value (an "opening"/"prev" Balance/Equity takes the first day's).
    USD columns: each daily amount at its own day's rate, summed; Balance and
    Equity at the last day's rate.
  * Daily: every daily row stacked, with a Date column and the USD columns.
  * Verifikasi: files, column mapping and every check that needs a look
    (missing FX rates, duplicate dates, accounts missing on the last day,
    Balance/Equity blank or 0 on the last day, ...).

Rows are streamed: memory grows with the number of accounts, not rows. Daily
rows are spooled to a temp file next to the output and written in a second pass.
"""
import argparse
import calendar
import datetime
import hashlib
import pickle
import re
import sys
import tempfile
from array import array
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import deal_segregator as ds
from angka import DOT_DECIMAL, parse_number, version_label
from upload_limits import MAX_FILES

JUDUL = "DPM Segregated"
EKSTENSI = (".csv", ".xlsx", ".xlsm")
HEADER_CARI = 10          # header row may sit below a title, within these rows
DAFTAR_MAKS = 20          # accounts named in one Verifikasi cell

# Header keys: upper case, only letters, digits and %.
AKUN = ("LOGIN", "ACCOUNT", "ACCOUNTID", "ACCOUNTNO", "ACCOUNTNUMBER", "LOGINID",
        "MT5LOGIN", "MT5ACCOUNT", "TRADINGACCOUNT", "CLIENTACCOUNT")
MATA_UANG = ("CURRENCY", "CCY", "ACCOUNTCURRENCY", "CURR")
TANGGAL = ("DATE", "REPORTDATE", "DATADATE", "TRADEDATE", "BUSINESSDATE", "DAY",
           "DATETIME", "TIMESTAMP", "TIME")
# Never parsed as numbers: identifiers and free text.
ID_KUNCI = {"DEAL", "ORDER", "POSITION", "ID", "TICKET", "NO", "NUMBER", "AGENT",
            "PHONE", "ZIP", "ZIPCODE", "POSTCODE"}
TEKS_KUNCI = {"NAME", "GROUP", "COMMENT", "COUNTRY", "CITY", "EMAIL", "TYPE", "SYMBOL",
              "ENTRY", "STATUS", "REASON", "DESK", "PLATFORM", "SERVER", "SHEET", "IB",
              "LEADSOURCE", "SOURCE", "STATE", "ADDRESS", "COLOR", "COLOUR"}
# Numeric but not additive: shown in Daily, left out of the monthly sums.
TIDAK_DIJUMLAH = ("PRICE", "RATE", "LEVERAGE", "LEVEL", "%", "PERCENT", "PCT", "RATIO")
# Numeric, additive, but not money: summed, never converted to USD.
BUKAN_UANG = ("VOLUME", "LOT", "DEALS", "TRADES", "COUNT", "ORDERS", "POSITIONS", "DAYS",
              "CLIENTS", "CUSTOMERS", "QTY", "QUANTITY")
SNAPSHOT = ("BALANCE", "EQUITY")
SNAPSHOT_AWAL = ("PREV", "OPENING", "OPEN", "START", "BEGIN", "BOD", "SOD")
LEMBAR_LEWAT = ("VERIFIKASI", "VERIFICATION", "GUIDE", "LEGEND", "README", "INSTRUCTION",
                "MONTHLY", "SUMMARY", "QUERYRESULT")
TOTAL = {"TOTAL", "GRANDTOTAL", "SUBTOTAL", "SUM", "JUMLAH"}
BULAN = {m.upper(): i for i, m in enumerate(calendar.month_abbr) if m}
BULAN.update({m.upper(): i for i, m in enumerate(calendar.month_name) if m})
BULAN["SEPT"] = 9

KUNING = ds.KUNING
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
_BATCH = 2000


def kunci(nama) -> str:
    return re.sub(r"[^A-Z0-9%]", "", ds._teks(nama).upper())


def _kata(k: str, daftar) -> bool:
    return any(w in k for w in daftar)


# ------------------------------------------------------------------ dates
_FMT_TANGGAL = ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S",
                "%Y.%m.%d", "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M:%S.%f", "%Y/%m/%d",
                "%Y/%m/%d %H:%M:%S", "%Y%m%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y",
                "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d %b %Y", "%d %B %Y",
                "%d-%b-%Y", "%b %d, %Y")


def tanggal_sel(v):
    """A date from a cell, or None when it is not one."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not v.is_integer():
            v = int(v) if 20000 < v < 80000 else v
        if isinstance(v, (int, float)) and float(v).is_integer():
            n = int(v)
            if 19000101 <= n <= 21001231:
                return tanggal_sel(str(n))
            if 20000 < n < 80000:      # Excel serial day
                return datetime.date(1899, 12, 30) + datetime.timedelta(days=n)
        return None
    teks = ds._teks(v)
    if not teks:
        return None
    for fmt in _FMT_TANGGAL:
        try:
            return datetime.datetime.strptime(teks, fmt).date()
        except ValueError:
            pass
    return None


def tanggal_nama_file(nama: str, tahun_cadangan=None):
    """Date in a file name, or None. Year-less '1 Sept' takes tahun_cadangan."""
    stem = Path(nama).stem
    m = re.search(r"(?<!\d)(20\d{2})[-_. ]?(\d{2})[-_. ]?(\d{2})(?!\d)", stem)
    if m:
        try:
            return datetime.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d{1,2})[-_. ](\d{1,2})[-_. ](20\d{2})(?!\d)", stem)
    if m:
        try:
            return datetime.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d{1,2})[-_. ]*([A-Za-z]{3,9})(?:[-_. ,]*(20\d{2}))?(?!\d)", stem)
    if m and m[2].upper() in BULAN:
        tahun = int(m[3]) if m[3] else tahun_cadangan
        if tahun:
            try:
                return datetime.date(tahun, BULAN[m[2].upper()], int(m[1]))
            except ValueError:
                pass
    return None


# ------------------------------------------------------------------ reading
def _cari_kolom(keys, daftar):
    for nama in daftar:
        if nama in keys:
            return keys.index(nama)
    return None


def _header_akun(row):
    keys = [kunci(v) for v in row]
    return _cari_kolom(keys, AKUN) is not None


def _baris_csv(path):
    import csv
    enc = ds._deteksi_encoding(path)
    with ds._buka_teks(path, enc) as f:
        awal = f.readline()
        pemisah = ds._deteksi_pemisah(awal)
        f.seek(0)
        reader = csv.reader(f, delimiter=pemisah)
        for nomor, row in enumerate(reader, start=1):
            yield nomor, [v if v != "" else None for v in row]


def _lembar(path):
    """Yield (sheet name, numbered rows) for every sheet that may hold data."""
    if path.suffix.lower() == ".csv":
        yield "", _baris_csv(path)
        return
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if _kata(kunci(ws.title), LEMBAR_LEWAT):
                continue
            yield ws.title, ((n, list(r)) for n, r in enumerate(ws.iter_rows(values_only=True), 1))
    finally:
        wb.close()


def buka_harian(path):
    """Yield (sheet, header, header row number, numbered data rows) per data sheet."""
    for nama, rows in _lembar(path):
        for nomor, row in rows:
            if nomor > HEADER_CARI:
                break
            if row and _header_akun(row):
                header = [ds._teks(v) for v in row]
                yield nama, header, nomor, rows
                break


def jenis_file(path):
    """'daily', 'fx' or None (no account header and not an FX file)."""
    try:
        for _ in buka_harian(path):
            return "daily"
    except Exception as exc:
        sys.exit(f"STOP: could not read '{path.name}': {type(exc).__name__}: {exc}")
    return "fx" if ds._adalah_file_fx(path) else None


def _sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for blok in iter(lambda: f.read(1 << 20), b""):
            h.update(blok)
    return h.hexdigest()


# ------------------------------------------------------------------ columns
class Kolom:
    """Union of every daily file's columns, keyed by normalised header."""

    def __init__(self):
        self.nama, self.keys, self.ok, self.bad, self.contoh_bad = [], [], [], [], []

    def posisi(self, header):
        hasil = []
        for h in header:
            k = kunci(h)
            if not k:
                hasil.append(None)
                continue
            if k not in self.keys:
                self.keys.append(k)
                self.nama.append(h)
                for daftar in (self.ok, self.bad):
                    daftar.append(0)
                self.contoh_bad.append(None)
            hasil.append(self.keys.index(k))
        return hasil

    def __len__(self):
        return len(self.keys)


def peran(k, i, kolom, peran_tetap):
    """What a column is: account/currency/date/text/id/skip/numeric/rate/snap_last/snap_first."""
    if i in peran_tetap:
        return peran_tetap[i]
    if k in ID_KUNCI or k.endswith("ID") or k in TEKS_KUNCI:
        return "text"
    if _cari_kolom([k], TANGGAL) is not None:
        return "text"
    if _kata(k, SNAPSHOT) and not kolom.bad[i]:
        # By name, even when every cell is blank: a blank Balance must be flagged.
        return "snap_first" if _kata(k, SNAPSHOT_AWAL) else "snap_last"
    if kolom.ok[i] == 0 or kolom.bad[i]:
        return "text"
    if _kata(k, TIDAK_DIJUMLAH):
        return "rate"
    return "numeric"


def uang(k, p):
    if p in ("snap_last", "snap_first"):
        return True
    return p == "numeric" and not _kata(k, BUKAN_UANG) and not k.endswith("USD")


def _boleh_angka(k):
    return not (k in ID_KUNCI or k.endswith("ID") or k in TEKS_KUNCI
                or _cari_kolom([k], TANGGAL) is not None)


# ------------------------------------------------------------------ main work
class Akun:
    __slots__ = ("jumlah", "usd", "hari", "awal", "akhir", "fx_kurang", "snap", "teks", "pecah")

    def __init__(self, n):
        self.jumlah = array("d", bytes(8 * n))
        self.usd = array("d", bytes(8 * n))
        self.hari = set()
        self.awal = self.akhir = None
        self.fx_kurang = False
        self.snap = {}     # col -> [date, value, usd divisor or None]
        self.teks = {}     # col -> [date, value]
        self.pecah = set()

    def lebar(self, n):
        while len(self.jumlah) < n:
            self.jumlah.append(0.0)
            self.usd.append(0.0)


def _angka(v):
    """float, None for empty, or the string itself when it is not a number."""
    if v is None or isinstance(v, bool):
        return None if v is None else v
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v
    teks = ds._teks(v)
    if teks in ("", "-", "--"):
        return None
    n = parse_number(teks, ambiguous=DOT_DECIMAL, default=False)
    return teks if n is False or n is None else n


@ds._jaga
def proses(paths, path_out):
    for p in paths:
        if p.suffix.lower() not in EKSTENSI:
            sys.exit(f"STOP: '{p.name}' is not a .csv, .xlsx or .xlsm file. "
                     "Upload the daily segregated files and the Client Equity FX file.")
    harian, fx_files, asing = [], [], []
    for p in paths:
        ds._ingat_berkas(p)
        j = jenis_file(p)
        (harian if j == "daily" else fx_files if j == "fx" else asing).append(p)
    if asing:
        sys.exit("STOP: these files have neither an account column (Login / Account) in the "
                 "first rows nor the FX columns date, currency, rate: "
                 + ", ".join(f"'{p.name}'" for p in asing))
    if not fx_files:
        sys.exit("STOP: the Client Equity FX file is required (columns date, currency, rate). "
                 "Upload the month's Client Equity FX together with the daily files.")
    if not harian:
        sys.exit("STOP: no daily segregated file found. Each daily file needs an account "
                 "column (Login / Account) in its header.")
    fx = {}
    for p in fx_files:
        fx.update(ds.baca_fx(p))
    tahun_fx = sorted({d.year for d, _ in fx}) or [None]

    catatan, kuning = {}, set()
    def flag(label, teks):
        catatan[label] = teks
        kuning.add(label)

    # ---- pass 0: identical files, header, file date
    kolom = Kolom()
    berkas, sidik = [], {}
    kembar = []
    for p in harian:
        h = _sha(p)
        if h in sidik:
            kembar.append((p.name, sidik[h]))
            continue
        sidik[h] = p.name
        ds._ingat_berkas(p)
        tgl_nama = tanggal_nama_file(p.name, tahun_fx[-1] if len(tahun_fx) == 1 else None)
        lembar, berkas_tgl = 0, None
        for sheet, header, nomor_h, rows in buka_harian(p):
            lembar += 1
            kolom.posisi(header)
            i_tgl = _cari_kolom([kunci(x) for x in header], TANGGAL)
            if i_tgl is None or berkas_tgl is not None:
                continue
            for nomor, row in rows:
                if any(v is not None and ds._teks(v) for v in row):
                    berkas_tgl = tanggal_sel(row[i_tgl] if i_tgl < len(row) else None)
                    break
        if not lembar:
            continue
        urut = berkas_tgl or tgl_nama
        if urut is None:
            sys.exit(f"STOP: cannot tell which day '{p.name}' is. Add a Date column, or put the "
                     "date in the file name (for example 2026-09-01 or 20260901).")
        berkas.append({"path": p, "tgl_nama": tgl_nama, "urut": urut, "lembar": lembar})
    if kembar:
        flag("Identical files skipped",
             "; ".join(f"'{a}' is byte-for-byte the same as '{b}'" for a, b in kembar))
    berkas.sort(key=lambda b: (b["urut"], b["path"].name))
    multi_lembar = any(b["lembar"] > 1 for b in berkas)

    # Fixed roles found by header name.
    def i_pertama(daftar):
        return _cari_kolom(kolom.keys, daftar)
    i_akun = i_pertama(AKUN)
    i_cur = i_pertama(MATA_UANG)
    peran_tetap = {}
    for i, k in enumerate(kolom.keys):
        if k in AKUN:
            peran_tetap[i] = "account" if i == i_akun else "text"
        elif k in MATA_UANG:
            peran_tetap[i] = "currency" if i == i_cur else "text"
        elif k.endswith("USD") and k[:-3] in kolom.keys:
            peran_tetap[i] = "skip"
    i_tgl_kolom = [i for i, k in enumerate(kolom.keys) if k == "DATE"]
    for i in i_tgl_kolom:
        peran_tetap[i] = "skip_date"
    if i_cur is None:
        flag("No currency column", "No Currency / CCY column in the daily files: every amount "
             "is treated as USD (rate 1).")

    # ---- pass 1: stream rows, aggregate, spool daily rows
    akun = {}
    tanggal_set = {}
    kurs_kurang = {}
    tanpa_akun = total_lewat = 0
    beda_tanggal = []
    baris_total = 0
    per_file = []
    path_out.parent.mkdir(parents=True, exist_ok=True)
    spool = tempfile.NamedTemporaryFile(dir=path_out.parent, prefix=".segregated-", delete=False)
    spool_path = Path(spool.name)
    try:
        batch = []
        for b in berkas:
            p = b["path"]
            ds._ingat_berkas(p)
            n_file = 0
            tgl_file = set()
            dari_kolom = False
            for sheet, header, nomor_h, rows in buka_harian(p):
                pos = kolom.posisi(header)
                keys = [kunci(x) for x in header]
                j_akun = next((j for j, q in enumerate(pos) if q == i_akun), None)
                j_cur = next((j for j, q in enumerate(pos) if q == i_cur), None) if i_cur is not None else None
                j_tgl = _cari_kolom(keys, TANGGAL)
                dari_kolom = dari_kolom or j_tgl is not None
                if j_akun is None:
                    continue
                boleh = [_boleh_angka(kolom.keys[q]) if q is not None else False for q in pos]
                for nomor, row in rows:
                    ds._ingat_baris(nomor)
                    if not any(v is not None and ds._teks(v) for v in row):
                        continue
                    raw_akun = row[j_akun] if j_akun < len(row) else None
                    akun_teks = ds._teks(raw_akun)
                    if isinstance(raw_akun, float) and raw_akun.is_integer():
                        akun_teks = str(int(raw_akun))
                    if not akun_teks:
                        tanpa_akun += 1
                        continue
                    if kunci(akun_teks) in TOTAL:
                        total_lewat += 1
                        continue
                    if ds._NOTASI.match(akun_teks):
                        sys.exit(f"STOP: account {akun_teks!r} in '{p.name}', row {nomor} is in "
                                 "scientific notation (Excel re-saved the file). Export it again.")
                    if j_tgl is not None:
                        v = row[j_tgl] if j_tgl < len(row) else None
                        tgl = tanggal_sel(v)
                        if tgl is None:
                            sys.exit(f"STOP: invalid date in '{p.name}', row {nomor}: "
                                     f"{ds._teks(v)!r} (column {header[j_tgl]!r})")
                    else:
                        tgl = b["tgl_nama"]
                    tgl_file.add(tgl)
                    cur = ds._norm(row[j_cur]) if j_cur is not None and j_cur < len(row) else ""
                    if i_cur is None:
                        cur = "USD"
                    if i_cur is None:
                        pembagi = 1.0
                    else:
                        pembagi = ds.pembagi_usd(fx, tgl, cur) if cur else None
                    if pembagi is None:
                        kk = (tgl, cur or "(blank)")
                        kurs_kurang[kk] = kurs_kurang.get(kk, 0) + 1
                    lembar_nama = sheet if multi_lembar else ""
                    key = (akun_teks, cur, lembar_nama)
                    a = akun.get(key)
                    if a is None:
                        a = akun[key] = Akun(len(kolom))
                    a.lebar(len(kolom))
                    a.hari.add(tgl)
                    a.awal = tgl if a.awal is None or tgl < a.awal else a.awal
                    a.akhir = tgl if a.akhir is None or tgl > a.akhir else a.akhir
                    if pembagi is None:
                        a.fx_kurang = True
                    nilai = [None] * len(kolom)
                    for j, q in enumerate(pos):
                        if q is None or j >= len(row):
                            continue
                        v = row[j]
                        if q in peran_tetap or not boleh[j]:
                            if v is not None and ds._teks(v) != "":
                                nilai[q] = v
                                if q not in peran_tetap or peran_tetap[q] == "text":
                                    lama = a.teks.get(q)
                                    tv = ds._teks(v)
                                    if lama is not None and lama[0] == tgl and lama[1] != tv:
                                        a.pecah.add(q)
                                    if lama is None or tgl >= lama[0]:
                                        a.teks[q] = [tgl, tv]
                            continue
                        x = _angka(v)
                        if x is None:
                            continue
                        if isinstance(x, float):
                            kolom.ok[q] += 1
                            a.jumlah[q] += x
                            if pembagi is not None:
                                a.usd[q] += x / pembagi
                            k = kolom.keys[q]
                            if _kata(k, SNAPSHOT):
                                lama = a.snap.get(q)
                                awal = _kata(k, SNAPSHOT_AWAL)
                                if (lama is None or (tgl < lama[0] if awal else tgl >= lama[0])):
                                    a.snap[q] = [tgl, x, pembagi]
                        else:
                            kolom.bad[q] += 1
                            if kolom.contoh_bad[q] is None:
                                kolom.contoh_bad[q] = f"'{p.name}' row {nomor}: {ds._teks(v)!r}"
                            lama = a.teks.get(q)
                            tv = ds._teks(v)
                            if lama is not None and lama[0] == tgl and lama[1] != tv:
                                a.pecah.add(q)
                            if lama is None or tgl >= lama[0]:
                                a.teks[q] = [tgl, tv]
                        nilai[q] = x
                    batch.append((tgl, akun_teks, cur, lembar_nama, pembagi, nilai))
                    n_file += 1
                    if len(batch) >= _BATCH:
                        pickle.dump(batch, spool, protocol=pickle.HIGHEST_PROTOCOL)
                        batch = []
            baris_total += n_file
            if b["tgl_nama"] and tgl_file and tgl_file != {b["tgl_nama"]}:
                beda_tanggal.append(f"'{p.name}': name says {b['tgl_nama']}, rows say "
                                    + ", ".join(str(t) for t in sorted(tgl_file)))
            for t in tgl_file:
                tanggal_set.setdefault(t, []).append(p.name)
            per_file.append((p.name, sorted(tgl_file), n_file,
                             "Date column" if dari_kolom else "file name"))
        if batch:
            pickle.dump(batch, spool, protocol=pickle.HIGHEST_PROTOCOL)
        spool.close()
        if not akun:
            sys.exit("STOP: the daily files have a header but no account rows.")
        ds._ingat_berkas(None)      # errors from here on are not about an input row

        ringkasan = _tulis(path_out, spool_path, kolom, peran_tetap, akun, tanggal_set,
                           kurs_kurang, catatan, kuning, flag, berkas, fx_files, per_file,
                           baris_total, tanpa_akun, total_lewat, beda_tanggal, multi_lembar, i_akun,
                           i_cur)
    finally:
        try:
            spool.close()
        except Exception:
            pass
        spool_path.unlink(missing_ok=True)
    return ringkasan


def _sel(ws, v, fill=None, fmt=None, bold=False):
    """A plain value when no style is needed (much faster in write-only mode)."""
    tidak_aman = isinstance(v, str) and v.startswith(("=", "+", "-", "@"))
    if not (tidak_aman or fill is not None or bold
            or (fmt and not isinstance(v, datetime.date))):
        return v
    c = WriteOnlyCell(ws, value=v)
    if tidak_aman:
        c.data_type = "s"
        c.quotePrefix = True
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = Font(bold=True)
    return c


def _header(ws, nama, komentar=None):
    komentar = komentar or {}
    sel = []
    for j, n in enumerate(nama, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, len(str(n)) + 2) + 4
        c = WriteOnlyCell(ws, value=n)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        if n in komentar:
            c.comment = Comment(komentar[n], JUDUL)
        sel.append(c)
    ws.freeze_panes = "A2"
    ws.append(sel)


def _bulat(x, uang_):
    if not isinstance(x, float):
        return x
    return round(x, 2) if uang_ else round(x, 6)


def _urut_akun(key):
    a = key[0]
    return (0, int(a), "") if a.isdigit() else (1, 0, a), key[1], key[2]


def _daftar(items):
    items = list(items)
    teks = ", ".join(items[:DAFTAR_MAKS])
    return teks + (f" ... and {len(items) - DAFTAR_MAKS} more (see the Check column)"
                   if len(items) > DAFTAR_MAKS else "")


def _tulis(path_out, spool_path, kolom, peran_tetap, akun, tanggal_set, kurs_kurang,
           catatan, kuning, flag, berkas, fx_files, per_file, baris_total, tanpa_akun,
           total_lewat, beda_tanggal, multi_lembar, i_akun, i_cur):
    n = len(kolom)
    peran_k = [peran(kolom.keys[i], i, kolom, peran_tetap) for i in range(n)]
    hari = sorted(tanggal_set)
    hari_akhir, hari_awal = hari[-1], hari[0]
    angka_i = [i for i in range(n) if peran_k[i] in ("numeric", "snap_last", "snap_first")]
    uang_i = [i for i in angka_i if uang(kolom.keys[i], peran_k[i])]
    snap_i = [i for i in angka_i if peran_k[i] in ("snap_last", "snap_first")]
    snap_akhir = [i for i in snap_i if peran_k[i] == "snap_last"]
    pecah = set()
    for a in akun.values():
        pecah |= a.pecah
    # Text kept in the summary: attributes with one value per account per day.
    # A numeric column spoiled by a few non-numbers is flagged, not shown as text.
    teks_i = [i for i in range(n) if peran_k[i] == "text" and i not in pecah
              and not (kolom.ok[i] and kolom.bad[i])]
    nama = kolom.nama

    # ---- last-day checks
    cek = {}
    tidak_hadir, kosong = [], {i: [] for i in snap_akhir}
    for key, a in akun.items():
        alasan = []
        if a.akhir != hari_akhir:
            alasan.append(f"not in the last day's file ({hari_akhir}); last seen {a.akhir}")
            tidak_hadir.append(key[0])
        for i in snap_akhir:
            s = a.snap.get(i)
            if a.akhir == hari_akhir and (s is None or s[0] != hari_akhir):
                alasan.append(f"{nama[i]} blank on {hari_akhir}")
                kosong[i].append(key[0])
            elif s is not None and s[0] == hari_akhir and s[1] == 0:
                alasan.append(f"{nama[i]} is 0 on {hari_akhir}")
                kosong[i].append(key[0])
        if a.fx_kurang:
            alasan.append("FX rate missing for some days: USD left blank")
        if alasan:
            cek[key] = "; ".join(alasan)

    # ---- workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Monthly Summary")
    kol1 = ["Account", "Currency"] + (["Sheet"] if multi_lembar else [])
    kol1 += [nama[i] for i in teks_i]
    kol1 += ["Days", "First date", "Last date"]
    kol1 += [nama[i] for i in angka_i]
    kol1 += [f"{nama[i]} (USD)" for i in uang_i]
    kol1 += ["Check"]
    komentar = {f"{nama[i]} (USD)": "Each day's amount at that day's Client Equity FX rate, summed. "
                "Yellow/blank: a rate was missing for this account's currency on some day."
                for i in uang_i}
    for i in snap_i:
        hari_snap = "last" if peran_k[i] == "snap_last" else "first"
        komentar[nama[i]] = (f"Not summed: the {hari_snap} day's value. Yellow: the account is "
                             "missing that day, or the value is blank/0.")
        komentar[f"{nama[i]} (USD)"] = f"The {hari_snap} day's value at that day's rate."
    _header(ws, kol1, komentar)
    for key in sorted(akun, key=_urut_akun):
        a = akun[key]
        baris = [_sel(ws, key[0]), _sel(ws, key[1])]
        if multi_lembar:
            baris.append(_sel(ws, key[2]))
        for i in teks_i:
            t = a.teks.get(i)
            baris.append(_sel(ws, t[1] if t else None))
        baris += [_sel(ws, len(a.hari)), _sel(ws, a.awal, fmt="yyyy-mm-dd"),
                  _sel(ws, a.akhir, fmt="yyyy-mm-dd")]
        for i in angka_i:
            k = kolom.keys[i]
            fill = None
            if peran_k[i] in ("snap_last", "snap_first"):
                s = a.snap.get(i)
                v = s[1] if s is not None and (peran_k[i] == "snap_first" or s[0] == hari_akhir) else None
                if peran_k[i] == "snap_last" and (v is None or v == 0):
                    fill = KUNING
            else:
                v = a.jumlah[i] if i < len(a.jumlah) else 0.0
            baris.append(_sel(ws, _bulat(v, uang(k, peran_k[i])), fill))
        for i in uang_i:
            fill = None
            if peran_k[i] in ("snap_last", "snap_first"):
                s = a.snap.get(i)
                ok = s is not None and (peran_k[i] == "snap_first" or s[0] == hari_akhir)
                if ok and s[2] is not None:
                    v = s[1] / s[2]
                else:
                    v, fill = None, KUNING
            elif a.fx_kurang:
                v, fill = None, KUNING
            else:
                v = a.usd[i]
            baris.append(_sel(ws, _bulat(v, True), fill))
        baris.append(_sel(ws, cek.get(key), KUNING if key in cek else None))
        ws.append(baris)

    # ---- Daily (second pass over the spool)
    kol2 = ["Date", "Account", "Currency"] + (["Sheet"] if multi_lembar else [])
    lain = [i for i in range(n) if peran_k[i] not in ("account", "currency", "skip", "skip_date")]
    kol2 += [nama[i] for i in lain]
    kol2 += [f"{nama[i]} (USD)" for i in uang_i]
    jumlah_lembar = max(1, -(-baris_total // ds.EXCEL_DATA_MAKS))
    if jumlah_lembar > 1:
        catatan["Excel row limit"] = (f"Daily has {baris_total:,} rows; each sheet holds at most "
                                      f"{ds.EXCEL_DATA_MAKS:,}, so it continues on Daily 2, ...")

    def lembar_harian(idx):
        w = wb.create_sheet(ds._nama_lembar("Daily", idx, jumlah_lembar))
        _header(w, kol2, {f"{nama[i]} (USD)": "Amount / that day's Client Equity FX rate. "
                          "Yellow: no rate for this date and currency." for i in uang_i})
        return w

    idx_l, isi = 0, 0
    w2 = lembar_harian(0)
    with open(spool_path, "rb") as f:
        while True:
            try:
                batch = pickle.load(f)
            except EOFError:
                break
            for tgl, ak, cur, lb, pembagi, nilai in batch:
                if isi >= ds.EXCEL_DATA_MAKS:
                    idx_l += 1
                    isi = 0
                    w2 = lembar_harian(idx_l)
                nilai = list(nilai) + [None] * (n - len(nilai))
                baris = [_sel(w2, tgl, fmt="yyyy-mm-dd"), _sel(w2, ak), _sel(w2, cur)]
                if multi_lembar:
                    baris.append(_sel(w2, lb))
                for i in lain:
                    baris.append(_sel(w2, _bulat(nilai[i], i in uang_i)))
                for i in uang_i:
                    x = nilai[i]
                    if pembagi is None:
                        baris.append(_sel(w2, None, KUNING if isinstance(x, float) else None))
                    else:
                        baris.append(_sel(w2, _bulat(x / pembagi, True) if isinstance(x, float) else None))
                w2.append(baris)
                isi += 1

    # ---- Verifikasi
    bulan = sorted({(d.year, d.month) for d in hari})
    r = {}
    r["Version"] = version_label()
    r["Daily files read"] = len(berkas)
    r["Client Equity FX file"] = ", ".join(p.name for p in fx_files)
    r["Period"] = f"{hari_awal} to {hari_akhir} ({len(hari)} days with data)"
    r["Last day (Balance / Equity taken from)"] = str(hari_akhir)
    r["Daily rows"] = baris_total
    r["Accounts (Monthly Summary rows)"] = len(akun)
    r["Account column"] = nama[i_akun]
    r["Currency column"] = nama[i_cur] if i_cur is not None else "(none: treated as USD)"
    r["Date taken from"] = ("Date column" if any(p[3] == "Date column" for p in per_file)
                            else "file name")
    r["Summed over the month"] = ", ".join(nama[i] for i in angka_i if peran_k[i] == "numeric") or "-"
    r["Last day's value (not summed)"] = ", ".join(nama[i] for i in snap_akhir) or "-"
    awal_i = [i for i in snap_i if peran_k[i] == "snap_first"]
    if awal_i:
        r["First day's value (not summed)"] = ", ".join(nama[i] for i in awal_i)
    r["Converted to USD"] = ", ".join(nama[i] for i in uang_i) or "-"
    tak_uang = [nama[i] for i in angka_i if i not in uang_i]
    if tak_uang:
        r["Summed, not converted (counts / volume / already USD)"] = ", ".join(tak_uang)
    rate_i = [nama[i] for i in range(n) if peran_k[i] == "rate"]
    if rate_i:
        r["Numeric, not summed (prices / rates / levels): Daily only"] = ", ".join(rate_i)
    if teks_i:
        r["Text kept (latest value)"] = ", ".join(nama[i] for i in teks_i)
    if pecah:
        r["Text that differs within one day: Daily only"] = ", ".join(nama[i] for i in sorted(pecah))
    skip = [nama[i] for i in range(n) if peran_k[i] == "skip"]
    if skip:
        r["USD columns in the files: recalculated from the base column"] = ", ".join(skip)
    for label, teks in catatan.items():
        r[label] = teks
    if kurs_kurang:
        flag("Missing FX rate (USD left blank)",
             "; ".join(f"{d} {c}: {m:,} rows" for (d, c), m in sorted(kurs_kurang.items(),
                                                                      key=lambda x: (str(x[0][0]), x[0][1]))))
        r["Missing FX rate (USD left blank)"] = catatan["Missing FX rate (USD left blank)"]
    dobel = {d: f for d, f in tanggal_set.items() if len(f) > 1}
    if dobel:
        teks = "; ".join(f"{d}: {', '.join(f)}" for d, f in sorted(dobel.items()))
        flag("Same date in more than one file (rows were added together)", teks)
        r["Same date in more than one file (rows were added together)"] = teks
    if len(bulan) > 1:
        teks = ", ".join(f"{calendar.month_abbr[m]} {y}" for y, m in bulan)
        flag("More than one month", f"The files cover {teks}; the summary adds them all.")
        r["More than one month"] = catatan["More than one month"]
    y, m = hari_akhir.year, hari_akhir.month
    semua = {datetime.date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)}
    hilang = sorted(semua - set(hari))
    if hilang and len(bulan) == 1:
        teks = ", ".join(str(d.day) for d in hilang)
        flag(f"Days of {calendar.month_abbr[m]} {y} with no file", teks)
        r[f"Days of {calendar.month_abbr[m]} {y} with no file"] = teks
    if tidak_hadir:
        teks = f"{len(tidak_hadir):,} accounts: " + _daftar(tidak_hadir)
        flag("Accounts missing on the last day (Balance / Equity blank)", teks)
        r["Accounts missing on the last day (Balance / Equity blank)"] = teks
    if not snap_akhir:
        flag("No Balance / Equity column", "No column named Balance or Equity was found; "
             "every numeric column was summed.")
        r["No Balance / Equity column"] = catatan["No Balance / Equity column"]
    for i, daftar in kosong.items():
        if daftar:
            label = f"{nama[i]} blank or 0 on the last day"
            teks = f"{len(daftar):,} accounts: " + _daftar(daftar)
            flag(label, teks)
            r[label] = teks
    for i in range(n):
        if kolom.bad[i] and kolom.ok[i] and _boleh_angka(kolom.keys[i]) and i not in peran_tetap:
            label = f"Column {nama[i]} not summed"
            teks = (f"{kolom.ok[i]:,} numbers but {kolom.bad[i]:,} values that are not numbers "
                    f"(first: {kolom.contoh_bad[i]}). Fix the file or ask for the column to be treated as text.")
            flag(label, teks)
            r[label] = teks
    if beda_tanggal:
        teks = "; ".join(beda_tanggal)
        flag("File name date differs from the Date column (Date column used)", teks)
        r["File name date differs from the Date column (Date column used)"] = teks
    if tanpa_akun:
        flag("Rows without an account (skipped)", f"{tanpa_akun:,}")
        r["Rows without an account (skipped)"] = f"{tanpa_akun:,}"
    if total_lewat:
        r["Total rows skipped"] = f"{total_lewat:,}"
    for nama_f, tgls, jml, sumber in per_file:
        r[f"File: {nama_f}"] = (f"{', '.join(str(t) for t in tgls) or '-'} ({sumber}); {jml:,} rows")
    if "Identical files skipped" in catatan:
        r["Identical files skipped"] = catatan["Identical files skipped"]

    wv = wb.create_sheet("Verifikasi")
    wv.column_dimensions["A"].width = 60
    wv.column_dimensions["B"].width = 90
    judul = WriteOnlyCell(wv, value=f"Summary - {JUDUL}")
    judul.font = Font(bold=True, size=13)
    wv.append([judul])
    wv.append([])
    for label, val in r.items():
        fill = KUNING if label in kuning else None
        wv.append([_sel(wv, label, fill), _sel(wv, val, fill)])
    wb.save(path_out)
    r["_flags"] = sorted(l for l in r if l in kuning)
    r["_periode"] = f"{hari_akhir.year:04d}-{hari_akhir.month:02d}"
    return r


def main():
    ap = argparse.ArgumentParser(description="DPM Segregated: monthly roll-up of daily segregated files")
    ap.add_argument("input", nargs="+", help="daily segregated files and the Client Equity FX file")
    ap.add_argument("-o", "--output", help="result .xlsx (default: Segregated.xlsx next to the first input)")
    ap.add_argument("--delete-input", action="store_true",
                    help="delete the input files after the result is written (the web app passes this)")
    args = ap.parse_args()
    paths = [Path(p) for p in args.input]
    if len(paths) > MAX_FILES:
        sys.exit(f"STOP: {len(paths)} files given; select no more than {MAX_FILES} files per run.")
    for p in paths:
        if not p.is_file():
            sys.exit(f"STOP: file not found: {p}")
    out = Path(args.output) if args.output else paths[0].with_name("Segregated.xlsx")
    r = proses(paths, out)
    if args.delete_input:
        for p in paths:
            p.unlink(missing_ok=True)
    print(f"=== {JUDUL} DONE ===")
    for k in ("Daily files read", "Period", "Daily rows", "Accounts (Monthly Summary rows)"):
        print(f"  {k}: {r[k]}")
    for k in r["_flags"]:
        print(f"  CHECK: {k}")
    print(f"PERIOD: {r['_periode']}")
    print(f"Output: {out.resolve()}")


if __name__ == "__main__":
    main()
