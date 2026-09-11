from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from deal_segregator import proses


SRC = Path("/Users/bayudarmawan/Documents/Dupoin/zern/31 Jul_Deals History 2026_08_07 12_07_27 (1).csv")

assert SRC.is_file()
with TemporaryDirectory() as tmp:
    out = Path(tmp) / "hasil.xlsx"
    summary = proses([SRC], out)
    wb = load_workbook(out, read_only=True, data_only=True)
    assert wb.sheetnames == ["Daily", "Monthly Summary", "Verifikasi"]
    assert wb["Daily"].max_row - 1 == 9418
    assert wb["Monthly Summary"].max_row - 1 == 9418
    assert summary["dipakai"] == 209838
    wb.close()

with TemporaryDirectory() as tmp:
    out = Path(tmp) / "dedupe.xlsx"
    summary = proses([SRC, SRC], out)
    assert summary["dipakai"] == 209838

print("OK deal segregator engine")
