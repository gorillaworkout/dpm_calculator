"""KVB v1.30: D&W Report columns on the Settlement Date basis (--settlement-rows in
isi_template.py, --settlement-columns in hitung_dw.py).

Covers: paid next month but settled in the month, paid in the month but settled the
month before, blank Settlement Date, a gateway that only exists on the Settlement
basis, the finish/success status filter, totals and averages, column positions /
headers / fills, Missing Data + Date Check, and that every other sheet is identical
to the same KVB run without the flag. The DPM snapshot test lives in
test_kvb_channel_usd.py (test_dpm_fixture_unchanged)."""
import datetime
import importlib.util
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
PY = sys.executable
DT = datetime.datetime


def muat(nama):
    spec = importlib.util.spec_from_file_location(nama + "_s", HERE / f"{nama}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


hd = muat("hitung_dw")
it = muat("isi_template")
tu = muat("tanggal_util")

# ------------------------------------------------------------------ unit checks
assert hd.SETTLEMENT_COLUMNS is False, "flag must default to off (DPM)"
r = tu.Ringkasan("W")
h_luar = tu.tentukan_tanggal({"COMPLETED DATE": DT(2026, 9, 1)}, "COMPLETED DATE")
r.catat(h_luar, "COMPLETED DATE", False, settle_only=True)
r.catat(h_luar, "COMPLETED DATE", False)
assert r.settle_only == Counter({"2026-09": 1}) and sum(r.luar_bulan.values()) == 1
kat = [x[1] for x in r.baris_ringkasan()]
assert "kept for Settlement Date columns only" in kat, kat

# isi_template helper: status + month
assert it.baris_settlement_saja([DT(2026, 8, 31), "finish", None], 0, 1, 2, (2026, 8))
assert it.baris_settlement_saja([DT(2026, 8, 31), "SUCCESS", None], 0, 1, 2, (2026, 8))
assert not it.baris_settlement_saja([DT(2026, 8, 31), "refuse", None], 0, 1, 2, (2026, 8))
assert not it.baris_settlement_saja([DT(2026, 7, 31), "finish", None], 0, 1, 2, (2026, 8))
assert not it.baris_settlement_saja([None, "finish", None], 0, 1, 2, (2026, 8))
assert it.baris_settlement_saja([DT(2026, 8, 31), None, "success"], 0, 1, 2, (2026, 8))
assert it.baris_settlement_saja([DT(2026, 8, 31)], 0, None, None, (2026, 8))   # D: no Status

# ------------------------------------------------------------------ fixture
tmp = Path(tempfile.mkdtemp())
src = Workbook()
src.remove(src.active)
t = src.create_sheet("Transfer")
t.append(["汇出"] + [None] * 7 + ["收款"])
t.append(["汇款日期", "汇出部门", "支付渠道", "汇出币种", "Xero Rate", "金额", "USD (Xero)", "手续费",
          "收款日期", "收款部门", "收款人", "收入币种", "金额", "手续费", "USD (Xero)"])
t.append([DT(2026, 8, 4), "KVB", "ChipPay", "CNY", None, 1000, None, None,
          DT(2026, 8, 4), "KVB", "J Wallet", "USDT", 140, None, None])
d = src.create_sheet("D")
d.append(["Source.Name", "user_id", "account", "Paid Date", "Settlement Date", "method", "ticket",
          "USD (CRM)", "rate", "transaction", "receive_company_account"])
d.append(["Dep", 1, 1, DT(2026, 8, 3), DT(2026, 8, 3), "CNY", "D1", 140, 0.14, 1000, "ChipPay"])
d.append(["Dep", 2, 2, DT(2026, 8, 5), DT(2026, 7, 31), "CNY", "D2", 70, 0.14, 500, "ChipPay"])
d.append(["Dep", 3, 3, DT(2026, 9, 1), DT(2026, 8, 31), "CNY", "D3", 28, 0.14, 200, "ChipPay"])
d.append(["Dep", 4, 4, DT(2026, 8, 10), None, "CNY", "D4", 14, 0.14, 100, "ChipPay"])
d.append(["Dep", 5, 5, DT(2026, 9, 2), DT(2026, 8, 30), "USDT", "D5", 50, 1, 50, "USDT-TRC20"])
d.append(["Dep", 6, 6, DT(2026, 9, 3), DT(2026, 9, 3), "CNY", "D6", 7, 0.14, 50, "ChipPay"])
w = src.create_sheet("W")
w.append(["Source.Name", "User ID", "Account", "Apply Date", "Payment Currency", "Payment Gateway",
          "Reference", "USD (CRM)", "rate", "charges", "transaction", "Completed Date",
          "Settlement Date", "status"])
w.append(["Wd", 1, 1, DT(2026, 8, 6), "CNY", "ChipPay", "W1", 140, 0.14, 0, 1000,
          DT(2026, 8, 6), DT(2026, 8, 6), "finish"])
w.append(["Wd", 2, 2, DT(2026, 8, 31), "CNY", "ChipPay", "W2", 14, 0.14, 0, 100,
          DT(2026, 9, 1), DT(2026, 8, 31), "finish"])        # completed Sep, settled Aug
w.append(["Wd", 3, 3, DT(2026, 8, 31), "CNY", "ChipPay", "W3", 7, 0.14, 0, 50,
          DT(2026, 9, 1), DT(2026, 8, 31), "refuse"])        # status filter
w.append(["Wd", 4, 4, DT(2026, 8, 7), "CNY", "ChipPay", "W4", 21, 0.14, 0, 150,
          DT(2026, 8, 7), DT(2026, 8, 7), "refuse"])         # status filter, both bases
w.append(["Wd", 5, 5, DT(2026, 8, 29), "USDT", "USDT-ERC20", "W5", 25, 1, 0, 25,
          DT(2026, 9, 3), DT(2026, 8, 29), "success"])
w.append(["Wd", 6, 6, DT(2026, 8, 8), "USDT", "USDT-ERC20", "W6", 10, 1, 0, 10,
          DT(2026, 8, 8), DT(2026, 8, 8), "finish"])
x = src.create_sheet("Xero")
x.append(["Date", "CNY", "USDT", "USD"])
for day in range(1, 32):
    x.append([DT(2026, 8, day), 7.0, 1.0, 1.0])
f = src.create_sheet("Handing Fee")
f.append(["Payment Channel", "Currency", "Deposit", "Withdrawal"])
f.append(["ChipPay", "CNY", 0.004, 0.007])
f.append(["USDT-ERC20", "USDT", 0.01, 0.02])
f.append(["USDT-TRC20", "USDT", 0.01, 0.02])
src.save(tmp / "kvb.xlsx")


def jalan(*args):
    p = subprocess.run([PY, *map(str, args)], cwd=HERE, capture_output=True, text=True)
    assert p.returncode == 0, p.stdout + p.stderr
    return p.stdout


KVB_HD = ["--tanpa-extra-fees", "--usd-shadow", "CNY:CHIPPAY,EP,UPAY,ZPAY",
          "--fee-currency-suffix", "--warn-high-pct", "--cny-rate-check"]
jalan("siapkan_kvb.py", tmp / "kvb.xlsx", "-o", tmp / "t.xlsx")
jalan("buat_template.py", tmp / "t.xlsx", "-o", tmp / "tpl0.xlsx")
out_on = jalan("isi_template.py", tmp / "t.xlsx", "-t", tmp / "tpl0.xlsx", "--bulan", "2026-08",
               "--tanpa-extra-fees", "--settlement-rows", "-o", tmp / "tpl_on.xlsx")
out_off = jalan("isi_template.py", tmp / "t.xlsx", "-t", tmp / "tpl0.xlsx", "--bulan", "2026-08",
                "--tanpa-extra-fees", "-o", tmp / "tpl_off.xlsx")
assert "sheet D: 5 baris" in out_on and "sheet W: 5 baris" in out_on, out_on   # +2 / +2
assert "sheet D: 3 baris" in out_off and "sheet W: 3 baris" in out_off, out_off
assert out_on.count("HANYA untuk kolom Settlement Date") == 2, out_on

log_on = jalan("hitung_dw.py", tmp / "tpl_on.xlsx", "--period", "2026-08", "-o", tmp / "on.xlsx",
               *KVB_HD, "--settlement-columns")
jalan("hitung_dw.py", tmp / "tpl_off.xlsx", "--period", "2026-08", "-o", tmp / "off.xlsx", *KVB_HD)
assert "Settle : 4 baris" in log_on, log_on
assert "PERINGATAN Settle" not in log_on, log_on

on = load_workbook(tmp / "on.xlsx", data_only=True)
off = load_workbook(tmp / "off.xlsx", data_only=True)

# ------------------------------------------------------------------ every other sheet unchanged
LAIN_LAYOUT = {"D&W Report", "Missing Data", "Legend", "D", "W", "Date Check"}
assert on.sheetnames == off.sheetnames, (on.sheetnames, off.sheetnames)
for nama in on.sheetnames:
    if nama in LAIN_LAYOUT:
        continue
    a = [list(r) for r in on[nama].iter_rows(values_only=True)]
    b = [list(r) for r in off[nama].iter_rows(values_only=True)]
    assert a == b, f"sheet {nama!r} changed with --settlement-columns"


# D / W: the in-month rows are identical; the 2 settle-only rows have blank calculated cells
def baris(ws):
    rows = [r for r in ws.iter_rows(values_only=True)]
    return rows[0], [r for r in rows[1:] if any(v not in (None, "") for v in r)]


for nama in ("D", "W"):
    h_on, r_on = baris(on[nama])
    h_off, r_off = baris(off[nama])
    assert h_on == h_off
    assert len(r_on) == len(r_off) + 2, (nama, len(r_on), len(r_off))
    sisa = Counter(r_on) - Counter(r_off)                 # copied in source order
    assert sum(sisa.values()) == 2 and not (Counter(r_off) - Counter(r_on)), nama
    i_xu = list(h_on).index("Xero USD")
    assert all(r[i_xu] is None for r in sisa), nama


# ------------------------------------------------------------------ D&W Report layout
def bagian(ws):
    """{jenis: (header_row_idx, header, {(cur, gw): row}, grand_total_row)}"""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hasil, jenis, cur = {}, None, None
    for i, r in enumerate(rows):
        if r[0] == "Currency" and r[1] == "Payment Gateway":
            jenis = "deposit" if "deposit" not in hasil else "withdrawal"
            hasil[jenis] = [i + 1, r, {}, None]
            cur = None
            continue
        if jenis and hasil[jenis][3] is None:
            if r[0] == "Grand Total":
                hasil[jenis][3] = r
                continue
            if r[0]:
                cur = r[0]
            if r[1]:
                hasil[jenis][2][(cur, r[1])] = r
    return hasil


rp_on, rp_off = bagian(on["D&W Report"]), bagian(off["D&W Report"])
hdr_d = rp_on["deposit"][1][:13]
hdr_w = rp_on["withdrawal"][1][:14]
assert hdr_d == hd.SETTLE_KOL_D, hdr_d
assert hdr_w == hd.SETTLE_KOL_W, hdr_w
assert hdr_d[2] == "CRM USD (Paid Date)" and hdr_d[3] == "CRM USD (Settlement Date)"
assert hdr_d[4:6] == ["TRANSACTIONS (Paid Date)", "TRANSACTIONS (Settlement Date)"]
assert hdr_d[8:10] == ["CRM Average Rate (Paid Date)", "CRM Average Rate (Settlement Date)"]
assert hdr_w[2] == "CRM USD (Completed Date)" and hdr_w[3] == "CRM USD (Settlement Date)"
assert hdr_w[5:7] == ["Transactions (Completed Date)", "Transactions (Settlement Date)"]
assert hdr_w[9:11] == ["CRM average (Completed Date)", "CRM average (Settlement Date)"]

ws_r = on["D&W Report"]
for jenis, n_col, n_data in (("deposit", 13, 7), ("withdrawal", 14, 8)):
    hr = rp_on[jenis][0]
    for c in range(1, n_col + 1):
        fill = ws_r.cell(hr, c).fill.fgColor.rgb[-6:]
        assert fill == ("ACB9CA" if c <= n_data else "A9D08E"), (jenis, c, fill)
    # settlement columns carry a header comment
    hdr = rp_on[jenis][1]
    c_s = hdr.index("CRM USD (Settlement Date)") + 1
    assert "Settlement Date" in ws_r.cell(hr, c_s).comment.text
    # data cells: settlement average (calculated) is green, CRM USD (Settlement) is not
    c_avg = (10 if jenis == "deposit" else 11)
    assert ws_r.cell(hr + 1, c_avg).fill.fgColor.rgb[-6:] == "E2EFDA"
    assert ws_r.cell(hr + 1, c_s).fill.fill_type in (None, "none") or \
        ws_r.cell(hr + 1, c_s).fill.fgColor.rgb[-6:] != "E2EFDA"
    assert ws_r.cell(hr + 1, c_avg).number_format == hd.FMT_ACC5

# ------------------------------------------------------------------ numbers
D = rp_on["deposit"][2]
W = rp_on["withdrawal"][2]
ap = lambda a, b: a is not None and abs(a - b) < 1e-9     # noqa: E731
cp = D[("CNY", "ChipPay")]
# old basis: D1 + D2 + D4 (paid in Aug); settlement: D1 + D3 (settled in Aug)
assert ap(cp[2], 224) and ap(cp[3], 168) and ap(cp[4], 1600) and ap(cp[5], 1200), cp
assert ap(cp[8], 1600 / 224) and ap(cp[9], 1200 / 168), cp
trc = D[("USDT", "USDT-TRC20")]                       # settlement-only gateway
assert trc[2] == 0 and ap(trc[3], 50) and trc[4] == 0 and ap(trc[5], 50), trc
assert trc[8] is None and ap(trc[9], 1.0) and trc[10] is None, trc
gt = rp_on["deposit"][3]
assert ap(gt[2], 224) and ap(gt[3], 218) and ap(gt[4], 1600) and ap(gt[5], 1250), gt
assert ap(gt[8], 1600 / 224) and ap(gt[9], 1250 / 218), gt

wcp = W[("CNY", "ChipPay")]
# old: W1 only (W4 refuse). settlement: W1 + W2 (W3 refuse excluded)
assert ap(wcp[2], 140) and ap(wcp[3], 154) and ap(wcp[5], 1000) and ap(wcp[6], 1100), wcp
assert ap(wcp[9], 1000 / 140) and ap(wcp[10], 1100 / 154), wcp
erc = W[("USDT", "USDT-ERC20")]
assert ap(erc[2], 10) and ap(erc[3], 35) and ap(erc[5], 10) and ap(erc[6], 35), erc
gtw = rp_on["withdrawal"][3]
assert ap(gtw[2], 150) and ap(gtw[3], 189) and ap(gtw[5], 1010) and ap(gtw[6], 1135), gtw
assert ap(gtw[10], 1135 / 189), gtw

# old-basis columns are identical to the run without the flag (only shifted)
MAP_D = {2: 2, 4: 3, 6: 4, 7: 5, 8: 6, 10: 7, 11: 8, 12: 9}      # on-col -> off-col (0-based)
MAP_W = {2: 2, 4: 3, 5: 4, 7: 5, 8: 6, 9: 7, 11: 8, 12: 9, 13: 10}
for jenis, peta in (("deposit", MAP_D), ("withdrawal", MAP_W)):
    hdr_off = rp_off[jenis][1]
    for a, b in peta.items():
        assert rp_on[jenis][1][a].replace(" (Paid Date)", "").replace(
            " (Completed Date)", "").strip().lower() == \
            hdr_off[b].strip().lower(), (jenis, rp_on[jenis][1][a], hdr_off[b])
    for k, r_off in rp_off[jenis][2].items():
        r_on = rp_on[jenis][2][k]
        for a, b in peta.items():
            assert r_on[a] == r_off[b], (jenis, k, a, r_on[a], r_off[b])
    for a, b in peta.items():
        assert rp_on[jenis][3][a] == rp_off[jenis][3][b], (jenis, a)
    # the only extra line is the settlement-only gateway
    assert set(rp_on[jenis][2]) - set(rp_off[jenis][2]) == (
        {("USDT", "USDT-TRC20")} if jenis == "deposit" else set())

# ------------------------------------------------------------------ Missing Data / Legend / Date Check
md = [str(c) for r in on["Missing Data"].iter_rows(values_only=True) for c in r if c is not None]
teks = "\n".join(md)
assert "7. SETTLEMENT DATE COLUMNS (D&W REPORT)" in teks
assert "Settlement Date inside: 4 rows" in teks, teks
for tiket in ("D3", "D5", "W2", "W5"):
    assert tiket in md, tiket
assert "W3" not in md                                   # refuse never listed
assert "only): 1 rows" in teks                          # D4 blank Settlement Date
assert "2026-07" in md                                  # D2 settled in July
assert "All currencies reconcile." in teks, teks
md_off = "\n".join(str(c) for r in off["Missing Data"].iter_rows(values_only=True)
                   for c in r if c is not None)
assert "SETTLEMENT DATE COLUMNS" not in md_off
lg = "\n".join(str(c) for r in on["Legend"].iter_rows(values_only=True) for c in r if c)
assert "D&W REPORT: SETTLEMENT DATE COLUMNS" in lg
assert "D&W REPORT: SETTLEMENT DATE COLUMNS" not in "\n".join(
    str(c) for r in off["Legend"].iter_rows(values_only=True) for c in r if c)
dc = [r for r in on["Date Check"].iter_rows(values_only=True)
      if r and r[1] == "kept for Settlement Date columns only"]
assert sorted((r[0], r[3]) for r in dc) == [("D", 2), ("W", 2)], dc
print("KVB settlement columns: layout, totals, settle-only rows, status, blank, "
      "other sheets unchanged: OK")
