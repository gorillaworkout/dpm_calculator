#!/usr/bin/env python3
"""KVB Payment Channel Balance: USD blocks, CNY openings, fee fallback, rate warning.

Run: python3 test_kvb_channel_usd.py
"""
import datetime
import itertools
import json
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

import hitung_dw as h

ROOT = Path(__file__).resolve().parent
FLAGS = [
    "--tanpa-extra-fees",
    "--usd-shadow", "CNY:CHIPPAY,EP,UPAY,ZPAY",
    "--fee-currency-suffix",
    "--warn-high-pct",
    "--cny-rate-check",
]


def _bal(teks):
    out, n = h.baca_tabel_saldo_channel(teks)
    return out, n


def test_cny_opening_and_usd_shadow_paste():
    fused, n = _bal("Beckpay\tVND\t100\nChipPay CNY\t1,000.50")
    assert n == 2, n
    assert fused[("CNY", h.kunci_channel("CNY", "ChipPay"))] == 1000.50
    assert ("CHIPPAY CNY", h.kunci_channel("CHIPPAY CNY", "Beckpay")) not in fused

    tiga, _n = _bal("ChipPay\tUSD(CNY)\t12,345.67\nChipPay\tUSD\t999")
    shadow = ("USD(CNY)", h.kunci_channel("CNY", "ChipPay"))
    nyata = ("USD", h.kunci_channel("USD", "ChipPay"))
    assert tiga[shadow] == 12345.67
    assert tiga[nyata] == 999
    assert shadow != nyata

    satu, _n = _bal("ChipPay USD(CNY) 12,345.67")
    assert satu[shadow] == 12345.67
    spasi, _n = _bal("ChipPay USD (CNY) 1,000")
    assert spasi[shadow] == 1000

    assert h.kunci_channel("CNY", "ChipPay-CNY") == h.kunci_gw("ChipPay")
    for kode in ("CNY", "IDR", "THB", "VND", "MYR", "PHP", "PKR", "INR", "USDT"):
        assert kode in h.KODE_CURRENCY_DIKENAL, kode
        assert kode in h.CURRENCY_SUFFIXES, kode


def test_lookup_fee_currency_suffix():
    row = {
        "deposit": {"pct": 0.02, "fixed": 0.0, "min": None},
        "withdrawal": {"pct": 0.03, "fixed": 0.0, "min": None},
        "fixed": 0.0,
        "nama": "PA-PKR",
    }
    table = {("PKR", h.kunci_gw("PA-PKR")): row}
    h.FEE_CURRENCY_SUFFIX = False
    komp, _fixed, _key, how = h.lookup_fee(table, "PKR", "PA", "withdrawal")
    assert komp is None and how == "NOT FOUND", (komp, how)

    h.FEE_CURRENCY_SUFFIX = True
    komp, _fixed, key, how = h.lookup_fee(table, "PKR", "PA", "withdrawal")
    assert how == "currency-suffix", how
    assert komp["pct"] == 0.03
    assert key[1] == "PA-PKR"

    table[("PKR", h.kunci_gw("PA"))] = {
        "deposit": {"pct": 0.01, "fixed": 0.0, "min": None},
        "withdrawal": {"pct": 0.05, "fixed": 0.0, "min": None},
        "fixed": 0.0,
        "nama": "PA",
    }
    komp, _fixed, key, how = h.lookup_fee(table, "PKR", "PA", "withdrawal")
    assert how == "exact" and komp["pct"] == 0.05 and key[1] == "PA"
    h.FEE_CURRENCY_SUFFIX = False

    assert h.baca_rate(8) == {"pct": 8.0, "fixed": 0.0, "min": None}
    assert h.baca_rate("+8") == {"pct": 0.0, "fixed": 8.0, "min": None}
    campur = h.baca_rate("1.5%+50")
    assert abs(campur["pct"] - 0.015) < 1e-12 and campur["fixed"] == 50
    minim = h.baca_rate("0.45%,min 90")
    assert minim["min"] == 90


def test_rate_warning_text_does_not_change_pct():
    table = {
        ("PHP", h.kunci_gw("SHUNFAPAY")): {
            "deposit": {"pct": 0.01, "fixed": 0.0, "min": None},
            "withdrawal": {"pct": 8.0, "fixed": 0.0, "min": None},
            "fixed": 0.0,
            "nama": "SHUNFAPAY",
            "dep_raw": 0.01,
            "wd_raw": 8,
        }
    }
    teks = h.kumpulkan_rate_tinggi(table)
    assert teks == [
        "PHP-SHUNFAPAY withdrawal rate 8 = 800%. "
        "If this is a flat fee per transaction, write +8."
    ]
    assert h.baca_rate(8)["pct"] == 8.0


def _rec_cny(tgl, gateway, tx, usd, ticket):
    return {
        "Currency": "CNY", "Payment Gateway": gateway, "Date": tgl,
        "Transaction": tx, "USD": usd, "Ticket": ticket,
        "Type": "deposit", "Handling Fee": 0,
    }


def _tf(**kw):
    dasar = {
        "tgl": datetime.date(2026, 7, 29),
        "channel": "ChipPay",
        "currency": "CNY",
        "usd": None,
        "_xero_usd_computed": False,
        "jumlah": 1000.0,
        "fee": 10.0,
        "penerima": None,
        "tujuan_cur": "",
        "tujuan_jumlah": None,
    }
    dasar.update(kw)
    return dasar


def test_cny_transfer_rate_sources():
    """Each USD fund-transfer rate: Xero, same-day, monthly, then the run median."""
    h.USD_SHADOW = h.parse_usd_shadow("CNY:CHIPPAY,EP,UPAY")
    jul29 = datetime.date(2026, 7, 29)
    jul31 = datetime.date(2026, 7, 31)
    jul20 = datetime.date(2026, 7, 20)
    # ChipPay 29 Jul: 1200/8000 and 200/2000. Same-day and monthly = 0.14.
    # Median of the two row ratios = 0.125. UPay and EP have no D/W rows.
    records = [
        _rec_cny(jul29, "ChipPay", 8000, 1200, "A"),
        _rec_cny(jul29, "ChipPay", 2000, 200, "B"),
    ]
    chip = (None, "USD(CNY)", h.kunci_channel("CNY", "ChipPay"))
    try:
        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": list(records)}
        xero = _tf(usd=145.0, _xero_usd_computed=True, penerima="EP",
                   tujuan_cur="CNY", tujuan_jumlah=900.0)
        h.tambah_usd_shadow(agg, report, [xero])
        out_k = (jul29,) + chip[1:]
        in_k = (jul29, "USD(CNY)", h.kunci_channel("CNY", "EP"))
        assert abs(agg[out_k]["ft"] - 145.0) < 1e-9, agg[out_k]
        assert abs(agg[out_k]["ft_fee"] - 1.45) < 1e-9, agg[out_k]
        assert abs(agg[in_k]["ft"] - -130.5) < 1e-9, agg[in_k]
        assert not agg[in_k]["ft_fee"]
        sumber = {(x["channel"], x["side"]): x["source"] for x in report["usd_shadow_ft"]}
        assert sumber == {("ChipPay", "out"): "xero", ("EP", "in"): "xero"}, sumber
        assert all(x["posted"] for x in report["usd_shadow_ft"])

        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": list(records)}
        h.tambah_usd_shadow(agg, report, [_tf(usd=999.0)])
        assert abs(agg[out_k]["ft"] - 140.0) < 1e-9, agg[out_k]
        assert abs(agg[out_k]["ft_fee"] - 1.4) < 1e-9, agg[out_k]
        assert report["usd_shadow_ft"][0]["source"] == "same-day"
        assert "via same-day weighted USD/Transaction" in h.baris_legend_ft(report["usd_shadow_ft"][0])

        # Each side uses its own gateway ratio when Xero USD was not computed.
        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": records + [_rec_cny(jul29, "EP", 500, 100, "E")]}
        beda = _tf(penerima="EP", tujuan_cur="CNY", tujuan_jumlah=800.0)
        h.tambah_usd_shadow(agg, report, [beda])
        assert abs(agg[out_k]["ft"] - 140.0) < 1e-9
        assert abs(agg[in_k]["ft"] - -160.0) < 1e-9, agg[in_k]["ft"]
        assert [x["source"] for x in report["usd_shadow_ft"]] == ["same-day", "same-day"]

        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": list(records)}
        h.tambah_usd_shadow(agg, report, [_tf(tgl=jul31, jumlah=500.0, fee=5.0)])
        k31 = (jul31,) + chip[1:]
        assert abs(agg[k31]["ft"] - 70.0) < 1e-9, agg[k31]
        assert abs(agg[k31]["ft_fee"] - 0.7) < 1e-9
        assert report["usd_shadow_ft"][0]["source"] == "monthly"
        assert "via monthly weighted USD/Transaction" in h.baris_legend_ft(report["usd_shadow_ft"][0])

        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": list(records)}
        h.tambah_usd_shadow(agg, report, [_tf(
            tgl=jul20, channel="UPay", jumlah=250.0, fee=2.5,
            penerima="EP", tujuan_cur="CNY", tujuan_jumlah=200.0)])
        k_up = (jul20, "USD(CNY)", h.kunci_channel("CNY", "UPay"))
        k_ep = (jul20, "USD(CNY)", h.kunci_channel("CNY", "EP"))
        assert abs(agg[k_up]["ft"] - 31.25) < 1e-9, agg[k_up]
        assert abs(agg[k_up]["ft_fee"] - 0.3125) < 1e-9
        assert abs(agg[k_ep]["ft"] - -25.0) < 1e-9, agg[k_ep]
        by = {(x["channel"], x["side"]): x for x in report["usd_shadow_ft"]}
        assert by[("UPay", "out")]["source"] == "median"
        assert by[("EP", "in")]["source"] == "median"
        assert abs(by[("UPay", "out")]["rate"] - 0.125) < 1e-12
        assert "via run median USD/Transaction" in h.baris_legend_ft(by[("UPay", "out")])

        agg = defaultdict(lambda: defaultdict(float))
        report = {"records": []}
        h.tambah_usd_shadow(agg, report, [_tf(
            tgl=datetime.date(2026, 7, 15), jumlah=400.0, fee=1.0, usd=50.0)])
        assert report["usd_shadow_ft"][0]["posted"] is False
        assert report["usd_shadow_ft"][0]["source"] is None
        assert not any(h.shadow_sumber(k[1]) for k in agg)
        assert "NOT posted" in h.baris_legend_ft(report["usd_shadow_ft"][0])
    finally:
        h.USD_SHADOW = {}


def _buku_kvb(path):
    wb = Workbook()
    d = wb.active
    d.title = "D"
    d.append(["Source Name", "Currency", "Paid Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket"])
    d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 30), 10000, "ChipPay",
              1450.56, 14.5056, "D30"])
    d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 30), 1000, "ChipPay",
              None, 14.5, "MISS"])
    d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 31), 20000, "ChipPay",
              2901.12, 0.145056, "D31"])
    d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 30), 1000, "EP",
              145.06, 14.506, "EP30"])
    d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 30), 1000, "OddPay",
              200, 14.5, "OUT1"])
    d.append(["Deposit", "PKR", datetime.datetime(2026, 7, 30), 5000, "PA",
              18, 278, "PAD"])
    d.append(["Deposit", "USD", datetime.datetime(2026, 7, 30), 50, "ChipPay",
              50, 1, "USD50"])
    w = wb.create_sheet("W")
    w.append(["Source Name", "Currency", "Completed Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket", "Status"])
    w.append(["Withdrawal", "CNY", datetime.datetime(2026, 7, 31), 5000, "ChipPay",
              725.28, 14.5056, "W31", "finish"])
    w.append(["Withdrawal", "PKR", datetime.datetime(2026, 7, 31), 800, "PA",
              3, 278, "PAW", "finish"])
    w.append(["Withdrawal", "PHP", datetime.datetime(2026, 7, 30), 100, "SHUNFAPAY",
              1.8, 55, "P1", "finish"])
    x = wb.create_sheet("Xero")
    x.append(["Date", "CNY", "PKR", "PHP", "USD"])
    x.append([datetime.datetime(2026, 7, 30), 6.9, 278, 55, 1])
    x.append([datetime.datetime(2026, 7, 31), 6.9, 278, 55, 1])
    fee = wb.create_sheet("D&W FEE")
    fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
    fee.append(["CNY", "ChipPay", 0.004, 0.007])
    fee.append(["CNY", "EP", 0.01, 0.01])
    fee.append(["CNY", "OddPay", 0.01, 0.01])
    fee.append(["PKR", "PA-PKR", 0.02, 0.03])
    fee.append(["PHP", "SHUNFAPAY", 0.01, 8])
    fee.append(["USD", "ChipPay", 0, 0])
    ft = wb.create_sheet("Fund Transfer Table")
    ft.append(["Date", "Department", "Payment Gateway", "Currency", "Amount", "Fee",
               "Xero (USD)", "收款日期", "收款部门", "收款人", "收入币种", "金额"])
    ft.append([datetime.datetime(2026, 7, 30), "TD", "ChipPay", "CNY", 1000, 10,
               None, None, None, None, None, None])
    wb.save(path)


def _baris(ws):
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        item = {}
        for nama, val in zip(header, row):
            if isinstance(val, datetime.datetime):
                val = val.date()
            item[nama] = val
        out.append(item)
    return out


def _teks(ws):
    bagian = []
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if isinstance(val, str) and val.strip():
                bagian.append(val)
    return "\n".join(bagian)


def test_payment_channel_balance_usd_block():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src = tmp / "kvb.xlsx"
        dst = tmp / "out.xlsx"
        opening = tmp / "open.txt"
        opening.write_text(
            "ChipPay\tCNY\t1000.50\n"
            "ChipPay\tUSD(CNY)\t12345.67\n"
            "ChipPay\tUSD\t999\n"
            "EP\tCNY\t10\n",
            encoding="utf-8",
        )
        _buku_kvb(src)
        proc = subprocess.run(
            [sys.executable, str(ROOT / "hitung_dw.py"), str(src), "-o", str(dst),
             "--period", "2026-07", "--channel-opening", str(opening), *FLAGS],
            cwd=ROOT, capture_output=True, text=True)
        assert proc.returncode == 0, proc.stderr[-2000:] + proc.stdout[-2000:]
        book = load_workbook(dst, data_only=True)

        wsheet = book["W"]
        header = [c.value for c in wsheet[1]]
        fee_col = header.index("Handling Fee") + 1
        ticket_col = header.index("Ticket") + 1
        fee_by_ticket = {}
        for row in range(2, wsheet.max_row + 1):
            fee_by_ticket[wsheet.cell(row, ticket_col).value] = wsheet.cell(row, fee_col).value
        assert abs(fee_by_ticket["P1"] - 800) < 1e-6, fee_by_ticket
        assert abs(fee_by_ticket["PAW"] - 24) < 1e-6, fee_by_ticket
        assert abs(fee_by_ticket["W31"] - 35) < 1e-6, fee_by_ticket

        pcb = _baris(book["Payment Channel Balance"])

        def gw(row):
            return h.norm(row["Payment Gateway"])

        cny_idx = [i for i, row in enumerate(pcb)
                   if row["Currency"] == "CNY" and gw(row) == "CHIPPAY"]
        assert cny_idx, "CNY ChipPay block missing"
        assert cny_idx == list(range(cny_idx[0], cny_idx[-1] + 1))
        nxt = pcb[cny_idx[-1] + 1]
        assert nxt["Currency"] == "USD" and gw(nxt) == "CHIPPAY", nxt
        assert nxt["Date"] == datetime.date(2026, 6, 30)
        assert abs(nxt["Balance"] - 12345.67) < 0.01

        shadow = []
        for row in pcb[cny_idx[-1] + 1:]:
            if gw(row) != "CHIPPAY" or row["Currency"] != "USD":
                break
            if (row["Date"] == datetime.date(2026, 6, 30)
                    and row["Balance"] is not None
                    and abs(row["Balance"] - 999) < 0.01
                    and shadow):
                break
            shadow.append(row)
        jul30 = next(row for row in shadow if row["Date"] == datetime.date(2026, 7, 30))
        jul31 = next(row for row in shadow if row["Date"] == datetime.date(2026, 7, 31))
        assert abs(jul30["Deposit"] - 1450.56) < 0.01, jul30
        assert abs(jul30["Deposit charges"] - 5.80224) < 0.001, jul30
        assert abs(jul30["Fund transfer"] - (1000 / 6.9)) < 0.001, jul30
        assert abs(jul30["Fund transfer charges"] - (10 / 6.9)) < 0.001, jul30
        assert abs(jul31["Deposit"] - 2901.12) < 0.01, jul31
        assert abs(jul31["Withdrawal"] - 725.28) < 0.01, jul31
        assert abs(jul31["Withdrawal charges"] - 5.07696) < 0.001, jul31
        expected_30 = 12345.67 + 1450.56 - 5.80224 - (1000 / 6.9) - (10 / 6.9)
        expected_31 = expected_30 + 2901.12 - (80 * 2901.12 / 20000) - 725.28 - (35 * 725.28 / 5000)
        assert abs(jul30["Balance"] - expected_30) < 0.01, (jul30["Balance"], expected_30)
        assert abs(jul31["Balance"] - expected_31) < 0.01, (jul31["Balance"], expected_31)

        after = pcb[cny_idx[-1] + 1 + len(shadow)]
        assert after["Currency"] == "USD" and gw(after) == "CHIPPAY", after
        assert abs(after["Balance"] - 999) < 0.01, after
        real_jul30 = next(row for row in pcb[cny_idx[-1] + 1 + len(shadow):]
                          if row["Date"] == datetime.date(2026, 7, 30)
                          and row["Currency"] == "USD" and gw(row) == "CHIPPAY")
        assert abs(real_jul30["Deposit"] - 50) < 0.01, real_jul30

        cny_jul30 = next(row for row in pcb
                         if row["Date"] == datetime.date(2026, 7, 30)
                         and row["Currency"] == "CNY" and gw(row) == "CHIPPAY")
        assert abs(cny_jul30["Deposit"] - 11000) < 0.01, cny_jul30
        assert abs(cny_jul30["Deposit charges"] - 44) < 0.01, cny_jul30
        assert abs(cny_jul30["Fund transfer"] - 1000) < 0.01
        assert abs(cny_jul30["Fund transfer charges"] - 10) < 0.01
        assert abs(cny_jul30["Balance"] - 10946.50) < 0.01, cny_jul30

        ep_cny = [i for i, row in enumerate(pcb) if row["Currency"] == "CNY" and gw(row) == "EP"]
        assert ep_cny == list(range(ep_cny[0], ep_cny[-1] + 1))
        ep_next = pcb[ep_cny[-1] + 1]
        assert ep_next["Currency"] == "USD" and gw(ep_next) == "EP", ep_next
        assert not any(row["Currency"] == "USD" and gw(row) == "ODDPAY" for row in pcb)

        hilang = _teks(book["Missing Data"])
        assert "PA (PKR) used fee row PA-PKR" not in hilang
        assert "2c. FEE ROW" not in hilang
        assert "PHP-SHUNFAPAY withdrawal rate 8 = 800%" in hilang
        assert "write +8" in hilang
        assert "OUT1" in hilang
        assert "MISS" in hilang
        assert "x100 (~14.5): 5" in hilang
        assert "like ~0.145: 1" in hilang
        assert "could not be converted" not in hilang
        legenda = _teks(book["Legend"])
        assert "ChipPay USD(CNY) 12,345.67" in legenda
        assert "via Xero USD" in legenda
        assert "2026-07-30 ChipPay CNY out" in legenda
        assert "PA (PKR) used fee row PA-PKR" in legenda
        assert "not a missing fee" in legenda
        assert "+8" in legenda and "1.5%+50" in legenda and "min" in legenda.lower()
        book.close()


def test_cny_transfer_sources_on_sheet():
    """No CNY Xero rate: Legend names same-day, monthly, and median conversions."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src, dst = tmp / "kvb.xlsx", tmp / "out.xlsx"
        wb = Workbook()
        d = wb.active
        d.title = "D"
        d.append(["Source Name", "Currency", "Paid Date", "Transaction", "Payment Gateway",
                  "USD", "Rate", "Ticket"])
        d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 29), 8000, "ChipPay",
                  1200, 0.15, "A"])
        d.append(["Deposit", "CNY", datetime.datetime(2026, 7, 29), 2000, "ChipPay",
                  200, 0.10, "B"])
        w = wb.create_sheet("W")
        w.append(["Source Name", "Currency", "Completed Date", "Transaction", "Payment Gateway",
                  "USD", "Rate", "Ticket", "Status"])
        x = wb.create_sheet("Xero")
        x.append(["Date", "USD"])
        x.append([datetime.datetime(2026, 7, 29), 1])
        fee = wb.create_sheet("D&W FEE")
        fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
        fee.append(["CNY", "ChipPay", 0, 0])
        fee.append(["CNY", "UPay", 0, 0])
        fee.append(["CNY", "EP", 0, 0])
        ft = wb.create_sheet("Fund Transfer Table")
        ft.append(["Date", "Department", "Payment Gateway", "Currency", "Amount", "Fee",
                   "Xero (USD)", "收款日期", "收款部门", "收款人", "收入币种", "金额"])
        ft.append([datetime.datetime(2026, 7, 29), "TD", "ChipPay", "CNY", 1000, 10,
                   None, None, None, None, None, None])
        ft.append([datetime.datetime(2026, 7, 31), "TD", "ChipPay", "CNY", 500, 5,
                   None, None, None, None, None, None])
        ft.append([datetime.datetime(2026, 7, 20), "TD", "UPay", "CNY", 250, 2.5,
                   None, None, None, "EP", "CNY", 200])
        wb.save(src)
        proc = subprocess.run(
            [sys.executable, str(ROOT / "hitung_dw.py"), str(src), "-o", str(dst),
             "--period", "2026-07", *FLAGS],
            cwd=ROOT, capture_output=True, text=True)
        assert proc.returncode == 0, proc.stderr[-2000:] + proc.stdout[-2000:]
        book = load_workbook(dst, data_only=True)
        pcb = _baris(book["Payment Channel Balance"])

        def ambil(tgl, cur, gw):
            return next(row for row in pcb
                        if row["Date"] == tgl and row["Currency"] == cur
                        and h.norm(row["Payment Gateway"]) == gw)

        jul29 = ambil(datetime.date(2026, 7, 29), "USD", "CHIPPAY")
        jul31 = ambil(datetime.date(2026, 7, 31), "USD", "CHIPPAY")
        upay = ambil(datetime.date(2026, 7, 20), "USD", "UPAY")
        ep = ambil(datetime.date(2026, 7, 20), "USD", "EP")
        cny = ambil(datetime.date(2026, 7, 29), "CNY", "CHIPPAY")
        assert abs(jul29["Fund transfer"] - 140) < 0.01, jul29
        assert abs(jul29["Fund transfer charges"] - 1.4) < 0.01, jul29
        assert abs(jul31["Fund transfer"] - 70) < 0.01, jul31
        assert abs(jul31["Fund transfer charges"] - 0.7) < 0.01, jul31
        assert abs(upay["Fund transfer"] - 31.25) < 0.01, upay
        assert abs(upay["Fund transfer charges"] - 0.3125) < 0.01, upay
        assert abs(ep["Fund transfer"] - -25) < 0.01, ep
        assert not ep["Fund transfer charges"], ep
        assert abs(cny["Fund transfer"] - 1000) < 0.01, cny
        assert abs(cny["Fund transfer charges"] - 10) < 0.01, cny
        legenda = _teks(book["Legend"])
        assert "via same-day weighted USD/Transaction" in legenda
        assert "via monthly weighted USD/Transaction" in legenda
        assert "via run median USD/Transaction" in legenda
        assert "NOT posted" not in legenda
        hilang = _teks(book["Missing Data"])
        assert "could not be converted" not in hilang
        book.close()


def test_dpm_fixture_unchanged():
    """A DPM-shaped workbook run without KVB flags matches the pre-change snapshot."""
    baseline_path = ROOT / "fixtures" / "dpm_unchanged_baseline.json"
    assert baseline_path.is_file(), "DPM baseline snapshot is missing"
    baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src, dst = tmp / "dpm.xlsx", tmp / "out.xlsx"
        opening = tmp / "open.txt"
        opening.write_text("77Pay\tVND\t1000\nPA\tPHP\t50\n", encoding="utf-8")
        wb = Workbook()
        d = wb.active
        d.title = "D"
        d.append(["Source Name", "Currency", "Paid Date", "Transaction", "Payment Gateway",
                  "USD", "Rate", "Ticket"])
        d.append(["Deposit", "VND", datetime.datetime(2026, 7, 30), 100000, "77Pay", 4.0, 25000, "D1"])
        d.append(["Deposit", "PHP", datetime.datetime(2026, 7, 30), 1000, "PA-PHP", 18, 55, "D2"])
        d.append(["Deposit", "PKR", datetime.datetime(2026, 7, 31), 5000, "PA", 18, 278, "D3"])
        w = wb.create_sheet("W")
        w.append(["Source Name", "Currency", "Completed Date", "Transaction", "Payment Gateway",
                  "USD", "Rate", "Ticket", "Status"])
        w.append(["Withdrawal", "VND", datetime.datetime(2026, 7, 31), 20000, "77Pay", 0.8, 25000, "W1", "finish"])
        w.append(["Withdrawal", "PHP", datetime.datetime(2026, 7, 30), 100, "SHUNFAPAY", 1.8, 55, "W2", "finish"])
        w.append(["Withdrawal", "PKR", datetime.datetime(2026, 7, 31), 800, "PA", 3, 278, "W3", "finish"])
        x = wb.create_sheet("Xero")
        x.append(["Date", "VND", "PHP", "PKR", "USD"])
        x.append([datetime.datetime(2026, 7, 30), 25000, 55, 278, 1])
        x.append([datetime.datetime(2026, 7, 31), 25000, 55, 278, 1])
        fee = wb.create_sheet("D&W FEE")
        fee.append([None, None])
        fee.append([None, None])
        fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
        fee.append(["VND", "77Pay", 0.025, 0.0])
        fee.append(["PHP", "PA", 0.01, 0.005])
        fee.append(["PHP", "SHUNFAPAY", 0.01, 8])
        fee.append(["PKR", "PA", 0.02, 0.01])
        fee.append(["PKR", "PA-PKR", 0.09, 0.08])
        ft = wb.create_sheet("Fund Transfer Table")
        ft.append(["Date", "Department", "Payment Gateway", "Currency", "Amount", "Fee",
                   "Xero (USD)", "收款日期", "收款部门", "收款人", "收入币种", "金额"])
        ft.append([datetime.datetime(2026, 7, 30), "TD", "77Pay", "VND", 5000, 10,
                   None, None, None, None, None, None])
        wb.save(src)
        proc = subprocess.run(
            [sys.executable, str(ROOT / "hitung_dw.py"), str(src), "-o", str(dst),
             "--period", "2026-07", "--channel-opening", str(opening)],
            cwd=ROOT, capture_output=True, text=True)
        assert proc.returncode == 0, proc.stderr[-2000:]
        book = load_workbook(dst, data_only=True)
        assert book.sheetnames == baseline["sheets"]
        for name in book.sheetnames:
            rows = []
            for row in book[name].iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    val = cell.value
                    if isinstance(val, str) and val.startswith("GENERATED "):
                        val = "GENERATED <time>"
                    if isinstance(val, datetime.datetime):
                        val = val.strftime("%Y-%m-%d %H:%M:%S")
                    elif isinstance(val, datetime.date):
                        val = val.isoformat()
                    elif isinstance(val, float):
                        val = round(val, 6)
                    rows.append([cell.coordinate, val])
            if rows != baseline["cells"][name]:
                for left, right in itertools.zip_longest(rows, baseline["cells"][name]):
                    if left != right:
                        raise AssertionError(f"{name}: {left} != {right}")
        book.close()


def test_plain_usd_opening_seeds_shadow_block():
    """Regression: 'ChipPay USD 28,598' must seed the USD(CNY) block, not an orphan."""
    lama = h.USD_SHADOW
    h.USD_SHADOW = h.parse_usd_shadow("CNY:CHIPPAY,EP,UPAY,ZPAY")
    try:
        jul1 = datetime.date(2026, 7, 1)
        agg = {(jul1, "USD(CNY)", h.kunci_channel("CNY", g)): {}
               for g in ("ChipPay", "EP", "UPay", "ZPay")}
        for teks in ("CHIPPAY USD 28,598", "chippay   usd   28,598",
                     "ChipPay\tUSD\t28,598", "USD ChipPay 28,598",
                     "USD\tChipPay\t28,598", "ChipPay USD(CNY) 28,598",
                     " ChipPay  usd (cny)  28,598 ", "Chip Pay USD 28,598"):
            saldo, n = _bal(teks + "\nEP USD 24,779\nUPAY USD 13,619\nZPay\tUSD\t10,897")
            assert n == 4, (teks, n, saldo)
            hasil = h.pindah_opening_usd_ke_shadow(saldo, agg)
            for g, v in (("CHIPPAY", 28598), ("EP", 24779), ("UPAY", 13619), ("ZPAY", 10897)):
                assert hasil.get(("USD(CNY)", g)) == v, (teks, g, hasil)
                assert ("USD", g) not in hasil, (teks, g, hasil)
        # Explicit USD(CNY) wins; a real USD channel with real USD movement is kept.
        saldo, _n = _bal("ChipPay\tUSD(CNY)\t5\nChipPay\tUSD\t999")
        hasil = h.pindah_opening_usd_ke_shadow(saldo, agg)
        assert hasil[("USD(CNY)", "CHIPPAY")] == 5 and hasil[("USD", "CHIPPAY")] == 999
        saldo, _n = _bal("EP USD 7")
        agg2 = dict(agg); agg2[(jul1, "USD", "EP")] = {}
        assert h.pindah_opening_usd_ke_shadow(saldo, agg2) == {("USD", "EP"): 7}
        # Non-shadow USD channels untouched.
        saldo, _n = _bal("BuziPay\tUSD\t16,277")
        assert h.pindah_opening_usd_ke_shadow(saldo, agg) == {("USD", "BUZIPAY"): 16277}
    finally:
        h.USD_SHADOW = lama


if __name__ == "__main__":
    test_plain_usd_opening_seeds_shadow_block()
    test_cny_opening_and_usd_shadow_paste()
    test_lookup_fee_currency_suffix()
    test_rate_warning_text_does_not_change_pct()
    test_cny_transfer_rate_sources()
    test_payment_channel_balance_usd_block()
    test_cny_transfer_sources_on_sheet()
    test_dpm_fixture_unchanged()
    print("OK KVB CNY USD channel, fee fallback, rate warning, rate check; DPM unchanged")
