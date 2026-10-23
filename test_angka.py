#!/usr/bin/env python3
"""Shared number parser (angka.py, PR #26) and every place that now uses it.

python3 test_angka.py
"""
import datetime
import os
import subprocess
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

import angka  # noqa: E402
from angka import DOT_DECIMAL, THOUSANDS, NumberError, parse_number  # noqa: E402

PY = sys.executable
LABEL = angka.version_label()


def near(a, b):
    return a is not None and abs(a - b) <= 1e-9 * max(1.0, abs(b))


def ok(teks, harapan, mode=DOT_DECIMAL):
    hasil = parse_number(teks, ambiguous=mode)
    assert near(hasil, harapan), (teks, mode, hasil, harapan)


def tolak(teks, mode=DOT_DECIMAL):
    try:
        hasil = parse_number(teks, ambiguous=mode)
    except NumberError as e:
        assert repr(teks) in str(e), (teks, str(e))
        return
    raise AssertionError(f"{teks!r} ({mode}) should be refused, got {hasil!r}")


BOTH = (DOT_DECIMAL, THOUSANDS)

# ---------------------------------------------------------------- formats
for m in BOTH:
    # scientific notation (the reported bug)
    ok("3.5E+09", 3.5e9, m); ok("3.5e9", 3.5e9, m); ok("3,5E+09", 3.5e9, m)
    ok("3,5e+09", 3.5e9, m); ok("-1.2E-03", -0.0012, m); ok("1E6", 1e6, m)
    ok("1.234e3", 1234.0, m); ok("2,5E3", 2500.0, m)
    # plain numbers and decimals
    ok("0", 0, m); ok("1234", 1234, m); ok("1234.5", 1234.5, m); ok("1234,5", 1234.5, m)
    ok(".5", 0.5, m); ok(",5", 0.5, m); ok("1234.", 1234, m); ok("377,23", 377.23, m)
    ok("0.500", 0.5, m); ok("0,500", 0.5, m); ok("1.2345", 1.2345, m); ok("0,0275", 0.0275, m)
    # Indonesian thousands
    ok("3.500.000.000", 3.5e9, m); ok("3.500.000,50", 3500000.5, m); ok("377.233,36", 377233.36, m)
    ok("1.000", 1000 if m == THOUSANDS else 1.0, m)
    # English thousands
    ok("3,500,000,000", 3.5e9, m); ok("3,500,000.50", 3500000.5, m); ok("1,234", 1234, m)
    ok("1,00,00,000.50", 10000000.5, m)            # Indian lakh grouping
    # negatives
    ok("-1234", -1234, m); ok("(1,234)", -1234, m); ok("(1.234,50)", -1234.5, m)
    ok("1234-", -1234, m); ok("\u22121.5", -1.5, m); ok("- 1,234.50", -1234.5, m)
    ok("+8", 8, m)
    # spaces, NBSP, currencies
    ok("  1234  ", 1234, m); ok("1\u00a0234\u00a0567,89", 1234567.89, m); ok("1 234 567.89", 1234567.89, m)
    ok("1\u202f234", 1234, m); ok("1'234'567.5", 1234567.5, m); ok("1_000_000", 1e6, m)
    ok("Rp 3.500.000", 3500000, m); ok("Rp3.500.000,50", 3500000.5, m); ok("Rp. 3.500.000", 3500000, m)
    ok("IDR 3,5E+09", 3.5e9, m); ok("3.500.000 IDR", 3500000, m); ok("USD 12.50", 12.5, m)
    ok("12.50 USD", 12.5, m); ok("12.50USDT", 12.5, m); ok("$1,234.56", 1234.56, m)
    ok("US$ 5", 5, m); ok("€ 5,5", 5.5, m); ok("-$1,234", -1234, m); ok("($1,234)", -1234, m)
    ok("Rp -1.000", -1000, m); ok("-Rp 1.000", -1000, m); ok("(Rp 1.234,50)", -1234.5, m)
    ok("\u00a0Rp\u00a03.500.000,50\u00a0", 3500000.5, m)

# ---------------------------------------------------------------- 3-digit ambiguity
ok("1.234", 1.234, DOT_DECIMAL); ok("1.234", 1234, THOUSANDS)
ok("377.233", 377.233, DOT_DECIMAL); ok("377.233", 377233, THOUSANDS)
ok("1,234", 1234, DOT_DECIMAL); ok("1,234", 1234, THOUSANDS)
ok("Rp 1.234", 1234, DOT_DECIMAL)          # rupiah marker -> thousands even in files
ok("IDR 1.234", 1234, DOT_DECIMAL)
ok("USD 1.234", 1.234, DOT_DECIMAL)
ok("1.234e2", 123.4, THOUSANDS)            # mantissa separator is always decimal
assert parse_number("1.234") == float("1.234")   # DOT_DECIMAL is the default

# ---------------------------------------------------------------- refused
for m in BOTH:
    for bad in ["abc", "1.2.3", "12,34.5", "1,234.5.6", "1.234,5,6", "1,2,3", "1234,567",
                "12.34,5", "5%", "inf", "-inf", "nan", "--5", "(-5)", "1.5k", "e5", "Rp",
                "1e", "1-2", "$", "-", "()", "1..2", "1,,2", "0x10", "1/2", "1 2.3.4",
                "12abc", "12 xyz", "abc 12", "k1.5"]:
        tolak(bad, m)
assert parse_number(None) is None and parse_number("") is None and parse_number("  \u00a0 ") is None
assert parse_number("abc", default=None) is None and parse_number("abc", default=0.0) == 0.0
try:
    parse_number(True)
    raise AssertionError("bool must be refused")
except NumberError:
    pass
try:
    parse_number(float("inf"))
    raise AssertionError("inf must be refused")
except NumberError:
    pass
assert issubclass(NumberError, ValueError)

# ---------------------------------------------------------------- numbers, Decimal
assert parse_number(5) == 5.0 and isinstance(parse_number(5), float)
assert parse_number(Decimal("1.25")) == 1.25
assert parse_number("1,234.505", as_decimal=True) == Decimal("1234.505")
assert parse_number("(1.234,50)", as_decimal=True) == Decimal("-1234.50")
assert parse_number(2.5, as_decimal=True) == Decimal("2.5")

# ---------------------------------------------------------------- old results unchanged
# Everything float() read before still gives the same float.
for teks in ["0", "1", "-1", "1.5", "-0.25", "1e5", "1E-5", "+3", ".5", "5.", "1234567.891",
             " 42 ", "3.5E+09", "100.000", "0.001", "12345678901234567890"]:
    assert parse_number(teks) == float(teks), teks
# old to_float: commas stripped
for teks in ["1,234", "1,234.5", "12,345,678.90", "-1,234.56"]:
    assert parse_number(teks) == float(teks.replace(",", "")), teks
# old baca_angka_saldo (typed balance)
for teks, v in [("377233.36", 377233.36), ("377.233,36", 377233.36), ("377,233.36", 377233.36),
                ("377,23", 377.23), ("377.233", 377233), ("377,233", 377233), ("-1.000", -1000),
                ("1 234 567,89", 1234567.89), ("3.500.000.000", 3.5e9), ("+5", 5)]:
    ok(teks, v, THOUSANDS)

assert LABEL.startswith("DW Calculator v"), LABEL
assert "v1.27 · 2026-10-23" in LABEL, LABEL

# ---------------------------------------------------------------- hitung_dw
os.environ.setdefault("PYTHONIOENCODING", "utf-8")
import hitung_dw as H  # noqa: E402

assert H.baca_angka_saldo("3.5E+09") == 3.5e9
assert H.baca_angka_saldo("3,5E+09") == 3.5e9
assert H.baca_angka_saldo("Rp 3.500.000,50") == 3500000.5
assert H.baca_angka_saldo("377.233") == 377233.0
assert H.baca_angka_saldo("") is None and H.baca_angka_saldo(None) is None
try:
    H.baca_angka_saldo("bukan angka")
    raise AssertionError("must exit")
except SystemExit as e:
    pesan = str(e.code)
    assert "--jwallet-opening bukan angka yang bisa dibaca: 'bukan angka'" in pesan, pesan
    assert LABEL in pesan, pesan

assert H.to_float("1,234.5") == 1234.5 and H.to_float("1.234,5") == 1234.5
assert H.to_float("3.5E+09") == 3.5e9 and H.to_float("(12.50)") == -12.5
assert H.to_float("1,5") == 1.5 and H.to_float("abc") is None and H.to_float("") is None
assert H.to_float(7) == 7.0 and H.to_float("Rp 1.000") == 1000.0
assert H.angka_sel("3.5E+09") == 3.5e9 and H.angka_sel(True) is None
assert H.angka_sel("Opening Balance") is None and H.angka_sel(None) is None

r = H.baca_rate("1,5%")
assert near(r["pct"], 0.015) and r["fixed"] == 0.0, r
r = H.baca_rate("1.5%+50")
assert near(r["pct"], 0.015) and r["fixed"] == 50.0, r
r = H.baca_rate("0.45%,min 90")
assert near(r["pct"], 0.0045) and r["min"] == 90.0, r
r = H.baca_rate("2,5% + 1.000,50")
assert near(r["pct"], 0.025) and r["fixed"] == 1000.5, r
assert H.baca_rate("8") == {"pct": 8.0, "fixed": 0.0, "min": None}
assert H._angka_persen_polos("8") == 8.0 and H._angka_persen_polos("8,5") == 8.5
assert H._angka_persen_polos("8%") is None and H._angka_persen_polos("+8") is None
assert H._angka_persen_polos("1.5%+50") is None and H._angka_persen_polos("min 5") is None

tabel, n = H.baca_tabel_saldo_channel(
    "PA\tVND\tVND 1.058.257,96\n"
    "\tTHB\t(4.976,10)\n"
    "\tZAR\t3,5E+03\n"
    "\tBRL\t-\n")
assert n == 4, (n, tabel)
vals = sorted(tabel.values())
assert near(vals[0], -4976.10) and vals[1] == 0.0 and near(vals[2], 3500.0) \
    and near(vals[3], 1058257.96), tabel
try:
    H.baca_tabel_saldo_channel("PA\tVND\t12abc34\n")
    raise AssertionError("must exit")
except SystemExit as e:
    assert "'12abc34'" in str(e.code) and LABEL in str(e.code), e.code

# Balance sheets with numbers stored as TEXT (Excel turns long numbers into '3.5E+09').
from openpyxl import Workbook  # noqa: E402

wb = Workbook()
ws = wb.active
ws.title = "Opening Balance"
ws.append(["As of date", "Currency", "Channel", "Opening Balance"])
ws.append([datetime.date(2026, 9, 30), "USDT", "J Wallet", "3.5E+09"])
ws.append([datetime.date(2026, 9, 30), "VND", "BeckPay", "1.234.567,89"])
ws.append([datetime.date(2026, 9, 30), "THB", "Other", 250.5])
ws.append([datetime.date(2026, 9, 30), "PHP", "Junk", "n/a"])
saldo, tgl, jw, tgl_jw = H.baca_opening_balance(wb)
assert jw == 3.5e9 and tgl_jw == datetime.date(2026, 9, 30), (jw, tgl_jw)
assert sorted(saldo.values()) == [250.5, 1234567.89], saldo
H.SALDO_AWAL_JW = None
assert H.saldo_awal_jwallet(wb, datetime.date(2026, 10, 1)) == (3.5e9, datetime.date(2026, 9, 30))

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "J Wallet"
ws2.append(["Date", "Currency", "转入", "转出", "余额"])
ws2.append([datetime.date(2026, 9, 29), "USDT", 10, None, "3,4E+09"])
ws2.append([datetime.date(2026, 9, 30), "USDT", 10, None, "3.5E+09"])
ws2.append([datetime.date(2026, 10, 1), "USDT", 10, None, 1])
assert H.saldo_awal_jwallet(wb2, datetime.date(2026, 10, 1)) == (3.5e9, datetime.date(2026, 9, 30))

# isi_template uses the same readers -> no more "J Wallet per None"
import isi_template as IT  # noqa: E402
hd = IT.muat_hitung_dw()
assert hd.saldo_awal_jwallet(wb, datetime.date(2026, 10, 1))[0] == 3.5e9

# CLI: --jwallet-opening 3.5E+09 is accepted (the input file check comes after)
for arg in (["--jwallet-opening", "3.5E+09"], ["--jwallet-opening=3,5E+09"],
            ["--jwallet-opening=-Rp 1.000"]):
    p = subprocess.run([PY, str(ROOT / "hitung_dw.py"), str(ROOT / "tidak-ada.xlsx"), *arg],
                       capture_output=True, text=True, cwd=ROOT,
                       env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    out = p.stdout + p.stderr
    assert "bukan angka" not in out, out
    assert ("3,500,000,000.00" in out) or ("-1,000.00" in out), out
p = subprocess.run([PY, str(ROOT / "hitung_dw.py"), "x.xlsx", "--jwallet-opening", "12abc"],
                   capture_output=True, text=True, cwd=ROOT,
                   env={**os.environ, "PYTHONIOENCODING": "utf-8"})
assert "--jwallet-opening bukan angka yang bisa dibaca: '12abc'" in p.stderr, p.stderr
assert LABEL in p.stderr, p.stderr

# ---------------------------------------------------------------- other tools
import gabung_fee as G  # noqa: E402
assert G.baca_rate("1,5%")[0] == 0.015
assert G.baca_rate("1.5%+50")[:2] == (0.015, 50.0)
assert G.baca_rate("8") == (8.0, 0.0, None, None)
assert G.baca_rate("1,000") == (1000.0, 0.0, None, None)
assert G.baca_rate("abc") == (None, 0.0, None, "abc")

import tambah_rate as TR  # noqa: E402
assert near(TR.pct("1.6%"), 0.016) and near(TR.pct("1,6%"), 0.016) and TR.pct("0") == 0.0
assert near(TR.pct("0.016"), 0.016) and near(TR.pct(" 2 % "), 0.02)
try:
    TR.pct("x%")
    raise AssertionError("must exit")
except SystemExit as e:
    assert "'x%'" in str(e.code) and LABEL in str(e.code), e.code

import pl_desk as P  # noqa: E402
assert P.money("(5,276,735.09)") == -5276735.09
assert P.money("-5,276,735.09") == -5276735.09
assert P.money("10.00") == 10.0 and P.money("1.005") == 1.01
assert P.money("5.276.735,09") == 5276735.09 and P.money("USD 1,234.5") == 1234.5
try:
    P.money("abc")
    raise AssertionError("must raise")
except ValueError:
    pass

import deal_segregator as D  # noqa: E402
assert D._angka("1.5", Path("x.csv"), 2, "Volume") == 1.5
assert D._angka(" -12.30 ", Path("x.csv"), 2, "Profit") == -12.3
assert D._angka("1,234.50", Path("x.csv"), 2, "Profit") == 1234.5
assert D._angka("3.5E+09", Path("x.csv"), 2, "Profit") == 3.5e9
assert D._angka(2.0, Path("x.csv"), 2, "Volume") == 2.0
try:
    D._angka("abc", Path("x.csv"), 7, "Volume")
    raise AssertionError("must exit")
except SystemExit as e:
    assert "STOP: invalid Volume value in 'x.csv', row 7: 'abc'" in str(e.code), e.code
    assert LABEL in str(e.code), e.code
try:
    D._angka("", Path("x.csv"), 7, "Volume")
    raise AssertionError("must exit")
except SystemExit:
    pass
assert D._angka_longgar("abc") == 0.0 and D._angka_longgar("1,5") == 1.5
assert D._angka_longgar("") == 0.0 and D._angka_longgar("inf") == 0.0

# ---------------------------------------------------------------- web form
TMP = Path(tempfile.mkdtemp())
os.environ["DW_AUTH_DB"] = str(TMP / "t.db")
os.environ["DW_JOB_DIR"] = str(TMP / "jobs")
os.environ["DW_ADMIN_USER"] = "boss"
os.environ["DW_ADMIN_PASS"] = "bosspassword1"
import app as webapp  # noqa: E402

import base64  # noqa: E402
import io  # noqa: E402
import time  # noqa: E402

webapp.app.config["TESTING"] = True
c = webapp.app.test_client()
c.environ_base["HTTP_AUTHORIZATION"] = "Basic " + base64.b64encode(b"boss:bosspassword1").decode()
buf = io.BytesIO()
wbx = Workbook()
wbx.active.append(["x"])
wbx.save(buf)
XLSX = buf.getvalue()


def post_saldo(nilai):
    return c.post("/tool/dw", data={"file": (io.BytesIO(XLSX), "kecil.xlsx"),
                                    "saldo_jw": nilai, "bulan": "6", "tahun": "2026"},
                  content_type="multipart/form-data")


for bad in ["bukan angka!!", "12abc", "1.2.3", "5%", "1" * 41]:
    r = post_saldo(bad)
    assert r.status_code == 400 and b"must be a number" in r.data, (bad, r.status_code)
    assert LABEL.encode() in r.data, r.data[-2000:]
    if len(bad) < 30:
        assert bad.encode() in r.data, (bad, r.data[-2000:])
from types import SimpleNamespace  # noqa: E402

perintah = []


def fake_run(command, **_kwargs):
    perintah.append(command)
    Path(command[command.index("-o") + 1]).write_bytes(b"PK fake xlsx")
    return SimpleNamespace(returncode=0, stdout="ok\n", stderr="")


real_run = webapp.subprocess.run
webapp.subprocess.run = fake_run
try:
    for good in ["3,5E+09", "3.5E+09", "Rp 3.500.000,50", "3,500,000.50", "(1,234)",
                 "-Rp 1.000", "377.233,36", "\u00a0377233.36"]:
        perintah.clear()
        r = post_saldo(good)
        assert r.status_code == 303, (good, r.status_code, r.data[-1500:])
        job_id = r.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
        for _ in range(400):
            info = webapp._read_state(job_id)
            if info and info["state"] in ("done", "failed"):
                break
            time.sleep(0.05)
        assert info["state"] == "done", info
        assert info["saldo_jw"] == good.strip(), info["saldo_jw"]
        hd_cmd = next(x for x in perintah if Path(x[1]).name == "hitung_dw.py")
        arg = f"--jwallet-opening={good.strip()}"
        assert arg in hd_cmd, hd_cmd
        # the real hitung_dw reads exactly that argument
        p = real_run([PY, str(ROOT / "hitung_dw.py"), str(ROOT / "tidak-ada.xlsx"), arg],
                           capture_output=True, text=True, cwd=ROOT,
                           env={**os.environ, "PYTHONIOENCODING": "utf-8"})
        assert "J Wallet mulai dari" in p.stdout and "bukan angka" not in p.stderr, \
            (good, p.stdout, p.stderr)
finally:
    webapp.subprocess.run = real_run

src = (Path(webapp.__file__).read_text(encoding="utf-8"))
assert "SALDO_POLA" not in src
assert 'f"{flag_saldo}={saldo_jw}"' in src

print("test_angka OK")
