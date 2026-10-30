#!/usr/bin/env python3
"""Bikin 'Template KVB Plus.xlsx' -- template upload KVB, gaya sama dengan
template DPM (buat_template.py): sheet Guide di depan, header berwarna per
sheet, komentar kolom wajib, format tanggal, header dibekukan.

Nama header SENGAJA nama asli export KVB Plus (Source.Name, method,
receive_company_account, ...). siapkan_kvb.py yang menerjemahkannya ke nama
DPM -- jadi jangan diganti di sini (dicek oleh test_kvb_template.py).

Sheet:
    Guide         penjelasan (English)
    Transfer      fund transfer: baris 1 汇出 / 收款, baris 2 header
    D             deposit
    W             withdrawal
    Xero          kurs harian (local currency per 1 USD)
    Handling Fee  tabel fee
    CNY RATE      v1.32: Date | TD | Rate (CNY per 1 USDT) per hari per gateway

Pakai:
    python3 buat_template_kvb.py                       -> Template KVB Plus.xlsx
    python3 buat_template_kvb.py -o /tmp/template.xlsx
"""

import argparse
import datetime
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl belum terpasang:  python3 -m pip install --user openpyxl")

NAMA_FILE = "Template KVB Plus.xlsx"

BIRU, HIJAU, UNGU, ABU = "4472C4", "2E7D32", "7030A0", "808080"
ORANYE, EMAS, EMAS_TUA, KVB = "C55A11", "BF8F00", "7F6000", "2EB5C4"
FMT_TGL, FMT_JAM = "yyyy-mm-dd", "hh:mm:ss"

# --- header asli export KVB Plus (urutan = urutan kolom di export) ----------
TRANSFER_R1 = {1: "汇出", 9: "收款"}
TRANSFER_R2 = ["汇款日期", "汇出部门", "支付渠道", "汇出币种", "Xero Rate", "金额",
               "USD (Xero)", "手续费", "收款日期", "收款部门", "收款人", "收入币种", "金额",
               "手续费", "USD (Xero)", "USD (Xero)", "收款银行或方式", "汇率", "备注"]
TRANSFER_EN = {
    "汇款日期": "remittance date (sending side)", "汇出部门": "sending department",
    "支付渠道": "payment channel the money leaves", "汇出币种": "sending currency",
    "金额": "amount", "手续费": "transfer fee", "收款日期": "receiving date",
    "收款部门": "receiving department",
    "收款人": "recipient channel - write 'J Wallet' when the money goes into the USDT wallet",
    "收入币种": "receiving currency", "USD (Xero)": "value in USD at the Xero rate",
    "Xero Rate": "Xero rate of the sending currency", "收款银行或方式": "receiving bank or method",
    "汇率": "exchange rate", "备注": "remark",
}

HDR_D = ["Source.Name", "sales_uid", "sales_name", "sales_dept", "sales_area", "user_id",
         "account", "Paid Date", "Paid Time", "Settlement Date", "method", "mt_order",
         "ticket", "USD (CRM)", "rate", "transaction", "receive_company_account",
         "referenace", "hash"]
HDR_W = ["Source.Name", "Sales UID", "Sales Name", "Sales Dept", "Sales Area", "User ID",
         "Account", "Apply Date", "Apply Time", "Payment Currency", "Payment Gateway",
         "Reference", "USD (CRM)", "rate", "charges", "transaction", "Trading Charges",
         "Pay Channel", "Completed Date", "Completed Time", "Settlement Date", "status",
         "Remark"]
# Kolom yang dibaca kalkulator -> artinya. Komentar "Required column".
WAJIB_D = {
    "Source.Name": "Source Name (Deposit / Rebate ...)",
    "Paid Date": "deposit date - report month, Payment Channel Balance and the Xero rate",
    "Settlement Date": "used for the Settlement Date columns of the D&W Report",
    "method": "becomes 'Currency' (CNY, VND, ...)",
    "USD (CRM)": "becomes 'USD'",
    "transaction": "becomes 'Transaction' - amount in the original currency",
    "receive_company_account": "becomes 'Payment Gateway' (ChipPay, EP, ...)",
}
WAJIB_W = {
    "Source.Name": "Source Name (Withdrawal / Wallet withdrawal ...)",
    "Apply Date": "withdrawal date on Payment Channel Balance",
    "Payment Currency": "becomes 'Currency'",
    "Payment Gateway": "payment gateway",
    "USD (CRM)": "becomes 'USD'",
    "transaction": "becomes 'Transaction' - amount in the original currency",
    "Completed Date": "withdrawal date for the D&W Report, Detail and the Xero rate",
    "Settlement Date": "used for the Settlement Date columns of the D&W Report",
    "status": "only finish / success rows are counted",
}
XERO_KOLOM = ["Date", "AED", "AUD", "BND", "BRL", "CNY", "EGP", "EUR", "GBP", "HKD", "IDR",
              "INR", "JPY", "KES", "KHR", "KRW", "LAK", "MMK", "MXN", "MYR", "NGN", "PHP",
              "PKR", "SGD", "THB", "TRY", "TWD", "VND", "ZAR", "USD", "USDT", "USDC"]
FEE_KOLOM = ["Payment Channel", "Currency", "Deposit", "Withdrawal", "Remark"]
FEE_PANDUAN = ("Rate guide: +8 = flat 8 per transaction (the plus sign is required). "
               "1.5%+50 = 1.5% of the transaction plus a flat 50. 0.45%,min 90 = 0.45%, "
               "but at least 90. A plain number 8 is read as 800%, not as a flat fee. "
               "Write +8 for a flat fee.")
CNY_KOLOM = ["Date", "TD", "Rate"]
CNY_EN = {
    "Date": "one row per day (2026-07-01). Real dates or text like 2026-07-01 are both read.",
    "TD": "CNY gateway, e.g. CHIPPAY, EP, UPAY, ZPAY. Case and spaces are ignored: "
          "ChipPay = CHIPPAY.",
    "Rate": "CNY per 1 USDT, e.g. 6.74. Balance (USDT) = CNY balance / this rate.",
}
CNY_CONTOH = [(datetime.date(2026, 7, 1), "CHIPPAY", 6.74),
              (datetime.date(2026, 7, 1), "EP", 6.82),
              (datetime.date(2026, 7, 1), "UPAY", 6.88),
              (datetime.date(2026, 7, 1), "ZPAY", 6.8),
              (datetime.date(2026, 7, 2), "CHIPPAY", 6.74)]

TIPIS = Side(style="thin", color="D9D9D9")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)


def _sel_header(sel, warna):
    sel.font = Font(bold=True, color="FFFFFF", size=10)
    sel.fill = PatternFill("solid", fgColor=warna)
    sel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    sel.border = KOTAK


def _komentar(teks, tinggi=80, lebar=260):
    return Comment(teks, "template", height=tinggi, width=lebar)


def tulis_header(ws, kolom, warna, baris=1, wajib=None, lebar=None):
    for c, t in enumerate(kolom, start=1):
        sel = ws.cell(baris, c, t)
        _sel_header(sel, warna)
        if wajib and t in wajib:
            sel.comment = _komentar(f"Required column - do not rename or delete.\n{wajib[t]}",
                                    90, 260)
        w = (lebar or {}).get(t)
        if w is None:
            w = 20 if len(t) > 14 else 13 if "Date" in t or "DATE" in t.upper() else 11
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[baris].height = 30


def format_baris_data(ws, kolom, baris):
    """Format baris pertama di bawah header, seperti template DPM."""
    for c, t in enumerate(kolom, start=1):
        u = t.upper()
        if "DATE" in u or t in ("汇款日期", "收款日期"):
            ws.cell(baris, c).number_format = FMT_TGL
        elif "TIME" in u:
            ws.cell(baris, c).number_format = FMT_JAM


def sheet_guide(ws):
    ws.title = "Guide"
    ws.sheet_properties.tabColor = ABU
    for kol, w in zip("ABC", (3, 26, 96)):
        ws.column_dimensions[kol].width = w
    r = 2
    ws.cell(r, 2, "KVB Plus D&W - Upload Template").font = Font(bold=True, size=18)
    r += 1
    ws.cell(r, 2, "Paste the KVB Plus exports into these sheets, upload this one file in "
                  "KVB Tools > Generate D&W, pick the month. Nothing else to set up."
            ).font = Font(italic=True, color=ABU, size=10)
    baris = [
        ("", ""),
        ("HOW TO USE", ""),
        ("1. Sheet 'Transfer'", "Fund transfers between channels, from row 3. Row 1 (汇出 = "
                                "money out, 收款 = money in) and row 2 are headers - keep them. "
                                "Write 'J Wallet' in 收款人 when the money goes into the USDT "
                                "wallet."),
        ("2. Sheet 'D'", "Paste the DEPOSIT export from row 2. Keep the header row as it is - "
                         "the KVB column names (method, receive_company_account, ...) are "
                         "translated by the calculator."),
        ("3. Sheet 'W'", "Paste the WITHDRAWAL export from row 2. Keep the header row."),
        ("4. Sheet 'Xero'", "One row per date, one column per currency, local currency per "
                            "1 USD (CNY 7.18 = 1 USD)."),
        ("5. Sheet 'Handling Fee'", "One row per Payment Channel + Currency with the deposit "
                                    "and withdrawal rate. See the rate guide in cell F1."),
        ("6. Sheet 'CNY RATE'", "NEW. One row per day per CNY gateway: Date | TD | Rate, "
                                "Rate = CNY per 1 USDT. Fills 'CNY Rate' and "
                                "'Balance (USDT)' on Payment Channel Balance."),
        ("7. Upload", "KVB Tools > Generate D&W: choose this file and the report month."),
        ("8. Result", "A new finished workbook is downloaded. Your file is never modified."),
        ("", ""),
        ("SHEET 'CNY RATE'", ""),
        ("Columns", "A Date, B TD, C Rate - from row 2. Only columns A to C are read."),
        ("Example (grey)", "Columns E to G show an example. They are NOT read; delete them or "
                           "leave them."),
        ("Missing days", "A day with no row uses the nearest earlier day of the same gateway, "
                         "marked '(carried)'. A gateway with no rows is left blank. Missing "
                         "Data section 8 lists every carried or blank day."),
        ("Without the sheet", "The report still runs; the USDT columns stay blank and Missing "
                              "Data says so."),
        ("", ""),
        ("FEE RATES", ""),
        ("+8", "Flat 8 per transaction (the plus sign is required)."),
        ("1.5%+50", "1.5% of the transaction plus a flat 50."),
        ("0.45%,min 90", "0.45% of the transaction, but at least 90."),
        ("8", "A plain number is a percentage: 8 = 800%. Write +8 for a flat fee."),
        ("", ""),
        ("WHAT YOU GET", ""),
        ("Payment Channel Balance", "Per date, currency and channel: deposits, charges, fund "
                                    "transfers, withdrawals and a running balance. CNY "
                                    "gateways also get a USD block and the Balance (USDT) "
                                    "columns."),
        ("D&W Report", "Currency x Payment Gateway summary, on the Paid / Completed Date and "
                       "the Settlement Date basis."),
        ("MTOATD, J Wallet, D&W Detail", "Daily reconciliation, USDT wallet ledger, one row "
                                         "per transaction."),
        ("Missing Data / Legend", "What could not be calculated, and what every column and "
                                  "colour means."),
    ]
    for judul, isi in baris:
        r += 1
        if judul and not isi:
            ws.cell(r, 2, judul).font = Font(bold=True, size=12, color=BIRU)
        elif judul or isi:
            ws.cell(r, 2, judul).font = Font(bold=True, size=10)
            ws.cell(r, 2).alignment = Alignment(vertical="top")
            sel = ws.cell(r, 3, isi)
            sel.font = Font(size=10)
            sel.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 28 if len(isi) > 95 else 15


def sheet_transfer(ws):
    ws.sheet_properties.tabColor = EMAS
    akhir = len(TRANSFER_R2)
    for c in range(1, akhir + 1):
        sel = ws.cell(1, c, TRANSFER_R1.get(c))
        _sel_header(sel, EMAS_TUA)
    for c, t in enumerate(TRANSFER_R2, start=1):
        sel = ws.cell(2, c, t)
        _sel_header(sel, EMAS)
        if t in TRANSFER_EN:
            sel.comment = _komentar(TRANSFER_EN[t])
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.cell(1, 1).comment = _komentar("汇出 = money going OUT of a channel (columns A-H).", 60)
    ws.cell(1, 9).comment = _komentar("收款 = where the money LANDS (columns I-S).", 60)
    ws.row_dimensions[2].height = 30
    format_baris_data(ws, TRANSFER_R2, 3)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(akhir)}2"


def sheet_data(ws, kolom, warna, wajib):
    ws.sheet_properties.tabColor = warna
    lebar = {t: 20 for t in ("sales_name", "sales_dept", "sales_area", "Sales Name",
                             "Sales Dept", "Sales Area", "Payment Gateway", "Pay Channel",
                             "receive_company_account")}
    lebar.update({t: 18 for t in ("hash", "referenace", "Reference", "Remark")})
    tulis_header(ws, kolom, warna, wajib=wajib, lebar=lebar)
    format_baris_data(ws, kolom, 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolom))}1"


def sheet_xero(ws):
    ws.sheet_properties.tabColor = ORANYE
    tulis_header(ws, XERO_KOLOM, ORANYE, lebar={"Date": 12})
    ws.cell(1, 1).comment = _komentar("One row per date. Each currency column = local "
                                      "currency per 1 USD.", 60)
    ws.cell(2, 1).number_format = FMT_TGL
    for c in range(2, len(XERO_KOLOM) + 1):
        ws.cell(2, c).number_format = "#,##0.######"
    ws.freeze_panes = "B2"


def sheet_fee(ws):
    ws.sheet_properties.tabColor = UNGU
    tulis_header(ws, FEE_KOLOM, UNGU, lebar={"Payment Channel": 24, "Remark": 30,
                                             "Deposit": 14, "Withdrawal": 14})
    for c in (3, 4):
        ws.cell(1, c).comment = _komentar(FEE_PANDUAN, 120, 300)
    ws.cell(1, 1).comment = _komentar("Required column - do not rename or delete.\n"
                                      "Gateway name as in D / W. Case, spaces and "
                                      "punctuation are ignored.", 80)
    sel = ws.cell(1, 6, FEE_PANDUAN)
    sel.font = Font(italic=True, size=9, color=ABU)
    sel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["F"].width = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:E1"


def sheet_cny_rate(ws):
    ws.sheet_properties.tabColor = KVB
    tulis_header(ws, CNY_KOLOM, KVB, lebar={"Date": 13, "TD": 16, "Rate": 11})
    for c, t in enumerate(CNY_KOLOM, start=1):
        ws.cell(1, c).comment = _komentar(f"Required column - do not rename or delete.\n"
                                          f"{CNY_EN[t]}", 100, 280)
    ws.cell(2, 1).number_format = FMT_TGL
    ws.cell(2, 3).number_format = "0.0000"
    # Contoh di kolom E-G, BUKAN di A-C: kalkulator hanya membaca kolom yang
    # header-nya Date | TD | Rate, jadi contoh ini tidak pernah terbaca sebagai rate.
    abu_muda = PatternFill("solid", fgColor="F2F2F2")
    judul = ws.cell(1, 5, "EXAMPLE - not read. Fill columns A-C.")
    judul.font = Font(bold=True, italic=True, size=9, color="C00000")
    ws.merge_cells("E1:G1")
    for c, t in enumerate(("Example date", "Example TD", "Example rate"), start=5):
        sel = ws.cell(2, c, t)
        sel.font = Font(bold=True, italic=True, size=9, color=ABU)
        sel.fill = abu_muda
        sel.border = KOTAK
    for i, (d, td, rate) in enumerate(CNY_CONTOH, start=3):
        for c, v in enumerate((d, td, rate), start=5):
            sel = ws.cell(i, c, v)
            sel.font = Font(italic=True, size=9, color=ABU)
            sel.fill = abu_muda
            sel.border = KOTAK
            if c == 5:
                sel.number_format = FMT_TGL
            elif c == 7:
                sel.number_format = "0.0000"
    ws.cell(len(CNY_CONTOH) + 3, 5, "Rate = CNY per 1 USDT. One row per day per gateway."
            ).font = Font(italic=True, size=9, color=ABU)
    ws.column_dimensions["D"].width = 4
    for kol, w in zip("EFG", (14, 13, 12)):
        ws.column_dimensions[kol].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:C1"


def buat(dst: Path) -> Path:
    out = openpyxl.Workbook()
    sheet_guide(out.active)
    sheet_transfer(out.create_sheet("Transfer"))
    sheet_data(out.create_sheet("D"), HDR_D, BIRU, WAJIB_D)
    sheet_data(out.create_sheet("W"), HDR_W, HIJAU, WAJIB_W)
    sheet_xero(out.create_sheet("Xero"))
    sheet_fee(out.create_sheet("Handling Fee"))
    sheet_cny_rate(out.create_sheet("CNY RATE"))
    out.active = 0
    tetap = datetime.datetime(2026, 10, 30)
    out.properties.creator = "DW Calculator"
    out.properties.created = tetap
    out.properties.modified = tetap
    dst.parent.mkdir(parents=True, exist_ok=True)
    out.save(dst)
    return dst


def main():
    ap = argparse.ArgumentParser(description="Bikin Template KVB Plus.xlsx")
    ap.add_argument("-o", "--output", help=f"default: {NAMA_FILE} di folder skrip")
    args = ap.parse_args()
    dst = (Path(args.output).expanduser() if args.output
           else Path(__file__).resolve().parent / NAMA_FILE)
    buat(dst)
    wb = openpyxl.load_workbook(dst, read_only=True)
    print(f"Simpan : {dst}  ({dst.stat().st_size / 1024:,.0f} KB)")
    print(f"         sheet: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
