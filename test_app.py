"""Smoke test: python3 test_app.py

Self-contained: every input is synthetic and generated at test time by testkit.py
(D&W source, Deals History CSV, Client Equity FX, and an empty 'Template D&W.xlsx'
built with buat_template.py). Nothing outside the repo is read, the real users.db /
session.key are never touched, and no real client data is needed.

Model app-nya ASINKRON sejak 28 Aug 2026: POST -> 303 ke /job/<id>, kerjanya di
thread, lalu /job/<id>/download. Jadi tes ini POST, ambil job id dari header
Location, tunggu state 'done'/'failed', baru unduh.
"""
import io
import shutil
import subprocess
import time
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

# Every route is behind Basic Auth. Use a throwaway user database and send the
# credentials on every request, so this file keeps testing the tools themselves.
# Assigned, never setdefault: an exported server DW_* must not leak in.
import testkit as T

TEST_ROOT = T.hermetic_env("tester", "tester-password-1")

import app as A  # noqa: E402
from app import JOBS, TOOLS, app  # noqa: E402

T.assert_isolated(TEST_ROOT)

assert app.config["MAX_CONTENT_LENGTH"] == 1024 * 1024 * 1024
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024
c = app.test_client()
c.environ_base["HTTP_AUTHORIZATION"] = "Basic " + __import__("base64").b64encode(
    b"tester:tester-password-1").decode()
FIXTURES = TEST_ROOT / "fixtures"
FIXTURES.mkdir()
DEALS_FIXTURE = T.write_deals_csv(FIXTURES / "31 Jul_Deals History 2026_08_07 12_07_27.csv")
EQUITY_FIXTURE = T.write_equity_fx(FIXTURES / "Client Equity - FX.xlsx")
DAILY_ROWS = T.DEALS_EXPECTED["daily_rows"]
KEPT = T.DEALS_EXPECTED["kept"]
DW_SOURCE = T.write_dpm_dw_source(FIXTURES / "D&W JUN 2026.xlsx", 2026, 6)
TEMPLATE = T.build_template(DW_SOURCE, FIXTURES / "Template D&W.xlsx")

# The private 'Template D&W.xlsx' is not in the repo. isi_template.py defaults to
# BASE/'Template D&W.xlsx', so hand every job the generated template instead.
_real_run = subprocess.run
COMMANDS = []          # every script command line the jobs ran, for argument checks


def _run_with_template(cmd, *args, **kwargs):
    if isinstance(cmd, list):
        COMMANDS.append([str(part) for part in cmd])
    if (isinstance(cmd, list) and len(cmd) > 1 and str(cmd[1]).endswith("isi_template.py")
            and "-t" not in cmd and "--template" not in cmd):
        cmd = [*cmd, "-t", str(TEMPLATE)]
    return _real_run(cmd, *args, **kwargs)


A.subprocess.run = _run_with_template

# Setiap upload WAJIB menyertakan bulan laporan (28 Aug 2026 -> 31 Aug 2026).
BULAN = {"bulan": "6", "tahun": "2026"}


def kirim(nama_file, fh, **tambahan):
    """POST lalu tunggu sampai job selesai. -> (job_id, state, info)"""
    data = {"file": (fh, nama_file)}
    data.update(tambahan)
    r = c.post("/tool/dw", data=data)
    if r.status_code != 303:
        return None, r.status_code, r
    job_id = r.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
    for _ in range(600):                       # maksimal 5 menit
        info = A._read_state(job_id)
        if info and info["state"] in ("done", "failed"):
            return job_id, info["state"], info
        time.sleep(0.5)
    raise AssertionError("job tidak pernah selesai")


landing = c.get("/")
assert landing.status_code == 200
assert set(TOOLS) == {"dw", "segregate", "pl-desk", "segregated"}
assert b"Dupoin DPM Tools" in landing.data
assert b"Generate D&amp;W" in landing.data
assert b"Deal Segregator" in landing.data
assert b"PL by Desk" in landing.data
assert b"one ZIP containing two XLSX workbooks" in landing.data
assert c.get("/tool/mtoatd").status_code == 404, "menu yang tidak ada harus 404"
assert c.get("/tool/hitung").status_code == 404, "menu lama sudah dihapus"
assert c.get("/tool/../../etc/passwd").status_code == 404

# Halaman upload harus memuat pemilih bulan, dan default-nya BULAN LALU.
halaman = c.get("/tool/dw")
assert halaman.status_code == 200
assert b'name="bulan"' in halaman.data and b'name="tahun"' in halaman.data, \
    "dropdown bulan laporan hilang dari halaman upload"
assert b'name="saldo_jw"' in halaman.data
assert b"multiple" not in halaman.data

halaman_segregate = c.get("/tool/segregate")
assert halaman_segregate.status_code == 200
class FilePickerParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.pickers = []
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "input" and a.get("type") == "file": self.pickers.append(a)
file_picker_parser = FilePickerParser()
file_picker_parser.feed(halaman_segregate.data.decode())
assert [p.get("id") for p in file_picker_parser.pickers] == ["deals-file", "equity-file"], \
    "Deals History and Client Equity FX need separate file pickers"
assert "multiple" in file_picker_parser.pickers[0]
assert "multiple" not in file_picker_parser.pickers[1]
assert b'id="deals-drop"' in halaman_segregate.data
assert b'id="equity-drop"' in halaman_segregate.data
assert b"Drop Deals History files here" in halaman_segregate.data
assert b"Drop Client Equity FX here" in halaman_segregate.data
assert b"var CHUNK = 10 * 1024 * 1024" in halaman_segregate.data
assert b"var DIRECT = 0" in halaman_segregate.data
assert b"DataTransfer" in halaman_segregate.data
assert halaman_segregate.data.count(b'class="drop-files"') == 2
assert b'className = "file-remove"' in halaman_segregate.data
assert b"button.textContent = 'Remove'" in halaman_segregate.data
assert b"unique(selected.concat(incoming))" in halaman_segregate.data
assert b"selected.splice(index, 1)" in halaman_segregate.data
assert b"input.files = transfer.files" in halaman_segregate.data
assert b'id="upload-progress"' in halaman_segregate.data
assert b'id="upload-progress-bar"' in halaman_segregate.data
assert b'id="upload-progress-title"' in halaman_segregate.data
assert b'id="upload-progress-detail"' in halaman_segregate.data
assert b"function updateProgress" in halaman_segregate.data
assert b"function uploadError" in halaman_segregate.data
assert b"File ' + (fileIndex + 1) + ' of ' + files.length" in halaman_segregate.data
assert b"retrying ' + info.n + '/' +" in halaman_segregate.data
assert b"var RETRY = {max: 8," in halaman_segregate.data
assert b"MB uploaded" in halaman_segregate.data
assert b"office firewall/antivirus or VPN/WARP, or try another network" in halaman_segregate.data
assert b"multiple" in halaman_segregate.data
assert b'name="bulan"' not in halaman_segregate.data
assert b'name="saldo_jw"' not in halaman_segregate.data
assert b"Download the template" not in halaman_segregate.data
assert b"12 GB total upload limit" in halaman_segregate.data
assert b"deleted from the server" in halaman_segregate.data
assert b"70 files maximum" in halaman_segregate.data
assert b"in pieces automatically" in halaman_segregate.data
assert b"fileCount > MAX_FILES" in halaman_segregate.data
assert b"file.size === 0" in halaman_segregate.data

missing = c.post("/tool/segregate", data={}, content_type="multipart/form-data")
assert missing.status_code == 400
assert b".csv, .xlsx, or .xlsm" in missing.data
invalid = c.post(
    "/tool/segregate",
    data={"file": [(io.BytesIO(b"x"), "a.txt")]},
    content_type="multipart/form-data",
)
assert invalid.status_code == 400
assert b".csv, .xlsx, or .xlsm" in invalid.data

too_many = c.post("/tool/segregate", data={
    "file": [(io.BytesIO(b"x"), f"{i}.csv") for i in range(71)]
}, content_type="multipart/form-data")
assert too_many.status_code == 400 and b"no more than 70 files" in too_many.data

empty = c.post("/tool/segregate", data={
    "file": (io.BytesIO(b""), "empty.csv")
}, content_type="multipart/form-data")
assert empty.status_code == 400 and b"Empty files cannot be processed" in empty.data


def kirim_segregate(files):
    r = c.post("/tool/segregate", data={"file": files},
               content_type="multipart/form-data")
    assert r.status_code == 303, (r.status_code, r.data[:500])
    job_id = r.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
    for _ in range(600):
        info = A._read_state(job_id)
        if info and info["state"] in ("done", "failed"):
            return job_id, info
        time.sleep(0.5)
    raise AssertionError("Deal Segregator job did not finish")


# Unicode-only stems are stripped by secure_filename; the validated suffix must survive.
captured = {}
unicode_job_id = None
real_thread = A.threading.Thread
class CapturingThread:
    def __init__(self, target, args, **kwargs):
        captured["sources"] = args[4]
    def start(self):
        pass
A.threading.Thread = CapturingThread
try:
    r = c.post("/tool/segregate", data={"file": (io.BytesIO(b"x"), "交易.xlsx")},
               content_type="multipart/form-data")
    assert r.status_code == 303
    unicode_job_id = r.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
    assert captured["sources"][0].suffix == ".xlsx", captured["sources"][0]
finally:
    A.threading.Thread = real_thread
    if unicode_job_id:
        shutil.rmtree(JOBS / unicode_job_id, ignore_errors=True)
        A._state_path(unicode_job_id).unlink(missing_ok=True)

assert DEALS_FIXTURE.is_file(), f"synthetic Deals History fixture missing: {DEALS_FIXTURE}"
job_id, info = kirim_segregate([(DEALS_FIXTURE.open("rb"), DEALS_FIXTURE.name)])
assert info["state"] == "done", info
assert info["name"].endswith("-hasil.zip"), info["name"]
r = c.get(f"/job/{job_id}/download")
assert r.status_code == 200
assert r.mimetype == "application/zip"
with zipfile.ZipFile(io.BytesIO(r.data)) as archive:
    assert archive.namelist() == ["Deals - Daily.xlsx", "Deals - Monthly Summary.xlsx"]
    wb = load_workbook(io.BytesIO(archive.read("Deals - Daily.xlsx")), read_only=True, data_only=True)
    try:
        assert wb.sheetnames == ["Daily", "Daily - In", "Verifikasi"]
        assert sum(1 for _ in wb["Daily"].iter_rows(values_only=True)) - 1 == DAILY_ROWS
        assert sum(1 for _ in wb["Daily - In"].iter_rows(values_only=True)) > 1
    finally:
        wb.close()

assert EQUITY_FIXTURE.is_file(), f"Client Equity FX fixture missing: {EQUITY_FIXTURE}"
job_id, info = kirim_segregate([
    (DEALS_FIXTURE.open("rb"), DEALS_FIXTURE.name),
    (EQUITY_FIXTURE.open("rb"), EQUITY_FIXTURE.name),
])
assert info["state"] == "done", info
assert "FX conversion: True" in (info["log"] or "") or "pakai_fx: True" in (info["log"] or "")
fx_response = c.get(f"/job/{job_id}/download")
assert fx_response.status_code == 200 and fx_response.mimetype == "application/zip"
with zipfile.ZipFile(io.BytesIO(fx_response.data)) as archive:
    wb = load_workbook(io.BytesIO(archive.read("Deals - Daily.xlsx")), read_only=True, data_only=True)
    try:
        headers = [cell.value for cell in wb["Daily"][1]]
        assert headers[-4:] == ["Commission (USD)", "Fee (USD)", "Swap (USD)", "Profit (USD)"]
        assert sum(1 for _ in wb["Daily"].iter_rows(values_only=True)) - 1 == DAILY_ROWS
    finally:
        wb.close()

job_id, info = kirim_segregate([
    (DEALS_FIXTURE.open("rb"), DEALS_FIXTURE.name),
    (DEALS_FIXTURE.open("rb"), DEALS_FIXTURE.name),
])
assert info["state"] == "done", info
# the same file twice: duplicates are dropped, so 'kept' equals ONE file's unique 'out' rows
assert f"kept: {KEPT}" in (info["log"] or ""), info["log"]
assert c.get(f"/job/{job_id}/download").data[:2] == b"PK"

# --- yang harus DITOLAK sebelum job dibuat ---------------------------------
assert c.post("/tool/dw", data={"file": (io.BytesIO(b"x"), "a.txt"), **BULAN}
              ).status_code == 400, "file bukan .xlsx harus ditolak"

src = DW_SOURCE
assert src.is_file(), "synthetic 'D&W JUN 2026.xlsx' missing"

r = c.post("/tool/dw", data={"file": (src.open("rb"), src.name)})
assert r.status_code == 400 and b"report month" in r.data, \
    "upload tanpa bulan laporan harus ditolak"
r = c.post("/tool/dw", data={"file": (src.open("rb"), src.name),
                             "bulan": "13", "tahun": "2026"})
assert r.status_code == 400, "bulan 13 harus ditolak"

# --- template KOSONG -> pesan jelas, bukan 500 mentah ----------------------
tpl = TEMPLATE
_id, state, info = kirim(tpl.name, tpl.open("rb"), **BULAN)
assert state == "failed", state
assert "TIDAK ADA BARIS DATA" in (info["log"] or ""), (info["log"] or "")[:400]

# --- file yang ADA datanya -> jadi, dan bulannya masuk nama file -----------
COMMANDS.clear()
job_id, state, info = kirim(src.name, src.open("rb"), **BULAN)
assert state == "done", (state, info.get("error"), (info.get("log") or "")[-600:])
assert info["periode"] == "2026-06", info["periode"]
assert "Jun 2026" in info["name"], info["name"]
# The chosen month must reach both scripts as an argument, and hitung_dw must say so.
ran = {Path(cmd[1]).name: cmd for cmd in COMMANDS}
assert ran["isi_template.py"][ran["isi_template.py"].index("--bulan") + 1] == "2026-06", ran
assert ran["hitung_dw.py"][ran["hitung_dw.py"].index("--period") + 1] == "2026-06", ran
assert "Periode: Jun 2026" in info["log"], info["log"]

r = c.get(f"/job/{job_id}/download")
assert r.status_code == 200, r.status_code
assert r.data[:2] == b"PK", r.data[:200]                       # xlsx, bukan HTML
assert "attachment" in r.headers["Content-Disposition"]
assert not (JOBS / job_id).exists(), "job dir harus dibersihkan setelah diunduh"
hasil_jun = len(r.data)

# --- saldo pembuka J Wallet yang diketik di halaman upload -------------------
halaman = c.get("/tool/dw")
assert b'name="saldo_jw"' in halaman.data, "field saldo pembuka hilang dari halaman"
assert b'class="prevlbl"' in halaman.data, "label bulan sebelumnya hilang"
assert b'name="saldo_channel"' in halaman.data, "textarea saldo channel hilang"
assert b"out.innerHTML" not in halaman.data, "upload error DOM XSS"
# v1.24 moved chunk retries into makeChunkSender (exponential backoff, RETRY.max tries).
assert b"var sendChunk = makeChunkSender(" in halaman.data, "chunk sender missing"
assert b"var RETRY = {max: 8," in halaman.data and b"if (n > RETRY.max) return fail(err);" in halaman.data, \
    "chunk upload must retry transient failures"

r = c.post("/tool/dw", data={"file": (src.open("rb"), src.name),
                             "saldo_jw": "bukan angka!!", **BULAN})
assert r.status_code == 400 and b"must be a number" in r.data, "saldo ngawur harus ditolak"
assert b"bukan angka!!" in r.data and A.VERSI_LABEL.encode() in r.data, \
    "the refused value and the version must be shown"

job_id, state, info = kirim(src.name, src.open("rb"), saldo_jw="377.233,36", **BULAN)
assert state == "done", (state, info.get("error"), (info.get("log") or "")[-500:])
assert info["saldo_jw"] == "377.233,36", info["saldo_jw"]
# angka Eropa/Indonesia harus dibaca 377233.36, dan tanggalnya akhir bulan SEBELUMNYA
assert "377,233.36" in info["log"], (info["log"] or "")[:1500]
assert "2026-05-31" in info["log"], (info["log"] or "")[:1500]
assert c.get(f"/job/{job_id}/download").data[:2] == b"PK"

# tanpa saldo -> tidak boleh ada baris 'Saldo :' sama sekali
job_id, state, info = kirim(src.name, src.open("rb"), **BULAN)
assert state == "done", state
assert info["saldo_jw"] is None
assert "diketik di halaman upload" not in (info["log"] or "")
assert c.get(f"/job/{job_id}/download").data[:2] == b"PK"

# --- bulan yang SALAH -> berhenti dengan pesan yang menyebut bulannya --------
# Sengaja GAGAL, bukan menghasilkan workbook kosong: file kosong yang kelihatan
# normal jauh lebih berbahaya daripada pesan error.
job_id, state, info = kirim(src.name, src.open("rb"), bulan="1", tahun="2026")
assert state == "failed", state
assert "BULAN LAPORANNYA SALAH" in (info["error"] or ""), info["error"]
assert "2026-06" in (info["log"] or ""), (info["log"] or "")[-800:]

# semua menu yang terdaftar harus bisa dibuka
for slug in TOOLS:
    assert c.get(f"/tool/{slug}").status_code == 200, slug

# 'staging' holds in-flight chunked uploads and is permanent; it must however be
# empty once every run has finished, or a large upload was left behind on disk.
sisa = [p for p in JOBS.iterdir() if p.is_dir() and p.name != "staging"]
assert not sisa, f"job dir bocor: {sisa}"
staging = JOBS / "staging"
if staging.is_dir():
    tertinggal = list(staging.iterdir())
    assert not tertinggal, f"staging bocor: {tertinggal}"

subprocess.run = _real_run
print("OK", len(TOOLS), "menu,", f"{hasil_jun:,}", "bytes hasil Jun 2026")
