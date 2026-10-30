"""Regression check for the downloadable blank KVB workbook template."""
import base64
import io
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(tempfile.mkdtemp())
os.environ["DW_AUTH_DB"] = str(ROOT / "users.db")
os.environ["DW_JOB_DIR"] = str(ROOT / "jobs")
os.environ["DW_ADMIN_USER"] = "tester"
os.environ["DW_ADMIN_PASS"] = "tester-password-1"

from app import app

client = app.test_client()
client.environ_base["HTTP_AUTHORIZATION"] = "Basic " + base64.b64encode(
    b"tester:tester-password-1"
).decode()

page = client.get("/kvb/tool/dw")
assert page.status_code == 200
assert b'href="/kvb/template"' in page.data
assert b"Download Template KVB Plus.xlsx" in page.data

response = client.get("/kvb/template")
assert response.status_code == 200
assert "Template KVB Plus.xlsx" in response.headers["Content-Disposition"]

workbook = load_workbook(io.BytesIO(response.data), data_only=True)
assert workbook.sheetnames == ["Guide", "Transfer", "D", "W", "Xero", "Handling Fee",
                               "CNY RATE"], workbook.sheetnames


def values_below(ws, first_row):
    return [c.value for row in ws.iter_rows(min_row=first_row) for c in row
            if c.value is not None]


assert workbook["Transfer"]["A1"].value == "汇出"
assert workbook["Transfer"]["I1"].value == "收款"
assert workbook["Transfer"]["A2"].value == "汇款日期"
assert not values_below(workbook["Transfer"], 3), "Transfer must contain headers only"
assert workbook["Transfer"].freeze_panes == "A3"
for sheet in ("D", "W", "Xero", "Handling Fee"):
    assert not values_below(workbook[sheet], 2), f"{sheet} must contain headers only"
    assert workbook[sheet].freeze_panes in ("A2", "B2"), sheet
assert workbook["D"]["A1"].value == "Source.Name"
assert workbook["W"]["A1"].value == "Source.Name"
assert workbook["Xero"]["A1"].value == "Date"
assert workbook["Handling Fee"]["A1"].value == "Payment Channel"
guide = str(workbook["Handling Fee"]["F1"].value)
assert "+8" in guide and "flat" in guide
assert "1.5%+50" in guide and "min" in guide.lower()

# Same header names as the KVB Plus export (siapkan_kvb.py translates them).
import buat_template_kvb as btk  # noqa: E402
import siapkan_kvb as sk  # noqa: E402

for sheet, kolom in (("D", btk.HDR_D), ("W", btk.HDR_W), ("Xero", btk.XERO_KOLOM),
                     ("Handling Fee", btk.FEE_KOLOM)):
    assert [c.value for c in workbook[sheet][1]][:len(kolom)] == kolom, sheet
assert [c.value for c in workbook["Transfer"][2]] == btk.TRANSFER_R2
for h_ in ("METHOD", "RECEIVE_COMPANY_ACCOUNT", "USD (CRM)", "TRANSACTION", "SOURCE.NAME"):
    assert h_ in sk.D_ALIAS, h_
assert workbook["D"]["H2"].number_format == "yyyy-mm-dd"
assert workbook["D"]["I2"].number_format == "hh:mm:ss"
assert workbook["W"]["S2"].number_format == "yyyy-mm-dd"
assert "Required column" in workbook["D"]["K1"].comment.text

# CNY RATE: header only in A-C, the example block sits in E-G and is not read.
cny = workbook["CNY RATE"]
assert [cny.cell(1, c).value for c in (1, 2, 3)] == ["Date", "TD", "Rate"]
assert all(cny.cell(r, c).value is None for r in range(2, cny.max_row + 1) for c in (1, 2, 3))
assert "EXAMPLE" in str(cny["E1"].value) and "not read" in str(cny["E1"].value)
assert cny["E2"].value == "Example date" and cny["F3"].value == "CHIPPAY"
assert cny.freeze_panes == "A2"
assert workbook["Guide"]["B2"].value.startswith("KVB Plus D&W")
workbook.close()

# The committed file equals what buat_template_kvb.py writes.
import tempfile as _tf  # noqa: E402

with _tf.TemporaryDirectory() as _d:
    fresh = btk.buat(Path(_d) / "t.xlsx")
    a = load_workbook(fresh)
    b = load_workbook(io.BytesIO(response.data))
    assert a.sheetnames == b.sheetnames
    for ws_a in a:
        ws_b = b[ws_a.title]
        sel_a = {(c.coordinate, c.value, c.number_format, c.fill.fgColor.rgb)
                 for row in ws_a.iter_rows() for c in row}
        sel_b = {(c.coordinate, c.value, c.number_format, c.fill.fgColor.rgb)
                 for row in ws_b.iter_rows() for c in row}
        assert sel_a == sel_b, (ws_a.title, sel_a ^ sel_b)
        assert ws_a.freeze_panes == ws_b.freeze_panes, ws_a.title
    a.close()
    b.close()

# The downloaded template, filled with synthetic rows, runs end to end the way
# the web app runs a KVB upload (siapkan_kvb -> isi_template -> hitung_dw).
import datetime  # noqa: E402
import subprocess  # noqa: E402
import sys  # noqa: E402

import app as app_mod  # noqa: E402

HERE = Path(__file__).resolve().parent
with _tf.TemporaryDirectory() as _d:
    tmp = Path(_d)
    wb = load_workbook(io.BytesIO(response.data))
    d = wb["D"]
    hd = [c.value for c in d[1]]
    hw = [c.value for c in wb["W"][1]]

    def put(ws, header, values):
        r = ws.max_row + 1 if any(c.value is not None for c in ws[ws.max_row]) else ws.max_row
        r = max(r, 2)
        for k, v in values.items():
            ws.cell(r, header.index(k) + 1, v)

    for day, tx, t in ((1, 6900, "T1"), (2, 13800, "T2")):
        put(d, hd, {"Source.Name": "07-2026 Deposit", "Paid Date": datetime.datetime(2026, 7, day),
                    "Settlement Date": datetime.datetime(2026, 7, day), "method": "CNY",
                    "ticket": t, "USD (CRM)": tx / 6.9, "rate": 1 / 6.9, "transaction": tx,
                    "receive_company_account": "ChipPay"})
    put(wb["W"], hw, {"Source.Name": "07-2026 Withdrawal",
                      "Apply Date": datetime.datetime(2026, 7, 2),
                      "Payment Currency": "CNY", "Payment Gateway": "ChipPay",
                      "Reference": "W1", "USD (CRM)": 100, "rate": 1 / 6.9, "charges": 0,
                      "transaction": 690, "Completed Date": datetime.datetime(2026, 7, 2),
                      "Settlement Date": datetime.datetime(2026, 7, 2), "status": "finish"})
    x = wb["Xero"]
    hx = [c.value for c in x[1]]
    for day in range(1, 32):
        x.cell(day + 1, 1, datetime.datetime(2026, 7, day))
        for cur, v in (("CNY", 6.9), ("USD", 1), ("USDT", 1)):
            x.cell(day + 1, hx.index(cur) + 1, v)
    fee = wb["Handling Fee"]
    for i, v in enumerate(("ChipPay", "CNY", 0.004, 0.007), start=1):
        fee.cell(2, i, v)
    cr = wb["CNY RATE"]
    cr["A2"], cr["B2"], cr["C2"] = datetime.datetime(2026, 7, 1), "ChipPay", 7.0
    cr["A3"], cr["B3"], cr["C3"] = "2026-07-02", "CHIPPAY", 6.9
    wb.save(tmp / "kvb.xlsx")

    def jalan(*args):
        r = subprocess.run([sys.executable, *map(str, args)], cwd=HERE,
                           capture_output=True, text=True)
        assert r.returncode == 0, (args[0], r.stdout[-3000:], r.stderr[-3000:])
        return r.stdout

    out1 = jalan("siapkan_kvb.py", tmp / "kvb.xlsx", "-o", tmp / "t.xlsx")
    assert "CNY RATE" in out1 and "2 baris" in out1, out1
    jalan("buat_template.py", tmp / "t.xlsx", "-o", tmp / "tpl0.xlsx")
    jalan("isi_template.py", tmp / "t.xlsx", "-t", tmp / "tpl0.xlsx", "--bulan", "2026-07",
          *app_mod.KVB_ARG["isi_template.py"], "-o", tmp / "tpl.xlsx")
    jalan("hitung_dw.py", tmp / "tpl.xlsx", "--period", "2026-07", "-o", tmp / "hasil.xlsx",
          *app_mod.KVB_ARG["hitung_dw.py"])
    hasil = load_workbook(tmp / "hasil.xlsx", read_only=True, data_only=True)
    rows = list(hasil["Payment Channel Balance"].iter_rows(values_only=True))
    header = list(rows[0])
    i_d, i_c, i_g = header.index("Date"), header.index("Currency"), header.index("Payment Gateway")
    i_b, i_r, i_u = header.index("Balance"), header.index("CNY Rate"), header.index("Balance (USDT)")
    by_day = {r[i_d].date(): r for r in rows[1:]
              if r[i_c] == "CNY" and str(r[i_g]).upper() == "CHIPPAY" and r[i_d]}
    jul1, jul2 = by_day[datetime.date(2026, 7, 1)], by_day[datetime.date(2026, 7, 2)]
    assert jul1[i_r] == 7.0 and abs(jul1[i_u] - jul1[i_b] / 7.0) < 1e-9, jul1
    assert jul2[i_r] == 6.9 and abs(jul2[i_u] - jul2[i_b] / 6.9) < 1e-9, jul2
    d_rows = list(hasil["D&W Detail"].iter_rows(values_only=True))
    assert len([r for r in d_rows[1:] if r and r[0] is not None]) >= 3, d_rows[:5]
    hasil.close()

print("KVB template route, sidebar link, sheets, headers, CNY RATE, generator match "
      "and end-to-end run: OK")
