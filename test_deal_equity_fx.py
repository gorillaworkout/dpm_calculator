"""Focused Deal + Client Equity FX regression: python3 test_deal_equity_fx.py"""
import csv
import datetime
import io
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from deal_segregator import proses

HEADERS = ["Deal", "Position", "Login", "Group", "Country", "Time", "Type", "Entry", "Symbol",
           "Volume", "Commission", "Fee", "Swap", "Profit", "Currency"]
USD_HEADERS = ["Commission (USD)", "Fee (USD)", "Swap (USD)", "Profit (USD)"]


def write_deals(path):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        # Fokus tes ini murni konversi FX. Baris 'in' dan 'out' tidak digabung.
        w.writerow(["1", "P1", "100", r"real\DPMKT-15", "MY", "2026.07.01 01:02:03.004",
                    "buy", "out", "EURUSD", "1", "-2", "-1", "-4", "20", "JPY"])
        w.writerow(["2", "P2", "100", r"real\DPMKT-15", "MY", "2026.07.02 01:02:03.004",
                    "buy", "out", "EURUSD", "1", "-3", "-2", "-5", "30", "JPY"])


def write_fx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    ws.append([datetime.date(2026, 7, 1), "JPY", 2])
    ws.append([datetime.date(2026, 7, 2), "JPY", None])
    wb.save(path)
    wb.close()


with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    deals, fx, out = tmp / "deals.csv", tmp / "Client Equity - FX.xlsx", tmp / "result.zip"
    write_deals(deals)
    write_fx(fx)
    summary = proses([deals, fx], out)

    assert zipfile.is_zipfile(out), "Deal+Equity output must be a ZIP"
    with zipfile.ZipFile(out) as z:
        assert z.namelist() == ["Deals - Daily.xlsx", "Deals - Monthly Summary.xlsx"]
        wb = load_workbook(io.BytesIO(z.read("Deals - Daily.xlsx")))
        ws = wb["Daily"]
        headers = [c.value for c in ws[1]]
        assert headers[-4:] == USD_HEADERS
        first = {headers[i]: c.value for i, c in enumerate(ws[2])}
        assert first["Commission (USD)"] == -1
        assert first["Fee (USD)"] == -0.5
        assert first["Swap (USD)"] == -2
        assert first["Profit (USD)"] == 10
        second = {headers[i]: c for i, c in enumerate(ws[3])}
        for name in USD_HEADERS:
            assert second[name].value is None
            assert second[name].fill.fgColor.rgb[-6:] == "FFFF00"
        verification = dict(row for row in wb["Verifikasi"].iter_rows(values_only=True)
                            if row and row[0])
        assert "Missing FX rates" in next(k for k in verification if "Missing FX rates" in k)
        wb.close()

        monthly = load_workbook(io.BytesIO(z.read("Deals - Monthly Summary.xlsx")))
        monthly_ws = monthly["Monthly Summary"]
        monthly_headers = [c.value for c in monthly_ws[1]]
        monthly_row = {monthly_headers[i]: c for i, c in enumerate(monthly_ws[2])}
        for name in USD_HEADERS:
            assert monthly_row[name].value is None, f"partial aggregate leaked in {name}"
            assert monthly_row[name].fill.fgColor.rgb[-6:] == "FFFF00"
        monthly.close()
    assert summary["pakai_fx"] is True
    assert summary["kurs_kurang"] == 1

print("OK Deal + Client Equity FX")

# Each row is converted with the FX rate of ITS OWN date: the 'in' row (opening leg,
# own sheet 'Daily - In') uses the opening date's rate, the 'out' row uses the close rate.
with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    deals, fx, out = tmp / "deals.csv", tmp / "fx.xlsx", tmp / "result.zip"
    with deals.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerow(["in-1", "PX", "100", r"real\DPMKT-15", "MY",
                    "2026.07.01 01:02:03.004", "buy", "in", "EURUSD", "1",
                    "-10", "-2", "-4", "0", "JPY"])
        w.writerow(["out-1", "PX", "100", r"real\DPMKT-15", "MY",
                    "2026.07.02 01:02:03.004", "buy", "out", "EURUSD", "1",
                    "-8", "-4", "-12", "40", "JPY"])
    wb = Workbook(); ws = wb.active; ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    ws.append([datetime.date(2026, 7, 1), "JPY", 2])
    ws.append([datetime.date(2026, 7, 2), "JPY", 4])
    wb.save(fx); wb.close()
    proses([deals, fx], out)
    with zipfile.ZipFile(out) as z:
        wb = load_workbook(io.BytesIO(z.read("Deals - Daily.xlsx")), data_only=True)
        headers = [c.value for c in wb["Daily"][1]]
        row_out = {headers[i]: c.value for i, c in enumerate(wb["Daily"][2])}
        row_in = {headers[i]: c.value for i, c in enumerate(wb["Daily - In"][2])}
        wb.close()
    assert row_out["Commission"] == -8 and row_out["Commission (USD)"] == -2
    assert row_out["Fee (USD)"] == -1 and row_out["Swap (USD)"] == -3
    assert row_out["Profit (USD)"] == 10
    assert row_in["Commission"] == -10 and row_in["Commission (USD)"] == -5
    assert row_in["Fee (USD)"] == -1 and row_in["Swap (USD)"] == -2

print("OK per-row-date FX")


def expect_bad_fx(rows, *parts):
    with TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        deals, fx = tmp / "deals.csv", tmp / "fx.xlsx"
        write_deals(deals)
        wb = Workbook()
        ws = wb.active
        ws.title = "Query result"
        ws.append(["date", "Currency", "rate"])
        for item in rows:
            ws.append(item)
        wb.save(fx)
        wb.close()
        try:
            proses([deals, fx], tmp / "out.zip")
        except SystemExit as exc:
            message = str(exc)
            assert all(part in message for part in parts), message
        else:
            raise AssertionError("invalid FX data was accepted")


expect_bad_fx([[datetime.date(2026, 7, 1), "JPY", -2]], "positive", "row 2")
expect_bad_fx([
    [datetime.date(2026, 7, 1), "JPY", 2],
    [datetime.date(2026, 7, 1), "JPY", 3],
], "duplicate", "JPY")

with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    deals, first, second = tmp / "deals.csv", tmp / "fx-1.xlsx", tmp / "fx-2.xlsx"
    write_deals(deals)
    write_fx(first)
    write_fx(second)
    try:
        proses([deals, first, second], tmp / "out.zip")
    except SystemExit as exc:
        assert "one Client Equity FX" in str(exc), exc
    else:
        raise AssertionError("multiple Client Equity FX workbooks were accepted")

print("OK FX validation")
