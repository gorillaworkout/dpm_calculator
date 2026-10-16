#!/usr/bin/env python3
"""Bad bytes in an uploaded text file must not crash the Deal Segregator.

python3 test_deal_segregator_decode.py
"""
import contextlib
import io
import sys
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

import deal_segregator
from deal_segregator import proses

HEADER = ["Deal", "Login", "Name", "Time", "Type", "Entry", "Symbol",
          "Volume", "Commission", "Fee", "Swap", "Profit", "Currency"]
SURROGATE = b"\x3d\xd8"  # lone high surrogate, U+D83D
CHUNK = io.DEFAULT_BUFFER_SIZE


def _u16(teks):
    return teks.encode("utf-16-le")


def _baris(*sel):
    return _u16(",".join(sel) + "\n")


def _tulis(path, blob):
    path.write_bytes(blob)


def _verifikasi(zip_path, folder):
    with zipfile.ZipFile(zip_path) as archive:
        archive.extract("Deals - Daily.xlsx", folder)
    book = load_workbook(Path(folder) / "Deals - Daily.xlsx", read_only=True, data_only=True)
    try:
        return {row[0]: row[1] for row in book["Verifikasi"].iter_rows(values_only=True)
                if row and len(row) > 1 and row[0]}
    finally:
        book.close()


def _harian(zip_path, folder):
    with zipfile.ZipFile(zip_path) as archive:
        archive.extract("Deals - Daily.xlsx", folder)
    book = load_workbook(Path(folder) / "Deals - Daily.xlsx", read_only=True, data_only=True)
    try:
        return list(book["Daily"].iter_rows(values_only=True))
    finally:
        book.close()


def _file_utf16_rusak(path):
    """Lone surrogate in Name, and another on the first read-buffer boundary."""
    head = _baris(*HEADER)
    # Early in the file: surrogate in Name, then a normal character.
    satu = (_u16("1,501,") + SURROGATE
            + _u16("Z,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,10,USD\n"))
    # Second row's name is padded so the surrogate is the last unit of the
    # first decoder chunk, then a normal character.
    lead = _u16("2,502,")
    pos = len(head) + len(satu) + len(lead)
    target = CHUNK - 2  # 8190: last UTF-16 unit that fits in an 8192-byte read
    # File layout is BOM + body. The decoder's first chunk is the first
    # CHUNK bytes of the file, BOM included.
    pos_in_file = 2 + pos
    gap = target - pos_in_file
    assert gap % 2 == 0 and gap >= 0, (pos_in_file, gap, CHUNK)
    nama = (b"A\x00" * (gap // 2)) + SURROGATE + _u16("Z")
    dua_awal = lead + nama
    ekor = _u16(",2026.07.02 00:00:00,buy,out,EURUSD,2,-2,0,0,20,USD\n")
    blob = b"\xff\xfe" + head + satu + dua_awal + ekor
    assert blob[target:target + 2] == SURROGATE, blob[target - 4:target + 6]
    assert blob[target + 2:target + 4] == _u16("Z")
    _tulis(path, blob)


def _file_bersih(path, pad_name):
    head = _baris(*HEADER)
    satu = _baris("1", "501", "Client Z", "2026.07.01 00:00:00", "buy", "out",
                  "EURUSD", "1", "-1", "0", "0", "10", "USD")
    dua = _baris("2", "502", pad_name, "2026.07.02 00:00:00", "buy", "out",
                 "EURUSD", "2", "-2", "0", "0", "20", "USD")
    _tulis(path, b"\xff\xfe" + head + satu + dua)


def _nama_panjang(path_rusak):
    """The padded name from the damaged file, without the surrogate."""
    blob = path_rusak.read_bytes()
    # Recover the name text the clean file must use: A's plus Z, no surrogate.
    target = CHUNK - 2
    # Walk back over A\x00 until the comma before the name.
    start = target
    while blob[start - 2:start] == b"A\x00":
        start -= 2
    banyak = (target - start) // 2
    return ("A" * banyak) + "Z"


def _berhenti(func):
    try:
        func()
    except SystemExit as exc:
        pesan = exc.code if isinstance(exc.code, str) else str(exc)
        return pesan
    raise AssertionError("the run should have stopped")


with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    rusak = tmp / "rusak.csv"
    _file_utf16_rusak(rusak)
    bersih = tmp / "bersih.csv"
    _file_bersih(bersih, _nama_panjang(rusak))
    out_r, out_b = tmp / "rusak.zip", tmp / "bersih.zip"
    proses([rusak], out_r)
    proses([bersih], out_b)
    assert _harian(out_r, tmp / "r") == _harian(out_b, tmp / "b"), "damaged names changed the totals"
    v = _verifikasi(out_r, tmp / "vr")
    kunci = next(k for k in v if "rows with unreadable characters (replaced)" in k)
    assert v[kunci].startswith("2;"), v[kunci]
    assert "501" in v[kunci] and "502" in v[kunci]
    assert v["Total Profit ('out' rows)"] == 30
    bersih_v = _verifikasi(out_b, tmp / "vb")
    assert not any("unreadable" in str(k) for k in bersih_v)
print("OK utf-16 surrogate in Name, including the read-buffer boundary")

with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    path = tmp / "profit.csv"
    head = _baris(*HEADER)
    # Surrogate replaces the profit amount. A following digit keeps it off EOF.
    baris = _u16("1,501,Client,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,") + SURROGATE + _u16("0,USD\n")
    _tulis(path, b"\xff\xfe" + head + baris)
    pesan = _berhenti(lambda: proses([path], tmp / "out.zip"))
    assert pesan.startswith("STOP:"), pesan
    assert "Profit" in pesan and "row" in pesan
    assert "Traceback" not in pesan
print("OK utf-16 surrogate in Profit stops cleanly")

with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    head = b"Deal,Login,Name,Time,Type,Entry,Symbol,Volume,Commission,Fee,Swap,Profit,Currency\n"
    baik = (head
            + b"1,501,Client Name,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,10,USD\n")
    rusak = (head
             + b"1,501,Client \xff Name,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,10,USD\n")
    (tmp / "baik.csv").write_bytes(baik)
    (tmp / "rusak.csv").write_bytes(rusak)
    proses([tmp / "baik.csv"], tmp / "baik.zip")
    proses([tmp / "rusak.csv"], tmp / "rusak.zip")
    assert _harian(tmp / "rusak.zip", tmp / "r") == _harian(tmp / "baik.zip", tmp / "b")
    v = _verifikasi(tmp / "rusak.zip", tmp / "vr")
    kunci = next(k for k in v if "rows with unreadable characters (replaced)" in k)
    assert v[kunci].startswith("1;"), v[kunci]
    assert "501" in v[kunci]
    jelek = (head
             + b"1,501,Client,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,\xff,USD\n")
    (tmp / "profit.csv").write_bytes(jelek)
    pesan = _berhenti(lambda: proses([tmp / "profit.csv"], tmp / "profit.zip"))
    assert pesan.startswith("STOP:") and "Profit" in pesan and "Traceback" not in pesan, pesan
print("OK utf-8 invalid bytes in Name are counted, and in Profit they stop")

# Header-first MT4: a cp1252 comment byte must not crash, and a UTF-16
# surrogate in Comment is reported without changing the trade totals.
def _peta_kosong():
    """Header-only account-type file: MT4 requires one; every login is then USD."""
    import tempfile as _tf
    from pathlib import Path as _P
    p = _P(_tf.mkdtemp()) / "account_types.csv"
    p.write_text("account,account_type\n", encoding="utf-8")
    return p


with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    header = ("Deal,Login,Open Time,Type,Symbol,Volume,Open Price,S/L,T/P,"
              "Close Time,Close Price,Reason,Gateway Order,Gateway Volume,"
              "Open Price Delta,Close Price Delta,Agent,Commission,Taxes,Swap,"
              "Profit,Points,Comment\n")
    baris = ("9,700,2026.07.01 00:00:00,buy,EURUSD,1,1,0,0,"
             "2026.07.02 00:00:00,1,0,0,0,0,0,0,-1,0,0,15,0,")
    (tmp / "cp.csv").write_bytes((header + baris + "catatan\xd0\n").encode("latin1"))
    out = tmp / "cp.zip"
    ringkas = proses([tmp / "cp.csv"], out, allow_mt4=True, mt4_accounts=_peta_kosong())
    assert ringkas["mt4_dipakai"] == 1, ringkas
    v = _verifikasi(out, tmp / "cpv")
    assert v["MT4 Total Profit (closed buy/sell)"] == 15

    blob = b"\xff\xfe" + _u16(header) + _u16(baris) + SURROGATE + _u16("Z\n")
    (tmp / "sur.csv").write_bytes(blob)
    out2 = tmp / "sur.zip"
    proses([tmp / "sur.csv"], out2, allow_mt4=True, mt4_accounts=_peta_kosong())
    v2 = _verifikasi(out2, tmp / "surv")
    assert v2["MT4 Total Profit (closed buy/sell)"] == 15
    kunci = next(k for k in v2 if "rows with unreadable characters (replaced)" in k)
    assert "700" in v2[kunci]
print("OK MT4 comment bytes do not crash")

with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    # Invalid UTF-8 in an unused column is read (cp1252 or replacement) and
    # does not hide the account. A surrogate in the account id stops cleanly.
    (tmp / "akun.csv").write_bytes(
        b"account,account_type,note\n100,CentAccount,ok\xff\n")
    (tmp / "deals.csv").write_bytes(
        b"Deal,Login,Time,Type,Entry,Symbol,Volume,Commission,Fee,Swap,Profit,Currency\n"
        b"1,100,2026.07.01 00:00:00,buy,out,EURUSD,1,-1,0,0,5,USD\n")
    # DPM refuses MT4, and this deals file is MT5. The account file is only
    # consulted with --allow-mt4, including when it was passed as --mt4-accounts.
    ringkas = proses([tmp / "deals.csv"], tmp / "out.zip", allow_mt4=True,
                     mt4_accounts=str(tmp / "akun.csv"))
    assert ringkas["dipakai"] == 1, ringkas

    peta = b"\xff\xfe" + _u16("account,account_type\n") + SURROGATE + _u16("1,CentAccount\n")
    (tmp / "peta.csv").write_bytes(peta)
    pesan = _berhenti(lambda: proses(
        [tmp / "deals.csv"], tmp / "out2.zip", allow_mt4=True,
        mt4_accounts=str(tmp / "peta.csv")))
    assert pesan.startswith("STOP:") and "account" in pesan and "Traceback" not in pesan, pesan
print("OK account-type text does not crash, and a damaged account id stops")

# An unexpected exception is a STOP line naming the file and row. The traceback
# stays on stderr after that line, for the server log.
with TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    src = tmp / "deals.csv"
    src.write_text("Deal,Login\n", encoding="utf-8")
    asli = deal_segregator.proses

    def meledak(*_args, **_kwargs):
        deal_segregator._lokasi["path"] = str(src)
        deal_segregator._lokasi["row"] = 9
        raise RuntimeError("sensor failed")

    deal_segregator.proses = meledak
    lama = sys.argv
    sys.argv = ["deal_segregator.py", str(src), "-o", str(tmp / "out.zip")]
    galat = io.StringIO()
    try:
        with contextlib.redirect_stderr(galat):
            kode = _berhenti(deal_segregator.main)
    finally:
        deal_segregator.proses = asli
        sys.argv = lama
    assert kode == "1", kode
    teks = galat.getvalue()
    assert teks.startswith("STOP: RuntimeError"), teks
    assert src.name in teks and "row 9" in teks and "sensor failed" in teks
    assert "Traceback (most recent call last):" in teks
    assert teks.index("STOP:") < teks.index("Traceback")
print("OK unexpected errors become STOP with file and row")

import os
import tempfile
_akar = Path(tempfile.mkdtemp())
os.environ.setdefault("DW_AUTH_DB", str(_akar / "auth.db"))
os.environ.setdefault("DW_JOB_DIR", str(_akar / "jobs"))
from app import _pesan_gagal_job, _pisah_traceback

jejak = (
    "STOP: RuntimeError in 'deals.csv', row 9: sensor failed\n"
    "Traceback (most recent call last):\n"
    "  File \"deal_segregator.py\", line 1, in proses\n"
    "RuntimeError: sensor failed\n"
)
pengguna, server = _pisah_traceback(jejak)
assert pengguna.startswith("STOP:") and "Traceback" not in pengguna
assert server.startswith("Traceback")
assert _pesan_gagal_job(jejak, "").startswith("STOP: RuntimeError")
assert _pesan_gagal_job("Traceback (most recent call last):\nValueError: x\n", "") == "ValueError: x"
print("OK web job hides the traceback from the user and keeps the STOP line")
