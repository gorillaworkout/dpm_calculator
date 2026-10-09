"""Downloadable Deal Segregator templates must match the parser and run cleanly.

python3 test_segregator_templates.py
"""
import base64
import io
import os
import tempfile
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(tempfile.mkdtemp())
os.environ["DW_AUTH_DB"] = str(ROOT / "users.db")
os.environ["DW_JOB_DIR"] = str(ROOT / "jobs")
os.environ["DW_ADMIN_USER"] = "template-tester"
os.environ["DW_ADMIN_PASS"] = "template-tester-password-1"

from deal_segregator import (  # noqa: E402
    FX_SHEET, KOLOM_WAJIB, _KOLOM_PETA, _MT4_WAJIB, baca_fx, baca_peta_akun, proses,
)
from segregator_templates import NAMA_BERKAS  # noqa: E402

MT5_HEADER = [
    "Deal", "ID", "Order", "Position", "Login", "Name", "Group", "Country",
    "Time", "Type", "Entry", "Symbol", "Volume", "Price", "S / L", "T / P",
    "Reason", "Commission", "Fee", "Swap", "Profit", "Dealer", "Currency",
    "Comment",
]
TEKS_MT5 = ("Deal", "ID", "Order", "Position", "Login")

import app as A  # noqa: E402

A.app.config["TESTING"] = True
client = A.app.test_client()
client.environ_base["HTTP_AUTHORIZATION"] = "Basic " + base64.b64encode(
    b"template-tester:template-tester-password-1").decode()


def sheet_header(ws):
    return [cell.value for cell in next(ws.iter_rows(max_row=1))]


def guide_marks(book):
    assert "Guide" in book.sheetnames
    marks = {}
    for row in book["Guide"].iter_rows(values_only=True):
        if not row or not row[0] or row[0] in ("Column", "Format", "Layout", "Sheet"):
            continue
        marks[row[0]] = row[1]
    return marks


def verification(path):
    with zipfile.ZipFile(path) as archive:
        blob = archive.read("Deals - Daily.xlsx")
    book = load_workbook(io.BytesIO(blob), data_only=True)
    try:
        rows = list(book["Verifikasi"].iter_rows(values_only=True))
    finally:
        book.close()
    return {row[0]: row[1] for row in rows if row and row[0]}


anon = A.app.test_client()
assert anon.get("/tool/segregate/template/mt5").status_code == 401
assert anon.get("/kvb/tool/segregate/template/mt4").status_code == 401
print("OK template download requires a login")

assert client.get("/tool/segregate/template/mt4").status_code == 404
assert client.get("/tool/segregate/template/accounts").status_code == 404
assert client.get("/kvb/tool/segregate/template/nope").status_code == 404
print("OK DPM does not serve the MT4 templates")

dpm = client.get("/tool/segregate")
assert dpm.status_code == 200
assert b"Templates &amp; format guide" in dpm.data
assert b'href="/tool/segregate/template/mt5"' in dpm.data
assert b'href="/tool/segregate/template/fx"' in dpm.data
assert b'href="/tool/segregate/template/mt4"' not in dpm.data
assert b'href="/tool/segregate/template/accounts"' not in dpm.data
assert b"Upload the original MT5/MT4 export (CSV) directly, or the template." in dpm.data
assert b"Don't open and re-save a CSV in Excel." in dpm.data
assert b"Deal and Login must be full numbers, not 1.41E+08." in dpm.data
assert b"MT4 and account-type files are optional for KVB." in dpm.data
assert b"DPM is MT5 only." in dpm.data
assert b"Download the template" not in dpm.data
kvb = client.get("/kvb/tool/segregate")
assert kvb.status_code == 200
for href in ("/kvb/tool/segregate/template/mt5",
             "/kvb/tool/segregate/template/mt4",
             "/kvb/tool/segregate/template/accounts",
             "/kvb/tool/segregate/template/fx"):
    assert f'href="{href}"'.encode() in kvb.data, href
assert b"DPM is MT5 only." in kvb.data
print("OK both Deal Segregator pages link the templates they can use")

blobs = {}
for company, names in (("dpm", ("mt5", "fx")), ("kvb", ("mt5", "mt4", "accounts", "fx"))):
    prefix = "" if company == "dpm" else "/kvb"
    for name in names:
        response = client.get(f"{prefix}/tool/segregate/template/{name}")
        assert response.status_code == 200, (company, name, response.status_code)
        assert NAMA_BERKAS[name] in response.headers["Content-Disposition"]
        blobs[name] = response.data
print("OK every template downloads")

mt5_book = load_workbook(io.BytesIO(blobs["mt5"]))
assert sheet_header(mt5_book["Deals History"]) == MT5_HEADER
assert mt5_book.sheetnames[0] == "Deals History"
for column in KOLOM_WAJIB:
    assert column in MT5_HEADER
marks = guide_marks(mt5_book)
for column in KOLOM_WAJIB:
    assert marks[column] == "required", column
for column in MT5_HEADER:
    if column not in KOLOM_WAJIB:
        assert marks[column] == "optional", column
ws = mt5_book["Deals History"]
for name in TEKS_MT5:
    col = MT5_HEADER.index(name) + 1
    for excel_row in range(2, 5):
        cell = ws.cell(excel_row, col)
        assert cell.number_format == "@", (name, cell.number_format)
        assert isinstance(cell.value, str) and "E+" not in cell.value.upper()
mt5_book.close()
print("OK MT5 template header, guide, and text ids")

mt4_book = load_workbook(io.BytesIO(blobs["mt4"]))
mt4_header = sheet_header(mt4_book["Raw Report"])
for column in _MT4_WAJIB:
    assert column in mt4_header, column
assert mt4_header[-1] == "Currency"
marks = guide_marks(mt4_book)
for column in _MT4_WAJIB:
    assert marks[column] == "required", column
assert marks["Currency"] == "optional"
deal_col = mt4_header.index("Deal") + 1
login_col = mt4_header.index("Login") + 1
for excel_row in range(2, 5):
    for col in (deal_col, login_col):
        cell = mt4_book["Raw Report"].cell(excel_row, col)
        assert cell.number_format == "@" and isinstance(cell.value, str)
mt4_book.close()
print("OK MT4 template header and guide")

akun_book = load_workbook(io.BytesIO(blobs["accounts"]))
assert sheet_header(akun_book["Accounts"]) == list(_KOLOM_PETA)
types = [akun_book["Accounts"].cell(row, 2).value for row in range(2, 6)]
assert types == ["Classic", "CentAccount", "Plus", "Pro"]
marks = guide_marks(akun_book)
assert marks["account"] == "required" and marks["account_type"] == "required"
assert akun_book["Accounts"].cell(2, 1).number_format == "@"
akun_book.close()
print("OK account-type template")

fx_book = load_workbook(io.BytesIO(blobs["fx"]))
assert fx_book.sheetnames[0] == FX_SHEET
assert sheet_header(fx_book[FX_SHEET]) == ["date", "Currency", "rate"]
marks = guide_marks(fx_book)
assert marks["date"] == "required"
assert marks["Currency"] == "required"
assert marks["rate"] == "required"
assert FX_SHEET in "".join(str(c) for row in fx_book["Guide"].iter_rows(values_only=True) for c in row)
fx_book.close()
print("OK Client Equity FX template")

folder = ROOT / "runs"
folder.mkdir()
paths = {}
for name, blob in blobs.items():
    path = folder / NAMA_BERKAS[name]
    path.write_bytes(blob)
    paths[name] = path

fx = baca_fx(paths["fx"])
assert len(fx) == 3
akun = baca_peta_akun(paths["accounts"])
assert akun["100002"] == "USC"
assert akun["100001"] == "USD" and akun["100003"] == "USD" and akun["100004"] == "USD"

mt5_out = folder / "mt5.zip"
proses([paths["mt5"]], mt5_out)
summary = verification(mt5_out)
assert summary["Unique 'out' rows kept"] == 2
assert "rows repaired (line break inside a field)" not in "".join(map(str, summary))
both = folder / "mt5-fx.zip"
proses([paths["mt5"], paths["fx"]], both)
both_summary = verification(both)
assert "Client Equity FX" in "".join(both_summary)
assert both_summary["Unique 'out' rows kept"] == 2
mt4_out = folder / "mt4.zip"
proses([paths["mt4"]], mt4_out, allow_mt4=True, mt4_accounts=paths["accounts"])
mt4_summary = verification(mt4_out)
assert mt4_summary["MT4 closed buy/sell kept"] == 2
assert "1 CentAccount" in mt4_summary[f"  - {paths['accounts'].name} (MT4 account types)"]
print("OK each template passes through the parser")


def finish(response):
    assert response.status_code == 303, (response.status_code, response.data[:400])
    job_id = response.headers["Location"].rstrip("/").rsplit("/", 1)[-1]
    state = None
    for _ in range(200):
        state = A._read_state(job_id)
        if state and state["state"] in ("done", "failed"):
            break
        time.sleep(0.05)
    assert state and state["state"] == "done", state
    return job_id


with paths["mt5"].open("rb") as mt5_file, paths["fx"].open("rb") as fx_file:
    posted = client.post("/tool/segregate", data={
        "file": [(mt5_file, paths["mt5"].name), (fx_file, paths["fx"].name)],
    }, content_type="multipart/form-data")
finish(posted)
with paths["mt5"].open("rb") as mt5_file, paths["mt4"].open("rb") as mt4_file, \
        paths["fx"].open("rb") as fx_file, paths["accounts"].open("rb") as akun_file:
    posted = client.post("/kvb/tool/segregate", data={
        "file": [
            (mt5_file, paths["mt5"].name),
            (mt4_file, paths["mt4"].name),
            (fx_file, paths["fx"].name),
        ],
        "mt4_accounts": (akun_file, paths["accounts"].name),
    }, content_type="multipart/form-data")
finish(posted)
print("OK downloaded templates upload and the run succeeds")
print("OK templates")
