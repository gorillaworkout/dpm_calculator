#!/usr/bin/env python3
"""Cek integritas sheet 'D&W Report' hasil hitung_dw.py (layout lama DAN layout
KVB --settlement-columns v1.30).

Semua kolom hitungan D&W Report ditulis sebagai NILAI (bukan rumus Excel). Skrip
ini memastikan:
  1. tidak ada rumus di workbook yang menunjuk ke 'D&W Report' (kalau ada, setiap
     referensi dilaporkan dengan judul kolom yang ditunjuk), dan D&W Report sendiri
     tidak berisi rumus;
  2. setiap "rumus" yang diwujudkan sebagai nilai cocok dengan kolom yang BENAR
     menurut judulnya (bukan posisinya): FE Gain/Loss, CRM average (basis lama &
     Settlement Date), Xero average, Rounding Adjustment, Grand Total (jumlah /
     rasio), Recorded in Xero;
  3. format angka (rate = 5 desimal) dan warna kolom data / hitungan mengikuti
     judul kolom.

Pakai:  python3 verifikasi/cek_dw_report.py hasil.xlsx [hasil2.xlsx ...]
Import: from verifikasi.cek_dw_report import periksa   -> list masalah (kosong = OK)
"""
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries

REPORT_SHEET = "D&W Report"
FILL_DATA, FILL_CALC_HDR, FILL_CALC = "ACB9CA", "A9D08E", "E2EFDA"
FMT_ACC5_MARK = "0.00000"

_SUFFIX = re.compile(r"\s*\((paid|completed) date\)$")


def semantik(judul):
    """Judul kolom D&W Report -> nama semantik (None kalau tidak dikenal)."""
    if judul is None:
        return None
    h = " ".join(str(judul).split()).lower()
    settle = h.endswith("(settlement date)")
    h = _SUFFIX.sub("", h.replace(" (settlement date)", ""))
    peta = {"currency": "cur", "payment gateway": "gw", "crm usd": "crm",
            "transactions": "tx", "xero usd": "xu", "total fe gain/-loss (usd)": "fe",
            "crm average rate": "avg", "crm average": "avg",
            "xero average rate": "xavg", "xero average": "xavg",
            "recorded in xero": "rec", "rounding adjustment": "round",
            "handling charges income": "ch"}
    s = peta.get(h)
    if s and settle:
        s += "_s" if s in ("crm", "tx", "avg") else "?"
    return s


DIHITUNG = {"fe", "avg", "avg_s", "xavg", "rec", "round"}
RATE = {"avg", "avg_s", "xavg"}
JUMLAH = {"crm", "crm_s", "tx", "tx_s", "xu", "ch", "fe", "round"}


def _bagian(rows):
    """[(jenis, idx_header(0-based), header, [(idx, row) data], (idx, row) grand total)]"""
    out, cur = [], None
    for i, r in enumerate(rows):
        if len(r) > 1 and r[0] == "Currency" and r[1] == "Payment Gateway":
            cur = ["deposit" if not out else "withdrawal", i, list(r), [], None]
            out.append(cur)
            continue
        if cur is not None and cur[4] is None:
            if r and r[0] == "Grand Total":
                cur[4] = (i, list(r))
            elif r and any(isinstance(v, (int, float)) and not isinstance(v, bool)
                           for v in r[2:]):
                cur[3].append((i, list(r)))      # termasuk baris Payment Gateway kosong
    return out


def _referensi(formula, sheet_formula):
    """[(sheet, kolom_awal, kolom_akhir, baris_awal, baris_akhir)] dari teks rumus."""
    hasil = []
    for m in re.finditer(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_\.]*))!"
                         r"(\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?)", formula):
        sh = (m.group(1) or m.group(2)).replace("''", "'")
        hasil.append((sh, m.group(3)))
    if sheet_formula == REPORT_SHEET:            # referensi tanpa nama sheet
        tanpa = re.sub(r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_\.]*)!"
                       r"\$?[A-Z]{1,3}\$?\d*(?::\$?[A-Z]{1,3}\$?\d*)?", "", formula)
        tanpa = re.sub(r'"[^"]*"', "", tanpa)
        for m in re.finditer(r"(?<![A-Za-z0-9_])(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"
                             r"(?![A-Za-z0-9_(])", tanpa):
            hasil.append((REPORT_SHEET, m.group(1)))
    return hasil


def _judul_kolom(bag, baris0, kolom1):
    """Judul (dan jenis tabel) yang ditunjuk sel D&W Report (baris 0-based, kolom 1-based)."""
    pilih = None
    for b in bag:
        if baris0 >= b[1]:
            pilih = b
    if pilih is None:
        return None, None
    h = pilih[2]
    return pilih[0], (h[kolom1 - 1] if kolom1 - 1 < len(h) else None)


def periksa(sumber, layout=None):
    """sumber: path atau Workbook. layout: None / 'old' / 'settlement' (opsional:
    pastikan layout memang itu). Return list teks masalah."""
    wb = sumber if not isinstance(sumber, (str, bytes)) and hasattr(sumber, "sheetnames") \
        else openpyxl.load_workbook(sumber, read_only=True)
    masalah = []
    if REPORT_SHEET not in wb.sheetnames:
        return [f"sheet '{REPORT_SHEET}' tidak ada"]
    ws = wb[REPORT_SHEET]
    sel = [list(r) for r in ws.iter_rows()]
    rows = [[c.value for c in r] for r in sel]
    bag = _bagian(rows)
    if [b[0] for b in bag] != ["deposit", "withdrawal"]:
        return [f"D&W Report: bagian deposit/withdrawal tidak ditemukan ({[b[0] for b in bag]})"]

    # ---- 1. rumus yang menunjuk D&W Report
    for w in wb.worksheets:
        for r in w.iter_rows():
            for c in r:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                if w.title == REPORT_SHEET:
                    masalah.append(f"D&W Report!{c.coordinate}: berisi rumus {v!r} "
                                   f"(D&W Report harus nilai statis)")
                for sh, ref in _referensi(v, w.title):
                    if sh != REPORT_SHEET:
                        continue
                    a, b = (ref.replace("$", "").split(":") + [None])[:2]
                    m = re.match(r"([A-Z]+)(\d*)", a)
                    kol = column_index_from_string(m.group(1))
                    brs = int(m.group(2)) - 1 if m.group(2) else 0
                    jenis, judul = _judul_kolom(bag, brs, kol)
                    masalah.append(f"{w.title}!{c.coordinate}: rumus {v!r} menunjuk "
                                   f"'D&W Report'!{ref} = {jenis} kolom {judul!r} -- periksa "
                                   f"setelah kolom bergeser")

    # ---- 2/3. nilai hitungan per judul kolom
    def dekat(a, b):
        if a is None or b is None:
            return a is None and b is None
        return abs(a - b) <= 1e-6 * max(1.0, abs(a), abs(b))

    for jenis, i_hdr, hdr, data, gt in bag:
        sem = {}
        for j, h in enumerate(hdr):
            s = semantik(h)
            if s is None:
                if h not in (None, ""):
                    masalah.append(f"{jenis}: judul kolom tak dikenal {h!r}")
                continue
            if s in sem:
                masalah.append(f"{jenis}: judul kolom ganda {h!r}")
            sem[s] = j
        wajib = {"cur", "gw", "crm", "tx", "xu", "fe", "avg", "xavg", "rec", "round"}
        if jenis == "withdrawal":
            wajib.add("ch")
        punya_s = {"crm_s", "tx_s", "avg_s"} & set(sem)
        if punya_s and punya_s != {"crm_s", "tx_s", "avg_s"}:
            masalah.append(f"{jenis}: kolom Settlement Date tidak lengkap: {sorted(punya_s)}")
        if layout == "settlement" and not punya_s:
            masalah.append(f"{jenis}: kolom Settlement Date tidak ada")
        if layout == "old" and punya_s:
            masalah.append(f"{jenis}: kolom Settlement Date muncul di layout lama")
        if wajib - set(sem):
            masalah.append(f"{jenis}: kolom hilang {sorted(wajib - set(sem))}")
            continue

        def g(r, s):
            return r[sem[s]] if s in sem and sem[s] < len(r) else None

        tanda = 1 if jenis == "deposit" else -1   # FE: deposit Xero-CRM, withdrawal CRM-Xero

        def cek_baris(lbl, r):
            crm, tx, xu = g(r, "crm") or 0.0, g(r, "tx") or 0.0, g(r, "xu") or 0.0
            harap = {"fe": tanda * (xu - crm), "round": -tanda * xu,
                     "avg": tx / crm if crm else None, "xavg": tx / xu if xu else None}
            if punya_s:
                cs, ts = g(r, "crm_s") or 0.0, g(r, "tx_s") or 0.0
                harap["avg_s"] = ts / cs if cs else None
            for s, e in harap.items():
                if not dekat(g(r, s), e):
                    masalah.append(f"{jenis} {lbl}: kolom {hdr[sem[s]]!r} = {g(r, s)!r}, "
                                   f"seharusnya {e!r} (dari kolom menurut judulnya)")

        for i, r in data:
            cek_baris(f"baris {i + 1} ({r[0]}/{r[1]})", r)
        if gt is None:
            masalah.append(f"{jenis}: Grand Total tidak ada")
            continue
        i_gt, rg = gt
        cek_baris(f"Grand Total baris {i_gt + 1}", rg)
        for s in JUMLAH & set(sem):
            tot = sum((g(r, s) or 0.0) for _i, r in data)
            if not dekat(g(rg, s) or 0.0, tot):
                masalah.append(f"{jenis} Grand Total {hdr[sem[s]]!r} = {g(rg, s)!r}, "
                               f"jumlah baris = {tot!r}")
        isi_rec = [g(r, "rec") for _i, r in data if isinstance(g(r, "rec"), (int, float))]
        if not dekat(g(rg, "rec"), sum(isi_rec) if isi_rec else None):
            masalah.append(f"{jenis} Grand Total 'Recorded in Xero' tidak sama dengan jumlahnya")

        # format & warna menurut judul
        hdr_cells = sel[i_hdr]
        for s, j in sem.items():
            if s in ("cur", "gw"):
                continue
            want_hdr = FILL_CALC_HDR if s in DIHITUNG else FILL_DATA
            got = (hdr_cells[j].fill.fgColor.rgb or "")[-6:] if hdr_cells[j].fill else ""
            if got != want_hdr:
                masalah.append(f"{jenis}: warna judul {hdr[j]!r} = {got}, seharusnya {want_hdr}")
            for i, _r in data + [gt]:
                c = sel[i][j] if j < len(sel[i]) else None
                if c is None:
                    continue
                fmt5 = FMT_ACC5_MARK in (c.number_format or "")
                if fmt5 != (s in RATE):
                    masalah.append(f"{jenis} {c.coordinate}: format angka {hdr[j]!r} salah "
                                   f"({c.number_format!r})")
                if i != i_gt:
                    hijau = (c.fill.fgColor.rgb or "")[-6:] == FILL_CALC if c.fill and \
                        c.fill.fill_type == "solid" else False
                    if hijau != (s in DIHITUNG):
                        masalah.append(f"{jenis} {c.coordinate}: warna sel {hdr[j]!r} salah")
    return masalah


def main():
    gagal = False
    for p in sys.argv[1:]:
        m = periksa(p)
        print(f"{p}: {'OK' if not m else f'{len(m)} masalah'}")
        for x in m[:200]:
            print("  -", x)
        gagal |= bool(m)
    sys.exit(1 if gagal else 0)


if __name__ == "__main__":
    main()
