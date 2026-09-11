#!/usr/bin/env python3
"""Bandingkan sheet 'Channel Balance' hasil vs blok 'Payment Channel Balance' sumber.

Pakai:  python3 banding_pcb.py "file-hasil.xlsx" "workbook-sumber.xlsx"
"""
import os, sys
HASIL = sys.argv[1] if len(sys.argv) > 1 else "../D&W JUN 2026-hasil.xlsx"
SUMBER = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser(
    "~/Downloads/22.08 Bayu - D&W-Dupoin Markets-JUN 2026v3.xlsx")

import openpyxl, importlib.util
spec=importlib.util.spec_from_file_location('hd', os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','hitung_dw.py'))
hd=importlib.util.module_from_spec(spec); spec.loader.exec_module(hd)

src=openpyxl.load_workbook(SUMBER,read_only=True,data_only=True)
ws=src['Payment Channel Balance']; data=[list(r) for r in ws.iter_rows(values_only=True)]
r4,r5=data[3],data[4]
def peta(a,b):
    cur=None; out={}
    for j in range(a,b+1):
        if j-1<len(r4) and r4[j-1] not in (None,''): cur=str(r4[j-1]).strip()
        if j-1<len(r5) and r5[j-1] not in (None,'') and cur: out[(cur,str(r5[j-1]).strip())]=j-1
    return out
BLOK={'Deposit':(2,69),'Deposit charges':(70,138),'Fund transfer':(139,214),
      'Fund transfer charges':(215,283),'Withdrawal':(284,357),
      'Withdrawal charges':(358,431),'Balance':(432,498)}
row=next(r for r in data[5:] if hasattr(r[0],'year') and str(r[0])[:10]=='2026-06-01')
EXTRA={'SHUNFAPA':'SHUNFAPAY'}
def key(cur,ch):
    k=hd.kunci_channel(cur,ch)
    return (hd.norm(cur), EXTRA.get(k,k))
mereka={}
for nama,(a,b) in BLOK.items():
    for (cur,ch),j in peta(a,b).items():
        v=row[j] if j<len(row) else None
        if isinstance(v,(int,float)):
            mereka.setdefault(key(cur,ch),{})[nama]=float(v)

out=openpyxl.load_workbook(HASIL,read_only=True,data_only=True)
cb = out['Payment Channel Balance'] if 'Payment Channel Balance' in out.sheetnames \
     else out['Channel Balance']   # nama sheet hasil diganti 28 Aug 2026
hdr=[c for c in next(cb.iter_rows(max_row=1,values_only=True))]
KOL=['Deposit','Deposit charges','Fund transfer','Fund transfer charges','Withdrawal','Withdrawal charges','Balance']
print(f"{'channel':22} {'kolom':22} {'KITA':>18} {'MEREKA':>18}   selisih")
n_ok=n_beda=0
for r in cb.iter_rows(min_row=2,values_only=True):
    if not r[1]: continue
    k=key(r[1],r[2])
    m=mereka.get(k)
    if not m:
        print(f"{r[1]+'/'+str(r[2]):22} {'(tidak ada di sheet mereka)':22}")
        continue
    for nm in KOL:
        kita=r[hdr.index(nm)]
        kita=0.0 if kita is None else float(kita)
        thm=m.get(nm)
        if thm is None: continue
        if abs(kita-thm)<0.02: n_ok+=1
        else:
            n_beda+=1
            print(f"{r[1]+'/'+str(r[2]):22} {nm:22} {kita:>18,.2f} {thm:>18,.2f}   {kita-thm:>+15,.2f}")
print(f"\ncocok {n_ok}, beda {n_beda}")
