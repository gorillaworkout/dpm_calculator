#!/usr/bin/env python3
"""Tambahkan baris rate baru ke sheet 'D&W FEE' sebuah workbook.

Pakai:
    python3 tambah_rate.py "DPM-D&W.xlsx" VND BeckPay 1.6% 0%
    python3 tambah_rate.py "DPM-D&W.xlsx" THB Other 0 0

File harus DITUTUP dulu dari Excel. Backup otomatis dibuat sebelum menulis.
Kalau kombinasi Currency + Payment Gateway sudah ada, barisnya di-update.
"""
import shutil
import sys
from pathlib import Path

import openpyxl

FEE_SHEET, HEADER_ROW = "D&W FEE", 3


def pct(v):
    v = str(v).strip()
    return float(v.rstrip("%")) / 100 if v.endswith("%") else float(v)


def main():
    if len(sys.argv) != 6:
        sys.exit(__doc__)
    path = Path(sys.argv[1]).expanduser()
    cur, gw = sys.argv[2].strip(), sys.argv[3].strip()
    dep, wd = pct(sys.argv[4]), pct(sys.argv[5])

    if not path.exists():
        sys.exit(f"File tidak ditemukan: {path}")
    backup = path.with_name(f"{path.stem}-backup{path.suffix}")
    shutil.copy2(path, backup)
    print(f"Backup : {backup}")

    wb = openpyxl.load_workbook(path)
    ws = wb[FEE_SHEET]

    target = None
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        a, b = ws.cell(row, 1).value, ws.cell(row, 2).value
        if a and b and str(a).strip().upper() == cur.upper() and str(b).strip().upper() == gw.upper():
            target = row
            print(f"Update : baris {row} (sudah ada) {a} / {b}")
            break
    if target is None:
        target = ws.max_row + 1
        while ws.cell(target - 1, 1).value in (None, "") and target - 1 > HEADER_ROW:
            target -= 1
        print(f"Tambah : baris baru {target}")

    ws.cell(target, 1, cur)
    ws.cell(target, 2, gw)
    ws.cell(target, 3, dep).number_format = "0.00%"
    ws.cell(target, 4, wd).number_format = "0.00%"

    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Gagal simpan: '{path}' sedang dibuka di Excel. Tutup dulu, jalankan ulang.")
    print(f"Simpan : {path}")
    print(f"         {cur} / {gw}  deposit {dep*100:g}%  withdrawal {wd*100:g}%")


if __name__ == "__main__":
    main()
