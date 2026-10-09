"""MT4 Raw Report support for Deal Segregator. python3 test_deal_segregator_mt4.py"""
import base64
import csv
import datetime
import io
import os
import tempfile
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(tempfile.mkdtemp())
os.environ["DW_AUTH_DB"] = str(ROOT / "users.db")
os.environ["DW_JOB_DIR"] = str(ROOT / "jobs")
os.environ["DW_ADMIN_USER"] = "mt4-tester"
os.environ["DW_ADMIN_PASS"] = "mt4-tester-password-1"

from deal_segregator import proses  # noqa: E402

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "kvb"
MT5_FIXTURE = FIXTURE / "KVB_MT5_Deals_2026-10-04.csv"
MT4_FIXTURE = FIXTURE / "KVB_MT4_RawReport_2026-10-04.csv"

MT5_HEADERS = [
    "Deal", "Login", "Time", "Type", "Entry", "Symbol", "Volume",
    "Commission", "Fee", "Swap", "Profit", "Currency",
]
MT4_HEADER = [
    "Deal", "Login", "Open Time", "Type", "Symbol", "Volume", "Open Price",
    "S/L", "T/P", "Close Time", "Close Price", "Reason", "Gateway Order",
    "Gateway Volume", "Open Price Delta", "Close Price Delta", "Agent",
    "Commission", "Taxes", "Swap", "Profit", "Points", "Comment",
]
COMBINED = [
    "Platform", "Login", "Date", "Type", "Symbol", "Deals", "Volume",
    "Commission", "Fee", "Swap", "Profit", "Currency",
]

K_BAL = "MT4 balance/credit rows dropped (same rule as MT5 non-trade Entry)"
K_CAN = "MT4 cancelled pending orders dropped"
K_FOOT = "MT4 footer/summary lines dropped"
K_OUT = "MT4 closed buy/sell kept"
K_DUP = "MT4 duplicate Deal IDs dropped"
K_OTHER = "MT4 other rows dropped"
K_PROFIT = "MT4 Total Profit (closed buy/sell)"
K_COMM = "MT4 Total Commission (closed buy/sell)"
K_FEE = "MT4 Total Fee (Taxes, closed buy/sell)"
K_AGENT = "MT4 Agent column (sum of kept rows; not included in Profit, Commission, or Fee)"
K_DAILY = "Daily 'out' rows (Platform+Login+Date+Type+Symbol+Currency)"
K_DAILY_IN = "Daily 'in' rows (Platform+Login+Date+Type+Symbol+Currency, sheet 'Daily - In')"
K_COMB_P = "Combined Total Profit (MT5 + MT4 out rows)"
K_COMB_C = "Combined Total Commission (MT5 + MT4 out rows)"
K_WHY = "MT4 currency is chosen per file"
SHEETS = ["Daily", "Daily - In", "Verifikasi"]


def write_mt5(path, rows):
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(MT5_HEADERS)
        w.writerows(rows)


def mt5_row(deal="1", login="100", time="2026.10.04 01:02:03", type_="buy",
            entry="out", symbol="EURUSD", volume="1", commission="0", fee="0",
            swap="0", profit="10", currency="USD"):
    return [deal, login, time, type_, entry, symbol, volume, commission, fee,
            swap, profit, currency]


def write_mt4(path, rows, start="2026.10.04", end="2026.10.04", footer=True):
    """Content is what identifies the file. The name is chosen by the caller."""
    with path.open("w", encoding="utf-8", newline="") as f:
        f.write(f"Raw Report for 'abcc' from {start} to {end}\n")
        w = csv.writer(f, delimiter=";")
        w.writerow(MT4_HEADER)
        for row in rows:
            w.writerow([row.get(h, "") for h in MT4_HEADER])
        if footer:
            w.writerow([""] * 16 + ["0.00", "-1.00", "0.00", "0.00", "10.00", "", ""])
            w.writerow([""] * 18 + ["Profit:", "", "10.00", "", ""])
            w.writerow([""] * 18 + ["Deposit:", "", "1.00", "", ""])
            w.writerow([""] * 18 + ["Withdrawal:", "", "-1.00", "", ""])
            w.writerow([""] * 18 + ["Credit:", "", "0.00", "", ""])


def mt4_trade(deal, login="100", open_time="2026.10.04 01:00:00", type_="buy",
              symbol="eurusd", volume="1", close_time="2026.10.04 02:00:00",
              agent="0", commission="0", taxes="0", swap="0", profit="0",
              comment=""):
    return {
        "Deal": deal, "Login": login, "Open Time": open_time, "Type": type_,
        "Symbol": symbol, "Volume": volume, "Close Time": close_time,
        "Agent": agent, "Commission": commission, "Taxes": taxes, "Swap": swap,
        "Profit": profit, "Comment": comment,
    }


def books(path):
    out = {}
    with zipfile.ZipFile(path) as z:
        assert z.namelist() == ["Deals - Daily.xlsx", "Deals - Monthly Summary.xlsx"]
        for name in z.namelist():
            wb = load_workbook(io.BytesIO(z.read(name)), data_only=True)
            out[name] = {s: [tuple(r) for r in wb[s].iter_rows(values_only=True)]
                         for s in wb.sheetnames}
            out[name]["__sheets__"] = list(wb.sheetnames)
            wb.close()
    return out


def verification(rows):
    return {r[0]: r[1] for r in rows if r and r[0]}


def run(paths, allow_mt4=False, mt4_currency=None):
    folder = Path(tempfile.mkdtemp())
    out = folder / "hasil.zip"
    summary = proses([Path(p) for p in paths], out, allow_mt4=allow_mt4,
                     mt4_currency=mt4_currency)
    return summary, books(out), out


def row_map(headers, values):
    return {headers[i]: values[i] for i in range(len(headers))}


# The attached MT5 export is UTF-16, tab-separated, and ends with a Total row
# plus per-currency summary rows (including USC). Entry is blank on those
# rows, so the existing parser drops them and keeps USC as a currency.
assert MT5_FIXTURE.is_file() and MT4_FIXTURE.is_file()
summary, hasil, _ = run([MT5_FIXTURE], allow_mt4=False)
daily = hasil["Deals - Daily.xlsx"]
assert daily["__sheets__"] == SHEETS
assert daily["Daily"][0][0] == "Login"
assert "Platform" not in daily["Daily"][0]
currencies = {r[10] for r in daily["Daily"][1:]}
assert "USC" in currencies and "USD" in currencies
v = verification(daily["Verifikasi"])
assert v["Total source rows (all files)"] == 7
assert v["Entry empty/other (dropped)"] == 3
assert v["Unique 'out' rows kept"] == 2
assert v["Unique 'in' rows kept"] == 2
assert v["Daily 'out' rows (Login+Date+Type+Symbol+Currency)"] == summary["harian"]
assert K_OUT not in v and K_WHY not in v and "Platform" not in "".join(v)
assert summary["dipakai"] == 2 and summary["dipakai_in"] == 2
assert "mt4_dipakai" not in summary
print("OK MT5 fixture already parsed (USC kept, footer dropped, no Platform column)")

# Same MT5 file, KVB flag on, must not grow sheets or change verification text.
_, dengan_flag, _ = run([MT5_FIXTURE], allow_mt4=True)
assert dengan_flag["Deals - Daily.xlsx"]["Verifikasi"] == daily["Verifikasi"]
assert dengan_flag["Deals - Daily.xlsx"]["__sheets__"] == daily["__sheets__"]
assert dengan_flag["Deals - Daily.xlsx"]["Daily"] == daily["Daily"]
assert dengan_flag["Deals - Daily.xlsx"]["Daily - In"] == daily["Daily - In"]
assert dengan_flag["Deals - Monthly Summary.xlsx"]["__sheets__"] == [
    "Monthly Summary", "Monthly Summary - In", "Verifikasi"]
print("OK MT5-only output unchanged when MT4 is allowed")

# DPM default refuses MT4 by content, including when the name looks like MT5.
with tempfile.TemporaryDirectory() as tmp:
    mt4 = Path(tmp) / "Deals History.csv"
    mt5 = Path(tmp) / "Raw Report.csv"
    write_mt4(mt4, [mt4_trade("9")])
    write_mt5(mt5, [mt5_row()])
    try:
        proses([mt4, mt5], Path(tmp) / "out.zip")
    except SystemExit as exc:
        message = str(exc)
        assert "STOP:" in message and "MT5 Deals History only" in message, message
        assert "Deals History.csv" in message and "KVB" in message, message
    else:
        raise AssertionError("DPM accepted an MT4 Raw Report")
    summary, hasil, _ = run([mt5], allow_mt4=False)
    assert hasil["Deals - Daily.xlsx"]["__sheets__"] == SHEETS
    assert hasil["Deals - Daily.xlsx"]["Daily"][1][0] == "100"
print("OK DPM rejects MT4 by content and still reads a misnamed MT5 file")

# Closed buy/sell -> out at Close Time only. Taxes land in Fee. Agent is not a
# column and is not added to Commission. Same Login+Date+Type+Symbol aggregates.
# Balance/credit, cancelled pending orders and the footer stay dropped.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    src = tmp / "notes.csv"
    write_mt4(src, [
        mt4_trade("1", login="100", open_time="2026.10.03 23:00:00",
                  close_time="2026.10.04 01:00:00", symbol="btcusd", volume="0.02",
                  commission="-0.17", taxes="-1.5", swap="-0.25", profit="3.98",
                  agent="2"),
        mt4_trade("2", login="200", open_time="2026.10.04 03:00:00",
                  close_time="2026.10.04 04:00:00", type_="sell", symbol="ethusd",
                  volume="0.5", commission="-0.4", taxes="0", swap="0", profit="-8",
                  agent="0.5"),
        mt4_trade("3", login="100", open_time="2026.10.01 00:00:00",
                  close_time="2026.10.04 05:00:00", symbol="btcusd", volume="1",
                  profit="1", agent="0"),
        {"Deal": "4", "Login": "300", "Open Time": "2026.10.04 00:05:00",
         "Type": "balance", "Symbol": "CRM-S-1,Swap", "Profit": "-2.35",
         "Comment": "CRM-S-1,Swap"},
        {"Deal": "5", "Login": "301", "Open Time": "2026.10.04 00:06:00",
         "Type": "credit", "Profit": "9", "Comment": "credit in"},
        {"Deal": "6", "Login": "302", "Open Time": "2026.10.03 01:00:00",
         "Type": "buy limit", "Symbol": "btcusd", "Volume": "0.2",
         "Close Time": "2026.10.04 02:00:00", "Comment": "cancelled"},
        {"Deal": "7", "Login": "303", "Open Time": "2026.10.04 01:00:00",
         "Type": "sell stop", "Symbol": "ethusd", "Volume": "0.1",
         "Close Time": "2026.10.04 02:00:00", "Comment": "cancelled"},
    ])
    summary, hasil, _ = run([src], allow_mt4=True)
    daily = hasil["Deals - Daily.xlsx"]
    assert daily["__sheets__"] == SHEETS
    assert "Daily - MT4" not in daily and "Daily - MT4 In" not in daily
    monthly = hasil["Deals - Monthly Summary.xlsx"]
    assert monthly["__sheets__"] == [
        "Monthly Summary", "Monthly Summary - In", "Verifikasi"]
    assert daily["Daily"][0] == tuple(COMBINED)
    assert "Agent" not in daily["Daily"][0]
    assert daily["Daily"][1:] == [
        ("MT4", "100", datetime.datetime(2026, 10, 4), "buy", "BTCUSD", 2, 1.02,
         -0.17, -1.5, -0.25, 4.98, "USD"),
        ("MT4", "200", datetime.datetime(2026, 10, 4), "sell", "ETHUSD", 1, 0.5,
         -0.4, 0, 0, -8, "USD"),
    ]
    assert daily["Daily - In"] == [tuple(COMBINED)]
    v = verification(daily["Verifikasi"])
    assert v[K_OUT] == 3
    assert v[K_BAL] == 2 and v[K_CAN] == 2 and v[K_FOOT] == 5
    assert v[K_DUP] == 0 and v[K_OTHER] == 0
    assert v[K_PROFIT] == round(3.98 + 1 - 8, 2)
    assert v[K_AGENT] == 2.5
    assert v[K_FEE] == -1.5
    assert v[K_COMM] == round(-0.17 + -0.4, 2)
    assert v[K_DAILY] == 2 and v[K_DAILY_IN] == 0
    assert v[K_COMB_P] == v[K_PROFIT]
    assert v[K_COMB_C] == v[K_COMM]
    assert "open question" not in v
    assert "cannot separate USD from USC" in v[K_WHY]
    assert "Balance and credit" in v[K_WHY]
    assert v["MT4 balance/credit Profit excluded"] == round(-2.35 + 9, 2)
    assert summary["mt4_dipakai"] == 3 and "mt4_dipakai_in" not in summary
    file_line = next(val for key, val in v.items() if "notes.csv" in key and "MT4" in key)
    assert "currency USD (default; no currency was selected for this file)" in file_line
    assert "balance/credit" in file_line and "cancelled" in file_line
    assert "opening" not in file_line
print("OK MT4 out-only rows, Agent excluded, balance rows dropped")

# A close is kept even when Open Time is outside the report period. No in row.
with tempfile.TemporaryDirectory() as tmp:
    src = Path(tmp) / "week.csv"
    write_mt4(src, [
        mt4_trade("a", login="1", open_time="2026.10.01 00:00:00",
                  close_time="2026.10.04 12:00:00", volume="1", profit="5"),
        mt4_trade("b", login="2", open_time="2026.09.30 23:59:59",
                  close_time="2026.10.04 12:00:00", volume="1", profit="7"),
    ], start="2026.10.01", end="2026.10.04", footer=False)
    _, hasil, _ = run([src], allow_mt4=True)
    daily = hasil["Deals - Daily.xlsx"]
    assert [r[1] for r in daily["Daily"][1:]] == ["1", "2"]
    assert daily["Daily - In"] == [tuple(COMBINED)]
    assert K_DAILY_IN not in verification(daily["Verifikasi"]) or \
        verification(daily["Verifikasi"])[K_DAILY_IN] == 0
print("OK closes outside the open window are still out rows")

# Duplicate Deal IDs across two MT4 files count once. Blank Login on a trade stops.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    first, second = tmp / "a.csv", tmp / "b.csv"
    write_mt4(first, [mt4_trade("same", login="5", profit="4", volume="1")], footer=False)
    write_mt4(second, [mt4_trade("same", login="5", profit="99", volume="9")], footer=False)
    _, hasil, _ = run([first, second], allow_mt4=True)
    v = verification(hasil["Deals - Daily.xlsx"]["Verifikasi"])
    assert v[K_DUP] == 1 and v[K_OUT] == 1
    assert hasil["Deals - Daily.xlsx"]["Daily"][1][10] == 4
    blank = tmp / "blank.csv"
    write_mt4(blank, [mt4_trade("z", login="", profit="1")], footer=False)
    try:
        proses([blank], tmp / "out.zip", allow_mt4=True)
    except SystemExit as exc:
        assert "blank Login" in str(exc) and "blank.csv" in str(exc)
    else:
        raise AssertionError("blank MT4 Login was accepted")
print("OK MT4 duplicate deals and blank Login")

# The same login and Deal ID on MT4 and MT5 stay two rows, grouped by Platform.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    mt5, mt4 = tmp / "mt5.csv", tmp / "mt4.csv"
    write_mt5(mt5, [mt5_row(deal="42", login="100", symbol="EURUSD", profit="10",
                            commission="-1")])
    write_mt4(mt4, [mt4_trade("42", login="100", symbol="eurusd", profit="3", volume="2",
                              agent="1.25", commission="-0.5")], footer=False)
    summary, hasil, _ = run([mt5, mt4], allow_mt4=True)
    daily = hasil["Deals - Daily.xlsx"]
    assert daily["__sheets__"] == SHEETS
    assert daily["Daily"][0] == tuple(COMBINED)
    assert "Agent" not in daily["Daily"][0]
    mt5_row_out = row_map(daily["Daily"][0], daily["Daily"][1])
    mt4_row_out = row_map(daily["Daily"][0], daily["Daily"][2])
    assert mt5_row_out["Platform"] == "MT5" and mt5_row_out["Profit"] == 10
    assert mt5_row_out["Symbol"] == "EURUSD" and mt5_row_out["Commission"] == -1
    assert mt4_row_out["Platform"] == "MT4" and mt4_row_out["Profit"] == 3
    assert mt4_row_out["Commission"] == -0.5
    assert daily["Daily - In"][0][0] == "Platform"
    v = verification(daily["Verifikasi"])
    assert v["Unique logins"] == 1 and v["Unique MT4 logins"] == 1
    assert v["Total Profit ('out' rows)"] == 10
    assert v[K_PROFIT] == 3
    assert v[K_COMB_P] == 13
    assert v[K_COMB_C] == round(-1 + -0.5, 2)
    assert v[K_AGENT] == 1.25
    assert summary["dipakai"] == 1 and summary["mt4_dipakai"] == 1
    assert summary["harian"] == 2
print("OK mixed MT4+MT5 stay apart by Platform and combine on one sheet")

# Per-file currency. MT4 USC is divided by 100 even when an FX file gives another
# USC rate. A missing USC rate is not yellow for that row. Agent is not converted.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    deals, fx = tmp / "mt4.csv", tmp / "fx.xlsx"
    write_mt4(deals, [mt4_trade(
        "1", login="8", open_time="2026.10.04 01:00:00",
        close_time="2026.10.04 09:00:00", profit="10", agent="4",
        commission="-2", taxes="-1", swap="-1", volume="1")], footer=False)
    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    ws.append([datetime.date(2026, 10, 4), "USC", 50])
    wb.save(fx)
    wb.close()
    _, hasil, out = run([deals, fx], allow_mt4=True, mt4_currency={deals.name: "USC"})
    wb = load_workbook(io.BytesIO(zipfile.ZipFile(out).read("Deals - Daily.xlsx")))
    try:
        headers = [c.value for c in wb["Daily"][1]]
        assert "Agent" not in headers and "Agent (USD)" not in headers
        assert "Profit (USD)" in headers
        row = {headers[i]: c.value for i, c in enumerate(wb["Daily"][2])}
        assert row["Currency"] == "USC"
        assert row["Platform"] == "MT4"
        assert row["Volume"] == 1
        assert row["Profit"] == 0.1 and row["Commission"] == -0.02
        assert row["Profit (USD)"] == 0.1
        assert row["Commission (USD)"] == -0.02
        assert row["Fee (USD)"] == -0.01
        assert row["Swap (USD)"] == -0.01
    finally:
        wb.close()
    v = verification(hasil["Deals - Daily.xlsx"]["Verifikasi"])
    line = next(val for key, val in v.items() if "mt4.csv" in key)
    assert "currency USC (selected for this file)" in line

    # Explicit USD is recorded as selected, and still converts with the USD rate.
    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    ws.append([datetime.date(2026, 10, 4), "USD", 2])
    usd_fx = tmp / "usd-fx.xlsx"
    wb.save(usd_fx)
    wb.close()
    _, hasil, out = run([deals, usd_fx], allow_mt4=True, mt4_currency={deals.name: "USD"})
    wb = load_workbook(io.BytesIO(zipfile.ZipFile(out).read("Deals - Daily.xlsx")))
    try:
        headers = [c.value for c in wb["Daily"][1]]
        row = {headers[i]: c.value for i, c in enumerate(wb["Daily"][2])}
        assert row["Currency"] == "USD" and row["Profit (USD)"] == 5
    finally:
        wb.close()
    chosen = verification(hasil["Deals - Daily.xlsx"]["Verifikasi"])
    assert "currency USD (selected for this file)" in next(
        val for key, val in chosen.items() if "mt4.csv" in key)

    # Blank/missing USC rate is not 1.
    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    ws.append([datetime.date(2026, 10, 4), "USD", None])
    missing = tmp / "missing-fx.xlsx"
    wb.save(missing)
    wb.close()
    _, _, out = run([deals, missing], allow_mt4=True, mt4_currency={deals.name: "USC"})
    wb = load_workbook(io.BytesIO(zipfile.ZipFile(out).read("Deals - Daily.xlsx")))
    try:
        headers = [c.value for c in wb["Daily"][1]]
        cells = list(wb["Daily"][2])
        profit_usd = cells[headers.index("Profit (USD)")]
        assert profit_usd.value == 0.1
        rgb = getattr(profit_usd.fill.fgColor, "rgb", None)
        assert not (isinstance(rgb, str) and rgb.endswith("FFFF00"))
    finally:
        wb.close()
print("OK MT4 USD/USC choice and FX")

# Two MT4 files keep their own currency, so the same login does not merge.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    usd, usc = tmp / "usd.csv", tmp / "usc.csv"
    trade = dict(login="77", symbol="btcusd", volume="1", profit="10",
                  close_time="2026.10.04 02:00:00")
    write_mt4(usd, [mt4_trade("1", **trade)], footer=False)
    write_mt4(usc, [mt4_trade("2", profit="1000", **{k: v for k, v in trade.items() if k != "profit"})],
              footer=False)
    _, hasil, _ = run([usd, usc], allow_mt4=True, mt4_currency={usc.name: "USC"})
    rows = hasil["Deals - Daily.xlsx"]["Daily"][1:]
    assert [(r[0], r[1], r[10], r[11]) for r in rows] == [
        ("MT4", "77", 10, "USD"),
        ("MT4", "77", 10, "USC"),
    ]
    v = verification(hasil["Deals - Daily.xlsx"]["Verifikasi"])
    usd_line = next(val for key, val in v.items() if key.strip().startswith("- usd.csv"))
    usc_line = next(val for key, val in v.items() if "usc.csv" in key)
    assert "currency USD (default; no currency was selected for this file)" in usd_line
    assert "currency USC (selected for this file)" in usc_line
print("OK two MT4 files do not share a currency")

# A currency meant for a non-MT4 file is ignored. A bad MT4 currency stops the run.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    mt5 = tmp / "deals.csv"
    write_mt5(mt5, [mt5_row(currency="USD")])
    _, hasil, _ = run([mt5], allow_mt4=True, mt4_currency={mt5.name: "USC"})
    assert hasil["Deals - Daily.xlsx"]["Daily"][1][10] == "USD"
    assert hasil["Deals - Daily.xlsx"]["Daily"][0][0] == "Login"
    bad = tmp / "raw.csv"
    write_mt4(bad, [mt4_trade("1", profit="1")], footer=False)
    try:
        proses([bad], tmp / "out.zip", allow_mt4=True, mt4_currency={bad.name: "EUR"})
    except SystemExit as exc:
        assert "USD or USC" in str(exc) and "EUR" in str(exc)
    else:
        raise AssertionError("EUR was accepted as an MT4 currency")
print("OK currency map ignores MT5 files and rejects a bad MT4 currency")

# Trimmed copies of the attached 2026-10-04 samples.
summary, hasil, _ = run([MT5_FIXTURE, MT4_FIXTURE], allow_mt4=True)
daily = hasil["Deals - Daily.xlsx"]
assert daily["__sheets__"] == SHEETS
mt4_rows = [r for r in daily["Daily"][1:] if r[0] == "MT4"]
assert {r[4] for r in mt4_rows} == {"BTCUSD"}
logins = {r[1] for r in mt4_rows}
assert "603477" in logins and "20033032" not in logins and "20023282" not in logins
opened_earlier = [r for r in mt4_rows if r[1] == "603477"]
assert opened_earlier and opened_earlier[0][10] == 3.98
same_day = [r for r in mt4_rows if r[1] == "20038203"]
assert same_day and same_day[0][10] == 2.77
assert all(r[0] != "MT4" for r in daily["Daily - In"][1:])
assert all(r[1] != "603477" and r[1] != "20038203" for r in daily["Daily - In"][1:])
v = verification(daily["Verifikasi"])
assert v[K_BAL] == 2 and v[K_CAN] == 2 and v[K_FOOT] == 7
assert v[K_OUT] == 3 and v[K_DAILY_IN] == 2  # the two MT5 in rows only
assert v["Entry empty/other (dropped)"] == 3
assert "USC" in {r[11] for r in daily["Daily"][1:] if r[0] == "MT5"}
assert v[K_COMB_P] == round(v["Total Profit ('out' rows)"] + v[K_PROFIT], 2)
assert summary["dipakai"] == 2 and summary["mt4_dipakai"] == 3
print("OK trimmed KVB MT4+MT5 fixture")

# Formula-like Login stays literal. Platform is the first column.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    src = tmp / "formula.csv"
    write_mt4(src, [mt4_trade("1", login="=1+1", symbol="=eurusd", profit="1")], footer=False)
    _, _, out = run([src], allow_mt4=True)
    wb = load_workbook(io.BytesIO(zipfile.ZipFile(out).read("Deals - Daily.xlsx")))
    try:
        platform, login, _date, _type, symbol = list(wb["Daily"][2])[:5]
        assert platform.value == "MT4"
        assert login.value == "=1+1" and symbol.value == "=EURUSD"
        assert login.quotePrefix and login.data_type == "s"
        assert symbol.quotePrefix and symbol.data_type == "s"
    finally:
        wb.close()
print("OK MT4 formula text stays literal")


def wait(client, location):
    job_id = location.rstrip("/").rsplit("/", 1)[-1]
    for _ in range(200):
        state = __import__("app")._read_state(job_id)
        if state and state["state"] in ("done", "failed"):
            return job_id, state
        time.sleep(0.02)
    raise AssertionError("job timeout")


import app as A  # noqa: E402

A.app.config["TESTING"] = True
token = base64.b64encode(b"mt4-tester:mt4-tester-password-1").decode()
client = A.app.test_client()
client.environ_base["HTTP_AUTHORIZATION"] = f"Basic {token}"
client.environ_base["HTTP_ACCEPT"] = "text/html"

kvb_page = client.get("/kvb/tool/segregate")
assert kvb_page.status_code == 200
assert b"MT4 Raw Report is optional" in kvb_page.data
assert b"mt4-currency" in kvb_page.data
assert b'value="USC"' in kvb_page.data
assert b"separate sheets" in kvb_page.data
assert b"deleted from the server" in kvb_page.data
dpm_page = client.get("/tool/segregate")
assert dpm_page.status_code == 200
assert b"MT4 Raw Report is optional" not in dpm_page.data
assert b"deleted from the server" in dpm_page.data
assert b"mt4-currency" not in dpm_page.data
assert b"Drop Deals History files here" in dpm_page.data
kvb_home = client.get("/kvb")
assert b"MT4 Raw Report is optional" in kvb_home.data
assert b"Platform column" in kvb_home.data
print("OK KVB tool page offers USD/USC per MT4 file; DPM page does not")

mt4_body = (
    "Raw Report for 'abcc' from 2026.10.04 to 2026.10.04\n"
    + ";".join(MT4_HEADER) + "\n"
    + ";".join(["9", "100", "2026.10.04 01:00:00", "buy", "btcusd", "1", "", "", "",
                "2026.10.04 02:00:00", "", "", "", "", "", "", "0", "-1", "0", "0", "4", "", ""])
    + "\n"
).encode()

bad_currency = client.post(
    "/kvb/tool/segregate",
    data={"file": (io.BytesIO(mt4_body), "whatever.csv"),
          "mt4_currency": "whatever.csv|EUR"},
    content_type="multipart/form-data")
assert bad_currency.status_code == 400, bad_currency.status_code
assert b"USD or USC" in bad_currency.data

dpm = client.post("/tool/segregate",
                  data={"file": (io.BytesIO(mt4_body), "Deals History.csv"),
                        "mt4_currency": "Deals History.csv|USC"},
                  content_type="multipart/form-data")
assert dpm.status_code == 303, (dpm.status_code, dpm.data[:300])
_, dpm_state = wait(client, dpm.headers["Location"])
assert dpm_state["state"] == "failed", dpm_state
assert "MT5 Deals History only" in (dpm_state.get("error") or "")
assert "KVB" in (dpm_state.get("error") or "")

commands = []
real_run = A.subprocess.run


def capture(command, **kwargs):
    commands.append(command)
    return real_run(command, **kwargs)


A.subprocess.run = capture
try:
    kvb = client.post("/kvb/tool/segregate",
                      data={"file": (io.BytesIO(mt4_body), "whatever.csv"),
                            "mt4_currency": "whatever.csv|USC"},
                      content_type="multipart/form-data")
    assert kvb.status_code == 303, (kvb.status_code, kvb.data[:300])
    _, kvb_state = wait(client, kvb.headers["Location"])
    assert kvb_state["state"] == "done", kvb_state
    assert commands and "--allow-mt4" in commands[-1]
    assert "--mt4-currency" in commands[-1]
    assert "whatever.csv=USC" in commands[-1]
    assert Path(commands[-1][1]).name == "deal_segregator.py"
finally:
    A.subprocess.run = real_run
print("OK KVB runs MT4 with the chosen currency and DPM's job rejects it")

print("OK deal segregator MT4")
