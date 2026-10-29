"""Shared helpers for the test_*.py scripts. Not a test itself.

Two jobs:

* ``hermetic_env()`` points DW_AUTH_DB / DW_JOB_DIR / DW_ADMIN_* at a fresh
  temporary directory, by ASSIGNMENT. A test must never inherit the server's
  exported DW_* values, and must never open the real users.db or session.key.
  ``assert_isolated()`` proves it after ``app``/``auth`` are imported.
* Synthetic fixture builders. Every test input is generated here at test time,
  so the suite runs the same on a laptop, on the box and on the server, and no
  real client data is ever committed. The numbers are made up.
"""
import atexit
import csv
import datetime as dt
import os
import shutil
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parent


# --------------------------------------------------------------------------
# Hermetic environment
# --------------------------------------------------------------------------
def hermetic_env(admin_user="tester", admin_pass="tester-password-1", admin=True):
    """Fresh temp users DB, session key and jobs dir for this test process.

    Call BEFORE ``import app``. Returns the temp root (removed at exit).
    ``admin=False`` removes DW_ADMIN_* so no admin is bootstrapped.
    """
    root = Path(tempfile.mkdtemp(prefix="dw-test-"))
    atexit.register(shutil.rmtree, root, True)
    os.environ["DW_AUTH_DB"] = str(root / "users.db")
    os.environ["DW_JOB_DIR"] = str(root / "jobs")
    if admin:
        os.environ["DW_ADMIN_USER"] = admin_user
        os.environ["DW_ADMIN_PASS"] = admin_pass
    else:
        os.environ.pop("DW_ADMIN_USER", None)
        os.environ.pop("DW_ADMIN_PASS", None)
    if "auth" in sys.modules or "app" in sys.modules:
        raise RuntimeError("hermetic_env() must run before app/auth are imported")
    return root


def assert_isolated(root):
    """The imported app must use only files under ``root``."""
    root = Path(root).resolve()
    auth = sys.modules.get("auth")
    if auth is not None:
        for path in (Path(auth.DB).resolve(), Path(auth.SECRET_FILE).resolve()):
            assert root in path.parents, f"test is using {path}, not its temp dir {root}"
            assert path not in (REPO / "users.db", REPO / "session.key"), path
    app = sys.modules.get("app")
    if app is not None:
        assert root in Path(app.JOBS).resolve().parents, f"jobs dir {app.JOBS} is outside {root}"


# --------------------------------------------------------------------------
# DPM D&W source workbook (input of isi_template.py / hitung_dw.py)
# --------------------------------------------------------------------------
def write_dpm_dw_source(path, year=2026, month=6):
    """A small DPM D&W export: D, W, Xero rates and the D&W FEE table."""
    from openpyxl import Workbook

    DT = dt.datetime
    wb = Workbook()
    d = wb.active
    d.title = "D"
    d.append(["Source Name", "Currency", "Paid Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket"])
    d.append(["Deposit", "VND", DT(year, month, 3), 1000000, "77Pay", 40.0, 25000, "D1"])
    d.append(["Deposit", "PHP", DT(year, month, 10), 5500, "PA-PHP", 100, 55, "D2"])
    d.append(["Deposit", "PKR", DT(year, month, 28), 27800, "PA", 100, 278, "D3"])
    w = wb.create_sheet("W")
    w.append(["Source Name", "Currency", "Completed Date", "Transaction", "Payment Gateway",
              "USD", "Rate", "Ticket", "Status"])
    w.append(["Withdrawal", "VND", DT(year, month, 15), 250000, "77Pay", 10, 25000, "W1", "finish"])
    w.append(["Withdrawal", "PHP", DT(year, month, 20), 1100, "SHUNFAPAY", 20, 55, "W2", "finish"])
    x = wb.create_sheet("Xero")
    x.append(["Date", "VND", "PHP", "PKR", "USD"])
    day = DT(year, month, 1) - dt.timedelta(days=31)
    end = DT(year, month, 1) + dt.timedelta(days=61)
    while day <= end:
        x.append([day, 25000, 55, 278, 1])
        day += dt.timedelta(days=1)
    fee = wb.create_sheet("D&W FEE")
    fee.append([None, None])
    fee.append([None, None])
    fee.append(["Currency", "Payment Gateway", "Deposit", "Withdrawal"])
    for row in (["VND", "77Pay", 0.025, 0.0], ["PHP", "PA", 0.01, 0.005],
                ["PHP", "SHUNFAPAY", 0.01, 8], ["PKR", "PA", 0.02, 0.01]):
        fee.append(row)
    wb.save(path)
    wb.close()
    return Path(path)


def build_template(source, out):
    """Generate an empty 'Template D&W.xlsx' with buat_template.py (the real template is private)."""
    import subprocess

    subprocess.run([sys.executable, str(REPO / "buat_template.py"), str(source), "-o", str(out)],
                   cwd=REPO, check=True, capture_output=True, text=True)
    return Path(out)


# --------------------------------------------------------------------------
# Deals History CSV + Client Equity FX
# --------------------------------------------------------------------------
DEALS_HEADERS = ["Deal", "Position", "Login", "Group", "Country", "Time", "Type", "Entry",
                 "Symbol", "Volume", "Commission", "Fee", "Swap", "Profit", "Currency"]
DEALS_LOGINS = (100, 101, 102)
DEALS_DAYS = 31                      # July 2026
DEALS_OUT_PER_DAY = 2                # two closing deals per login per day


def deals_expected(logins=DEALS_LOGINS):
    return {
        # one Daily row per login per day ('out' rows aggregated)
        "daily_rows": len(logins) * DEALS_DAYS,
        # unique 'out' deals kept; uploading the same file twice must not change it
        "kept": len(logins) * DEALS_DAYS * DEALS_OUT_PER_DAY,
    }


DEALS_EXPECTED = deals_expected()


def write_deals_csv(path, currency="USD", logins=DEALS_LOGINS):
    """Synthetic MT5 Deals History: per login per day one 'in' and two 'out' deals."""
    with Path(path).open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(DEALS_HEADERS)
        n = 0
        for day in range(1, DEALS_DAYS + 1):
            for login in logins:
                entries = ["in"] + ["out"] * DEALS_OUT_PER_DAY
                for k, entry in enumerate(entries):
                    n += 1
                    w.writerow([n, f"P{login}-{day}-{k}", login, r"real\DPMKT-15", "MY",
                                f"2026.07.{day:02d} 0{k + 1}:02:03.004", "buy", entry, "EURUSD",
                                "1", "-2", "-1", "-4", "20" if entry == "out" else "0", currency])
    return Path(path)


def write_equity_fx(path, currency="USD"):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Query result"
    ws.append(["date", "Currency", "rate"])
    for day in range(1, DEALS_DAYS + 1):
        ws.append([dt.date(2026, 7, day), currency, 1])
    wb.save(path)
    wb.close()
    return Path(path)


# --------------------------------------------------------------------------
# PL by Desk: FinanceOS reconciliation (PDF + Excel) and Metabase export
# --------------------------------------------------------------------------
# (section code, section name, [(level, label, amount)]) -- a made-up August 2026
# reconciliation that adds up exactly like a real FinanceOS export.
PL_COMPANY = "DPM Malaysia Sdn. Bhd."
PL_MANUAL = [
    "Manual rebate IB partner VN batch one",
    ("Manual rebate IB partner campaign for the August promotion period",
     "settlement for desk GEM"),           # wraps onto a second line in the PDF
    "2026-08-21 [USD]",
    "Manual rebate IB partner TH",
    "Manual rebate IB partner MY",
    "Manual rebate IB partner ID",
    "Manual rebate IB partner EG",
]
PL_MANUAL_AMOUNTS = [-1000.00, -1200.00, -500.00, -750.00, -300.00, -250.00, -1000.00]
PL_SECTIONS = [
    ("A", "Trade", -4628109.21, [
        (1, "Closed P/L", -5276735.09),
        (1, "Floating P/L Changes", 312450.55),
        (1, "Swap", -18230.40),
        (1, "Swap (deals)", -5406.04),
        (1, "Commissions", 358561.77),
        (1, "CRM Swap Refund", 1250.00),
    ]),
    ("B", "Cash Movement", 3955539.65, [
        (1, "Deposit to MT5 Trading Account", 6120500.10),
        (1, "Withdrawal MT5 Trading Account", -1553610.45),
        (1, "Withdrawal from Clients' Wallet", -210300.00),
        (1, "Clients' Wallet transfer to MT5 trading account (test)", 100.00),
        (1, "Rebates (calculated by system)", -401150.00),
        (2, "Rebate", -402115.30),
        (2, "Rollback Rebate", 1015.30),
        (2, "Test", -50.00),
    ]),
    ("C", "Adjustment", -48115.40, [
        (1, "MT5 Trading Account Changes", -15830.25),
        (2, "Profit Sharing", -12000.00),
        (2, "Negative Balance Compensation", -3450.25),
        (2, "Other Adjustment", -380.00),
        (3, "Adjustment: Inactive", 820.00),
        (3, "Adjustment: Written Off", -1200.00),
        (1, "Rewards Wallet Changes", -35500.50),
        (2, "Reward Other Promotions", -33400.50),
        (3, "Welcome Bonus", -25000.00),
        (3, "Deposit Bonus", -8400.50),
        (2, "Swap Free", -2100.00),
        (1, "Rewards (Clients' Wallet)", -9000.00),
        (2, "TransferOut", -4000.00),
        (2, "Manual Rebate (Reward)", -5000.00),
        *[(3, label, amount) for label, amount in zip(PL_MANUAL, PL_MANUAL_AMOUNTS)],
        (1, "MT5 Transfer IN", 50000.00),
        (1, "MT5 Transfer OUT", -50000.00),
        (1, "Profit by sharing (Wallet's Profit Sharing)", 12000.00),
        (1, "Currency Rate Diff(USD)", 215.33),
        (2, "System: FX translation of the opening balance", 215.33),
        (1, "Rounding Adjustment", 0.02),
    ]),
]
PL_EQUITY = {                      # previous, current, change
    "Total": (11450000.00, 11848488.72, 398488.72),
    "MT5": (10250000.00, 10600000.00, 350000.00),
    "Wallet": (1200000.00, 1248488.72, 48488.72),
}
PL_TOTAL_MOVEMENT = -720684.96
PL_DIFF = round(PL_TOTAL_MOVEMENT - PL_EQUITY["Total"][2], 2)   # -1,119,173.68

# Metabase rows: code, deposit, withdrawal, ndp, lots, closed, commission, rebate.
# EG is not on the Desks sheet on purpose (auto-assigned to GEM).
PL_METABASE = [
    ("VN", 2000000.00, 500000.00, 1500000.00, 1200.5, -2000000.00, 120000.00, 150000.00),
    ("TH", 1500000.00, 400000.00, 1100000.00, 900.25, -1500000.00, 90000.00, 110000.00),
    ("MY", 1000000.00, 300000.00, 700000.00, 600.0, -900000.00, 60000.00, 70000.00),
    ("ID", 900000.00, 200000.00, 700000.00, 500.0, -500000.00, 50000.00, 40000.00),
    ("EG", 520000.00, 103127.72, 416872.28, 210.75, -273899.21, 26000.00, 22000.00),
    ("NG", 200000.00, 50000.00, 150000.00, 90.5, -100000.00, 12526.93, 10115.30),
]
PL_COUNTRY_NAMES = {"VN": "Vietnam", "TH": "Thailand", "MY": "Malaysia", "ID": "Indonesia",
                    "EG": "Egypt", "NG": "Nigeria"}
# Desk names across row 1 starting in column B, so column F (6) is GEM. "Thai " keeps a
# trailing space on purpose: the loader must strip it.
PL_DESKS = [("VN", ["VN"]), ("Thai ", ["TH"]), ("MY", ["MY"]), ("ID", ["ID"]), ("GEM", ["NG"])]


def _fmt(amount):
    return f"{amount:,.2f}"


def _manual_text(label):
    return label if isinstance(label, str) else " ".join(label)


def pl_finance_lines():
    """[(x, text)] for the FinanceOS PDF, in reading order."""
    lines = [
        (40, "Trading Account Reconciliation"),
        (40, PL_COMPANY),
        (40, "Period: 2026-08-01 to 2026-08-31"),
        (40, "Currency: CONSOLIDATED USD"),
        (40, "First day this currency rate is used: 2026-08-01 (synthetic test data)"),
        (40, "Equity / Balance"),
    ]
    for name, (prev, cur, chg) in PL_EQUITY.items():
        lines.append((40, f"{name}  Previous: {_fmt(prev)}  Current: {_fmt(cur)}  Change: {_fmt(chg)}"))
    for code, name, subtotal, items in PL_SECTIONS:
        lines.append((40, f"{code}.) {name}"))
        for level, label, amount in items:
            x = 60 + 14 * (level - 1)
            if isinstance(label, tuple):
                lines.append((x, f"- {label[0]}"))
                lines.append((x + 8, f"{label[1]}  {_fmt(amount)}"))
            else:
                lines.append((x, f"- {label}  {_fmt(amount)}"))
        lines.append((40, f"Sub-Total:  {_fmt(subtotal)}"))
    lines.append((40, f"Total Movement:  {_fmt(PL_TOTAL_MOVEMENT)}"))
    lines.append((40, f"Total Movement - Equity Change (Diff):  {_fmt(PL_DIFF)}"))
    return lines


def _pdf_escape(text):
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(path, lines, size=8, leading=11, top=810):
    """Minimal text PDF (Helvetica, one Tj per line at its own x). Paginates."""
    pages, current, y = [], [], top
    for x, text in lines:
        if y < 40:
            pages.append(current)
            current, y = [], top
        current.append(f"BT /F1 {size} Tf {x} {y} Td ({_pdf_escape(text)}) Tj ET")
        y -= leading
    pages.append(current)
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"]
    kids = []
    for content in pages:
        stream = "\n".join(content).encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        content_id = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                       f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>")
        kids.append(len(objects))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        body = body if isinstance(body, bytes) else body.encode("latin-1")
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    Path(path).write_bytes(bytes(out))
    return Path(path)


def write_pl_finance_pdf(path):
    return write_pdf(path, pl_finance_lines())


def write_pl_finance_excel(path):
    """The same reconciliation as FinanceOS's Excel export (indent = level)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"
    ws.append(["Trading Account Reconciliation"])
    ws.append(["Company", PL_COMPANY])
    ws.append(["Currency", "CONSOLIDATED USD"])
    ws.append(["Period", "2026-08-01 to 2026-08-31"])
    ws.append(["Rate note", "First day this currency rate is used: 2026-08-01 (synthetic test data)"])
    ws.append(["Equity / Balance", "Previous", "Current", "Change"])
    for name, values in PL_EQUITY.items():
        ws.append([name, *values])
    for code, name, subtotal, items in PL_SECTIONS:
        ws.append([f"{code}.) {name} Sub-Total", subtotal])
        for level, label, amount in items:
            ws.append([_manual_text(label), amount])
            ws.cell(ws.max_row, 1).alignment = Alignment(indent=level)
    ws.append(["Total Movement", PL_TOTAL_MOVEMENT])
    ws.append(["Total Movement - Equity Change (Diff)", PL_DIFF])
    wb.save(path)
    wb.close()
    return Path(path)


def write_pl_metabase(path, data_date=20260815):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Query result-Desk Rebate"
    ws.append(["data_date", "country_code", "deposit_usd", "withdrawal_usd", "ndp_usd", "lots",
               "closed_pnl_usd", "commission_usd", "rebate_usd"])
    for row in PL_METABASE:
        ws.append([data_date, *row])
    desks = wb.create_sheet("Desks")
    for col, (name, codes) in enumerate(PL_DESKS, start=2):
        desks.cell(1, col).value = name
        for offset, code in enumerate(codes, start=2):
            desks.cell(offset, col).value = code
    target = wb.create_sheet("Target")
    target.append(["Metric", *[name.strip() for name, _ in PL_DESKS]])
    for metric in ("Closed P/L", "Commissions", "Rebate", "Deposit", "Withdrawal", "NDP", "Volume"):
        target.append([metric])
    names = wb.create_sheet("Country code")
    names.append(["code", "name"])
    for code, name in PL_COUNTRY_NAMES.items():
        names.append([code, name])
    wb.save(path)
    wb.close()
    return Path(path)


def pl_metabase_total(index):
    """Sum of one Metabase column (index into the PL_METABASE tuples), to the cent."""
    return round(sum(row[index] for row in PL_METABASE), 2)
