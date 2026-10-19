#!/usr/bin/env python3
"""MT4 Raw Report (';') opened in Excel with a ',' list separator and saved as CSV.

Excel splits each line only on ',' (Comments such as "CRM-S-1333365,261003,Swap"),
keeps every ';' inside the cells, pads every row with ',' to the widest row and
saves with ','. The result is the original ';' line plus ",,,,," at the end:

    Raw Report for 'abcc' from 2026.08.04 to 2026.08.04,,,,,
    Deal;Login;...;Points;Comment,,,,,
    92134628;601651;...;-213.00;;CRM-W-256811;,,,,,

Before the fix the last cell (",,,,,") counted as data, so the first data row
stopped with "line break inside a field ... at row 3 (Login ...). the row has
more columns than the header." Totals must match the original file.

The real KVB reports are used when present (MT4_REAL_RAW, MT4_FULL_RAW,
MT4_ACCOUNT_TYPES); the committed fixtures always run.

python3 test_deal_segregator_mt4_comma_padding.py
"""
import io
import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

import deal_segregator as ds
from deal_segregator import proses

ROOT = Path(__file__).resolve().parent
FIX = ROOT / "fixtures" / "kvb"
FIX_RAW = FIX / "KVB_MT4_RawReport_2026-10-04.csv"
FIX_TRAILING = FIX / "KVB_MT4_RawReport_2026-07-01_trailing_delim.csv"
FIX_TYPES = FIX / "KVB_MT4_Account_Types_sample.csv"


def _default_full():
    d = Path("/home/box/agent-data/agents/24ca3b67-52c5-4d69-aa5a-b9f914126756/attachments")
    hits = sorted(d.glob("8763a29c*.csv")) if d.exists() else []
    return str(hits[0]) if hits else str(d / "8763a29c.csv")


REAL_RAW = Path(os.environ.get("MT4_REAL_RAW", "/workspace/kvb-mt4/KVB_MT4_RawReport_2026-10-04.csv"))
FULL_RAW = Path(os.environ.get("MT4_FULL_RAW", _default_full()))
REAL_TYPES = Path(os.environ.get("MT4_ACCOUNT_TYPES", "/workspace/kvb-mt4/KVB_MT4_Account_Types.xlsx"))

HEADER = ("Deal;Login;Open Time;Type;Symbol;Volume;Open Price;S/L;T/P;Close Time;"
          "Close Price;Reason;Gateway Order;Gateway Volume;Open Price Delta;"
          "Close Price Delta;Agent;Commission;Taxes;Swap;Profit;Points;Comment")
TITLE = "Raw Report for 'abcc' from 2026.08.04 to 2026.08.04"
TRADE = ("97888615;20010437;2026.08.04 17:51:18;sell;btcusd;0.01;85121.71;0;0;"
         "2026.08.04 23:38:24;85651.47;;;;;;0.00;-0.09;0.00;0.00;-5.30;-52976;{c}")
FOOTER = ";;;;;;;;;;;;;;;;;;Profit:;;-5.30"


def _read(src):
    raw = Path(src).read_bytes()
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252"), "cp1252"


def comma_locale_resave(src, dst, lebar=None):
    """Excel with ',' list separator: split on ',' only, pad to the widest row
    (or `lebar` cells), save with ','. -> width used."""
    teks, enc = _read(src)
    rows = [b.split(",") for b in teks.splitlines()]
    lebar = max(lebar or 0, max(len(r) for r in rows))
    Path(dst).write_bytes(("\r\n".join(",".join(r + [""] * (lebar - len(r))) for r in rows)
                           + "\r\n").encode(enc))
    return lebar


def _verifikasi(zip_path):
    out = {}
    with zipfile.ZipFile(zip_path) as z:
        for nama in z.namelist():
            wb = load_workbook(io.BytesIO(z.read(nama)), read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
                if ws.title == "Verifikasi":
                    rows = [r for r in rows if r and r[0] is not None and str(r[0]).startswith("MT4")]
                out[(nama, ws.title)] = rows
            wb.close()
    return out


def _run(path, types, out):
    proses([path], out, allow_mt4=True, mt4_accounts=types)
    return _verifikasi(out)


def _kunci(hasil):
    rows = next(v for (n, t), v in hasil.items() if t == "Verifikasi" and "Daily" in n)
    v = {str(r[0]): (r[1] if len(r) > 1 else None) for r in rows}
    return {k: v[k] for k in ("MT4 closed buy/sell kept", "MT4 Total Profit (closed buy/sell)",
                              "MT4 Total Commission (closed buy/sell)",
                              "MT4 Total Volume (closed buy/sell)",
                              "MT4 Total Profit raw (before USC ÷100)")}


def _bandingkan(src, types, lebar_list=(None, 6, 9)):
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        asli = d / "asli.csv"
        shutil.copy(src, asli)
        ref = _run(asli, types, d / "asli.zip")
        for lebar in lebar_list:
            p = d / f"MT4_RawReport_20260804_{lebar}.csv"
            w = comma_locale_resave(src, p, lebar)
            baris = p.read_text(encoding=_read(src)[1]).splitlines()
            assert w == 1 or baris[0].endswith(","), baris[0]
            info = ds._info_mt4(p)
            assert info and info["pemisah"] == ";" and info["header"][-1] == "Comment", info
            res = _run(p, types, d / f"{lebar}.zip")
            assert _kunci(res) == _kunci(ref), (w, _kunci(res), _kunci(ref))
            assert res == ref, f"width {w}: workbook differs from the original"
        return _kunci(ref)


def _tulis(d, lines, nama="MT4_RawReport_20260804.csv"):
    p = Path(d) / nama
    p.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8"))
    return p


def _peta(d):
    p = Path(d) / "types.csv"
    p.write_text("account,account_type\n", encoding="utf-8")
    return p


def _stop(p, d):
    try:
        proses([p], Path(d) / "o.zip", allow_mt4=True, mt4_accounts=_peta(d))
    except SystemExit as e:
        return str(e)
    return None


# ----------------------------------------------------------------------- tests
def test_buang_padding_unit():
    f = ds._buang_padding_mt4
    assert f(["a", ",,,,,"], ";") == ["a", ""]
    assert f(["a", "-1899.71,,,,"], ";") == ["a", "-1899.71"]
    assert f(["a", "Comment,,,,,"], ";") == ["a", "Comment"]
    assert f(["a", "x;;;", ], ",") == ["a", "x"]
    assert f(["a", "\t\t"], ";") == ["a", ""]
    assert f(["a", "b"], ";") == ["a", "b"]
    assert f(["a", ",,x"], ";") == ["a", ",,x"]          # data after the padding stays
    assert f(["a", ";;"], ";") == ["a", ";;"]            # own delimiter is not touched here
    assert f([], ";") == []


def test_exact_user_error_variant():
    """Comma-padded ';' report: the row-3 STOP from the user is gone."""
    with tempfile.TemporaryDirectory() as d:
        lines = [TITLE + ",,,,,", HEADER + ",,,,,",
                 TRADE.format(c="[sl];") + ",,,,,",
                 TRADE.replace("97888615", "97888616").format(c="a,b;") + ",,,",
                 FOOTER + ",,,,,"]
        p = _tulis(d, lines)
        assert _stop(p, d) is None


def test_fixture_comma_locale_resave():
    _bandingkan(FIX_RAW, FIX_TYPES)


def test_trailing_fixture_comma_locale_resave():
    v = _bandingkan(FIX_TRAILING, FIX_TYPES)
    assert v["MT4 closed buy/sell kept"] == 5
    assert round(v["MT4 Total Profit (closed buy/sell)"], 2) == 77.84


def test_one_extra_trailing_empty_cell():
    with tempfile.TemporaryDirectory() as d:
        for extra in (";", ";;", ";;;;;", "; ;\t;", ";,,", ";;,,,"):
            p = _tulis(d, [TITLE, HEADER, TRADE.format(c="[sl]") + extra, FOOTER + extra])
            assert _stop(p, d) is None, extra


def test_extra_non_empty_cell_still_stops():
    with tempfile.TemporaryDirectory() as d:
        # Deal/Login not numbers: the extra cell cannot be the Comment
        p = _tulis(d, [TITLE, HEADER, TRADE.replace("20010437", "abc").format(c="[sl];x") + ",,,"])
        msg = _stop(p, d)
        assert msg and "more columns than the header" in msg, msg
        # Comment is not the last header column (currency layout): no rejoin
        hdr = HEADER.replace(";", ",") + ",,"
        row = TRADE.replace(";", ",").format(c="[sl]") + ",,USD,extra"
        p = _tulis(d, [hdr, row], "deals.csv")
        msg = _stop(p, d)
        assert msg and "more columns than the header" in msg, msg


def test_comment_with_delimiter_is_rejoined():
    with tempfile.TemporaryDirectory() as d:
        p = _tulis(d, [TITLE + ",,,", HEADER + ",,,", TRADE.format(c="a;b;c;") + ",,,", FOOTER + ",,,"])
        assert _stop(p, d) is None


def _real(path):
    if not path.exists() or not REAL_TYPES.exists():
        print(f"SKIP: {path} or {REAL_TYPES} not present")
        return False
    return True


def test_real_raw_report():
    if _real(REAL_RAW):
        v = _bandingkan(REAL_RAW, REAL_TYPES)
        print("  real 2026-10-04:", v)


def test_full_raw_report():
    if _real(FULL_RAW):
        v = _bandingkan(FULL_RAW, REAL_TYPES, (None, 6))
        assert v["MT4 closed buy/sell kept"] == 24325
        assert round(v["MT4 Total Profit (closed buy/sell)"], 2) == -286498.73
        print("  full 2026-07-01:", v)


if __name__ == "__main__":
    for nama, f in list(globals().items()):
        if nama.startswith("test_") and callable(f):
            f()
            print("OK", nama)
    print("OK")
